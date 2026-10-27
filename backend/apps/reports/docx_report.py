"""The Word (DOCX) report, laid out like the PDF (pdf_report.py): a title band, headline score table, pillar table
with status colours, gap and recommendation tables with priority tags, the action timeline and InGage support.
Short labels instead of paragraphs; every value comes from the report's frozen `report_data`.
"""
import io
from datetime import datetime

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

BRAND = RGBColor(99, 130, 0)
INK = RGBColor(52, 51, 53)
MUTED = RGBColor(110, 110, 115)
GREEN = RGBColor(0, 130, 54)
AMBER = RGBColor(166, 95, 0)
RED = RGBColor(193, 0, 7)
BRAND_HEX, LIGHT_HEX, ROW_HEX = '638200', 'F0F7DC', 'FAFAF8'

STATUS = {'strong': ('Strong', GREEN), 'developing': ('Developing', BRAND), 'at_risk': ('At risk', AMBER),
          'not_started': ('Not started', RED)}
PRIORITY = {'blocking': ('BLOCKING', RED), 'high': ('HIGH', AMBER), 'medium': ('MEDIUM', BRAND),
            'optional': ('OPTIONAL', MUTED)}


def _date(value):
    if not value:
        return ''
    try:
        return datetime.fromisoformat(str(value).replace('Z', '+00:00')).strftime('%d %b %Y')
    except ValueError:
        return str(value)[:10]


def _shade(cell, hex_colour):
    props = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_colour)
    props.append(shd)


def _write(cell, text, *, bold=False, colour=INK, size=9, align=None):
    cell.text = ''
    para = cell.paragraphs[0]
    run = para.add_run(str(text if text is not None else ''))
    run.bold, run.font.size, run.font.color.rgb = bold, Pt(size), colour
    if align:
        para.alignment = align
    return run


def _heading(doc, text, subtitle=''):
    para = doc.add_paragraph()
    para.paragraph_format.space_before, para.paragraph_format.space_after = Pt(12), Pt(2)
    run = para.add_run(text)
    run.bold, run.font.size, run.font.color.rgb = True, Pt(14), INK
    if subtitle:
        sub = doc.add_paragraph()
        sub.paragraph_format.space_after = Pt(4)
        r = sub.add_run(subtitle)
        r.font.size, r.font.color.rgb = Pt(8.5), MUTED


def _table(doc, headers, rows, widths=None):
    """rows: [[(text, colour, bold), ...]] -- a header band in the brand colour and light zebra rows."""
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, h in zip(table.rows[0].cells, headers):
        _shade(cell, LIGHT_HEX)
        _write(cell, h, bold=True, colour=BRAND, size=8)
    for n, row in enumerate(rows):
        cells = table.add_row().cells
        for cell, (text, colour, bold) in zip(cells, row):
            _write(cell, text, bold=bold, colour=colour)
            if n % 2:
                _shade(cell, ROW_HEX)
    if widths:
        for row in table.rows:
            for cell, w in zip(row.cells, widths):
                cell.width = w
    return table


def _note(doc, text, colour=MUTED, bold=False):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.font.size, run.font.color.rgb, run.bold = Pt(8.5), colour, bold
    return para


def build_docx(report):
    data = report.report_data
    doc = Document()
    for section in doc.sections:
        section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Pt(42)

    # Title band
    band = doc.add_table(rows=1, cols=2)
    left, right = band.rows[0].cells
    for cell in (left, right):
        _shade(cell, BRAND_HEX)
    _write(left, 'AI READINESS DISCOVERY REPORT', bold=True, colour=RGBColor(240, 247, 220), size=8)
    name = left.add_paragraph().add_run(data['institution'].get('name', ''))
    name.bold, name.font.size, name.font.color.rgb = True, Pt(20), RGBColor(255, 255, 255)
    meta = left.add_paragraph().add_run(f"Sprint {data['sprint'].get('sprint_code', '')}   |   Version {report.version}"
                                        f"   |   {_date(data.get('generated_at'))}")
    meta.font.size, meta.font.color.rgb = Pt(9), RGBColor(240, 247, 220)
    baseline = (data.get('baseline') or {}).get('status') or 'not submitted'
    _write(right, f'BASELINE: {str(baseline).upper()}', bold=True, colour=RGBColor(255, 255, 255), size=9,
           align=WD_ALIGN_PARAGRAPH.RIGHT)

    source = data.get('cri_source') or {}
    if source.get('baseline_scoring_run_id'):
        conf = source.get('baseline_confidence')
        _note(doc, f"Approved baseline ({source.get('baseline_status')}): CRI {source.get('baseline_cri') or 0:.1f}"
                   + (f', confidence {conf * 100:.0f}%' if conf is not None else '')
                   + f", calculated {_date(source.get('baseline_calculated_at'))}"
                   + f"   |   Live CRI in this report: {data.get('overall_cri', 0):.1f}", colour=BRAND, bold=True)
    processing = data.get('processing') or {}
    if processing and not processing.get('complete', True):
        _note(doc, '! Document processing incomplete: some facts may be missing.', colour=AMBER, bold=True)

    # Headline scores
    metrics = data.get('evidence_metrics') or {}
    adjusted = data.get('coverage_adjusted_cri')
    tiles = doc.add_table(rows=2, cols=4)
    labels = ['LIVE CRI', 'COVERAGE-ADJUSTED', 'EVIDENCE CONFIDENCE', 'OPEN DATA GAPS']
    values = [f"{data.get('overall_cri', 0):.1f} / 100", '-' if adjusted is None else f'{adjusted:.1f} / 100',
              f"{(data.get('confidence_score') or 0) * 100:.0f}%", str(len(data.get('missing_data_appendix') or []))]
    for i, (label, value) in enumerate(zip(labels, values)):
        _write(tiles.rows[0].cells[i], label, bold=True, colour=MUTED, size=7.5)
        _write(tiles.rows[1].cells[i], value, bold=True, size=16,
               colour=RED if label == 'OPEN DATA GAPS' and value != '0' else INK)
    strong = len(data.get('strengths') or [])
    facts = (metrics.get('confirmed_facts') or 0) + (metrics.get('corrected_facts') or 0)
    _note(doc, f"{strong} of {len(data.get('pillar_scorecards') or [])} pillars strong   |   {facts} confirmed facts   |   "
               f"{metrics.get('documents_processed', 0)} documents processed")

    # Pillars
    _heading(doc, 'Pillar scores', 'Score out of 100 per AI readiness pillar')
    pillars = data.get('pillar_scorecards') or []
    if pillars:
        rows = []
        for p in pillars:
            label, colour = STATUS.get(p.get('status'), (str(p.get('status') or ''), MUTED))
            score = p.get('raw_score') or 0
            bar = '█' * int(round(score / 10)) + '░' * (10 - int(round(score / 10)))
            rows.append([(p.get('label') or p.get('pillar'), INK, False), (bar, colour, False),
                         (f'{score:.0f}', INK, True), (label, colour, True)])
        _table(doc, ['PILLAR', 'SCORE', '', 'STATUS'], rows)
    else:
        _note(doc, 'No pillar scores yet.')

    # Gaps
    _heading(doc, 'Priority data gaps', 'Documents or data still needed')
    gaps = data.get('missing_data_appendix') or []
    if gaps:
        with_pillar = any(g.get('pillar_label') for g in gaps)
        rows = [[(PRIORITY.get(g.get('priority'), (str(g.get('priority', '')).upper(), MUTED))[0],
                  PRIORITY.get(g.get('priority'), ('', MUTED))[1], True), (g.get('title'), INK, False)]
                + ([(g.get('pillar_label') or '-', MUTED, False)] if with_pillar else []) for g in gaps]
        _table(doc, ['PRIORITY', 'GAP'] + (['PILLAR'] if with_pillar else []), rows)
    else:
        _note(doc, 'No open data gaps.')

    # Recommendations
    _heading(doc, 'Recommendations', 'Sorted by expected CRI lift')
    recs = sorted(data.get('recommendations') or [], key=lambda r: -(r.get('expected_cri_lift') or 0))
    if recs:
        _table(doc, ['PRIORITY', 'ACTION', 'CRI LIFT', 'TIMELINE'], [[
            (PRIORITY.get(r.get('priority'), (str(r.get('priority', '')).upper(), MUTED))[0],
             PRIORITY.get(r.get('priority'), ('', MUTED))[1], True),
            (r.get('title'), INK, False), (f"+{r.get('expected_cri_lift') or 0:.1f}", GREEN, True),
            (r.get('timeline') or '-', MUTED, False)] for r in recs])
    else:
        _note(doc, 'No recommendations generated yet.')

    # Timeline
    _heading(doc, 'Action timeline', 'What to do first, from the recommendations above')
    buckets = list(data.get('ninety_day_action_plan') or []) + list(data.get('twelve_month_roadmap') or [])
    if buckets:
        cols = buckets[:4]
        timeline = doc.add_table(rows=2, cols=len(cols))
        for i, b in enumerate(cols):
            head = timeline.rows[0].cells[i]
            _shade(head, BRAND_HEX)
            _write(head, b.get('timeline') or 'Unscheduled', bold=True, colour=RGBColor(255, 255, 255), size=9,
                   align=WD_ALIGN_PARAGRAPH.CENTER)
            body = timeline.rows[1].cells[i]
            body.text = ''
            for n, item in enumerate(b.get('items') or []):
                para = body.paragraphs[0] if n == 0 else body.add_paragraph()
                colour = PRIORITY.get(item.get('priority'), ('', MUTED))[1]
                mark = para.add_run('▌ ')
                mark.font.color.rgb, mark.font.size = colour, Pt(8)
                run = para.add_run(item.get('title', ''))
                run.font.size, run.font.color.rgb = Pt(8.5), INK
    else:
        _note(doc, 'No actions scheduled yet.')

    # Support
    _heading(doc, 'How InGage can help')
    offers = data.get('how_ingage_can_help') or []
    if offers:
        _table(doc, ['OFFERING', 'RECOMMENDATIONS'], [[(o.get('offering'), INK, True),
                                                      (str(o.get('recommendation_count', 0)), BRAND, False)]
                                                     for o in offers])
    else:
        _note(doc, 'No support offerings linked yet.')

    coverage = data.get('evidence_coverage') or {}
    if coverage:
        _note(doc, f"Coverage-adjusted CRI = CRI x AI-QIF evidence coverage. {coverage.get('summary', '')}")

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
