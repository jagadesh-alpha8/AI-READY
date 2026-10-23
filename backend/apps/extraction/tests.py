import io
import unittest
from datetime import timedelta
from unittest.mock import MagicMock, patch

import httpx
from fpdf import FPDF
from openai import APIConnectionError, APIStatusError, APITimeoutError, RateLimitError
import anthropic as anthropic_sdk
from PIL import Image

from django.conf import settings
from django.core.exceptions import ImproperlyConfigured
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase, TestCase, override_settings
from django.utils import timezone
from rest_framework import status
from rest_framework.test import APITestCase

from apps.accounts.models import User
from apps.documents.models import Document
from apps.facts.models import ExtractedFact
from apps.gaps.models import GapItem
from apps.institutions.models import Institution
from apps.sprints.models import Sprint

from .exceptions import AIResponseError, PermanentExtractionError, RecoverableExtractionError
from .models import ExtractionJob
from .services.ai_service import detect_provider, get_ai_service
from .services.anthropic_client import AnthropicExtractionService, AnthropicResponseError
from .services.conflict_checker import ConflictValidationError, OpenAIConflictChecker
from .services.gap_detector import RuleBasedGapDetector
from .services.openai_classifier import (
    CLASSIFICATION_SCHEMA, ClassificationValidationError, OpenAIDocumentClassifier,
)
from .services.openai_client import OpenAIExtractionService, OpenAIResponseError
from .services.openai_fact_extractor import OpenAIFactExtractor
from .services.pdf_reader import PDFPageReader
from .services.text_cleaning import clean_text, is_meaningful_text, remove_repeated_headers_footers
from .tasks import _handle_recoverable, run_extraction_job

PASSWORD = 'Str0ng!DevPassw0rd'

def route_combined_analysis(test):
    """The pipeline's first AI request now classifies AND extracts the first chunk's facts
    (services/document_analysis.py). Route it to the two existing mocks -- classification from
    the classifier's, facts from the fact extractor's -- so a test keeps steering each part the
    way it always has; later chunks go to the fact extractor's mock as before."""
    from apps.extraction.services import openai_classifier, openai_fact_extractor

    def answer(**kwargs):
        if kwargs.get('schema_name') == 'document_analysis':
            classification = openai_classifier.get_ai_service().extract_structured_data(**kwargs)
            facts = openai_fact_extractor.get_ai_service().extract_structured_data(**kwargs)
            return {'classification': classification, **(facts if isinstance(facts, dict) else {'facts': facts})}
        return openai_fact_extractor.get_ai_service().extract_structured_data(**kwargs)

    patcher = patch('apps.extraction.services.document_analysis.get_ai_service')
    test.analysis_service = patcher.start().return_value
    test.analysis_service.extract_structured_data.side_effect = answer
    test.addCleanup(patcher.stop)


#: A classification result the real OpenAIDocumentClassifier would accept --
#: used to fake a plausible AI response wherever a test just needs the
#: pipeline to get past the (now real, OpenAI-backed) classification step
#: without asserting anything about classification itself.
FAKE_CLASSIFICATION_RESULT = {
    'document_type': 'naac_ssr',
    'document_title': 'Self Study Report',
    'reporting_year': None,
    'institution_name': None,
    'confidence': 0.9,
    'reasoning': 'Test fixture classification result.',
}


def make_pdf_bytes(page_texts=('Sample document text for extraction pipeline tests.',)):
    """A real, pdfplumber-readable PDF -- one page per string in
    `page_texts`; an empty string produces a page with no text drawn on it
    (simulating a blank/scanned page)."""
    pdf = FPDF()
    for text in page_texts:
        pdf.add_page()
        if text:
            pdf.set_font('Helvetica', size=12)
            pdf.cell(0, 10, text)
    return bytes(pdf.output())


def make_scanned_pdf_bytes():
    """A PDF page with an embedded image and no real text layer -- what an
    actual scan looks like to a text extractor."""
    image = Image.new('RGB', (40, 40), color='white')
    buf = io.BytesIO()
    image.save(buf, format='PNG')
    buf.seek(0)
    pdf = FPDF()
    pdf.add_page()
    pdf.image(buf, x=10, y=10, w=20)
    return bytes(pdf.output())


def make_pdf(name='doc.pdf', page_texts=('Sample document text for extraction pipeline tests.',)):
    return SimpleUploadedFile(name, make_pdf_bytes(page_texts), content_type='application/pdf')


class ExtractionJobTestBase(APITestCase):
    def setUp(self):
        # These tests are about the review flow of unreviewed facts, and their fixture facts are 96% confident;
        # auto-confirmation (apps/facts/auto_confirm.py) is switched off here and tested in AutoConfirmTests.
        no_auto = override_settings(FACT_AUTO_CONFIRM_MIN_CONFIDENCE=None)
        no_auto.enable()
        self.addCleanup(no_auto.disable)
        self.institution = Institution.objects.create(name='MKCE', city='Karur', state='TN')
        self.other_institution = Institution.objects.create(name='Other', city='X', state='Y')

        self.iqac = self._make_user('iqac@test.edu', User.Role.IQAC_COORDINATOR, self.institution)
        self.outsider = self._make_user('outsider@test.edu', User.Role.HOD, self.other_institution)

        self.sprint = Sprint.objects.create(institution=self.institution, mode=Sprint.SprintMode.VERIFIED_CRI)

        # ExtractionPipeline's default classifier and fact extractor are both
        # real (OpenAI-backed) as of this task -- these tests exercise the
        # pipeline's orchestration, not classification/extraction themselves,
        # so both underlying OpenAI calls are mocked here once for every test
        # in this file that runs the real pipeline. Tests about classification/
        # extraction behavior itself live in OpenAIDocumentClassifierTests and
        # OpenAIFactExtractorTests, further down, with their own mocks.
        classifier_patcher = patch('apps.extraction.services.openai_classifier.get_ai_service')
        mock_classifier_openai_cls = classifier_patcher.start()
        mock_classifier_openai_cls.return_value.extract_structured_data.return_value = dict(FAKE_CLASSIFICATION_RESULT)
        self.addCleanup(classifier_patcher.stop)

        # Defaults to "the AI found nothing" -- individual tests that care
        # about persisted fact content override
        # self.fact_openai_mock.return_value.extract_structured_data.return_value.
        fact_patcher = patch('apps.extraction.services.openai_fact_extractor.get_ai_service')
        self.fact_openai_mock = fact_patcher.start()
        self.fact_openai_mock.return_value.extract_structured_data.return_value = {'facts': []}
        self.addCleanup(fact_patcher.stop)
        route_combined_analysis(self)

        self._auth(self.iqac)
        upload = self.client.post(
            f'/api/v1/sprints/{self.sprint.id}/upload-file/',
            {'file': make_pdf(), 'document_type': 'naac_ssr'}, format='multipart',
        )
        self.document = Document.objects.get(id=upload.data['id'])
        self.sprint.refresh_from_db()  # now 'collecting'
        self.client.credentials()  # tests authenticate explicitly as whichever persona they need

    @staticmethod
    def _make_user(email, role, institution):
        return User.objects.create_user(
            email=email, username=email.split('@')[0], password=PASSWORD,
            first_name='Test', last_name=role, role=role, institution=institution,
        )

    def _auth(self, user):
        login = self.client.post('/api/v1/auth/login/', {'email': user.email, 'password': PASSWORD})
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {login.data["access_token"]}')


@override_settings(CELERY_TASK_ALWAYS_EAGER=True, CELERY_TASK_EAGER_PROPAGATES=True)
class JobCreationTests(ExtractionJobTestBase):
    def test_post_creates_job_for_eligible_document(self):
        self._auth(self.iqac)
        response = self.client.post(f'/api/v1/sprints/{self.sprint.id}/extraction-jobs/')
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(len(response.data), 1)
        self.assertEqual(str(response.data[0]['document_id']), str(self.document.id))
        self.assertEqual(str(response.data[0]['sprint_id']), str(self.sprint.id))

    def test_post_without_trailing_slash_matches_frontend_client(self):
        self._auth(self.iqac)
        response = self.client.post(f'/api/v1/sprints/{self.sprint.id}/extraction-jobs')
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)

    def test_post_does_not_duplicate_jobs_for_already_processed_documents(self):
        self._auth(self.iqac)
        self.client.post(f'/api/v1/sprints/{self.sprint.id}/extraction-jobs/')
        response = self.client.post(f'/api/v1/sprints/{self.sprint.id}/extraction-jobs/')
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data, [])
        self.assertEqual(ExtractionJob.objects.filter(document=self.document).count(), 1)

    def test_explicit_repeat_after_completion_does_not_create_another_job(self):
        self._auth(self.iqac)
        payload = {'document_ids': [str(self.document.id)]}
        first = self.client.post(f'/api/v1/sprints/{self.sprint.id}/extraction-jobs/', payload, format='json')
        self.assertEqual(first.status_code, status.HTTP_201_CREATED, first.data)
        self.assertEqual(first.data[0]['status'], ExtractionJob.Status.COMPLETED)

        again = self.client.post(f'/api/v1/sprints/{self.sprint.id}/extraction-jobs/', payload, format='json')
        self.assertEqual(again.status_code, status.HTTP_201_CREATED, again.data)
        self.assertEqual(again.data, [])
        self.assertEqual(ExtractionJob.objects.filter(document=self.document).count(), 1)

    def test_legacy_processed_document_without_a_sprint_job_is_still_queued(self):
        self._auth(self.iqac)
        self.document.status = Document.Status.PROCESSED
        self.document.save(update_fields=['status', 'updated_at'])

        response = self.client.post(f'/api/v1/sprints/{self.sprint.id}/extraction-jobs/')
        self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.data)
        self.assertEqual(len(response.data), 1)
        self.assertEqual(str(response.data[0]['document_id']), str(self.document.id))

    def test_same_content_renamed_in_another_sprint_gets_its_own_job_and_facts(self):
        self._auth(self.iqac)
        self.fact_openai_mock.return_value.extract_structured_data.return_value = {'facts': [_valid_fact()]}

        first_response = self.client.post(
            f'/api/v1/sprints/{self.sprint.id}/extraction-jobs/', {'document_id': str(self.document.id)},
        )
        self.assertEqual(first_response.status_code, status.HTTP_201_CREATED, first_response.data)
        self.assertEqual(self.sprint.facts.filter(field_key='total_faculty').count(), 1)

        second_sprint = Sprint.objects.create(institution=self.institution, mode=Sprint.SprintMode.QUICK_CRI)
        self.document.file.open('rb')
        identical_bytes = self.document.file.read()
        self.document.file.close()
        upload = self.client.post(
            f'/api/v1/sprints/{second_sprint.id}/upload-file/',
            {'file': SimpleUploadedFile('renamed-ssr.pdf', identical_bytes, content_type='application/pdf'),
             'document_type': 'naac_ssr'}, format='multipart',
        )
        self.assertEqual(upload.status_code, status.HTTP_201_CREATED, upload.data)
        second_document = Document.objects.get(id=upload.data['id'])
        self.assertEqual(second_document.checksum, self.document.checksum)

        second_response = self.client.post(
            f'/api/v1/sprints/{second_sprint.id}/extraction-jobs/', {'document_id': str(second_document.id)},
        )
        self.assertEqual(second_response.status_code, status.HTTP_201_CREATED, second_response.data)
        second_job = ExtractionJob.objects.get(document=second_document)
        self.assertNotEqual(second_job.id, first_response.data[0]['id'])
        self.assertEqual(second_job.sprint_id, second_sprint.id)
        self.assertEqual(second_sprint.facts.filter(field_key='total_faculty').count(), 1)
        self.assertEqual(
            list(ExtractedFact.objects.filter(field_key='total_faculty').values_list('sprint_id', flat=True)),
            [second_sprint.id, self.sprint.id],
        )

    def test_retry_of_a_failed_document_reuses_its_processing_record(self):
        self._auth(self.iqac)
        initial = self.client.post(f'/api/v1/sprints/{self.sprint.id}/extraction-jobs/')
        self.assertEqual(initial.status_code, status.HTTP_201_CREATED, initial.data)
        job = ExtractionJob.objects.get(document=self.document)
        job.status = ExtractionJob.Status.FAILED
        job.save(update_fields=['status', 'updated_at'])

        retried = self.client.post(f'/api/v1/extraction-jobs/sprints/{self.sprint.id}/process-incomplete')
        self.assertEqual(retried.status_code, status.HTTP_200_OK, retried.data)
        self.assertEqual(len(retried.data), 1)
        self.assertEqual(retried.data[0]['id'], str(job.id))
        self.assertEqual(ExtractionJob.objects.filter(document=self.document).count(), 1)

    def test_post_with_explicit_document_id_targets_only_that_document(self):
        self._auth(self.iqac)
        response = self.client.post(
            f'/api/v1/sprints/{self.sprint.id}/extraction-jobs/', {'document_id': str(self.document.id)},
        )
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(len(response.data), 1)

    @patch('apps.extraction.views.run_extraction_job')
    def test_post_with_document_ids_extracts_only_the_selected_documents(self, mock_task):
        mock_task.delay.return_value = MagicMock(id='task-1')
        self._auth(self.iqac)
        unselected = Document.objects.create(
            sprint=self.sprint, document_type='aqar_report', original_filename='aqar.pdf',
            mime_type='application/pdf', status=Document.Status.UPLOADED,
        )
        response = self.client.post(
            f'/api/v1/sprints/{self.sprint.id}/extraction-jobs/',
            {'document_ids': [str(self.document.id)]}, format='json',
        )
        self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.data)
        self.assertEqual(len(response.data), 1)
        self.assertTrue(ExtractionJob.objects.filter(document=self.document).exists())
        self.assertFalse(ExtractionJob.objects.filter(document=unselected).exists())

        # Selecting it again while its job is still pending queues nothing new.
        again = self.client.post(
            f'/api/v1/sprints/{self.sprint.id}/extraction-jobs/',
            {'document_ids': [str(self.document.id)]}, format='json',
        )
        self.assertEqual(again.data, [])

    def test_post_with_document_ids_outside_sprint_is_not_found(self):
        self._auth(self.iqac)
        other_sprint = Sprint.objects.create(institution=self.institution, mode=Sprint.SprintMode.QUICK_CRI)
        response = self.client.post(
            f'/api/v1/sprints/{other_sprint.id}/extraction-jobs/',
            {'document_ids': [str(self.document.id)]}, format='json',
        )
        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)

    def test_post_with_document_id_not_in_sprint_is_not_found(self):
        self._auth(self.iqac)
        other_sprint = Sprint.objects.create(institution=self.institution, mode=Sprint.SprintMode.QUICK_CRI)
        response = self.client.post(
            f'/api/v1/sprints/{other_sprint.id}/extraction-jobs/', {'document_id': str(self.document.id)},
        )
        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)

    def test_job_listing_scoped_to_sprint(self):
        self._auth(self.iqac)
        self.client.post(f'/api/v1/sprints/{self.sprint.id}/extraction-jobs/')
        response = self.client.get(f'/api/v1/sprints/{self.sprint.id}/extraction-jobs/')
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(len(response.data), 1)

    # -- unauthorized access -------------------------------------------------

    def test_unauthenticated_create_rejected(self):
        response = self.client.post(f'/api/v1/sprints/{self.sprint.id}/extraction-jobs/')
        self.assertEqual(response.status_code, status.HTTP_401_UNAUTHORIZED)

    def test_create_from_outside_institution_rejected(self):
        self._auth(self.outsider)
        response = self.client.post(f'/api/v1/sprints/{self.sprint.id}/extraction-jobs/')
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(ExtractionJob.objects.count(), 0)

    def test_list_from_outside_institution_rejected(self):
        self._auth(self.outsider)
        response = self.client.get(f'/api/v1/sprints/{self.sprint.id}/extraction-jobs/')
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_detail_from_outside_institution_rejected(self):
        self._auth(self.iqac)
        create = self.client.post(f'/api/v1/sprints/{self.sprint.id}/extraction-jobs/')
        job_id = create.data[0]['id']

        self._auth(self.outsider)
        response = self.client.get(f'/api/v1/extraction-jobs/{job_id}/')
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_detail_accessible_to_own_institution(self):
        self._auth(self.iqac)
        create = self.client.post(f'/api/v1/sprints/{self.sprint.id}/extraction-jobs/')
        job_id = create.data[0]['id']
        response = self.client.get(f'/api/v1/extraction-jobs/{job_id}/')
        self.assertEqual(response.status_code, status.HTTP_200_OK)


@override_settings(CELERY_TASK_ALWAYS_EAGER=True, CELERY_TASK_EAGER_PROPAGATES=True)
class SuccessfulRunTests(ExtractionJobTestBase):
    """Exercises the real pipeline end to end (OpenAI calls mocked at the
    class/document level, everything else genuinely real) -- 'completed
    jobs' and 'task state transitions' coverage."""

    def test_job_completes_through_all_seven_steps(self):
        self._auth(self.iqac)
        response = self.client.post(f'/api/v1/sprints/{self.sprint.id}/extraction-jobs/')
        job = ExtractionJob.objects.get(id=response.data[0]['id'])

        self.assertEqual(job.status, ExtractionJob.Status.COMPLETED)
        self.assertEqual(job.current_step, ExtractionJob.Step.PREPARING_REVIEW_WORKSPACE)
        self.assertEqual(job.progress_percentage, 100)
        self.assertIsNotNone(job.started_at)
        self.assertIsNotNone(job.completed_at)
        self.assertEqual(job.error_message, '')
        self.assertEqual(job.retry_count, 0)

    def test_document_marked_processed_on_completion(self):
        self._auth(self.iqac)
        self.client.post(f'/api/v1/sprints/{self.sprint.id}/extraction-jobs/')
        self.document.refresh_from_db()
        self.assertEqual(self.document.status, Document.Status.PROCESSED)
        self.assertIsNotNone(self.document.processed_at)
        self.assertEqual(self.document.processing_status, 'classified')

    def test_ocr_required_is_corrected_from_the_upload_time_guess(self):
        """Every PDF is flagged ocr_required=True at upload time purely by
        file extension (apps.documents.constants.OCR_REQUIRED_EXTENSIONS),
        before any content is read. Once the real PDFPageReader actually
        reads this document's real, extractable text, that stale guess must
        be corrected to False -- not left misleadingly True forever."""
        self.assertTrue(self.document.ocr_required)  # the upload-time guess
        self._auth(self.iqac)
        self.client.post(f'/api/v1/sprints/{self.sprint.id}/extraction-jobs/')
        self.document.refresh_from_db()
        self.assertFalse(self.document.ocr_required)
        self.assertIsNotNone(self.document.page_count)

    def test_ocr_required_stays_true_for_a_genuinely_unreadable_document(self):
        """The other direction of the same fix: a document that genuinely
        has no extractable text must end up ocr_required=True from the real
        reader's finding, with real ocr_warnings explaining why -- not just
        coincidentally matching the upload-time guess for the wrong reason."""
        self._auth(self.iqac)
        blank_upload = self.client.post(
            f'/api/v1/sprints/{self.sprint.id}/upload-file/',
            {'file': make_pdf(name='blank.pdf', page_texts=('',)), 'document_type': 'aqar_report'},
            format='multipart',
        )
        blank_document = Document.objects.get(id=blank_upload.data['id'])

        self.client.post(
            f'/api/v1/sprints/{self.sprint.id}/extraction-jobs/', {'document_id': str(blank_document.id)},
        )

        blank_document.refresh_from_db()
        self.assertTrue(blank_document.ocr_required)
        self.assertTrue(blank_document.ocr_warnings)

    def test_sprint_advances_from_collecting_to_reviewing(self):
        self._auth(self.iqac)
        self.assertEqual(self.sprint.status, Sprint.Status.COLLECTING)
        self.client.post(f'/api/v1/sprints/{self.sprint.id}/extraction-jobs/')
        self.sprint.refresh_from_db()
        self.assertEqual(self.sprint.status, Sprint.Status.REVIEWING)

    def test_no_facts_are_fabricated_when_the_ai_finds_none(self):
        """Per 'no fake data': when the (mocked) AI genuinely returns no
        facts, a completed job must not have invented any."""
        self._auth(self.iqac)
        self.client.post(f'/api/v1/sprints/{self.sprint.id}/extraction-jobs/')
        self.assertEqual(self.document.extracted_facts.count(), 0)

    def test_facts_are_persisted_end_to_end_through_the_real_pipeline(self):
        """The core of 'connect OpenAIFactExtractor to ExtractedFact
        database': a fact the (mocked) AI returns must actually reach the
        database, not be dropped by an unmapped field_mapper stage."""
        self.fact_openai_mock.return_value.extract_structured_data.return_value = {
            'facts': [{
                'field_name': 'Total Faculty Count', 'field_key': 'total_faculty', 'value': '42',
                'data_type': 'number', 'pillar': 'aiqif_c4', 'owner_role': 'hr_officer',
                'source_page': '1', 'source_snippet': 'The institution has 42 faculty members.',
                'confidence_score': 0.96, 'confidence_reason': 'Explicitly stated in the source.',
            }],
        }
        self._auth(self.iqac)
        self.client.post(f'/api/v1/sprints/{self.sprint.id}/extraction-jobs/')

        fact = self.document.extracted_facts.get()
        self.assertEqual(fact.field_key, 'total_faculty')
        self.assertEqual(fact.value, 42)
        self.assertEqual(fact.normalized_value, 42)
        self.assertEqual(fact.data_type, 'number')
        self.assertEqual(fact.pillar, 'aiqif_c4')
        self.assertEqual(fact.owner_role, 'hr_officer')
        self.assertEqual(fact.source_page, '1')
        self.assertEqual(fact.source_document_id, self.document.id)
        self.assertEqual(fact.sprint_id, self.sprint.id)
        self.assertEqual(fact.extraction_method, 'openai')
        self.assertEqual(fact.status, 'extracted')

    def test_low_confidence_fact_creates_a_persisted_gap(self):
        """RuleBasedGapDetector's low_confidence check, exercised through
        the real pipeline (not just the unit tests further down) -- proves
        the gap it returns actually reaches the database via _persist_gaps'
        create_gap_if_new call, not just a returned-but-dropped dict."""
        self.fact_openai_mock.return_value.extract_structured_data.return_value = {
            'facts': [_valid_fact(confidence_score=0.3)],
        }
        self._auth(self.iqac)
        self.client.post(f'/api/v1/sprints/{self.sprint.id}/extraction-jobs/')

        fact = self.document.extracted_facts.get()
        gap = GapItem.objects.get(gap_type=GapItem.GapType.LOW_CONFIDENCE)
        self.assertEqual(gap.source_fact_id, fact.id)
        self.assertEqual(gap.sprint_id, self.sprint.id)
        self.assertEqual(gap.priority, GapItem.Priority.HIGH)  # below the "very low" threshold too

    def test_gaps_created_on_completion_are_real_not_fabricated(self):
        """With no facts extracted (the default mocked AI response is
        empty), RuleBasedGapDetector has nothing to flag either -- so once
        the sprint reaches 'reviewing', every gap present comes from
        apps.gaps.services' real, deterministic sprint-level generation
        (see apps/gaps/tests.py): genuinely missing required documents
        (everything but the uploaded naac_ssr), not fabricated data."""
        self._auth(self.iqac)
        self.client.post(f'/api/v1/sprints/{self.sprint.id}/extraction-jobs/')
        gaps = list(self.sprint.gaps.all())
        self.assertTrue(gaps)
        self.assertTrue(all(gap.gap_type == GapItem.GapType.MISSING_DOCUMENT for gap in gaps))
        missing_types = {gap.title for gap in gaps}
        self.assertNotIn('Missing NAAC SSR', missing_types)

    def test_job_status_exposes_full_lifecycle_via_api(self):
        self._auth(self.iqac)
        response = self.client.post(f'/api/v1/sprints/{self.sprint.id}/extraction-jobs/')
        detail = self.client.get(f'/api/v1/extraction-jobs/{response.data[0]["id"]}/')
        for field in ('status', 'current_step', 'progress_percentage', 'error_message', 'started_at', 'completed_at'):
            self.assertIn(field, detail.data)


@override_settings(CELERY_TASK_ALWAYS_EAGER=True, CELERY_TASK_EAGER_PROPAGATES=True)
class FailureTests(ExtractionJobTestBase):
    """Permanent and unexpected failures resolve in a single attempt --
    no retry-loop involved, so these are safe to exercise end to end."""

    @patch('apps.extraction.tasks.ExtractionPipeline')
    def test_permanent_error_fails_immediately_without_retry(self, mock_pipeline_cls):
        mock_pipeline_cls.return_value.run.side_effect = PermanentExtractionError('corrupt file')
        self._auth(self.iqac)

        response = self.client.post(f'/api/v1/sprints/{self.sprint.id}/extraction-jobs/')
        job = ExtractionJob.objects.get(id=response.data[0]['id'])

        self.assertEqual(job.status, ExtractionJob.Status.FAILED)
        self.assertEqual(job.retry_count, 0)
        self.assertIn('corrupt file', job.error_message)
        self.assertEqual(mock_pipeline_cls.return_value.run.call_count, 1)

        self.document.refresh_from_db()
        self.assertEqual(self.document.status, Document.Status.FAILED)

    @patch('apps.extraction.tasks.ExtractionPipeline')
    def test_unrecognized_error_fails_without_retry(self, mock_pipeline_cls):
        """An error the pipeline doesn't classify as recoverable/permanent
        must not be retried blindly."""
        mock_pipeline_cls.return_value.run.side_effect = ValueError('totally unexpected bug')
        self._auth(self.iqac)

        response = self.client.post(f'/api/v1/sprints/{self.sprint.id}/extraction-jobs/')
        job = ExtractionJob.objects.get(id=response.data[0]['id'])

        self.assertEqual(job.status, ExtractionJob.Status.FAILED)
        self.assertIn('Unexpected error', job.error_message)
        self.assertEqual(mock_pipeline_cls.return_value.run.call_count, 1)

    @patch('apps.extraction.tasks.ExtractionPipeline')
    def test_missing_job_id_does_not_crash_the_task(self, mock_pipeline_cls):
        import uuid

        from .tasks import run_extraction_job
        run_extraction_job.delay(str(uuid.uuid4()))
        mock_pipeline_cls.return_value.run.assert_not_called()

    @patch('apps.extraction.tasks.ExtractionPipeline')
    def test_first_recoverable_attempt_marks_job_retrying(self, mock_pipeline_cls):
        """One recoverable failure (before retries are exhausted) leaves the
        job visibly 'retrying', not silently stuck or marked failed."""
        mock_pipeline_cls.return_value.run.side_effect = RecoverableExtractionError('transient glitch')
        self._auth(self.iqac)

        response = self.client.post(f'/api/v1/sprints/{self.sprint.id}/extraction-jobs/')
        job = ExtractionJob.objects.get(id=response.data[0]['id'])

        self.assertEqual(job.status, ExtractionJob.Status.RETRYING)
        self.assertEqual(job.retry_count, 1)
        self.assertIn('transient glitch', job.error_message)


@override_settings(CELERY_TASK_ALWAYS_EAGER=True, CELERY_TASK_EAGER_PROPAGATES=True)
class StopAndProcessTests(ExtractionJobTestBase):
    """Per-row Stop / Process on the AI Processing Monitor, and the worker
    honouring a stop."""

    def _job(self, status_value, **extra):
        return ExtractionJob.objects.create(
            sprint=self.sprint, document=self.document, status=status_value, **extra,
        )

    def test_worker_never_starts_a_job_stopped_while_queued(self):
        job = self._job(ExtractionJob.Status.CANCELLED)
        run_extraction_job.apply(args=[str(job.id)])
        job.refresh_from_db()
        self.assertEqual(job.status, ExtractionJob.Status.CANCELLED)
        self.assertEqual(job.current_step, '')

    def test_running_job_stops_at_next_step_and_stays_cancelled(self):
        job = self._job(ExtractionJob.Status.PENDING)

        def stop_during_classification(**kwargs):
            # The first request classifies the document (and extracts the first chunk's facts).
            ExtractionJob.objects.filter(id=job.id).update(status=ExtractionJob.Status.CANCELLED)
            return {'classification': dict(FAKE_CLASSIFICATION_RESULT), 'facts': []}

        self.analysis_service.extract_structured_data.side_effect = stop_during_classification
        run_extraction_job.apply(args=[str(job.id)])

        job.refresh_from_db()
        self.assertEqual(job.status, ExtractionJob.Status.CANCELLED)
        self.assertEqual(job.current_step, ExtractionJob.Step.CLASSIFYING_DOCUMENTS)

    def test_stop_endpoint_stops_one_active_job(self):
        job = self._job(ExtractionJob.Status.RUNNING)
        self._auth(self.iqac)
        response = self.client.post(f'/api/v1/extraction-jobs/{job.id}/cancel')
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        self.assertEqual(response.data['status'], 'cancelled')

    def test_stop_endpoint_rejects_finished_job(self):
        job = self._job(ExtractionJob.Status.COMPLETED)
        self._auth(self.iqac)
        response = self.client.post(f'/api/v1/extraction-jobs/{job.id}/cancel')
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    @patch('apps.extraction.views.run_extraction_job')
    def test_process_endpoint_requeues_the_same_row(self, mock_task):
        mock_task.delay.return_value = MagicMock(id='task-2')
        job = self._job(ExtractionJob.Status.CANCELLED, progress_percentage=45, error_message='x')
        self._auth(self.iqac)
        response = self.client.post(f'/api/v1/extraction-jobs/{job.id}/process')
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        job.refresh_from_db()
        self.assertEqual(job.status, ExtractionJob.Status.PENDING)
        self.assertEqual(job.progress_percentage, 0)
        self.assertEqual(job.celery_task_id, 'task-2')
        self.assertEqual(ExtractionJob.objects.filter(document=self.document).count(), 1)

    def test_process_endpoint_rejects_completed_job(self):
        job = self._job(ExtractionJob.Status.COMPLETED)
        self._auth(self.iqac)
        response = self.client.post(f'/api/v1/extraction-jobs/{job.id}/process')
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    @patch('apps.extraction.views.run_extraction_job')
    def test_process_incomplete_requeues_only_unfinished_documents(self, mock_task):
        mock_task.delay.return_value = MagicMock(id='task-3')
        stopped = self._job(ExtractionJob.Status.CANCELLED)
        done_doc = Document.objects.create(
            sprint=self.sprint, document_type='aqar_report', original_filename='aqar.pdf',
            mime_type='application/pdf', status=Document.Status.PROCESSED,
        )
        ExtractionJob.objects.create(sprint=self.sprint, document=done_doc, status=ExtractionJob.Status.FAILED)
        ExtractionJob.objects.create(sprint=self.sprint, document=done_doc, status=ExtractionJob.Status.COMPLETED)

        self._auth(self.iqac)
        response = self.client.post(f'/api/v1/extraction-jobs/sprints/{self.sprint.id}/process-incomplete')
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        self.assertEqual([row['id'] for row in response.data], [str(stopped.id)])


class RetryDecisionLogicTests(TestCase):
    """Unit tests for the retry/backoff/exhaustion decision logic itself.

    Celery's eager mode (used above for end-to-end tests) runs a task
    exactly once per `.delay()` call and does not loop through retries --
    `self.retry()` just raises `celery.exceptions.Retry`, the same as it
    would be caught and rescheduled by a real worker. Rather than
    reimplementing Celery's own redelivery mechanics here, these tests
    drive `_handle_recoverable` directly with a stand-in task object,
    which is what actually encodes *this project's* retry policy.
    """

    def setUp(self):
        institution = Institution.objects.create(name='MKCE', city='Karur', state='TN')
        sprint = Sprint.objects.create(institution=institution, mode=Sprint.SprintMode.VERIFIED_CRI)
        document = Document.objects.create(
            sprint=sprint, document_type='naac_ssr', status=Document.Status.UPLOADED,
        )
        self.job = ExtractionJob.objects.create(sprint=sprint, document=document)

    @staticmethod
    def _fake_task(retries):
        task = MagicMock()
        task.request.retries = retries

        def _raise(*args, **kwargs):
            raise RuntimeError('celery would reschedule here')
        task.retry.side_effect = _raise
        return task

    def test_schedules_a_retry_while_attempts_remain(self):
        task = self._fake_task(retries=0)
        with self.assertRaises(RuntimeError):
            _handle_recoverable(task, self.job, RecoverableExtractionError('flaky'))

        self.job.refresh_from_db()
        self.assertEqual(self.job.status, ExtractionJob.Status.RETRYING)
        self.assertEqual(self.job.retry_count, 1)
        task.retry.assert_called_once()
        self.assertEqual(task.retry.call_args.kwargs['countdown'], settings.EXTRACTION_RETRY_BACKOFF_SECONDS)

    def test_backoff_grows_exponentially_with_attempt_number(self):
        task = self._fake_task(retries=2)
        with self.assertRaises(RuntimeError):
            _handle_recoverable(task, self.job, RecoverableExtractionError('still flaky'))

        expected = settings.EXTRACTION_RETRY_BACKOFF_SECONDS * (2 ** 2)
        self.assertEqual(task.retry.call_args.kwargs['countdown'], expected)

    def test_gives_up_once_retries_are_exhausted(self):
        task = self._fake_task(retries=settings.EXTRACTION_MAX_RETRIES)
        _handle_recoverable(task, self.job, RecoverableExtractionError('never recovers'))

        self.job.refresh_from_db()
        self.assertEqual(self.job.status, ExtractionJob.Status.FAILED)
        self.assertIn('Failed after', self.job.error_message)
        self.assertIn(str(settings.EXTRACTION_MAX_RETRIES), self.job.error_message)
        task.retry.assert_not_called()

    def test_does_not_endlessly_retry_past_the_configured_maximum(self):
        """A permanent-failure job never re-enters RETRYING once exhausted,
        no matter how many more times it's (hypothetically) invoked."""
        task = self._fake_task(retries=settings.EXTRACTION_MAX_RETRIES + 5)
        _handle_recoverable(task, self.job, RecoverableExtractionError('never recovers'))
        self.job.refresh_from_db()
        self.assertEqual(self.job.status, ExtractionJob.Status.FAILED)
        task.retry.assert_not_called()


def _fake_openai_request():
    return httpx.Request('POST', 'https://api.openai.com/v1/chat/completions')


def _fake_openai_response(status_code):
    return httpx.Response(status_code=status_code, request=_fake_openai_request())


def _fake_completion(content, finish_reason='stop'):
    """A stand-in for an OpenAI ChatCompletion, shaped just enough for
    OpenAIExtractionService._parse_response to read choices[0].finish_reason
    and choices[0].message.content."""
    choice = MagicMock()
    choice.finish_reason = finish_reason
    choice.message.content = content
    completion = MagicMock()
    completion.choices = [choice]
    return completion


@override_settings(AI_BASE_URL='')  # isolate from whatever the real .env happens to have set
class OpenAIExtractionServiceTests(SimpleTestCase):
    """Unit tests for the OpenAI service layer (apps.extraction.services.
    openai_client). No real network call is ever made here -- either the
    OpenAI client class itself is mocked, or a fake client double is
    injected via the constructor's `client=` param."""

    def _service(self, client=None, api_key='sk-test', model='gpt-test'):
        return OpenAIExtractionService(api_key=api_key, model=model, client=client)

    def _call(self, service):
        return service.extract_structured_data(
            system_prompt='Extract facts from this document.',
            user_content='document text...',
            response_schema={'type': 'object'},
        )

    # -- configuration --------------------------------------------------

    @override_settings(OPENAI_API_KEY='', OPENAI_EXTRACTION_MODEL='gpt-test')
    def test_missing_api_key_raises_configuration_error(self):
        with self.assertRaises(ImproperlyConfigured):
            OpenAIExtractionService()

    @override_settings(OPENAI_API_KEY='sk-test', OPENAI_EXTRACTION_MODEL='')
    def test_missing_model_raises_configuration_error(self):
        with self.assertRaises(ImproperlyConfigured):
            OpenAIExtractionService()

    @patch('apps.extraction.services.openai_client.OpenAI')
    def test_client_initializes_with_the_configured_api_key(self, mock_openai_cls):
        OpenAIExtractionService(api_key='sk-test', model='gpt-test')
        mock_openai_cls.assert_called_once_with(api_key='sk-test', base_url=None)

    @patch('apps.extraction.services.openai_client.OpenAI')
    def test_custom_base_url_is_passed_to_the_client(self, mock_openai_cls):
        OpenAIExtractionService(api_key='sk-test', model='my-combo', base_url='http://localhost:20128/v1')
        mock_openai_cls.assert_called_once_with(api_key='sk-test', base_url='http://localhost:20128/v1')

    def test_injected_client_is_used_instead_of_constructing_a_real_one(self):
        fake_client = MagicMock()
        service = self._service(client=fake_client)
        self.assertIs(service._client, fake_client)

    # -- OpenAI-compatible routers (AI_BASE_URL, e.g. 9router) -------------

    def _router_service(self, client):
        return OpenAIExtractionService(
            api_key='router-key', model='my-combo', base_url='http://localhost:20128/v1', client=client,
        )

    def test_router_gets_the_schema_in_the_system_prompt(self):
        """Routers often drop response_format, so the schema must also be in
        the prompt; direct OpenAI keeps the prompt untouched."""
        fake_client = MagicMock()
        fake_client.chat.completions.create.return_value = _fake_completion(content='{"a": 1}')
        self._call(self._router_service(fake_client))
        system = fake_client.chat.completions.create.call_args.kwargs['messages'][0]['content']
        self.assertTrue(system.startswith('Extract facts from this document.'))
        self.assertIn('ONLY one JSON value', system)
        self.assertIn('{"type":"object"}', system)

        direct_client = MagicMock()
        direct_client.chat.completions.create.return_value = _fake_completion(content='{"a": 1}')
        self._call(self._service(client=direct_client))
        direct_system = direct_client.chat.completions.create.call_args.kwargs['messages'][0]['content']
        self.assertEqual(direct_system, 'Extract facts from this document.')

    def test_fenced_or_wrapped_json_is_still_parsed(self):
        for content in ('```json\n{"a": 1}\n```', 'Here is the result:\n{"a": 1}\nHope this helps.'):
            fake_client = MagicMock()
            fake_client.chat.completions.create.return_value = _fake_completion(content=content)
            self.assertEqual(self._call(self._router_service(fake_client)), {'a': 1})

    def test_router_prose_reply_is_retried_once_then_succeeds(self):
        fake_client = MagicMock()
        fake_client.chat.completions.create.side_effect = [
            _fake_completion(content='Mumbai is the largest city.'),
            _fake_completion(content='{"a": 1}'),
        ]
        self.assertEqual(self._call(self._router_service(fake_client)), {'a': 1})
        self.assertEqual(fake_client.chat.completions.create.call_count, 2)

    def test_router_prose_reply_twice_raises(self):
        fake_client = MagicMock()
        fake_client.chat.completions.create.return_value = _fake_completion(content='Just prose, no JSON.')
        with self.assertRaises(OpenAIResponseError):
            self._call(self._router_service(fake_client))
        self.assertEqual(fake_client.chat.completions.create.call_count, 2)

    # -- successful response ---------------------------------------------

    def test_successful_response_returns_normalized_python_data(self):
        fake_client = MagicMock()
        fake_client.chat.completions.create.return_value = _fake_completion(
            content='{"institution_name": "MKCE", "confidence": 0.92}',
        )
        service = self._service(client=fake_client)

        result = self._call(service)

        self.assertEqual(result, {'institution_name': 'MKCE', 'confidence': 0.92})
        call_kwargs = fake_client.chat.completions.create.call_args.kwargs
        self.assertEqual(call_kwargs['model'], 'gpt-test')
        self.assertEqual(call_kwargs['response_format']['json_schema']['schema'], {'type': 'object'})

    # -- API errors -------------------------------------------------------

    def test_rate_limit_error_is_treated_as_recoverable(self):
        fake_client = MagicMock()
        fake_client.chat.completions.create.side_effect = RateLimitError(
            'rate limited', response=_fake_openai_response(429), body=None,
        )
        with self.assertRaises(RecoverableExtractionError):
            self._call(self._service(client=fake_client))

    def test_timeout_error_is_treated_as_recoverable(self):
        fake_client = MagicMock()
        fake_client.chat.completions.create.side_effect = APITimeoutError(request=_fake_openai_request())
        with self.assertRaises(RecoverableExtractionError):
            self._call(self._service(client=fake_client))

    def test_connection_error_is_treated_as_recoverable(self):
        fake_client = MagicMock()
        fake_client.chat.completions.create.side_effect = APIConnectionError(request=_fake_openai_request())
        with self.assertRaises(RecoverableExtractionError):
            self._call(self._service(client=fake_client))

    def test_server_side_status_error_is_treated_as_recoverable(self):
        fake_client = MagicMock()
        fake_client.chat.completions.create.side_effect = APIStatusError(
            'internal error', response=_fake_openai_response(500), body=None,
        )
        with self.assertRaises(RecoverableExtractionError):
            self._call(self._service(client=fake_client))

    def test_client_side_status_error_is_treated_as_permanent(self):
        fake_client = MagicMock()
        fake_client.chat.completions.create.side_effect = APIStatusError(
            'bad request', response=_fake_openai_response(400), body=None,
        )
        with self.assertRaises(PermanentExtractionError):
            self._call(self._service(client=fake_client))

    # -- invalid responses ------------------------------------------------

    def test_empty_content_raises_response_error(self):
        fake_client = MagicMock()
        fake_client.chat.completions.create.return_value = _fake_completion(content=None)
        with self.assertRaises(OpenAIResponseError):
            self._call(self._service(client=fake_client))

    def test_non_json_content_raises_response_error(self):
        fake_client = MagicMock()
        fake_client.chat.completions.create.return_value = _fake_completion(content='not valid json')
        with self.assertRaises(OpenAIResponseError):
            self._call(self._service(client=fake_client))

    def test_truncated_response_raises_response_error(self):
        fake_client = MagicMock()
        fake_client.chat.completions.create.return_value = _fake_completion(
            content='{"partial": tr', finish_reason='length',
        )
        with self.assertRaises(OpenAIResponseError):
            self._call(self._service(client=fake_client))

    def test_content_filtered_response_raises_response_error(self):
        fake_client = MagicMock()
        fake_client.chat.completions.create.return_value = _fake_completion(
            content=None, finish_reason='content_filter',
        )
        with self.assertRaises(OpenAIResponseError):
            self._call(self._service(client=fake_client))

    def test_no_choices_raises_response_error(self):
        fake_client = MagicMock()
        empty_response = MagicMock()
        empty_response.choices = []
        fake_client.chat.completions.create.return_value = empty_response
        with self.assertRaises(OpenAIResponseError):
            self._call(self._service(client=fake_client))


def _fake_anthropic_request():
    return httpx.Request('POST', 'https://api.anthropic.com/v1/messages')


def _fake_anthropic_response(status_code):
    return httpx.Response(status_code=status_code, request=_fake_anthropic_request())


def _fake_tool_use_message(input_data=None, name='extraction_result', stop_reason='tool_use'):
    """A stand-in for an Anthropic Message, shaped just enough for
    AnthropicExtractionService._parse_response to read stop_reason and find
    a matching tool_use content block."""
    message = MagicMock()
    message.stop_reason = stop_reason
    if input_data is None:
        message.content = []
    else:
        block = MagicMock()
        block.type = 'tool_use'
        block.name = name
        block.input = input_data
        message.content = [block]
    return message


class AnthropicExtractionServiceTests(SimpleTestCase):
    """Unit tests for the Claude counterpart of OpenAIExtractionService --
    same contract, same test shape, so the two stay verifiably equivalent.
    Every Anthropic call is mocked; no real network access."""

    def _service(self, client=None, api_key='sk-ant-test', model='claude-test'):
        return AnthropicExtractionService(api_key=api_key, model=model, client=client)

    def _call(self, service):
        return service.extract_structured_data(
            system_prompt='Extract facts.', user_content='document text...',
            response_schema={'type': 'object'},
        )

    # -- configuration --------------------------------------------------

    @override_settings(AI_API_KEY='', AI_MODEL='claude-test')
    def test_missing_api_key_raises_configuration_error(self):
        with self.assertRaises(ImproperlyConfigured):
            AnthropicExtractionService()

    @override_settings(AI_API_KEY='sk-ant-test', AI_MODEL='')
    def test_missing_model_raises_configuration_error(self):
        with self.assertRaises(ImproperlyConfigured):
            AnthropicExtractionService()

    @patch('apps.extraction.services.anthropic_client.Anthropic')
    def test_client_initializes_with_the_configured_api_key(self, mock_anthropic_cls):
        AnthropicExtractionService(api_key='sk-ant-test', model='claude-test')
        mock_anthropic_cls.assert_called_once_with(api_key='sk-ant-test')

    def test_injected_client_is_used_instead_of_constructing_a_real_one(self):
        fake_client = MagicMock()
        service = self._service(client=fake_client)
        self.assertIs(service._client, fake_client)

    # -- successful response ---------------------------------------------

    def test_successful_response_returns_the_tool_calls_input(self):
        fake_client = MagicMock()
        fake_client.messages.create.return_value = _fake_tool_use_message(
            input_data={'institution_name': 'MKCE', 'confidence': 0.92},
        )
        service = self._service(client=fake_client)

        result = self._call(service)

        self.assertEqual(result, {'institution_name': 'MKCE', 'confidence': 0.92})
        call_kwargs = fake_client.messages.create.call_args.kwargs
        self.assertEqual(call_kwargs['model'], 'claude-test')
        self.assertEqual(call_kwargs['tools'][0]['input_schema'], {'type': 'object'})
        self.assertTrue(call_kwargs['tools'][0]['strict'])
        self.assertEqual(call_kwargs['tool_choice']['name'], 'extraction_result')

    # -- API errors -------------------------------------------------------

    def test_rate_limit_error_is_treated_as_recoverable(self):
        fake_client = MagicMock()
        fake_client.messages.create.side_effect = anthropic_sdk.RateLimitError(
            'rate limited', response=_fake_anthropic_response(429), body=None,
        )
        with self.assertRaises(RecoverableExtractionError):
            self._call(self._service(client=fake_client))

    def test_timeout_error_is_treated_as_recoverable(self):
        fake_client = MagicMock()
        fake_client.messages.create.side_effect = anthropic_sdk.APITimeoutError(request=_fake_anthropic_request())
        with self.assertRaises(RecoverableExtractionError):
            self._call(self._service(client=fake_client))

    def test_connection_error_is_treated_as_recoverable(self):
        fake_client = MagicMock()
        fake_client.messages.create.side_effect = anthropic_sdk.APIConnectionError(
            request=_fake_anthropic_request(),
        )
        with self.assertRaises(RecoverableExtractionError):
            self._call(self._service(client=fake_client))

    def test_server_side_status_error_is_treated_as_recoverable(self):
        fake_client = MagicMock()
        fake_client.messages.create.side_effect = anthropic_sdk.APIStatusError(
            'internal error', response=_fake_anthropic_response(500), body=None,
        )
        with self.assertRaises(RecoverableExtractionError):
            self._call(self._service(client=fake_client))

    def test_client_side_status_error_is_treated_as_permanent(self):
        fake_client = MagicMock()
        fake_client.messages.create.side_effect = anthropic_sdk.APIStatusError(
            'bad request', response=_fake_anthropic_response(400), body=None,
        )
        with self.assertRaises(PermanentExtractionError):
            self._call(self._service(client=fake_client))

    # -- invalid responses ------------------------------------------------

    def test_no_matching_tool_call_raises_response_error(self):
        fake_client = MagicMock()
        fake_client.messages.create.return_value = _fake_tool_use_message(input_data=None)
        with self.assertRaises(AnthropicResponseError):
            self._call(self._service(client=fake_client))

    def test_truncated_response_raises_response_error(self):
        fake_client = MagicMock()
        fake_client.messages.create.return_value = _fake_tool_use_message(
            input_data={'a': 1}, stop_reason='max_tokens',
        )
        with self.assertRaises(AnthropicResponseError):
            self._call(self._service(client=fake_client))

    def test_refusal_raises_response_error(self):
        fake_client = MagicMock()
        fake_client.messages.create.return_value = _fake_tool_use_message(
            input_data=None, stop_reason='refusal',
        )
        with self.assertRaises(AnthropicResponseError):
            self._call(self._service(client=fake_client))

    def test_context_window_exceeded_raises_response_error(self):
        fake_client = MagicMock()
        fake_client.messages.create.return_value = _fake_tool_use_message(
            input_data=None, stop_reason='model_context_window_exceeded',
        )
        with self.assertRaises(AnthropicResponseError):
            self._call(self._service(client=fake_client))

    def test_response_error_is_also_a_generic_ai_response_error(self):
        """Code that doesn't care which provider answered (e.g.
        openai_fact_extractor's malformed-shape check) can catch just
        AIResponseError and still catch this."""
        fake_client = MagicMock()
        fake_client.messages.create.return_value = _fake_tool_use_message(input_data=None)
        with self.assertRaises(AIResponseError):
            self._call(self._service(client=fake_client))


@override_settings(AI_BASE_URL='')  # isolate from whatever the real .env happens to have set
class AIServiceFactoryTests(SimpleTestCase):
    """Unit tests for apps.extraction.services.ai_service -- the provider
    detection and factory every AI call site goes through."""

    # -- provider detection -----------------------------------------------

    def test_detects_openai_from_key_prefix(self):
        self.assertEqual(detect_provider('sk-proj-abc123'), 'openai')

    def test_detects_anthropic_from_key_prefix(self):
        self.assertEqual(detect_provider('sk-ant-api03-abc123'), 'anthropic')

    def test_unrecognized_key_format_raises(self):
        with self.assertRaises(ImproperlyConfigured):
            detect_provider('not-a-real-key-format')

    # -- factory ------------------------------------------------------------

    @override_settings(AI_API_KEY='', OPENAI_API_KEY='')
    def test_no_key_configured_raises(self):
        with self.assertRaises(ImproperlyConfigured):
            get_ai_service()

    @override_settings(AI_API_KEY='sk-test-key', AI_MODEL='gpt-test', OPENAI_API_KEY='')
    def test_returns_openai_service_for_an_openai_key(self):
        service = get_ai_service()
        self.assertIsInstance(service, OpenAIExtractionService)
        self.assertEqual(service.model, 'gpt-test')

    @override_settings(AI_API_KEY='sk-ant-test-key', AI_MODEL='claude-test', OPENAI_API_KEY='')
    def test_returns_anthropic_service_for_an_anthropic_key(self):
        service = get_ai_service()
        self.assertIsInstance(service, AnthropicExtractionService)
        self.assertEqual(service.model, 'claude-test')

    @override_settings(AI_API_KEY='', AI_MODEL='', OPENAI_API_KEY='sk-legacy-key', OPENAI_EXTRACTION_MODEL='gpt-legacy')
    def test_falls_back_to_openai_api_key_when_ai_api_key_unset(self):
        """Anyone with a .env from before multi-provider support existed
        keeps working unchanged."""
        service = get_ai_service()
        self.assertIsInstance(service, OpenAIExtractionService)
        self.assertEqual(service.api_key, 'sk-legacy-key')
        self.assertEqual(service.model, 'gpt-legacy')

    @override_settings(AI_API_KEY='sk-new-key', OPENAI_API_KEY='sk-old-key', AI_MODEL='')
    def test_ai_api_key_takes_precedence_over_openai_api_key(self):
        service = get_ai_service()
        self.assertEqual(service.api_key, 'sk-new-key')

    @override_settings(AI_API_KEY='sk-test', AI_MODEL='', OPENAI_API_KEY='', OPENAI_EXTRACTION_MODEL='')
    def test_default_model_used_when_none_configured(self):
        service = get_ai_service()
        self.assertEqual(service.model, 'gpt-4o-mini')

    @override_settings(AI_API_KEY='sk-ant-test', AI_MODEL='', OPENAI_API_KEY='', OPENAI_EXTRACTION_MODEL='')
    def test_default_model_used_for_anthropic_when_none_configured(self):
        service = get_ai_service()
        self.assertEqual(service.model, 'claude-haiku-4-5-20251001')

    @override_settings(AI_API_KEY='sk-ant-test', AI_MODEL='', OPENAI_API_KEY='', OPENAI_EXTRACTION_MODEL='gpt-4o-mini')
    def test_stale_openai_extraction_model_does_not_leak_into_anthropic(self):
        """The bug this test guards against: OPENAI_EXTRACTION_MODEL is left
        over in .env from before the key was switched to Anthropic -- it
        must never be handed to a Claude client as its model name."""
        service = get_ai_service()
        self.assertIsInstance(service, AnthropicExtractionService)
        self.assertEqual(service.model, 'claude-haiku-4-5-20251001')

    @override_settings(AI_API_KEY='sk-test', AI_MODEL='', OPENAI_API_KEY='', OPENAI_EXTRACTION_MODEL='gpt-4o')
    def test_openai_extraction_model_still_applies_when_provider_is_openai(self):
        service = get_ai_service()
        self.assertEqual(service.model, 'gpt-4o')

    @override_settings(AI_API_KEY='sk-test', AI_MODEL='gpt-explicit-override', OPENAI_EXTRACTION_MODEL='gpt-4o')
    def test_explicit_model_argument_wins_over_everything(self):
        service = get_ai_service(model='gpt-called-with')
        self.assertEqual(service.model, 'gpt-called-with')

    # -- custom OpenAI-compatible endpoint (AI_BASE_URL) ---------------------

    @override_settings(
        AI_API_KEY='sk-45c32fe83970c5b2-nihvmj-397d2a0c', AI_MODEL='my-combo',
        AI_BASE_URL='http://localhost:20128/v1', OPENAI_API_KEY='', OPENAI_EXTRACTION_MODEL='',
    )
    def test_custom_base_url_returns_an_openai_compatible_service(self):
        """A local router/gateway's key doesn't look like a real OpenAI or
        Anthropic key at all -- AI_BASE_URL must be enough on its own,
        without ever calling detect_provider() on that key."""
        service = get_ai_service()
        self.assertIsInstance(service, OpenAIExtractionService)
        self.assertEqual(service.api_key, 'sk-45c32fe83970c5b2-nihvmj-397d2a0c')
        self.assertEqual(service.model, 'my-combo')
        self.assertEqual(service.base_url, 'http://localhost:20128/v1')

    @override_settings(
        AI_API_KEY='sk-ant-looks-like-anthropic', AI_MODEL='my-combo',
        AI_BASE_URL='http://localhost:20128/v1', OPENAI_API_KEY='', OPENAI_EXTRACTION_MODEL='',
    )
    def test_custom_base_url_wins_even_over_an_anthropic_shaped_key(self):
        """The base URL is an explicit instruction, not a hint -- it must
        override key-format detection entirely, not just act as a fallback."""
        service = get_ai_service()
        self.assertIsInstance(service, OpenAIExtractionService)
        self.assertEqual(service.base_url, 'http://localhost:20128/v1')

    @override_settings(
        AI_API_KEY='sk-45c32fe83970c5b2-nihvmj-397d2a0c', AI_MODEL='',
        AI_BASE_URL='http://localhost:20128/v1', OPENAI_API_KEY='', OPENAI_EXTRACTION_MODEL='',
    )
    def test_custom_base_url_without_a_model_raises(self):
        """There's no sensible default model for an endpoint this app knows
        nothing about -- unlike OpenAI/Anthropic, which fall back to
        DEFAULT_MODELS, this must fail loudly instead of guessing."""
        with self.assertRaises(ImproperlyConfigured):
            get_ai_service()

    @override_settings(
        AI_API_KEY='sk-45c32fe83970c5b2-nihvmj-397d2a0c', AI_MODEL='',
        AI_BASE_URL='http://localhost:20128/v1', OPENAI_API_KEY='', OPENAI_EXTRACTION_MODEL='my-combo',
    )
    def test_openai_extraction_model_is_accepted_as_the_model_for_a_base_url_too(self):
        """OPENAI_EXTRACTION_MODEL is still a valid fallback for AI_BASE_URL
        setups, same as for plain OpenAI -- only AI_MODEL being unset should
        matter, not which specific settings name supplied it."""
        service = get_ai_service()
        self.assertEqual(service.model, 'my-combo')


class PDFPageReaderTests(TestCase):
    """Unit tests for the real PDFPageReader (apps.extraction.services.
    pdf_reader). Every PDF here is a real, valid, pdfplumber-readable file
    built with fpdf2 -- nothing about PDF parsing is mocked."""

    def setUp(self):
        institution = Institution.objects.create(name='MKCE', city='Karur', state='TN')
        self.sprint = Sprint.objects.create(institution=institution, mode=Sprint.SprintMode.VERIFIED_CRI)
        self.reader = PDFPageReader()

    def _make_document(self, content, filename='doc.pdf', document_type='naac_ssr', content_type='application/pdf'):
        upload = SimpleUploadedFile(filename, content, content_type=content_type)
        return Document.objects.create(
            sprint=self.sprint, document_type=document_type, file=upload,
            original_filename=filename, mime_type=content_type, status=Document.Status.UPLOADED,
        )

    def test_normal_text_pdf_extracts_real_text(self):
        document = self._make_document(
            make_pdf_bytes(['Hello, this is real NAAC SSR self-study report content for testing.']),
        )
        result = self.reader.read_pages(document)

        self.assertTrue(result['format_supported'])
        self.assertEqual(result['page_count'], 1)
        self.assertEqual(result['pages'][0]['page_number'], 1)
        self.assertIn('real NAAC SSR self-study report content', result['pages'][0]['text'])
        self.assertFalse(result['pages'][0]['requires_ocr'])
        self.assertFalse(result['requires_ocr'])

    def test_multi_page_pdf_preserves_page_numbers_and_text(self):
        document = self._make_document(
            make_pdf_bytes(['Page one content.', 'Page two content.', 'Page three content.']),
        )
        result = self.reader.read_pages(document)

        self.assertEqual(result['page_count'], 3)
        self.assertEqual([p['page_number'] for p in result['pages']], [1, 2, 3])
        self.assertIn('Page one', result['pages'][0]['text'])
        self.assertIn('Page two', result['pages'][1]['text'])
        self.assertIn('Page three', result['pages'][2]['text'])

    def test_max_pages_samples_without_changing_the_reported_page_count(self):
        document = self._make_document(make_pdf_bytes(['One', 'Two', 'Three']))
        result = self.reader.read_pages(document, max_pages=1)

        self.assertEqual(result['page_count'], 3)  # true total, not the sample size
        self.assertEqual(result['pages_read'], 1)
        self.assertEqual(len(result['pages']), 1)

    def test_empty_pdf_page_is_marked_as_requiring_ocr_not_silently_empty(self):
        document = self._make_document(make_pdf_bytes(['']))  # a page with nothing drawn on it
        result = self.reader.read_pages(document)

        self.assertEqual(result['page_count'], 1)
        self.assertEqual(result['pages'][0]['text'], '')
        self.assertTrue(result['pages'][0]['requires_ocr'])
        self.assertTrue(result['requires_ocr'])
        self.assertTrue(result['ocr_warnings'])  # the "why" isn't silent

    def test_scanned_pdf_is_detected_as_requiring_ocr_without_inventing_text(self):
        document = self._make_document(make_scanned_pdf_bytes())
        result = self.reader.read_pages(document)

        self.assertTrue(result['requires_ocr'])
        self.assertEqual(result['pages'][0]['text'], '')

    def test_corrupt_pdf_raises_permanent_error(self):
        document = self._make_document(b'this is not a pdf at all')
        with self.assertRaises(PermanentExtractionError):
            self.reader.read_pages(document)

    def test_unsupported_format_is_reported_honestly_not_silently_empty(self):
        document = self._make_document(
            b'\xd0\xcf\x11\xe0 old word file', filename='enrolment.doc',
            document_type='student_strength', content_type='application/msword',
        )
        result = self.reader.read_pages(document)

        self.assertFalse(result['format_supported'])
        self.assertEqual(result['pages'], [])
        self.assertIsNone(result['page_count'])
        self.assertIn('.doc', result['format_note'])
        self.assertIn('.docx or PDF', result['format_note'])  # tells the user what to do

    def test_configured_ocr_provider_is_consulted_for_low_text_pages(self):
        fake_ocr_provider = MagicMock()
        fake_ocr_provider.extract_text.return_value = 'This full paragraph of text was recovered by a real OCR backend.'
        reader = PDFPageReader(ocr_provider=fake_ocr_provider)
        document = self._make_document(make_pdf_bytes(['']))

        result = reader.read_pages(document)

        fake_ocr_provider.extract_text.assert_called_once()
        args, kwargs = fake_ocr_provider.extract_text.call_args
        self.assertEqual(args, (document, 1))
        # The page is rendered to a BGR image for the OCR engine.
        image = kwargs['image']
        self.assertEqual(image.ndim, 3)
        self.assertEqual(image.shape[2], 3)
        self.assertEqual(str(image.dtype), 'uint8')
        self.assertEqual(result['pages'][0]['text'], 'This full paragraph of text was recovered by a real OCR backend.')
        self.assertEqual(result['pages'][0]['text_source'], 'ocr')
        self.assertFalse(result['pages'][0]['requires_ocr'])
        self.assertEqual(result['ocr_pages'], 1)

    def test_digital_page_with_real_text_skips_ocr(self):
        fake_ocr_provider = MagicMock()
        reader = PDFPageReader(ocr_provider=fake_ocr_provider)
        document = self._make_document(make_pdf_bytes([
            'The institution has a total of 142 full-time faculty members across all departments.',
        ]))

        result = reader.read_pages(document)

        fake_ocr_provider.extract_text.assert_not_called()
        self.assertEqual(result['pages'][0]['text_source'], 'text')
        self.assertIn('142 full-time faculty', result['pages'][0]['text'])

    def test_page_ocr_cannot_read_either_stays_flagged(self):
        fake_ocr_provider = MagicMock()
        fake_ocr_provider.extract_text.return_value = ''
        reader = PDFPageReader(ocr_provider=fake_ocr_provider)
        document = self._make_document(make_pdf_bytes(['']))

        result = reader.read_pages(document)

        self.assertTrue(result['pages'][0]['requires_ocr'])
        self.assertEqual(result['pages'][0]['text_source'], 'none')
        self.assertIn('even after OCR', result['ocr_warnings'][0])

    def test_default_ocr_provider_never_fabricates_text(self):
        """With no real OCR backend configured, a low-text page stays
        honestly flagged rather than being filled with invented content."""
        document = self._make_document(make_pdf_bytes(['']))
        result = self.reader.read_pages(document)
        self.assertEqual(result['pages'][0]['text'], '')
        self.assertTrue(result['pages'][0]['requires_ocr'])


class TextCleaningTests(SimpleTestCase):
    def test_garbage_text_is_not_meaningful(self):
        self.assertFalse(is_meaningful_text('��� ### ~~~ ||| ��� ### ~~~ |||', 10))
        self.assertFalse(is_meaningful_text('short', 40))
        self.assertTrue(is_meaningful_text('Total sanctioned intake for B.E. CSE is 180 seats.', 40))

    def test_clean_text_normalises_whitespace_and_hyphenation(self):
        self.assertEqual(
            clean_text('Intro-\nduction   to\tAI\n\n\n\nAcademic year 2023-\n24\x0c'),
            'Introduction to AI\n\nAcademic year 2023-\n24',
        )

    def test_repeated_header_and_page_numbers_are_removed(self):
        pages = [f'ABC College of Engineering\nBody text for page {i}.\nPage {i} of 5' for i in range(1, 6)]
        cleaned = remove_repeated_headers_footers(pages)
        for i, text in enumerate(cleaned, start=1):
            self.assertEqual(text, f'Body text for page {i}.')

    def test_short_documents_keep_their_lines(self):
        pages = ['Heading\nBody one', 'Heading\nBody two']
        self.assertEqual(remove_repeated_headers_footers(pages), pages)


class OpenAIDocumentClassifierTests(TestCase):
    """Unit tests for the real, OpenAI-backed DocumentClassifier (apps.
    extraction.services.openai_classifier). The OpenAI service itself is
    always a fake double here -- see PDFPageReaderTests for real PDF I/O,
    and OpenAIExtractionServiceTests for the transport-level SDK mocking."""

    VALID_RESULT = {
        'document_type': 'naac_ssr',
        'document_title': 'Self Study Report',
        'reporting_year': '2025-26',
        'institution_name': 'M. Kumarasamy College of Engineering',
        'confidence': 0.92,
        'reasoning': 'The text contains NAAC SSR section headers and criteria.',
    }

    def setUp(self):
        institution = Institution.objects.create(name='MKCE', city='Karur', state='TN')
        self.sprint = Sprint.objects.create(institution=institution, mode=Sprint.SprintMode.VERIFIED_CRI)

    def _make_document(self, content=None, document_type='naac_ssr'):
        content = content if content is not None else make_pdf_bytes(['NAAC SSR self study report content.'])
        upload = SimpleUploadedFile('doc.pdf', content, content_type='application/pdf')
        return Document.objects.create(
            sprint=self.sprint, document_type=document_type, file=upload,
            original_filename='doc.pdf', mime_type='application/pdf', status=Document.Status.UPLOADED,
        )

    @staticmethod
    def _fake_service(result=None, side_effect=None):
        service = MagicMock()
        if side_effect is not None:
            service.extract_structured_data.side_effect = side_effect
        else:
            service.extract_structured_data.return_value = result
        return service

    def test_successful_classification_returns_the_validated_result(self):
        document = self._make_document()
        service = self._fake_service(dict(self.VALID_RESULT))
        classifier = OpenAIDocumentClassifier(openai_service=service)

        result = classifier.classify(document)

        self.assertEqual(result, self.VALID_RESULT)
        call_kwargs = service.extract_structured_data.call_args.kwargs
        self.assertIn('NAAC SSR self study report content', call_kwargs['user_content'])
        self.assertEqual(call_kwargs['response_schema'], CLASSIFICATION_SCHEMA)

    def test_uploaded_document_type_is_sent_as_a_hint(self):
        document = self._make_document(document_type='aqar_report')
        service = self._fake_service(dict(self.VALID_RESULT))
        classifier = OpenAIDocumentClassifier(openai_service=service)

        classifier.classify(document)

        call_kwargs = service.extract_structured_data.call_args.kwargs
        self.assertIn('aqar_report', call_kwargs['user_content'])

    def test_missing_information_is_accepted_as_null_not_an_error(self):
        document = self._make_document()
        service = self._fake_service({**self.VALID_RESULT, 'reporting_year': None, 'institution_name': None})
        classifier = OpenAIDocumentClassifier(openai_service=service)

        result = classifier.classify(document)

        self.assertIsNone(result['reporting_year'])
        self.assertIsNone(result['institution_name'])

    def test_document_with_no_extractable_text_still_gets_an_honest_prompt(self):
        document = self._make_document(content=make_pdf_bytes(['']))
        service = self._fake_service({**self.VALID_RESULT, 'confidence': 0.1})
        classifier = OpenAIDocumentClassifier(openai_service=service)

        classifier.classify(document)

        call_kwargs = service.extract_structured_data.call_args.kwargs
        self.assertIn('No extractable text was found', call_kwargs['user_content'])

    def test_confidence_out_of_range_is_rejected(self):
        document = self._make_document()
        service = self._fake_service({**self.VALID_RESULT, 'confidence': 1.5})
        classifier = OpenAIDocumentClassifier(openai_service=service)
        with self.assertRaises(ClassificationValidationError):
            classifier.classify(document)

    def test_negative_confidence_is_rejected(self):
        document = self._make_document()
        service = self._fake_service({**self.VALID_RESULT, 'confidence': -0.1})
        classifier = OpenAIDocumentClassifier(openai_service=service)
        with self.assertRaises(ClassificationValidationError):
            classifier.classify(document)

    def test_non_string_field_is_rejected(self):
        document = self._make_document()
        service = self._fake_service({**self.VALID_RESULT, 'document_title': 42})
        classifier = OpenAIDocumentClassifier(openai_service=service)
        with self.assertRaises(ClassificationValidationError):
            classifier.classify(document)

    def test_empty_reasoning_is_rejected(self):
        document = self._make_document()
        service = self._fake_service({**self.VALID_RESULT, 'reasoning': '   '})
        classifier = OpenAIDocumentClassifier(openai_service=service)
        with self.assertRaises(ClassificationValidationError):
            classifier.classify(document)

    def test_invalid_raw_ai_response_propagates_as_a_permanent_error(self):
        """A malformed raw response is already turned into OpenAIResponseError
        inside OpenAIExtractionService -- the classifier must let it
        propagate (it's already a PermanentExtractionError), not swallow it."""
        document = self._make_document()
        service = self._fake_service(side_effect=OpenAIResponseError('not valid JSON'))
        classifier = OpenAIDocumentClassifier(openai_service=service)
        with self.assertRaises(PermanentExtractionError):
            classifier.classify(document)


def _page(number, text):
    return {'page_number': number, 'text': text}


def _valid_fact(**overrides):
    fact = {
        'field_name': 'Total Faculty Count',
        'field_key': 'total_faculty',
        'value': '42',
        'data_type': 'number',
        'pillar': 'aiqif_c4',
        'owner_role': 'hr_officer',
        'source_page': '1',
        'source_snippet': 'The institution has 42 faculty members.',
        'confidence_score': 0.96,
        'confidence_reason': 'Explicitly stated in the source.',
    }
    fact.update(overrides)
    return fact


class OpenAIFactExtractorTests(TestCase):
    """Unit tests for the real, OpenAI-backed FactExtractor (apps.
    extraction.services.openai_fact_extractor). The OpenAI service is
    always a fake double here -- see PDFPageReaderTests for real PDF I/O."""

    def setUp(self):
        institution = Institution.objects.create(name='MKCE', city='Karur', state='TN')
        self.sprint = Sprint.objects.create(institution=institution, mode=Sprint.SprintMode.VERIFIED_CRI)
        self.document = Document.objects.create(
            sprint=self.sprint, document_type='naac_ssr', original_filename='ssr.pdf',
            mime_type='application/pdf', status=Document.Status.UPLOADED,
        )

    @staticmethod
    def _fake_service(results=None, side_effect=None):
        service = MagicMock()
        if side_effect is not None:
            service.extract_structured_data.side_effect = side_effect
        elif results is not None:
            service.extract_structured_data.side_effect = results
        return service

    # -- basic behavior -----------------------------------------------------

    def test_successful_extraction_returns_the_validated_fact(self):
        service = self._fake_service(results=[{'facts': [_valid_fact()]}])
        extractor = OpenAIFactExtractor(openai_service=service)
        pages = {'pages': [_page(1, 'The institution has 42 faculty members on its rolls.')]}

        facts = extractor.extract_facts(self.document, pages)

        self.assertEqual(len(facts), 1)
        fact = facts[0]
        self.assertEqual(fact['field_key'], 'total_faculty')
        self.assertEqual(fact['value'], 42)
        self.assertEqual(fact['normalized_value'], 42)
        self.assertEqual(fact['pillar'], 'aiqif_c4')
        self.assertEqual(fact['owner_role'], 'hr_officer')
        self.assertEqual(fact['source_page'], '1')
        self.assertEqual(fact['extraction_method'], 'openai')

    def test_empty_document_never_calls_openai(self):
        """A document with no extractable text on any page shouldn't waste
        an API call -- there's nothing to extract facts from."""
        service = self._fake_service(results=[])
        extractor = OpenAIFactExtractor(openai_service=service)

        facts = extractor.extract_facts(self.document, {'pages': [_page(1, ''), _page(2, '   ')]})

        self.assertEqual(facts, [])
        service.extract_structured_data.assert_not_called()

    def test_no_pages_at_all_never_calls_openai(self):
        service = self._fake_service(results=[])
        extractor = OpenAIFactExtractor(openai_service=service)
        facts = extractor.extract_facts(self.document, {'pages': []})
        self.assertEqual(facts, [])
        service.extract_structured_data.assert_not_called()

    # -- evidence requirements / hallucination prevention -------------------

    def test_fact_with_missing_evidence_snippet_is_dropped(self):
        service = self._fake_service(results=[{'facts': [_valid_fact(source_snippet='')]}])
        extractor = OpenAIFactExtractor(openai_service=service)
        facts = extractor.extract_facts(self.document, {'pages': [_page(1, 'Some real page text here.')]})
        self.assertEqual(facts, [])

    def test_fact_citing_a_page_outside_the_chunk_is_dropped(self):
        """The core hallucination guard: a source_page the model wasn't
        actually given (it only saw page 1) must not be trusted."""
        service = self._fake_service(results=[{'facts': [_valid_fact(source_page='99')]}])
        extractor = OpenAIFactExtractor(openai_service=service)
        facts = extractor.extract_facts(self.document, {'pages': [_page(1, 'Some real page text here.')]})
        self.assertEqual(facts, [])

    def test_fact_with_null_source_page_is_still_accepted(self):
        """'Prefer null over guessing' -- an honestly-unknown page shouldn't
        be penalized the way a fabricated one is."""
        service = self._fake_service(results=[{'facts': [_valid_fact(source_page=None)]}])
        extractor = OpenAIFactExtractor(openai_service=service)
        facts = extractor.extract_facts(self.document, {'pages': [_page(1, 'Some real page text here.')]})
        self.assertEqual(len(facts), 1)
        self.assertEqual(facts[0]['source_page'], '')

    # -- structured output / confidence validation ---------------------------

    def test_fact_with_invalid_data_type_is_dropped(self):
        service = self._fake_service(results=[{'facts': [_valid_fact(data_type='integer')]}])  # not a real choice
        extractor = OpenAIFactExtractor(openai_service=service)
        facts = extractor.extract_facts(self.document, {'pages': [_page(1, 'Some real page text here.')]})
        self.assertEqual(facts, [])

    def test_fact_with_invalid_pillar_is_dropped(self):
        service = self._fake_service(results=[{'facts': [_valid_fact(pillar='not_a_real_pillar')]}])
        extractor = OpenAIFactExtractor(openai_service=service)
        facts = extractor.extract_facts(self.document, {'pages': [_page(1, 'Some real page text here.')]})
        self.assertEqual(facts, [])

    def test_fact_with_invalid_owner_role_is_dropped(self):
        service = self._fake_service(results=[{'facts': [_valid_fact(owner_role='super_admin')]}])
        extractor = OpenAIFactExtractor(openai_service=service)
        facts = extractor.extract_facts(self.document, {'pages': [_page(1, 'Some real page text here.')]})
        self.assertEqual(facts, [])

    def test_malformed_top_level_response_is_treated_like_any_other_chunk_failure(self):
        """A response missing the "facts" list is turned into
        OpenAIResponseError inside _extract_chunk -- since that's a
        PermanentExtractionError, it's handled the same as any other
        per-chunk failure (logged, that chunk contributes no facts) rather
        than raised (see test_permanent_error_on_one_chunk_does_not_block_
        the_others for why single-chunk failures shouldn't be fatal)."""
        service = self._fake_service(results=[{'not_facts': []}])
        extractor = OpenAIFactExtractor(openai_service=service)
        facts = extractor.extract_facts(self.document, {'pages': [_page(1, 'Some real page text here.')]})
        self.assertEqual(facts, [])

    def test_confidence_out_of_range_is_dropped(self):
        service = self._fake_service(results=[{'facts': [_valid_fact(confidence_score=1.5)]}])
        extractor = OpenAIFactExtractor(openai_service=service)
        facts = extractor.extract_facts(self.document, {'pages': [_page(1, 'Some real page text here.')]})
        self.assertEqual(facts, [])

    def test_negative_confidence_is_dropped(self):
        service = self._fake_service(results=[{'facts': [_valid_fact(confidence_score=-0.2)]}])
        extractor = OpenAIFactExtractor(openai_service=service)
        facts = extractor.extract_facts(self.document, {'pages': [_page(1, 'Some real page text here.')]})
        self.assertEqual(facts, [])

    def test_boundary_confidence_values_are_accepted(self):
        service = self._fake_service(results=[{'facts': [
            _valid_fact(field_key='a', confidence_score=0.0),
            _valid_fact(field_key='b', confidence_score=1.0),
        ]}])
        extractor = OpenAIFactExtractor(openai_service=service)
        facts = extractor.extract_facts(self.document, {'pages': [_page(1, 'Some real page text here.')]})
        self.assertEqual({f['field_key'] for f in facts}, {'a', 'b'})

    # -- Django-side value typing --------------------------------------------

    def test_percentage_value_is_normalized_to_a_number(self):
        service = self._fake_service(results=[{'facts': [
            _valid_fact(field_key='placement_rate', data_type='percentage', value='87%'),
        ]}])
        extractor = OpenAIFactExtractor(openai_service=service)
        facts = extractor.extract_facts(self.document, {'pages': [_page(1, 'Placement rate: 87%.')]})
        self.assertEqual(facts[0]['normalized_value'], 87.0)

    def test_list_value_is_split_into_items(self):
        service = self._fake_service(results=[{'facts': [
            _valid_fact(
                field_key='ai_software', data_type='list', pillar='aiqif_c6',
                owner_role='lab_admin', value='TensorFlow; PyTorch; MATLAB',
            ),
        ]}])
        extractor = OpenAIFactExtractor(openai_service=service)
        facts = extractor.extract_facts(self.document, {'pages': [_page(1, 'Licensed AI software text.')]})
        self.assertEqual(facts[0]['value'], ['TensorFlow', 'PyTorch', 'MATLAB'])

    def test_unparseable_number_is_dropped_not_coerced(self):
        service = self._fake_service(results=[{'facts': [_valid_fact(value='approximately forty-two')]}])
        extractor = OpenAIFactExtractor(openai_service=service)
        facts = extractor.extract_facts(self.document, {'pages': [_page(1, 'Some real page text here.')]})
        self.assertEqual(facts, [])

    # -- multiple pages / multiple chunks / duplicates -----------------------

    def test_facts_from_different_pages_in_one_chunk_are_both_kept(self):
        service = self._fake_service(results=[{'facts': [
            _valid_fact(field_key='total_faculty', source_page='1'),
            _valid_fact(
                field_key='student_strength', pillar='aiqif_c5', owner_role='registrar',
                source_page='2', value='1200', source_snippet='Total student enrolment is 1200.',
            ),
        ]}])
        extractor = OpenAIFactExtractor(openai_service=service)
        pages = {'pages': [_page(1, 'Faculty section text.'), _page(2, 'Enrolment section text.')]}

        facts = extractor.extract_facts(self.document, pages)

        by_key = {f['field_key']: f for f in facts}
        self.assertEqual(by_key['total_faculty']['source_page'], '1')
        self.assertEqual(by_key['student_strength']['source_page'], '2')

    def test_large_document_is_split_into_multiple_chunks(self):
        pages = {'pages': [_page(i, f'Page {i} content. ' * 10) for i in range(1, 4)]}
        service = self._fake_service(results=[{'facts': []}, {'facts': []}, {'facts': []}])
        # A tiny max_chunk_chars forces one page per chunk.
        extractor = OpenAIFactExtractor(openai_service=service, max_chunk_chars=50)

        extractor.extract_facts(self.document, pages)

        self.assertEqual(service.extract_structured_data.call_count, 3)
        first_call_content = service.extract_structured_data.call_args_list[0].kwargs['user_content']
        self.assertIn('chunk 1 of 3', first_call_content)

    def test_duplicate_facts_across_chunks_keep_the_strongest_evidence(self):
        pages = {'pages': [_page(i, f'Page {i} content. ' * 10) for i in range(1, 3)]}
        service = self._fake_service(results=[
            {'facts': [_valid_fact(source_page='1', confidence_score=0.6, source_snippet='Weaker mention.')]},
            {'facts': [_valid_fact(source_page='2', confidence_score=0.95, source_snippet='Clear, explicit statement.')]},
        ])
        extractor = OpenAIFactExtractor(openai_service=service, max_chunk_chars=50)

        facts = extractor.extract_facts(self.document, pages)

        self.assertEqual(len(facts), 1)
        self.assertEqual(facts[0]['confidence_score'], 0.95)
        self.assertEqual(facts[0]['source_page'], '2')

    def test_too_many_chunks_are_capped_not_unbounded(self):
        pages = {'pages': [_page(i, f'Page {i} content. ' * 10) for i in range(1, 6)]}
        service = self._fake_service(results=[{'facts': []}] * 10)
        extractor = OpenAIFactExtractor(openai_service=service, max_chunk_chars=50, max_chunks=2)

        extractor.extract_facts(self.document, pages)

        self.assertEqual(service.extract_structured_data.call_count, 2)

    # -- OpenAI failure handling ----------------------------------------------

    def test_recoverable_error_on_a_chunk_propagates(self):
        service = self._fake_service(side_effect=RecoverableExtractionError('rate limited'))
        extractor = OpenAIFactExtractor(openai_service=service)
        with self.assertRaises(RecoverableExtractionError):
            extractor.extract_facts(self.document, {'pages': [_page(1, 'Some real page text here.')]})

    def test_permanent_error_on_one_chunk_does_not_block_the_others(self):
        pages = {'pages': [_page(i, f'Page {i} content. ' * 10) for i in range(1, 3)]}
        service = self._fake_service(side_effect=[
            PermanentExtractionError('malformed response for this chunk'),
            {'facts': [_valid_fact(source_page='2')]},
        ])
        extractor = OpenAIFactExtractor(openai_service=service, max_chunk_chars=50)

        facts = extractor.extract_facts(self.document, pages)

        self.assertEqual(len(facts), 1)
        self.assertEqual(facts[0]['source_page'], '2')


@unittest.skipUnless(
    settings.RUN_OPENAI_INTEGRATION_TESTS,
    'Set RUN_OPENAI_INTEGRATION_TESTS=true (and a real OPENAI_API_KEY) to run this against the real API.',
)
class OpenAIFactExtractorIntegrationTests(TestCase):
    """Opt-in only -- makes a real call to the real OpenAI API against a
    real PDF, spending real quota. Never runs as part of a normal
    `manage.py test`; see the skip condition above."""

    def test_real_extraction_against_a_real_pdf(self):
        institution = Institution.objects.create(name='MKCE', city='Karur', state='TN')
        sprint = Sprint.objects.create(institution=institution, mode=Sprint.SprintMode.VERIFIED_CRI)
        document = Document.objects.create(
            sprint=sprint, document_type='naac_ssr', original_filename='ssr.pdf',
            mime_type='application/pdf', status=Document.Status.UPLOADED,
        )
        pages = {'pages': [_page(1, (
            'AIOS Institute of Technology — Self Study Report\n'
            'The institution has 128 full-time faculty members, of whom 34 hold an AI/ML '
            'certification recognized by AICTE. Total student enrolment for AY 2025-26 is 2,450. '
            'The placement cell reported an 82% placement rate for the graduating batch, with 47 '
            'signed industry MoUs including partnerships with three AI research labs.'
        ))]}

        extractor = OpenAIFactExtractor()  # real OpenAIExtractionService, real API key from settings
        facts = extractor.extract_facts(document, pages)

        self.assertTrue(facts, 'Expected the real API to extract at least one fact from clearly-stated content.')
        for fact in facts:
            self.assertTrue(fact['source_snippet'])
            self.assertTrue(0.0 <= fact['confidence_score'] <= 1.0)


class RuleBasedGapDetectorTests(TestCase):
    """Unit tests for the real, deterministic GapDetector (apps.extraction.
    services.gap_detector). No AI involved -- nothing here is mocked."""

    def setUp(self):
        institution = Institution.objects.create(name='MKCE', city='Karur', state='TN')
        self.sprint = Sprint.objects.create(institution=institution, mode=Sprint.SprintMode.VERIFIED_CRI)
        self.document = Document.objects.create(
            sprint=self.sprint, document_type='naac_ssr', original_filename='ssr.pdf',
            mime_type='application/pdf', status=Document.Status.PROCESSED, uploaded_at=timezone.now(),
        )
        self.detector = RuleBasedGapDetector()

    def _make_fact(self, **overrides):
        defaults = {
            'sprint': self.sprint, 'document': self.document, 'source_document': self.document,
            'field_name': 'Total Faculty Count', 'field_key': 'total_faculty', 'value': 42,
            'normalized_value': 42, 'data_type': ExtractedFact.DataType.NUMBER,
            'pillar': 'aiqif_c4', 'owner_role': 'hr_officer',
            'source_snippet': '42 faculty members.', 'confidence_score': 0.9,
            'extraction_method': 'openai', 'status': ExtractedFact.Status.EXTRACTED,
        }
        defaults.update(overrides)
        return ExtractedFact.objects.create(**defaults)

    def test_low_confidence_fact_yields_low_confidence_gap(self):
        fact = self._make_fact(confidence_score=0.3)
        gaps = self.detector.detect_gaps(self.document, [])
        low_conf = [g for g in gaps if g['gap_type'] == GapItem.GapType.LOW_CONFIDENCE]
        self.assertEqual(len(low_conf), 1)
        self.assertEqual(low_conf[0]['source_fact'], fact)
        self.assertEqual(low_conf[0]['priority'], GapItem.Priority.HIGH)  # below the "very low" threshold

    def test_moderately_low_confidence_is_medium_priority(self):
        self._make_fact(confidence_score=0.6)
        gaps = self.detector.detect_gaps(self.document, [])
        low_conf = [g for g in gaps if g['gap_type'] == GapItem.GapType.LOW_CONFIDENCE]
        self.assertEqual(low_conf[0]['priority'], GapItem.Priority.MEDIUM)

    def test_high_confidence_fact_yields_unconfirmed_fact_gap(self):
        fact = self._make_fact(confidence_score=0.95)
        gaps = self.detector.detect_gaps(self.document, [])
        unconfirmed = [g for g in gaps if g['gap_type'] == GapItem.GapType.UNCONFIRMED_FACT]
        self.assertEqual(len(unconfirmed), 1)
        self.assertEqual(unconfirmed[0]['source_fact'], fact)

    def test_confirmed_fact_yields_no_gap(self):
        self._make_fact(confidence_score=0.95, status=ExtractedFact.Status.CONFIRMED)
        gaps = self.detector.detect_gaps(self.document, [])
        self.assertEqual(gaps, [])

    def test_stale_document_yields_stale_data_gap(self):
        self.document.uploaded_at = timezone.now() - timedelta(days=settings.GAP_STALE_DATA_DAYS + 1)
        self.document.save(update_fields=['uploaded_at'])
        gaps = self.detector.detect_gaps(self.document, [])
        stale = [g for g in gaps if g['gap_type'] == GapItem.GapType.STALE_DATA]
        self.assertEqual(len(stale), 1)
        self.assertEqual(stale[0]['related_document'], self.document)

    def test_recent_document_yields_no_stale_gap(self):
        gaps = self.detector.detect_gaps(self.document, [])
        self.assertEqual([g for g in gaps if g['gap_type'] == GapItem.GapType.STALE_DATA], [])

    def test_document_with_no_uploaded_at_yields_no_stale_gap(self):
        self.document.uploaded_at = None
        self.document.save(update_fields=['uploaded_at'])
        gaps = self.detector.detect_gaps(self.document, [])
        self.assertEqual([g for g in gaps if g['gap_type'] == GapItem.GapType.STALE_DATA], [])

    def test_never_detects_missing_document(self):
        """Deliberate scope boundary -- see gap_detector.py's module
        docstring for why missing_document stays a sprint-level concern
        rather than something this per-document detector attempts."""
        gaps = self.detector.detect_gaps(self.document, [])
        self.assertFalse(any(g['gap_type'] == GapItem.GapType.MISSING_DOCUMENT for g in gaps))


class OpenAIConflictCheckerTests(TestCase):
    """Unit tests for the real, OpenAI-backed ConflictChecker (apps.
    extraction.services.conflict_checker). The OpenAI service is always a
    fake double here -- never a real network call."""

    def setUp(self):
        institution = Institution.objects.create(name='MKCE', city='Karur', state='TN')
        self.sprint = Sprint.objects.create(institution=institution, mode=Sprint.SprintMode.VERIFIED_CRI)
        self.document_a = Document.objects.create(
            sprint=self.sprint, document_type='naac_ssr', original_filename='a.pdf',
            mime_type='application/pdf', status=Document.Status.PROCESSED,
        )
        self.document_b = Document.objects.create(
            sprint=self.sprint, document_type='aqar_report', original_filename='b.pdf',
            mime_type='application/pdf', status=Document.Status.PROCESSED,
        )

    def _make_fact(self, document, **overrides):
        defaults = {
            'sprint': self.sprint, 'document': document, 'source_document': document,
            'field_name': 'Total Faculty Count', 'field_key': 'total_faculty', 'value': 42,
            'normalized_value': 42, 'data_type': ExtractedFact.DataType.NUMBER,
            'pillar': 'aiqif_c4', 'owner_role': 'hr_officer',
            'source_snippet': '42 faculty members.', 'confidence_score': 0.9,
            'extraction_method': 'openai', 'status': ExtractedFact.Status.EXTRACTED,
        }
        defaults.update(overrides)
        return ExtractedFact.objects.create(**defaults)

    @staticmethod
    def _fake_service(results=None, side_effect=None):
        service = MagicMock()
        if side_effect is not None:
            service.extract_structured_data.side_effect = side_effect
        elif results is not None:
            service.extract_structured_data.side_effect = results
        return service

    def test_no_other_documents_never_calls_openai(self):
        self._make_fact(self.document_a)
        service = self._fake_service(results=[])
        checker = OpenAIConflictChecker(openai_service=service)
        conflicts = checker.check_conflicts(self.document_a, [])
        self.assertEqual(conflicts, [])
        service.extract_structured_data.assert_not_called()

    def test_matching_normalized_values_are_not_candidates(self):
        """The deterministic pre-filter: two facts that genuinely agree
        never reach the AI at all."""
        self._make_fact(self.document_a, value=42, normalized_value=42)
        self._make_fact(self.document_b, value=42, normalized_value=42)
        service = self._fake_service(results=[])
        checker = OpenAIConflictChecker(openai_service=service)
        conflicts = checker.check_conflicts(self.document_a, [])
        self.assertEqual(conflicts, [])
        service.extract_structured_data.assert_not_called()

    def test_confirmed_semantic_conflict_creates_a_conflict_dict(self):
        fact_a = self._make_fact(self.document_a, value=42, normalized_value=42)
        fact_b = self._make_fact(self.document_b, value=47, normalized_value=47)
        service = self._fake_service(results=[
            {'is_conflict': True, 'confidence': 0.88, 'explanation': 'Both describe total faculty headcount.'},
        ])
        checker = OpenAIConflictChecker(openai_service=service)

        conflicts = checker.check_conflicts(self.document_a, [])

        self.assertEqual(len(conflicts), 1)
        conflict = conflicts[0]
        self.assertEqual(conflict['gap_type'], GapItem.GapType.CONFLICT)
        self.assertEqual(conflict['source_fact'], fact_a)
        self.assertEqual(conflict['conflict_fact_b'], fact_b)
        self.assertEqual(conflict['conflict_value_a'], 42)
        self.assertEqual(conflict['conflict_value_b'], 47)
        self.assertEqual(conflict['conflict_confidence'], 0.88)
        self.assertIn('faculty headcount', conflict['description'])

    def test_semantic_non_conflict_creates_nothing(self):
        """The AI can determine the two values don't actually conflict
        (different populations) -- no gap, and neither fact's own value is
        touched (never 'silently resolved')."""
        fact_a = self._make_fact(self.document_a, value=42, normalized_value=42)
        fact_b = self._make_fact(self.document_b, value=47, normalized_value=47)
        service = self._fake_service(results=[{
            'is_conflict': False, 'confidence': 0.7,
            'explanation': 'One covers the CS department only, the other the whole institution.',
        }])
        checker = OpenAIConflictChecker(openai_service=service)

        conflicts = checker.check_conflicts(self.document_a, [])

        self.assertEqual(conflicts, [])
        fact_a.refresh_from_db()
        fact_b.refresh_from_db()
        self.assertEqual(fact_a.value, 42)
        self.assertEqual(fact_b.value, 47)

    def test_invalid_is_conflict_type_is_dropped(self):
        self._make_fact(self.document_a, value=42, normalized_value=42)
        self._make_fact(self.document_b, value=47, normalized_value=47)
        service = self._fake_service(results=[{'is_conflict': 'yes', 'confidence': 0.8, 'explanation': 'x'}])
        checker = OpenAIConflictChecker(openai_service=service)
        self.assertEqual(checker.check_conflicts(self.document_a, []), [])

    def test_confidence_out_of_range_is_dropped(self):
        self._make_fact(self.document_a, value=42, normalized_value=42)
        self._make_fact(self.document_b, value=47, normalized_value=47)
        service = self._fake_service(results=[{'is_conflict': True, 'confidence': 1.4, 'explanation': 'x'}])
        checker = OpenAIConflictChecker(openai_service=service)
        self.assertEqual(checker.check_conflicts(self.document_a, []), [])

    def test_empty_explanation_is_dropped(self):
        self._make_fact(self.document_a, value=42, normalized_value=42)
        self._make_fact(self.document_b, value=47, normalized_value=47)
        service = self._fake_service(results=[{'is_conflict': True, 'confidence': 0.8, 'explanation': '   '}])
        checker = OpenAIConflictChecker(openai_service=service)
        self.assertEqual(checker.check_conflicts(self.document_a, []), [])

    def test_rejected_facts_excluded_from_candidates(self):
        self._make_fact(self.document_a, value=42, normalized_value=42)
        self._make_fact(self.document_b, value=47, normalized_value=47, status=ExtractedFact.Status.REJECTED)
        service = self._fake_service(results=[])
        checker = OpenAIConflictChecker(openai_service=service)
        conflicts = checker.check_conflicts(self.document_a, [])
        self.assertEqual(conflicts, [])
        service.extract_structured_data.assert_not_called()

    def test_recoverable_error_propagates(self):
        self._make_fact(self.document_a, value=42, normalized_value=42)
        self._make_fact(self.document_b, value=47, normalized_value=47)
        service = self._fake_service(side_effect=RecoverableExtractionError('rate limited'))
        checker = OpenAIConflictChecker(openai_service=service)
        with self.assertRaises(RecoverableExtractionError):
            checker.check_conflicts(self.document_a, [])

    def test_permanent_error_on_one_pair_does_not_block_others(self):
        self._make_fact(self.document_a, value=42, normalized_value=42)
        self._make_fact(self.document_b, value=47, normalized_value=47)
        document_c = Document.objects.create(
            sprint=self.sprint, document_type='placement_report', original_filename='c.pdf',
            mime_type='application/pdf', status=Document.Status.PROCESSED,
        )
        self._make_fact(document_c, value=50, normalized_value=50)
        service = self._fake_service(side_effect=[
            PermanentExtractionError('malformed response'),
            {'is_conflict': True, 'confidence': 0.9, 'explanation': 'Genuinely disagree.'},
        ])
        checker = OpenAIConflictChecker(openai_service=service)

        conflicts = checker.check_conflicts(self.document_a, [])

        self.assertEqual(len(conflicts), 1)

    def test_too_many_pairs_are_capped(self):
        self._make_fact(self.document_a, value=42, normalized_value=42)
        for i in range(5):
            doc = Document.objects.create(
                sprint=self.sprint, document_type='placement_report', original_filename=f'extra{i}.pdf',
                mime_type='application/pdf', status=Document.Status.PROCESSED,
            )
            self._make_fact(doc, value=100 + i, normalized_value=100 + i)
        service = self._fake_service(results=[{'is_conflict': False, 'confidence': 0.5, 'explanation': 'x'}] * 10)
        checker = OpenAIConflictChecker(openai_service=service, max_pairs=2)

        checker.check_conflicts(self.document_a, [])

        self.assertEqual(service.extract_structured_data.call_count, 2)


@override_settings(CELERY_TASK_ALWAYS_EAGER=True, CELERY_TASK_EAGER_PROPAGATES=True)
class EndToEndGapAndConflictTests(ExtractionJobTestBase):
    """Proves the full documented flow -- PDF -> PageReader -> Classifier ->
    OpenAIFactExtractor -> AuditFieldMapper -> GapDetector -> ConflictChecker
    -> Review Workspace -- actually persists real GapItem rows, including a
    genuine cross-document conflict, rather than returning temporary
    results that never reach the database."""

    def setUp(self):
        super().setUp()
        self._auth(self.iqac)
        upload = self.client.post(
            f'/api/v1/sprints/{self.sprint.id}/upload-file/',
            {
                'file': make_pdf(name='doc2.pdf', page_texts=('A second, distinct document for conflict tests.',)),
                'document_type': 'aqar_report',
            },
            format='multipart',
        )
        self.document2 = Document.objects.get(id=upload.data['id'])
        self.client.credentials()

    def test_conflicting_facts_across_documents_are_persisted_as_a_conflict_gap(self):
        def fake_extract_facts(*, system_prompt, user_content, response_schema, schema_name='extraction_result', **kwargs):
            if 'doc2.pdf' in user_content:
                return {'facts': [_valid_fact(value='47', source_snippet='There are 47 teaching staff.')]}
            return {'facts': [_valid_fact(value='42', source_snippet='The institution has 42 faculty members.')]}

        self.fact_openai_mock.return_value.extract_structured_data.side_effect = fake_extract_facts

        conflict_patcher = patch('apps.extraction.services.conflict_checker.get_ai_service')
        mock_conflict_openai_cls = conflict_patcher.start()
        mock_conflict_openai_cls.return_value.extract_structured_data.return_value = {
            'is_conflict': True, 'confidence': 0.85,
            'explanation': 'Both figures describe the same total faculty headcount.',
        }
        self.addCleanup(conflict_patcher.stop)

        self._auth(self.iqac)
        self.client.post(f'/api/v1/sprints/{self.sprint.id}/extraction-jobs/')

        facts = list(ExtractedFact.objects.filter(sprint=self.sprint, field_key='total_faculty'))
        self.assertEqual(len(facts), 2)
        self.assertEqual({f.value for f in facts}, {42, 47})

        conflict = GapItem.objects.get(gap_type=GapItem.GapType.CONFLICT)
        self.assertEqual({conflict.conflict_value_a, conflict.conflict_value_b}, {42, 47})
        self.assertEqual(conflict.conflict_confidence, 0.85)
        self.assertIn('faculty headcount', conflict.description)
        self.assertEqual(conflict.status, GapItem.Status.OPEN)  # marked for human review, not auto-resolved

    def test_matching_facts_across_documents_create_no_conflict(self):
        self.fact_openai_mock.return_value.extract_structured_data.return_value = {
            'facts': [_valid_fact(value='42')],
        }
        conflict_patcher = patch('apps.extraction.services.conflict_checker.get_ai_service')
        mock_conflict_openai_cls = conflict_patcher.start()
        self.addCleanup(conflict_patcher.stop)

        self._auth(self.iqac)
        self.client.post(f'/api/v1/sprints/{self.sprint.id}/extraction-jobs/')

        self.assertEqual(GapItem.objects.filter(gap_type=GapItem.GapType.CONFLICT).count(), 0)
        mock_conflict_openai_cls.return_value.extract_structured_data.assert_not_called()



class AIPricingTests(SimpleTestCase):
    def test_cost_from_price_table(self):
        from decimal import Decimal
        from .pricing import cost_usd
        # Sonnet 5: $2 in / $10 out per million.
        self.assertEqual(cost_usd('claude-sonnet-5', input_tokens=1_000_000, output_tokens=500_000), Decimal('7.000000'))
        # Dated / prefixed names resolve to the base model.
        self.assertEqual(cost_usd('anthropic/claude-haiku-4-5-20251001', input_tokens=1_000_000), Decimal('1.000000'))
        # Longest match wins: opus-5-5 is not priced as opus-5.
        self.assertEqual(cost_usd('claude-opus-5-5', input_tokens=1_000_000), Decimal('4.000000'))

    def test_unknown_model_has_no_price_until_configured(self):
        from decimal import Decimal
        from .pricing import cost_usd
        self.assertIsNone(cost_usd('my-combo', input_tokens=1000, output_tokens=1000))
        with override_settings(AI_PRICING_JSON='{"my-combo": {"input": 1, "output": 2}}'):
            self.assertEqual(cost_usd('my-combo', input_tokens=1_000_000, output_tokens=1_000_000), Decimal('3.000000'))


@override_settings(CELERY_TASK_ALWAYS_EAGER=True, CELERY_TASK_EAGER_PROPAGATES=True, AI_BASE_URL='')
class AIUsageRecordingTests(ExtractionJobTestBase):
    """Every AI response made during an extraction is stored with its tokens
    and cost, attributed to the sprint / document / job / stage."""

    def _job(self):
        return ExtractionJob.objects.create(sprint=self.sprint, document=self.document)

    def test_openai_compatible_call_is_recorded_inside_a_scope_only(self):
        from .models import AIUsage
        from .usage import usage_scope
        completion = _fake_completion(content='{"a": 1}')
        completion.usage.prompt_tokens = 1200
        completion.usage.prompt_tokens_details.cached_tokens = 200
        completion.usage.completion_tokens = 300
        completion.usage.completion_tokens_details.reasoning_tokens = 100
        client = MagicMock()
        client.chat.completions.create.return_value = completion
        service = OpenAIExtractionService(api_key='k', model='claude-sonnet-5', client=client)

        service.extract_structured_data(system_prompt='s', user_content='u', response_schema={'type': 'object'})
        self.assertEqual(AIUsage.objects.count(), 0)  # outside an extraction: not recorded

        job = self._job()
        with usage_scope(job=job, stage=AIUsage.Stage.FACT_EXTRACTION):
            service.extract_structured_data(system_prompt='s', user_content='u', response_schema={'type': 'object'})
        row = AIUsage.objects.get()
        self.assertEqual((row.sprint, row.document, row.extraction_job), (self.sprint, self.document, job))
        self.assertEqual(row.institution_id, self.sprint.institution_id)
        self.assertEqual(row.stage, 'fact_extraction')
        self.assertEqual(row.provider, 'openai')
        self.assertEqual(
            (row.input_tokens, row.cache_read_tokens, row.output_tokens, row.reasoning_tokens), (1000, 200, 300, 100),
        )
        # 1000 x $2 + 200 x $0.2 + 300 x $10, per million tokens
        self.assertEqual(str(row.cost_usd), '0.005040')

    def test_anthropic_call_is_recorded_even_when_the_reply_is_unusable(self):
        from .models import AIUsage
        from .services.anthropic_client import AnthropicExtractionService
        from .usage import usage_scope
        message = _fake_tool_use_message(input_data=None)  # no tool call -> parse error
        message.usage.input_tokens = 900
        message.usage.output_tokens = 2100  # includes thinking
        message.usage.cache_read_input_tokens = 0
        message.usage.cache_creation_input_tokens = 0
        client = MagicMock()
        client.messages.create.return_value = message
        service = AnthropicExtractionService(api_key='sk-ant-x', model='claude-haiku-4-5', client=client)
        with usage_scope(job=self._job(), stage=AIUsage.Stage.CLASSIFICATION):
            with self.assertRaises(Exception):
                service.extract_structured_data(system_prompt='s', user_content='u', response_schema={'type': 'object'})
        row = AIUsage.objects.get()
        self.assertEqual(
            (row.provider, row.stage, row.input_tokens, row.output_tokens), ('anthropic', 'classification', 900, 2100),
        )
        self.assertEqual(str(row.cost_usd), '0.011400')  # 900 x $1 + 2100 x $5, per million

    def _usage_rows(self):
        from decimal import Decimal
        from .models import AIUsage
        job = self._job()
        common = dict(
            institution=self.sprint.institution, sprint=self.sprint, document=self.document, extraction_job=job,
        )
        AIUsage.objects.create(stage='classification', model='claude-sonnet-5', input_tokens=1000,
                               output_tokens=200, cost_usd=Decimal('0.004'), **common)
        AIUsage.objects.create(stage='fact_extraction', model='claude-sonnet-5', input_tokens=5000,
                               output_tokens=1500, cost_usd=Decimal('0.025'), **common)
        AIUsage.objects.create(stage='fact_extraction', model='my-combo', input_tokens=100,
                               output_tokens=50, cost_usd=None, **common)
        return job

    def test_sprint_usage_endpoint_totals_and_breakdowns(self):
        self._usage_rows()
        self._auth(self.iqac)
        response = self.client.get(f'/api/v1/sprints/{self.sprint.id}/ai-usage')
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        self.assertEqual(response.data['calls'], 3)
        self.assertEqual(response.data['input_tokens'], 6100)
        self.assertEqual(response.data['output_tokens'], 1750)
        self.assertEqual(response.data['cost_usd'], '0.029000')
        self.assertEqual(response.data['unpriced_calls'], 1)
        self.assertEqual(response.data['unpriced_models'], ['my-combo'])
        self.assertEqual(
            {r['stage']: r['calls'] for r in response.data['by_stage']}, {'classification': 1, 'fact_extraction': 2},
        )
        self.assertEqual(response.data['top_documents'][0]['filename'], self.document.original_filename)

    def test_job_list_includes_each_jobs_tokens_and_cost(self):
        job = self._usage_rows()
        self._auth(self.iqac)
        rows = self.client.get(f'/api/v1/sprints/{self.sprint.id}/extraction-jobs/').data
        row = next(r for r in rows if r['id'] == str(job.id))
        self.assertEqual(
            (row['ai_input_tokens'], row['ai_output_tokens'], row['ai_cost_usd']), (6100, 1750, '0.029000'),
        )

    def test_usage_survives_deleting_the_job(self):
        from .models import AIUsage
        job = self._usage_rows()
        job.delete()
        self.assertEqual(AIUsage.objects.filter(sprint=self.sprint).count(), 3)

    def test_billing_report_command(self):
        from io import StringIO
        from django.core.management import call_command
        self._usage_rows()
        out = StringIO()
        call_command('ai_usage_report', stdout=out)
        text = out.getvalue()
        self.assertIn(self.sprint.institution.name, text)
        self.assertIn('cost=$0.0290', text)

    def test_platform_usage_report_is_admin_only(self):
        self._usage_rows()
        self._auth(self.iqac)
        self.assertEqual(self.client.get('/api/v1/extraction-jobs/usage-report').status_code, status.HTTP_403_FORBIDDEN)

        admin_user = self._make_user('platform-admin@test.edu', User.Role.SUPER_ADMIN, None)
        self._auth(admin_user)
        response = self.client.get('/api/v1/extraction-jobs/usage-report')
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        self.assertEqual(response.data['calls'], 3)
        self.assertEqual(response.data['cost_usd'], '0.029000')
        row = response.data['rows'][0]
        self.assertEqual(row['institution'], self.sprint.institution.name)
        self.assertEqual((row['calls'], row['input_tokens'], row['output_tokens']), (3, 6100, 1750))
        self.assertEqual(row['unpriced_calls'], 1)

        # Date filter: nothing before today.
        empty = self.client.get('/api/v1/extraction-jobs/usage-report?until=2000-01-01')
        self.assertEqual((empty.data['calls'], empty.data['rows']), (0, []))
        bad = self.client.get('/api/v1/extraction-jobs/usage-report?since=yesterday')
        self.assertEqual(bad.status_code, status.HTTP_400_BAD_REQUEST)

    def test_import_from_dashboard_is_labelled_counted_and_not_doubled(self):
        from django.core.management import call_command
        from django.core.management.base import CommandError
        from .models import AIUsage
        args = ['import_ai_usage', '--sprint', str(self.sprint.id), '--requests', '711',
                '--input-tokens', '1369687', '--output-tokens', '227699', '--cost', '3.11',
                '--model', 'my-combo', '--date', '2026-10-07', '--note', '9Router 7-day total']
        call_command(*args, stdout=__import__('io').StringIO())
        row = AIUsage.objects.get(sprint=self.sprint)
        self.assertEqual((row.source, row.request_count, str(row.cost_usd)), ('imported', 711, '3.110000'))
        self.assertEqual(row.created_at.date().isoformat(), '2026-10-07')
        with self.assertRaises(CommandError):
            call_command(*args, stdout=__import__('io').StringIO())

        self._auth(self.iqac)
        usage = self.client.get(f'/api/v1/sprints/{self.sprint.id}/ai-usage').data
        self.assertEqual((usage['calls'], usage['imported_calls'], usage['cost_usd']), (711, 711, '3.110000'))
        self.assertEqual(usage['notes'], ['9Router 7-day total'])


# -- several AI routers (AI_ROUTER_URLS) --------------------------------------

class FakeRedis:
    """Just enough of redis-py for RouterPool: mget / set(nx, ex) / expire /
    eval (the release script)."""

    def __init__(self):
        self.data = {}

    def mget(self, keys):
        return [self.data.get(k) for k in keys]

    def set(self, key, value, nx=False, ex=None):
        if nx and key in self.data:
            return None
        self.data[key] = value
        return True

    def expire(self, key, seconds):
        return key in self.data

    def eval(self, script, numkeys, key, token):
        if self.data.get(key) == token:
            del self.data[key]
            return 1
        return 0


FOUR_ROUTERS = ','.join(f'http://localhost:{port}' for port in (20228, 20229, 20230, 20231))


@override_settings(AI_PROCESSING_MODE='multi_router', AI_ROUTER_MODEL='',
                   AI_ROUTER_URLS=FOUR_ROUTERS, AI_ROUTER_API_KEYS='k1,k2,k3,k4', AI_ROUTER_MAX_CONCURRENT=2,
                   AI_ROUTER_WAIT_SECONDS=0, AI_MODEL='my-combo')
class RouterPoolTests(SimpleTestCase):
    def _pool(self):
        from .services.router_pool import RouterPool, configured_routers
        return RouterPool(configured_routers(), redis_client=FakeRedis())

    def test_routers_are_read_from_settings_with_v1_added(self):
        from .services.router_pool import configured_routers
        routers = configured_routers()
        self.assertEqual([r.base_url for r in routers][0], 'http://localhost:20228/v1')
        self.assertEqual([r.api_key for r in routers], ['k1', 'k2', 'k3', 'k4'])
        self.assertEqual(routers[1].label, 'router2 (localhost:20229)')

    @override_settings(AI_ROUTER_API_KEYS='shared')
    def test_one_key_is_shared_by_every_router(self):
        from .services.router_pool import configured_routers
        self.assertEqual({r.api_key for r in configured_routers()}, {'shared'})

    @override_settings(AI_ROUTER_API_KEYS='k1,k2')
    def test_mismatched_key_count_is_rejected(self):
        from .services.router_pool import configured_routers
        with self.assertRaises(ImproperlyConfigured):
            configured_routers()

    @override_settings(AI_ROUTER_URLS='')
    def test_multi_router_mode_without_routers_is_a_config_error(self):
        from .services.router_pool import get_router_pool
        with self.assertRaisesMessage(ImproperlyConfigured, 'AI_ROUTER_URLS is empty'):
            get_router_pool()

    @override_settings(AI_PROCESSING_MODE='single')
    def test_single_mode_ignores_configured_routers(self):
        from .services.router_pool import get_router_pool
        self.assertIsNone(get_router_pool())

    def test_eight_jobs_spread_two_per_router_and_the_ninth_waits(self):
        pool = self._pool()
        leases = [pool.acquire() for _ in range(8)]
        counts = {}
        for lease in leases:
            counts[lease.router.index] = counts.get(lease.router.index, 0) + 1
        self.assertEqual(counts, {1: 2, 2: 2, 3: 2, 4: 2})

        with self.assertRaises(RecoverableExtractionError):
            pool.acquire()
        pool.release(leases[5])
        self.assertEqual(pool.acquire().router, leases[5].router)

    def test_retry_avoids_the_router_that_just_failed(self):
        pool = self._pool()
        failed = pool.acquire()
        pool.release(failed)
        pool.mark_unhealthy(failed.router)
        for _ in range(5):
            lease = pool.acquire(avoid=failed.router.label)
            self.assertNotEqual(lease.router, failed.router)
            pool.release(lease)

    def test_failed_router_is_still_used_when_nothing_else_is_free(self):
        pool = self._pool()
        held = [pool.acquire() for _ in range(7)]
        lease = pool.acquire(avoid='whatever')
        self.assertNotIn(lease.slot_key, {h.slot_key for h in held})

    def test_redis_down_still_returns_a_router(self):
        from .services.router_pool import RouterPool, configured_routers
        broken = MagicMock()
        broken.mget.side_effect = ConnectionError('redis down')
        lease = RouterPool(configured_routers(), redis_client=broken).acquire()
        self.assertIn(lease.router, configured_routers())

    def test_ai_calls_inside_a_lease_go_to_that_router(self):
        from .services.router_pool import use_lease
        lease = self._pool().acquire()
        with use_lease(lease):
            service = get_ai_service()
        self.assertEqual(service.base_url, lease.router.base_url)
        self.assertEqual(service.api_key, lease.router.api_key)
        self.assertEqual(service.model, 'my-combo')


@override_settings(CELERY_TASK_ALWAYS_EAGER=True, CELERY_TASK_EAGER_PROPAGATES=True,
                   AI_PROCESSING_MODE='multi_router', AI_ROUTER_MODEL='',
                   AI_ROUTER_URLS=FOUR_ROUTERS, AI_ROUTER_API_KEYS='k1,k2,k3,k4', AI_ROUTER_WAIT_SECONDS=0)
class RouterTaskTests(ExtractionJobTestBase):
    def setUp(self):
        super().setUp()
        self.redis = FakeRedis()
        patcher = patch('apps.extraction.services.router_pool.RouterPool._client', return_value=self.redis)
        patcher.start()
        self.addCleanup(patcher.stop)

    def _job(self, **extra):
        return ExtractionJob.objects.create(sprint=self.sprint, document=self.document, **extra)

    @patch('apps.extraction.tasks.ExtractionPipeline')
    def test_job_records_its_router_and_frees_the_slot(self, mock_pipeline_cls):
        job = self._job()
        run_extraction_job.apply(args=[str(job.id)])
        job.refresh_from_db()
        self.assertEqual(job.status, ExtractionJob.Status.COMPLETED)
        self.assertRegex(job.ai_router, r'^router[1-4] \(localhost:202(28|29|30|31)\)$')
        self.assertEqual([v for v in self.redis.data if ':slot:' in v], [])

    @patch('apps.extraction.tasks.ExtractionPipeline')
    def test_recoverable_failure_retries_on_another_router(self, mock_pipeline_cls):
        mock_pipeline_cls.return_value.run.side_effect = RecoverableExtractionError('rate limited')
        from celery.exceptions import Retry as CeleryRetry
        job = self._job()
        with self.assertRaises(CeleryRetry):  # eager mode re-raises the scheduled retry
            run_extraction_job.apply(args=[str(job.id)])
        job.refresh_from_db()
        first = job.ai_router
        self.assertEqual(job.status, ExtractionJob.Status.RETRYING)

        mock_pipeline_cls.return_value.run.side_effect = None
        run_extraction_job.apply(args=[str(job.id)])
        job.refresh_from_db()
        self.assertEqual(job.status, ExtractionJob.Status.COMPLETED)
        self.assertNotEqual(job.ai_router, first)

    @patch('apps.extraction.tasks.ExtractionPipeline')
    def test_duplicate_delivery_of_a_finished_job_is_skipped(self, mock_pipeline_cls):
        for status_value in (ExtractionJob.Status.COMPLETED, ExtractionJob.Status.FAILED):
            job = self._job(status=status_value)
            run_extraction_job.apply(args=[str(job.id)])
            job.refresh_from_db()
            self.assertEqual(job.status, status_value)
        mock_pipeline_cls.return_value.run.assert_not_called()

    @override_settings(AI_PROCESSING_MODE='single')  # routers still listed: single mode ignores them
    @patch('apps.extraction.tasks.ExtractionPipeline')
    def test_single_router_setup_is_unchanged(self, mock_pipeline_cls):
        job = self._job()
        run_extraction_job.apply(args=[str(job.id)])
        job.refresh_from_db()
        self.assertEqual(job.status, ExtractionJob.Status.COMPLETED)
        self.assertEqual(job.ai_router, '')

    @patch('apps.extraction.tasks.ExtractionPipeline')
    def test_stale_message_for_a_reprocessed_job_is_skipped(self, mock_pipeline_cls):
        job = self._job(celery_task_id='new-task-id')
        run_extraction_job.apply(args=[str(job.id)], task_id='old-task-id')
        job.refresh_from_db()
        self.assertEqual(job.status, ExtractionJob.Status.PENDING)
        mock_pipeline_cls.return_value.run.assert_not_called()

        run_extraction_job.apply(args=[str(job.id)], task_id='new-task-id')
        job.refresh_from_db()
        self.assertEqual(job.status, ExtractionJob.Status.COMPLETED)


class ResponseFormatFallbackTests(SimpleTestCase):
    """A gateway that rejects `response_format` gets the request again without
    it (the schema is in the prompt), and isn't sent it again afterwards."""

    def setUp(self):
        from .services import openai_client
        openai_client._NO_RESPONSE_FORMAT.clear()
        self.addCleanup(openai_client._NO_RESPONSE_FORMAT.clear)

    def _rejection(self):
        request = httpx.Request('POST', 'http://gateway.test/v1/chat/completions')
        response = httpx.Response(400, request=request)
        return APIStatusError(
            "The capability 'response_format' is not supported by this gateway profile.",
            response=response, body=None,
        )

    def _ok(self):
        message = MagicMock(content='{"answer": "yes"}')
        return MagicMock(choices=[MagicMock(finish_reason='stop', message=message)], usage=None)

    def test_resends_without_response_format_and_remembers(self):
        client = MagicMock()
        client.chat.completions.create.side_effect = [self._rejection(), self._ok(), self._ok()]
        service = OpenAIExtractionService(api_key='k', model='m', base_url='http://gateway.test/v1', client=client)
        schema = {'type': 'object', 'properties': {'answer': {'type': 'string'}}}

        self.assertEqual(service.extract_structured_data(system_prompt='s', user_content='u', response_schema=schema),
                         {'answer': 'yes'})
        service.extract_structured_data(system_prompt='s', user_content='u', response_schema=schema)

        calls = client.chat.completions.create.call_args_list
        self.assertIn('response_format', calls[0].kwargs)
        self.assertNotIn('response_format', calls[1].kwargs)
        self.assertNotIn('response_format', calls[2].kwargs)

    def test_other_400s_still_fail_permanently(self):
        client = MagicMock()
        request = httpx.Request('POST', 'http://gateway.test/v1/chat/completions')
        client.chat.completions.create.side_effect = APIStatusError(
            'bad model', response=httpx.Response(400, request=request), body=None,
        )
        service = OpenAIExtractionService(api_key='k', model='m', base_url='http://gateway.test/v1', client=client)
        with self.assertRaises(PermanentExtractionError):
            service.extract_structured_data(system_prompt='s', user_content='u', response_schema={'type': 'object'})


# -- AI_PROCESSING_MODE: single | multi_router -------------------------------

class ProcessingModeConfigTests(SimpleTestCase):
    def test_default_mode_is_single(self):
        from config.ai_processing_mode import parse_mode
        self.assertEqual(parse_mode(None), 'single')
        self.assertEqual(parse_mode(''), 'single')
        self.assertEqual(parse_mode(' Multi_Router '), 'multi_router')

    def test_invalid_mode_is_a_clear_config_error(self):
        from config.ai_processing_mode import parse_mode
        with self.assertRaisesMessage(ImproperlyConfigured, "AI_PROCESSING_MODE='turbo' is not valid"):
            parse_mode('turbo')

    @override_settings(AI_PROCESSING_MODE='turbo')
    def test_invalid_mode_is_reported_by_the_system_check(self):
        from .checks import check_ai_processing_mode
        self.assertEqual([e.id for e in check_ai_processing_mode(None)], ['extraction.E001'])

    def test_worker_slots_per_mode(self):
        from config.ai_processing_mode import default_worker_concurrency
        self.assertEqual(default_worker_concurrency('single', routers=4, per_router=2), 2)
        self.assertEqual(default_worker_concurrency('multi_router', routers=4, per_router=2), 8)
        self.assertEqual(default_worker_concurrency('multi_router', routers=3, per_router=2), 6)


@override_settings(AI_PROCESSING_MODE='single', AI_API_KEY='single-key', AI_MODEL='single-model',
                   AI_BASE_URL='http://gateway.test/v1', AI_ROUTER_URLS=FOUR_ROUTERS, AI_ROUTER_API_KEYS='k1,k2,k3,k4')
class SingleModeTests(SimpleTestCase):
    def test_uses_the_single_connection_even_with_routers_listed(self):
        from .services.router_pool import get_router_pool
        self.assertIsNone(get_router_pool())
        service = get_ai_service()
        self.assertEqual((service.base_url, service.api_key, service.model),
                         ('http://gateway.test/v1', 'single-key', 'single-model'))

    def test_status_describes_single_mode_without_secrets(self):
        from .services.processing_status import describe, summary_line
        info = describe()
        self.assertEqual((info['mode'], info['endpoint'], info['key']), ('single', 'gateway.test', 'set'))
        self.assertNotIn('single-key', summary_line(info))


@override_settings(AI_PROCESSING_MODE='multi_router', AI_MODEL='single-model', AI_ROUTER_MODEL='my-combo',
                   AI_BASE_URL='http://gateway.test/v1', AI_ROUTER_URLS=FOUR_ROUTERS,
                   AI_ROUTER_API_KEYS='k1,k2,k3,k4', AI_ROUTER_MAX_CONCURRENT=2, AI_ROUTER_WAIT_SECONDS=0,
                   CELERY_WORKER_CONCURRENCY=8)
class MultiRouterModeTests(SimpleTestCase):
    def _pool(self):
        from .services.router_pool import RouterPool, configured_routers
        return RouterPool(configured_routers(), redis_client=FakeRedis())

    def test_four_routers_with_two_worker_slots_each(self):
        from .services.processing_status import configuration_problems, describe
        from .services.router_pool import get_router_pool
        pool = get_router_pool()
        self.assertEqual([r.label for r in pool.routers], [
            'router1 (localhost:20228)', 'router2 (localhost:20229)',
            'router3 (localhost:20230)', 'router4 (localhost:20231)'])
        self.assertEqual(pool.max_concurrent, 2)
        info = describe()
        self.assertEqual((info['mode'], len(info['routers']), info['slots_per_router']), ('multi_router', 4, 2))
        self.assertEqual(configuration_problems(), ([], []))

    def test_requests_go_to_the_leased_router_with_the_router_model(self):
        from .services.router_pool import use_lease
        pool = self._pool()
        for _ in range(4):
            lease = pool.acquire()
            with use_lease(lease):
                service = get_ai_service()
            self.assertEqual((service.base_url, service.api_key, service.model),
                             (lease.router.base_url, lease.router.api_key, 'my-combo'))
            self.assertNotEqual(service.base_url, 'http://gateway.test/v1')

    def test_ai_call_outside_a_router_lease_is_refused(self):
        with self.assertRaisesMessage(ImproperlyConfigured, 'Refusing to fall back to AI_BASE_URL'):
            get_ai_service()

    def test_status_never_contains_router_keys(self):
        from .services.processing_status import describe, summary_line
        text = str(describe()) + summary_line()
        for key in ('k1', 'k2', 'k3', 'k4'):
            self.assertNotIn(f"'{key}'", text)

    @override_settings(CELERY_WORKER_CONCURRENCY=2)
    def test_too_few_workers_for_the_router_slots_is_a_warning(self):
        from .services.processing_status import configuration_problems
        errors, warnings = configuration_problems()
        self.assertEqual(errors, [])
        self.assertIn('some slots will stay unused', warnings[0])

    @override_settings(AI_ROUTER_API_KEYS='k1,k2')
    def test_wrong_number_of_keys_is_a_config_error(self):
        from .checks import check_ai_processing_mode
        self.assertEqual([e.id for e in check_ai_processing_mode(None)], ['extraction.E001'])

    def test_token_usage_and_duration_are_recorded_for_a_router_call(self):
        """Usage recording is mode-independent: a call on a router writes the
        same AIUsage row (tokens, duration, model) as a single-mode call."""
        from types import SimpleNamespace

        from .services.router_pool import use_lease
        from .usage import usage_scope

        usage = SimpleNamespace(prompt_tokens=1200, completion_tokens=300,
                                prompt_tokens_details=None, completion_tokens_details=None)
        message = SimpleNamespace(content='{"answer": "yes"}')
        response = SimpleNamespace(choices=[SimpleNamespace(finish_reason='stop', message=message)], usage=usage)
        lease = self._pool().acquire()
        with use_lease(lease):
            service = get_ai_service()
        service._client = MagicMock()
        service._client.chat.completions.create.return_value = response
        with patch('apps.extraction.models.AIUsage.objects.create') as create, usage_scope(stage='fact_extraction'):
            service.extract_structured_data(system_prompt='s', user_content='u',
                                            response_schema={'type': 'object'})
        row = create.call_args.kwargs
        self.assertEqual((row['input_tokens'], row['output_tokens'], row['model'], row['provider']),
                         (1200, 300, 'my-combo', 'openai_compatible'))
        self.assertGreaterEqual(row['duration_ms'], 0)


@override_settings(CELERY_TASK_ALWAYS_EAGER=True, CELERY_TASK_EAGER_PROPAGATES=True,
                   AI_PROCESSING_MODE='multi_router', AI_ROUTER_URLS='')
class MultiRouterMisconfiguredTaskTests(ExtractionJobTestBase):
    @patch('apps.extraction.tasks.ExtractionPipeline')
    def test_job_fails_clearly_instead_of_sticking_at_running(self, mock_pipeline_cls):
        job = ExtractionJob.objects.create(sprint=self.sprint, document=self.document)
        run_extraction_job.apply(args=[str(job.id)])
        job.refresh_from_db()
        self.assertEqual(job.status, ExtractionJob.Status.FAILED)
        self.assertIn('AI_ROUTER_URLS is empty', job.error_message)
        mock_pipeline_cls.return_value.run.assert_not_called()


class ModeSwitchAfterRestartTests(SimpleTestCase):
    """A fresh process = a restarted worker: the mode and worker slots come
    from the environment it starts with, and a bad value stops the start."""

    def _start(self, mode, extra=None):
        import os
        import subprocess
        import sys
        env = {**os.environ, 'AI_PROCESSING_MODE': mode, 'AI_ROUTER_URLS': FOUR_ROUTERS,
               'AI_ROUTER_MAX_CONCURRENT': '2', 'CELERY_WORKER_CONCURRENCY': '',
               'DJANGO_SETTINGS_MODULE': 'config.settings', **(extra or {})}
        code = ('import django; django.setup(); from django.conf import settings; '
                'print(settings.AI_PROCESSING_MODE, settings.CELERY_WORKER_CONCURRENCY)')
        return subprocess.run([sys.executable, '-c', code], cwd=settings.BASE_DIR, env=env,
                              capture_output=True, text=True, timeout=120)

    def test_switching_modes_takes_effect_on_restart(self):
        self.assertEqual(self._start('single').stdout.split()[-2:], ['single', '2'])
        self.assertEqual(self._start('multi_router').stdout.split()[-2:], ['multi_router', '8'])
        self.assertEqual(self._start('multi_router', {'CELERY_WORKER_CONCURRENCY': '6'}).stdout.split()[-2:],
                         ['multi_router', '6'])

    def test_invalid_mode_stops_startup(self):
        result = self._start('turbo')
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("AI_PROCESSING_MODE='turbo' is not valid", result.stderr)


class OfficeAndImageReaderTests(TestCase):
    """Word / Excel / CSV / image uploads are read into the same page shape as a PDF."""

    def setUp(self):
        institution = Institution.objects.create(name='MKCE', city='Karur', state='TN')
        self.sprint = Sprint.objects.create(institution=institution, mode=Sprint.SprintMode.VERIFIED_CRI)

    def _doc(self, content, filename):
        return Document.objects.create(
            sprint=self.sprint, document_type='faculty_master_list', file=SimpleUploadedFile(filename, content),
            original_filename=filename, status=Document.Status.UPLOADED)

    def test_docx_paragraphs_and_tables(self):
        import io
        from docx import Document as Docx
        d = Docx()
        d.add_paragraph('Total faculty strength is 312 as on 31 March 2026.')
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text = 'Department', 'Faculty'
        t.cell(1, 0).text, t.cell(1, 1).text = 'CSE', '64'
        buf = io.BytesIO()
        d.save(buf)
        result = PDFPageReader().read_pages(self._doc(buf.getvalue(), 'faculty.docx'))
        self.assertTrue(result['format_supported'])
        text = result['pages'][0]['text']
        self.assertIn('Total faculty strength is 312', text)
        self.assertIn('CSE | 64', text)
        self.assertEqual(result['pages'][0]['page_number'], 1)

    def test_xlsx_one_page_per_sheet(self):
        import io
        import openpyxl
        wb = openpyxl.Workbook()
        wb.active.title = 'Faculty'
        wb.active.append(['Name', 'Designation'])
        wb.active.append(['R. Priya', 'Assistant Professor'])
        wb.create_sheet('Labs').append(['AI Lab', 40])
        buf = io.BytesIO()
        wb.save(buf)
        result = PDFPageReader().read_pages(self._doc(buf.getvalue(), 'records.xlsx'))
        self.assertEqual(result['page_count'], 2)
        self.assertIn('Sheet: Faculty', result['pages'][0]['text'])
        self.assertIn('R. Priya | Assistant Professor', result['pages'][0]['text'])
        self.assertIn('AI Lab | 40', result['pages'][1]['text'])

    def test_csv(self):
        result = PDFPageReader().read_pages(self._doc(b'year,students\n2025-26,4120\n', 'strength.csv'))
        self.assertIn('2025-26 | 4120', result['pages'][0]['text'])

    def test_image_is_read_by_ocr(self):
        import io
        from PIL import Image
        from apps.extraction.services.base import OCRProvider

        class FakeOCR(OCRProvider):
            def extract_text(self, document, page_number, image=None):
                self.shape = image.shape
                return 'Accreditation certificate valid till 2029'

        buf = io.BytesIO()
        Image.new('RGB', (40, 20), 'white').save(buf, format='PNG')
        ocr = FakeOCR()
        result = PDFPageReader(ocr_provider=ocr).read_pages(self._doc(buf.getvalue(), 'certificate.png'))
        self.assertEqual(ocr.shape, (20, 40, 3))
        self.assertEqual(result['pages'][0]['text'], 'Accreditation certificate valid till 2029')
        self.assertEqual(result['pages'][0]['text_source'], 'ocr')

    def test_image_without_ocr_says_so(self):
        import io
        from PIL import Image
        from apps.extraction.services import stub
        buf = io.BytesIO()
        Image.new('RGB', (10, 10)).save(buf, format='JPEG')
        result = PDFPageReader(ocr_provider=stub.NullOCRProvider()).read_pages(self._doc(buf.getvalue(), 'scan.jpg'))
        self.assertIn('OCR is not available', result['ocr_warnings'][0])

    def test_unreadable_type_is_shown_on_the_document(self):
        from apps.extraction.models import ExtractionJob
        from apps.extraction.services import stub
        from apps.extraction.services.pipeline import ExtractionPipeline
        doc = self._doc(b'PK zip bytes', 'bundle.zip')
        job = ExtractionJob.objects.create(sprint=self.sprint, document=doc)
        pipe = ExtractionPipeline(classifier=stub.KnownMetadataClassifier(), fact_extractor=stub.NullFactExtractor(),
                                  conflict_checker=stub.NullConflictChecker())
        pipe.run(job)
        doc.refresh_from_db()
        self.assertEqual((doc.page_count, doc.processing_status), (0, 'unsupported_format'))
        self.assertIn('ZIP archives are not unpacked', doc.ocr_warnings[0])


@override_settings(CELERY_TASK_ALWAYS_EAGER=True, CELERY_TASK_EAGER_PROPAGATES=True)
class DocumentClassificationTests(ExtractionJobTestBase):
    """The AI classifier's verdict is saved on the document (separately from
    the assigned document_type), and reprocessing refreshes it without
    duplicating the document or its facts."""

    def setUp(self):
        super().setUp()
        patcher = patch('apps.extraction.services.openai_classifier.get_ai_service')
        self.classify_mock = patcher.start().return_value.extract_structured_data
        self.classify_mock.return_value = dict(FAKE_CLASSIFICATION_RESULT)
        self.addCleanup(patcher.stop)
        self._auth(self.iqac)

    def _run(self):
        """(Re)process self.document -- the same request the Upload Data Pack
        page sends for a selected document, whatever its status."""
        response = self.client.post(
            f'/api/v1/sprints/{self.sprint.id}/extraction-jobs/', {'document_id': str(self.document.id)},
        )
        self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.data)
        self.document.refresh_from_db()
        return ExtractionJob.objects.get(id=response.data[0]['id'])

    def _classify_as(self, document_type, confidence, reasoning='Grounded in the text.'):
        self.classify_mock.return_value = {
            **FAKE_CLASSIFICATION_RESULT, 'document_type': document_type, 'confidence': confidence,
            'reasoning': reasoning,
        }

    def test_confident_checklist_category_is_saved_and_the_assigned_type_kept(self):
        self._classify_as('aqar_report', 0.93, 'Titled Annual Quality Assurance Report 2024-25.')
        self._run()

        doc = self.document
        self.assertEqual(doc.classification_status, Document.ClassificationStatus.CLASSIFIED)
        self.assertEqual(doc.ai_document_type, 'aqar_report')
        self.assertEqual(doc.ai_confidence, 0.93)
        self.assertEqual(doc.ai_reasoning, 'Titled Annual Quality Assurance Report 2024-25.')
        self.assertIsNotNone(doc.classified_at)
        # The uploader chose naac_ssr; that original label is untouched.
        self.assertEqual((doc.document_type, doc.document_type_source), ('naac_ssr', 'uploader'))

        row = self.client.get(f'/api/v1/sprints/{self.sprint.id}/documents/').data[0]
        self.assertEqual(row['classification_status'], 'classified')
        self.assertEqual(row['ai_document_type_label'], 'AQAR / Annual Quality Assurance Report')
        self.assertEqual(row['document_type'], 'naac_ssr')

    @override_settings(AI_CLASSIFICATION_MIN_CONFIDENCE=0.7)
    def test_low_confidence_needs_review(self):
        self._classify_as('naac_ssr', 0.45)
        self._run()
        self.assertEqual(self.document.classification_status, Document.ClassificationStatus.NEEDS_REVIEW)
        self.assertEqual(self.document.ai_document_type, 'naac_ssr')

    def test_no_category_or_an_unknown_one_is_unclassified_not_a_default_type(self):
        for answer in (None, 'campus_brochure'):
            with self.subTest(answer=answer):
                self._classify_as(answer, 0.9, 'A marketing brochure, none of the categories.')
                self._run()
                self.assertEqual(self.document.classification_status, Document.ClassificationStatus.UNCLASSIFIED)
                self.assertEqual(self.document.ai_document_type, '')
                self.assertEqual(self.document.ai_reasoning, 'A marketing brochure, none of the categories.')

    def test_a_document_types_alias_maps_to_its_checklist_slot(self):
        self._classify_as('AQAR', 0.9)
        self._run()
        self.assertEqual(self.document.ai_document_type, 'aqar_report')

    def test_reprocessing_updates_the_classification_and_facts_in_place(self):
        self._classify_as('naac_ssr', 0.5)
        self.fact_openai_mock.return_value.extract_structured_data.return_value = {
            'facts': [_valid_fact(value='42', confidence_score=0.3)],
        }
        self._run()
        first_fact = self.document.extracted_facts.get()
        gap = GapItem.objects.get(gap_type=GapItem.GapType.LOW_CONFIDENCE, source_fact=first_fact)

        self._classify_as('naac_ssr', 0.95)
        self.fact_openai_mock.return_value.extract_structured_data.return_value = {
            'facts': [
                _valid_fact(value='45', confidence_score=0.3),
                _valid_fact(field_key='total_students', field_name='Total Students', value='900'),
            ],
        }
        self._run()

        self.assertEqual(self.document.classification_status, Document.ClassificationStatus.CLASSIFIED)
        self.assertEqual(self.document.ai_confidence, 0.95)
        self.assertEqual(Document.objects.filter(sprint=self.sprint).count(), 1)
        facts = {f.field_key: f for f in self.document.extracted_facts.all()}
        self.assertEqual(sorted(facts), ['total_faculty', 'total_students'])
        # Same row, new value -- so the gap raised against it still points at it, once.
        self.assertEqual(facts['total_faculty'].id, first_fact.id)
        self.assertEqual(facts['total_faculty'].value, 45)
        self.assertEqual(
            GapItem.objects.filter(gap_type=GapItem.GapType.LOW_CONFIDENCE, source_fact=first_fact).count(), 1,
        )
        gap.refresh_from_db()
        self.assertEqual(gap.source_fact_id, first_fact.id)

    def test_reprocessing_leaves_a_reviewed_fact_as_the_reviewer_left_it(self):
        self.fact_openai_mock.return_value.extract_structured_data.return_value = {'facts': [_valid_fact(value='42')]}
        self._run()
        fact = self.document.extracted_facts.get()
        fact.status = ExtractedFact.Status.CONFIRMED
        fact.save()

        self.fact_openai_mock.return_value.extract_structured_data.return_value = {'facts': [_valid_fact(value='99')]}
        self._run()

        fact.refresh_from_db()
        self.assertEqual(self.document.extracted_facts.count(), 1)
        self.assertEqual((fact.status, fact.value), (ExtractedFact.Status.CONFIRMED, 42))

    def test_a_rerun_whose_classification_fails_keeps_the_previous_result(self):
        self._classify_as('placement_report', 0.88)
        self._run()

        self.classify_mock.return_value = {**FAKE_CLASSIFICATION_RESULT, 'confidence': 7}  # invalid -> permanent failure
        job = self._run()

        self.assertEqual(job.status, ExtractionJob.Status.FAILED)
        self.assertEqual(self.document.classification_status, Document.ClassificationStatus.CLASSIFIED)
        self.assertEqual((self.document.ai_document_type, self.document.ai_confidence), ('placement_report', 0.88))

    def _confirm_category(self, category):
        response = self.client.post(
            f'/api/v1/documents/{self.document.id}/confirm-category', {'category': category}, format='json',
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        self.document.refresh_from_db()

    def test_reprocessing_keeps_a_category_a_person_confirmed(self):
        self._classify_as('naac_ssr', 0.45, 'Unsure.')
        self._run()
        self.assertEqual(self.document.classification_status, Document.ClassificationStatus.NEEDS_REVIEW)
        self._confirm_category('aqar_report')
        confirmed_at = self.document.category_confirmed_at

        self._classify_as('naac_ssr', 0.97, 'Now sure it is the SSR.')
        job = self._run()

        self.assertEqual(job.status, ExtractionJob.Status.COMPLETED)
        self.assertEqual(self.document.classification_status, Document.ClassificationStatus.CLASSIFIED)
        self.assertEqual(self.document.ai_document_type, 'aqar_report')
        self.assertEqual((self.document.category_confirmed_by, self.document.category_confirmed_at),
                         (self.iqac, confirmed_at))
        # The new verdict is dropped whole, like a reviewed fact's new value.
        self.assertEqual((self.document.ai_confidence, self.document.ai_reasoning), (0.45, 'Unsure.'))
        self.assertEqual(self.document.status, Document.Status.PROCESSED)  # the rest of the run still happened

    def test_reprocessing_keeps_a_not_a_checklist_document_decision(self):
        self._classify_as(None, 0.9)
        self._run()
        self._confirm_category(None)

        self._classify_as('placement_report', 0.95)
        self._run()

        self.assertEqual(self.document.classification_status, Document.ClassificationStatus.UNCLASSIFIED)
        self.assertEqual(self.document.ai_document_type, '')

    def test_a_confirmation_made_while_the_job_runs_still_wins(self):
        from .services.pipeline import ExtractionPipeline

        stale = Document.objects.get(id=self.document.id)  # what the pipeline loaded at the start
        Document.objects.filter(id=self.document.id).update(
            classification_status=Document.ClassificationStatus.CLASSIFIED, ai_document_type='lab_inventory',
            category_confirmed_by=self.iqac, category_confirmed_at=timezone.now(),
        )

        changed = ExtractionPipeline._apply_classification(
            stale, {**FAKE_CLASSIFICATION_RESULT, 'document_type': 'naac_ssr', 'confidence': 0.99},
        )

        self.assertEqual(changed, [])
        self.assertEqual((stale.classification_status, stale.ai_document_type), ('classified', 'lab_inventory'))

    def test_ai_fields_cannot_be_written_through_the_document_api(self):
        response = self.client.patch(
            f'/api/v1/documents/{self.document.id}/',
            {'ai_document_type': 'aqar_report', 'classification_status': 'classified', 'ai_confidence': 1},
            format='json',
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        self.document.refresh_from_db()
        self.assertEqual(
            (self.document.ai_document_type, self.document.classification_status, self.document.ai_confidence),
            ('', Document.ClassificationStatus.PENDING, None),
        )


class ClassifierPromptTests(TestCase):
    def setUp(self):
        institution = Institution.objects.create(name='MKCE', city='Karur', state='TN')
        self.sprint = Sprint.objects.create(institution=institution, mode=Sprint.SprintMode.VERIFIED_CRI)

    def _call_kwargs(self, **document_fields):
        document = Document.objects.create(
            sprint=self.sprint, original_filename='SSR_2025.pdf', mime_type='application/pdf',
            file=SimpleUploadedFile('SSR_2025.pdf', make_pdf_bytes(['Self Study Report']), 'application/pdf'),
            **document_fields,
        )
        service = MagicMock()
        service.extract_structured_data.return_value = dict(FAKE_CLASSIFICATION_RESULT)
        OpenAIDocumentClassifier(openai_service=service).classify(document)
        return service.extract_structured_data.call_args.kwargs

    def test_the_prompt_lists_every_checklist_category(self):
        from apps.documents.constants import CHECKLIST_TYPE_LABELS

        system_prompt = self._call_kwargs(document_type='naac_ssr')['system_prompt']
        for slug in CHECKLIST_TYPE_LABELS:
            self.assertIn(f'- {slug}:', system_prompt)

    def test_the_assigned_type_is_sent_as_a_hint_with_where_it_came_from(self):
        cases = [
            ({'document_type': 'naac_ssr', 'document_type_source': 'filename'},
             'Assigned document type: naac_ssr (guessed from filename keywords only)'),
            ({'document_type': 'naac_ssr', 'document_type_source': 'uploader'},
             'Assigned document type: naac_ssr (chosen by the uploader)'),
            ({'document_type': 'unclassified', 'document_type_source': 'none'},
             'Assigned document type: (none)'),
        ]
        for fields, expected in cases:
            with self.subTest(**fields):
                self.assertIn(expected, self._call_kwargs(**fields)['user_content'])


def _indicator(code, confidence=0.9, reason='The evidence states it directly.'):
    return {'indicator': code, 'reason': reason, 'confidence': confidence}


class FactAIQIFTaggingTests(TestCase):
    """OpenAIFactExtractor tags each fact with AI-QIF key indicators, from the
    approved 36 only; an unclear or invalid tag is dropped, never guessed."""

    def setUp(self):
        institution = Institution.objects.create(name='MKCE', city='Karur', state='TN')
        self.sprint = Sprint.objects.create(institution=institution, mode=Sprint.SprintMode.VERIFIED_CRI)
        self.document = Document.objects.create(
            sprint=self.sprint, document_type='naac_ssr', original_filename='ssr.pdf',
            mime_type='application/pdf', status=Document.Status.UPLOADED,
        )

    def _extract(self, **fact_overrides):
        service = MagicMock()
        service.extract_structured_data.return_value = {'facts': [_valid_fact(**fact_overrides)]}
        pages = {'pages': [_page(1, 'The institution has 42 faculty members on its rolls.')]}
        facts = OpenAIFactExtractor(openai_service=service).extract_facts(self.document, pages)
        return facts, service.extract_structured_data.call_args.kwargs

    def test_indicators_are_kept_with_their_criterion_reason_and_confidence(self):
        facts, _ = self._extract(aiqif_indicators=[
            _indicator('AIQIF-C4.1', 0.6, 'Counts faculty, says nothing of AI literacy.'),
            _indicator('C4.2', 0.92, 'Lists role-wise AI training.'),  # short code accepted
            _indicator('AIQIF-C4.2', 0.5),  # same indicator again: the stronger one is kept
        ])
        self.assertEqual(facts[0]['aiqif_mappings'], [
            {'indicator_code': 'AIQIF-C4.2', 'criterion': 'aiqif_c4', 'reason': 'Lists role-wise AI training.',
             'confidence': 0.92},
            {'indicator_code': 'AIQIF-C4.1', 'criterion': 'aiqif_c4',
             'reason': 'Counts faculty, says nothing of AI literacy.', 'confidence': 0.6},
        ])

    def test_a_fact_can_support_indicators_of_more_than_one_criterion(self):
        facts, _ = self._extract(aiqif_indicators=[_indicator('AIQIF-C4.1'), _indicator('AIQIF-C9.2', 0.8)])
        self.assertEqual(
            [(m['criterion'], m['indicator_code']) for m in facts[0]['aiqif_mappings']],
            [('aiqif_c4', 'AIQIF-C4.1'), ('aiqif_c9', 'AIQIF-C9.2')],
        )

    def test_invalid_tags_are_dropped_but_the_fact_is_kept(self):
        facts, _ = self._extract(aiqif_indicators=[
            _indicator('AIQIF-C10.1'),             # not an AI-QIF indicator
            _indicator('Faculty AI Literacy'),     # a name, not a code
            _indicator('AIQIF-C4.1', 1.4),         # confidence out of range
            _indicator('AIQIF-C4.3', 0.9, '  '),   # no reason
            {'indicator': 'AIQIF-C4.4', 'reason': 'x'},  # no confidence
            'AIQIF-C4.1',
        ])
        self.assertEqual(len(facts), 1)
        self.assertEqual(facts[0]['aiqif_mappings'], [])

    def test_missing_or_empty_tags_leave_the_fact_unmapped(self):
        for raw in ({}, {'aiqif_indicators': []}, {'aiqif_indicators': None}, {'aiqif_indicators': 'AIQIF-C4.1'}):
            with self.subTest(raw=raw):
                facts, _ = self._extract(**raw)
                self.assertEqual((len(facts), facts[0]['aiqif_mappings']), (1, []))

    def test_at_most_three_tags_strongest_first(self):
        codes = ['AIQIF-C4.1', 'AIQIF-C4.2', 'AIQIF-C4.3', 'AIQIF-C4.4', 'AIQIF-C9.2']
        facts, _ = self._extract(aiqif_indicators=[_indicator(c, 0.5 + i / 10) for i, c in enumerate(codes)])
        self.assertEqual(
            [m['indicator_code'] for m in facts[0]['aiqif_mappings']], ['AIQIF-C9.2', 'AIQIF-C4.4', 'AIQIF-C4.3'],
        )

    def test_the_prompt_and_schema_offer_exactly_the_36_approved_indicators(self):
        from apps.scoring.constants import AIQIF_INDICATORS

        _, kwargs = self._extract()
        self.assertEqual(len(AIQIF_INDICATORS), 36)
        item_schema = kwargs['response_schema']['properties']['facts']['items']
        self.assertEqual(
            item_schema['properties']['aiqif_indicators']['items']['properties']['indicator']['enum'],
            list(AIQIF_INDICATORS),
        )
        for code, (_criterion, name) in AIQIF_INDICATORS.items():
            self.assertIn(f'- {code}: {name}', kwargs['system_prompt'])
        self.assertIn('not from the document', kwargs['system_prompt'])
        self.assertIn('Never state or imply a maturity level', kwargs['system_prompt'])


@override_settings(CELERY_TASK_ALWAYS_EAGER=True, CELERY_TASK_EAGER_PROPAGATES=True,
                   AIQIF_MAPPING_MIN_CONFIDENCE=0.7)
class FactAIQIFMappingPersistenceTests(ExtractionJobTestBase):
    """The pipeline saves each fact's AI-QIF tags, the facts API returns them,
    and reprocessing replaces an unreviewed fact's tags without duplicating."""

    def setUp(self):
        super().setUp()
        self._auth(self.iqac)

    def _run(self, *facts):
        self.fact_openai_mock.return_value.extract_structured_data.return_value = {'facts': list(facts)}
        response = self.client.post(
            f'/api/v1/sprints/{self.sprint.id}/extraction-jobs/', {'document_id': str(self.document.id)},
        )
        self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.data)

    def _api_facts(self):
        return {f['field_key']: f for f in self.client.get(f'/api/v1/sprints/{self.sprint.id}/facts/').data}

    def test_tags_are_saved_and_shown_with_the_facts_evidence(self):
        self._run(
            _valid_fact(aiqif_indicators=[_indicator('AIQIF-C4.1', 0.9, 'Faculty count.'),
                                          _indicator('AIQIF-C9.2', 0.4, 'Cited figure.')]),
            _valid_fact(field_key='ai_labs', field_name='AI Labs', aiqif_indicators=[_indicator('AIQIF-C6.4', 0.5)]),
            _valid_fact(field_key='campus_area', field_name='Campus Area', aiqif_indicators=[]),
        )
        facts = self._api_facts()

        faculty = facts['total_faculty']
        UNREVIEWED = {'review_status': 'unreviewed', 'reviewed_by': None, 'reviewed_by_name': '', 'reviewed_at': None,
                      'review_note': ''}  # no person has looked at the AI's tags yet
        self.assertEqual(faculty['aiqif_mapping_status'], 'mapped')
        self.assertTrue(all(m.pop('id') for m in faculty['aiqif_mappings']))  # each tag has its own id
        self.assertEqual(faculty['aiqif_mappings'], [
            {'criterion': 'aiqif_c4', 'criterion_code': 'AIQIF-C4',
             'criterion_name': 'Faculty & Institutional Workforce AI Capability',
             'indicator_code': 'AIQIF-C4.1', 'indicator_name': 'Faculty AI Literacy & Awareness',
             'reason': 'Faculty count.', 'confidence': 0.9, 'status': 'mapped', **UNREVIEWED},
            {'criterion': 'aiqif_c9', 'criterion_code': 'AIQIF-C9',
             'criterion_name': 'AI Assurance, Measurement & Continuous Improvement',
             'indicator_code': 'AIQIF-C9.2', 'indicator_name': 'Evidence, Traceability & Auditability',
             'reason': 'Cited figure.', 'confidence': 0.4, 'status': 'needs_review', **UNREVIEWED},
        ])
        # The evidence stays on the fact; the CRI pillar is untouched by the tags.
        self.assertEqual(
            (faculty['source_document_id'], faculty['source_page'], faculty['source_snippet'], faculty['pillar']),
            (self.document.id, '1', 'The institution has 42 faculty members.', 'aiqif_c4'),
        )
        self.assertEqual(facts['ai_labs']['aiqif_mapping_status'], 'needs_review')
        self.assertEqual((facts['campus_area']['aiqif_mapping_status'], facts['campus_area']['aiqif_mappings']),
                         ('unmapped', []))

    def test_reprocessing_replaces_an_unreviewed_facts_tags_and_keeps_a_reviewed_facts(self):
        self._run(
            _valid_fact(aiqif_indicators=[_indicator('AIQIF-C4.1')]),
            _valid_fact(field_key='ai_labs', field_name='AI Labs', aiqif_indicators=[_indicator('AIQIF-C6.4')]),
        )
        reviewed = self.document.extracted_facts.get(field_key='ai_labs')
        reviewed.status = ExtractedFact.Status.CONFIRMED
        reviewed.save()

        self._run(
            _valid_fact(aiqif_indicators=[_indicator('AIQIF-C4.2'), _indicator('AIQIF-C4.1', 0.8)]),
            _valid_fact(field_key='ai_labs', field_name='AI Labs', aiqif_indicators=[_indicator('AIQIF-C6.1')]),
        )

        facts = self._api_facts()
        self.assertEqual(len(facts), 2)
        self.assertEqual([m['indicator_code'] for m in facts['total_faculty']['aiqif_mappings']],
                         ['AIQIF-C4.2', 'AIQIF-C4.1'])
        self.assertEqual([m['indicator_code'] for m in facts['ai_labs']['aiqif_mappings']], ['AIQIF-C6.4'])

    def test_a_fact_saved_without_any_tagging_stays_not_yet_mapped(self):
        """'Unmapped' means the AI looked and found no clear indicator; a fact
        from an extractor that doesn't tag at all stays 'pending'."""
        from .services.pipeline import ExtractionPipeline

        self._run()
        job = ExtractionJob.objects.filter(document=self.document).first()
        ExtractionPipeline._persist_facts(job, [{
            'field_name': 'Legacy', 'field_key': 'legacy_key', 'value': 1, 'normalized_value': 1,
            'data_type': 'number', 'pillar': 'aiqif_c4', 'owner_role': 'hr_officer', 'source_page': '1',
            'source_snippet': 'One.', 'confidence_score': 0.9, 'confidence_reason': 'Stated.',
            'extraction_method': 'manual',
        }])
        self.assertEqual(self._api_facts()['legacy_key']['aiqif_mapping_status'], 'pending')

    def test_facts_can_be_filtered_by_criterion_indicator_and_mapping_status(self):
        self._run(
            _valid_fact(aiqif_indicators=[_indicator('AIQIF-C4.1'), _indicator('AIQIF-C9.2')]),
            _valid_fact(field_key='ai_labs', field_name='AI Labs', aiqif_indicators=[_indicator('AIQIF-C6.4')]),
            _valid_fact(field_key='campus_area', field_name='Campus Area', aiqif_indicators=[]),
        )
        url = f'/api/v1/sprints/{self.sprint.id}/facts/'

        def keys(query):
            response = self.client.get(f'{url}?{query}')
            self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
            return sorted(f['field_key'] for f in response.data)

        self.assertEqual(keys('aiqif_criterion=aiqif_c9'), ['total_faculty'])
        self.assertEqual(keys('aiqif_criterion=aiqif_c4'), ['total_faculty'])  # once, despite two tags
        self.assertEqual(keys('aiqif_indicator=AIQIF-C6.4'), ['ai_labs'])
        self.assertEqual(keys('aiqif_mapping_status=unmapped'), ['campus_area'])



def _stage_now():
    from .usage import _scope
    return (_scope.get() or {}).get('stage')


class DocumentAnalyzerTests(TestCase):
    """services/document_analysis.py: the first chunk's request returns classification AND facts;
    later chunks are fact-only; there is never a separate classification request for the same text."""

    def setUp(self):
        from .services.openai_classifier import OpenAIDocumentClassifier
        institution = Institution.objects.create(name='MKCE', city='Karur', state='TN')
        sprint = Sprint.objects.create(institution=institution, mode=Sprint.SprintMode.VERIFIED_CRI)
        self.document = Document.objects.create(
            sprint=sprint, document_type='naac_ssr', original_filename='ssr.pdf',
            mime_type='application/pdf', status=Document.Status.UPLOADED,
        )
        self.classifier = MagicMock()
        self.classifier.validate.side_effect = OpenAIDocumentClassifier.validate  # the real checks
        self.calls = []  # (schema_name, usage stage) per request

    def _analyzer(self, answers, max_chunk_chars=60):
        from .services.document_analysis import DocumentAnalyzer
        from .services.openai_fact_extractor import OpenAIFactExtractor
        service = MagicMock()
        answers = list(answers)

        def answer(**kwargs):
            self.calls.append((kwargs.get('schema_name'), _stage_now()))
            result = answers.pop(0)
            if isinstance(result, Exception):
                raise result
            return result
        service.extract_structured_data.side_effect = answer
        extractor = OpenAIFactExtractor(openai_service=service, max_chunk_chars=max_chunk_chars, max_chunks=10)
        return DocumentAnalyzer(classifier=self.classifier, fact_extractor=extractor, ai_service=service)

    @staticmethod
    def _pages(*texts):
        return {'pages': [_page(i, t) for i, t in enumerate(texts, start=1)]}

    def _analyze(self, analyzer, pages):
        first = analyzer.analyze_first(self.document, pages)
        return first, analyzer.extract_rest(self.document, first)

    def test_one_request_returns_a_validated_classification_and_facts(self):
        analyzer = self._analyzer([{
            'classification': dict(FAKE_CLASSIFICATION_RESULT),
            'facts': [_valid_fact(), _valid_fact(field_key='bad', source_page='99')],  # page 99 was not sent
        }], max_chunk_chars=10_000)
        first, facts = self._analyze(analyzer, self._pages('The institution has 42 faculty members.'))

        self.assertEqual(first.classification['document_type'], 'naac_ssr')
        self.assertEqual([f['field_key'] for f in facts], ['total_faculty'])  # the invalid one dropped
        self.assertEqual(self.calls, [('document_analysis', 'document_analysis')])
        self.classifier.classify.assert_not_called()

    def test_an_invalid_classification_fails_the_combined_answer(self):
        from .services.openai_classifier import ClassificationValidationError
        for bad in (None, 'naac_ssr', {**FAKE_CLASSIFICATION_RESULT, 'confidence': 7}):
            with self.subTest(bad=bad):
                analyzer = self._analyzer([{'classification': bad, 'facts': []}], max_chunk_chars=10_000)
                with self.assertRaises(ClassificationValidationError):
                    analyzer.analyze_first(self.document, self._pages('Some text.'))

    def test_missing_facts_keep_the_classification(self):
        analyzer = self._analyzer([{'classification': dict(FAKE_CLASSIFICATION_RESULT), 'facts': None}],
                                  max_chunk_chars=10_000)
        first, facts = self._analyze(analyzer, self._pages('Some text.'))
        self.assertEqual(first.classification['confidence'], 0.9)
        self.assertEqual(facts, [])

    def test_a_long_document_classifies_once_with_the_first_chunk_then_extracts_facts_only(self):
        text = 'The institution has 42 faculty members. ' + 'x' * 30
        analyzer = self._analyzer([
            {'classification': dict(FAKE_CLASSIFICATION_RESULT), 'facts': [_valid_fact()]},
            {'facts': [_valid_fact(field_key='total_students', field_name='Total Students', value='900',
                                   source_page='2')]},
            {'facts': [_valid_fact(value='42', source_page='3', confidence_score=0.99)]},  # same key: merged
        ])
        first, facts = self._analyze(analyzer, self._pages(text, text, text))

        self.assertEqual(first.chunk_count, 3)
        self.assertEqual(self.calls, [
            ('document_analysis', 'document_analysis'),
            ('extracted_facts', 'fact_extraction'),
            ('extracted_facts', 'fact_extraction'),
        ])
        self.classifier.classify.assert_not_called()
        self.assertEqual(sorted(f['field_key'] for f in facts), ['total_faculty', 'total_students'])

    def test_a_document_without_text_asks_for_classification_only(self):
        def classify(document, pages=None):
            self.calls.append(('document_classification', _stage_now()))
            return dict(FAKE_CLASSIFICATION_RESULT)
        self.classifier.classify.side_effect = classify
        first, facts = self._analyze(self._analyzer([]), self._pages('', '   '))
        self.assertEqual(self.calls, [('document_classification', 'classification')])
        self.assertEqual((first.classification['document_type'], facts), ('naac_ssr', []))

    def test_a_recoverable_error_retries_the_job_whichever_request_it_hits(self):
        text = 'y' * 70
        analyzer = self._analyzer([RecoverableExtractionError('router busy')])
        with self.assertRaises(RecoverableExtractionError):
            analyzer.analyze_first(self.document, self._pages(text, text))

        analyzer = self._analyzer([
            {'classification': dict(FAKE_CLASSIFICATION_RESULT), 'facts': []},
            RecoverableExtractionError('timeout'),
        ])
        first = analyzer.analyze_first(self.document, self._pages(text, text))
        with self.assertRaises(RecoverableExtractionError):
            analyzer.extract_rest(self.document, first)

    def test_a_permanent_error_on_a_later_chunk_only_skips_that_chunk(self):
        text = 'The institution has 42 faculty members. ' + 'x' * 30
        analyzer = self._analyzer([
            {'classification': dict(FAKE_CLASSIFICATION_RESULT), 'facts': [_valid_fact()]},
            PermanentExtractionError('refused'),
        ])
        _, facts = self._analyze(analyzer, self._pages(text, text))
        self.assertEqual([f['field_key'] for f in facts], ['total_faculty'])

    def test_stop_is_honoured_before_the_first_request_and_between_chunks(self):
        from .exceptions import ExtractionCancelled
        text = 'y' * 70
        analyzer = self._analyzer([])
        analyzer.should_stop = lambda: True
        with self.assertRaises(ExtractionCancelled):
            analyzer.analyze_first(self.document, self._pages(text))
        self.assertEqual(self.calls, [])

        stopped = []
        analyzer = self._analyzer([{'classification': dict(FAKE_CLASSIFICATION_RESULT), 'facts': []}])
        analyzer.should_stop = lambda: bool(stopped)
        first = analyzer.analyze_first(self.document, self._pages(text, text, text))
        stopped.append(True)
        with self.assertRaises(ExtractionCancelled):
            analyzer.extract_rest(self.document, first)
        self.assertEqual(len(self.calls), 1)  # no request after Stop


@override_settings(CELERY_TASK_ALWAYS_EAGER=True, CELERY_TASK_EAGER_PROPAGATES=True)
class CombinedAnalysisPipelineTests(ExtractionJobTestBase):
    """The real pipeline makes ONE AI request for a one-chunk document; its classification still goes
    through the existing confidence / category rules and its facts through the existing persistence."""

    def setUp(self):
        super().setUp()
        patcher = patch('apps.extraction.services.openai_classifier.get_ai_service')
        self.classify_mock = patcher.start().return_value.extract_structured_data
        self.classify_mock.return_value = dict(FAKE_CLASSIFICATION_RESULT)
        self.addCleanup(patcher.stop)
        self._auth(self.iqac)

    _run = DocumentClassificationTests._run
    _classify_as = DocumentClassificationTests._classify_as
    _confirm_category = DocumentClassificationTests._confirm_category

    def test_a_one_chunk_document_needs_one_request(self):
        self.fact_openai_mock.return_value.extract_structured_data.return_value = {'facts': [_valid_fact()]}
        job = self._run()

        self.assertEqual(job.status, ExtractionJob.Status.COMPLETED)
        schemas = [c.kwargs.get('schema_name') for c in self.analysis_service.extract_structured_data.call_args_list]
        self.assertEqual(schemas, ['document_analysis'])
        self.assertEqual(self.document.classification_status, Document.ClassificationStatus.CLASSIFIED)
        self.assertEqual(self.document.extracted_facts.get().field_key, 'total_faculty')

    def test_usage_is_recorded_once_per_response_under_the_combined_stage(self):
        from .models import AIUsage
        from .usage import _save

        def answer(**kwargs):  # what a real AI client does with each response
            _save(provider='openai', model='m', input_tokens=1000, output_tokens=200)
            return {'classification': dict(FAKE_CLASSIFICATION_RESULT), 'facts': [_valid_fact()]}
        self.analysis_service.extract_structured_data.side_effect = answer
        job = self._run()

        row = AIUsage.objects.get()
        self.assertEqual((row.stage, row.extraction_job_id, row.document_id),
                         ('document_analysis', job.id, self.document.id))
        self.assertEqual((row.input_tokens, row.output_tokens), (1000, 200))  # one row, not split in two

    def test_a_low_confidence_combined_classification_needs_review(self):
        self._classify_as('naac_ssr', 0.3, 'Unsure.')
        self.fact_openai_mock.return_value.extract_structured_data.return_value = {'facts': [_valid_fact()]}
        self._run()
        self.assertEqual(self.document.classification_status, Document.ClassificationStatus.NEEDS_REVIEW)
        self.assertEqual(self.document.extracted_facts.count(), 1)  # the facts are still kept

    def test_an_unsupported_category_is_unclassified(self):
        self._classify_as('not_a_real_category', 0.95)
        self._run()
        self.assertEqual(self.document.classification_status, Document.ClassificationStatus.UNCLASSIFIED)

    def test_a_confirmed_category_and_a_reviewed_fact_survive_reprocessing(self):
        self._classify_as('naac_ssr', 0.45, 'Unsure.')
        self.fact_openai_mock.return_value.extract_structured_data.return_value = {'facts': [_valid_fact(value='42')]}
        self._run()
        self._confirm_category('aqar_report')
        fact = self.document.extracted_facts.get()
        fact.status = ExtractedFact.Status.CONFIRMED
        fact.save()

        self._classify_as('naac_ssr', 0.97)
        self.fact_openai_mock.return_value.extract_structured_data.return_value = {'facts': [_valid_fact(value='99')]}
        self._run()

        fact.refresh_from_db()
        self.assertEqual(self.document.ai_document_type, 'aqar_report')
        self.assertEqual(self.document.extracted_facts.count(), 1)
        self.assertEqual((fact.status, fact.value), (ExtractedFact.Status.CONFIRMED, 42))


# -- extraction completeness: partial documents are visible, retried in place, and gate the baseline ----------

PAGE_TEXTS = (
    'Page one marker. The institution has 42 faculty members on its rolls.',
    'Page two marker. Laboratories list twelve GPU workstations for students.',
    'Page three marker. The placement cell records 310 offers this academic year.',
)


@override_settings(CELERY_TASK_ALWAYS_EAGER=True, CELERY_TASK_EAGER_PROPAGATES=True,
                   OPENAI_FACT_EXTRACTION_MAX_CHUNK_CHARS=90)
class ExtractionCompletenessTests(ExtractionJobTestBase):
    """A document whose chunks failed or were skipped is "partial", never shown as fully processed."""

    def setUp(self):
        super().setUp()
        self.fail_page_two = True
        self.analysis_service.extract_structured_data.side_effect = self._answer
        self._auth(self.iqac)
        upload = self.client.post(f'/api/v1/sprints/{self.sprint.id}/upload-file/',
                                  {'file': make_pdf('three.pdf', PAGE_TEXTS), 'document_type': 'naac_ssr'},
                                  format='multipart')
        self.doc = Document.objects.get(id=upload.data['id'])
        Document.objects.filter(id=self.document.id).update(status=Document.Status.REJECTED)  # only self.doc counts
        self.admin = self._make_user('admin@test.edu', User.Role.INSTITUTION_ADMIN, self.institution)

    def _answer(self, **kwargs):
        content = kwargs.get('user_content', '')
        if kwargs.get('schema_name') == 'document_analysis':
            return {'classification': dict(FAKE_CLASSIFICATION_RESULT),
                    'facts': [_valid_fact(source_page='1', source_snippet='The institution has 42 faculty members')]}
        if self.fail_page_two and 'Page two marker' in content:
            raise PermanentExtractionError('The AI answer was not valid JSON.')
        return {'facts': []}

    def _process(self):
        r = self.client.post(f'/api/v1/sprints/{self.sprint.id}/extraction-jobs/', {'document_id': str(self.doc.id)})
        self.assertEqual(r.status_code, status.HTTP_201_CREATED, r.data)
        return ExtractionJob.objects.get(id=r.data[0]['id'])

    def _summary(self):
        return self.client.get(f'/api/v1/sprints/{self.sprint.id}/processing-summary').data

    def test_a_failed_chunk_makes_the_document_partial_everywhere(self):
        job = self._process()
        self.assertEqual((job.status, job.completeness), (ExtractionJob.Status.COMPLETED, 'partial'))
        self.assertEqual((job.chunks_total, job.chunks_completed), (3, 2))
        self.assertEqual([w['kind'] for w in job.extraction_warnings], ['chunk_failed'])
        self.assertIn('pages 2-2', job.extraction_warnings[0]['detail'])
        self.doc.refresh_from_db()
        self.assertEqual(self.doc.processing_status, 'processed_partial')
        listed = self.client.get(f'/api/v1/sprints/{self.sprint.id}/extraction-jobs').data
        rows = listed['results'] if isinstance(listed, dict) else listed
        self.assertEqual(rows[0]['completeness'], 'partial')
        summary = self._summary()
        self.assertFalse(summary['complete'])
        self.assertEqual((summary['counts']['partial'], summary['counts']['complete']), (1, 0))
        self.assertIn('1 partial', summary['description'])

    def test_chunks_over_the_limit_are_reported_as_skipped(self):
        self.fail_page_two = False
        with override_settings(OPENAI_FACT_EXTRACTION_MAX_CHUNKS=2):
            job = self._process()
        self.assertEqual(job.completeness, 'partial')
        self.assertEqual(job.extraction_warnings[0]['kind'], 'chunks_skipped')
        self.assertEqual(job.extraction_warnings[0]['chunks'], [3])

    def test_retrying_a_partial_document_reuses_its_job_and_never_duplicates_facts(self):
        job = self._process()
        facts_before = sorted(self.doc.extracted_facts.values_list('field_key', flat=True))
        self.fail_page_two = False
        r = self.client.post(f'/api/v1/extraction-jobs/{job.id}/process')
        self.assertEqual(r.status_code, status.HTTP_200_OK, r.data)
        job.refresh_from_db()
        self.assertEqual((r.data['id'], job.completeness, job.extraction_warnings), (str(job.id), 'complete', []))
        self.assertEqual(ExtractionJob.objects.filter(document=self.doc).count(), 1)
        self.assertEqual(sorted(self.doc.extracted_facts.values_list('field_key', flat=True)), facts_before)
        self.assertTrue(self._summary()['complete'])
        # A fully processed document is not queued again by "process unfinished" or a bulk selection.
        again = self.client.post(f'/api/v1/extraction-jobs/sprints/{self.sprint.id}/process-incomplete')
        self.assertEqual(again.data, [])
        bulk = self.client.post(f'/api/v1/sprints/{self.sprint.id}/extraction-jobs/',
                                {'document_ids': [str(self.doc.id)]}, format='json')
        self.assertEqual(bulk.data, [])

    def test_process_unfinished_picks_up_partial_documents(self):
        job = self._process()
        self.fail_page_two = False
        r = self.client.post(f'/api/v1/extraction-jobs/sprints/{self.sprint.id}/process-incomplete')
        self.assertEqual([row['id'] for row in r.data], [str(job.id)])
        job.refresh_from_db()
        self.assertEqual(job.completeness, 'complete')

    # -- the baseline gate ---------------------------------------------------------------------------
    def _pending_baseline(self):
        from apps.scoring.services import get_or_create_pending_baseline
        return get_or_create_pending_baseline(self.sprint, triggered_by=self.admin)

    def test_baseline_approval_is_refused_while_documents_are_partial_failed_or_unprocessed(self):
        self._process()
        upload = self.client.post(f'/api/v1/sprints/{self.sprint.id}/upload-file/',
                                  {'file': make_pdf('new.pdf', ('A newly added document, never processed yet.',)),
                                   'document_type': 'aqar_report'}, format='multipart')
        self.assertEqual(upload.status_code, status.HTTP_201_CREATED, upload.data)
        self._pending_baseline()
        from apps.gaps.models import GapItem
        GapItem.objects.filter(sprint=self.sprint, priority='blocking').update(status='resolved')  # gaps are another gate
        self._auth(self.admin)
        for path in ('approve', 'approve-provisional'):
            r = self.client.post(f'/api/v1/sprints/{self.sprint.id}/baseline/{path}', {'comments': 'ok'}, format='json')
            self.assertEqual(r.status_code, status.HTTP_400_BAD_REQUEST, r.data)
            self.assertIn('not fully processed', str(r.data['detail']))
            counts = r.data['processing']['counts']
            self.assertEqual((int(counts['partial']), int(counts['not_processed'])), (1, 1))
        self.assertEqual(self.sprint.baselines.get().status, 'pending')
        baseline_view = self.client.get(f'/api/v1/sprints/{self.sprint.id}/baseline').data
        self.assertFalse(baseline_view['processing']['complete'])

    def test_an_override_needs_a_reason_and_is_recorded_and_reported(self):
        from apps.reports.services import build_report_data
        from apps.scoring.models import BaselineDecisionHistory
        self._process()
        self._pending_baseline()
        self._auth(self.admin)
        r = self.client.post(f'/api/v1/sprints/{self.sprint.id}/baseline/approve-provisional',
                             {'comments': 'ok', 'processing_override_reason': '   '}, format='json')
        self.assertEqual(r.status_code, status.HTTP_400_BAD_REQUEST)  # blank is not a reason
        r = self.client.post(f'/api/v1/sprints/{self.sprint.id}/baseline/approve-provisional',
                             {'comments': 'ok', 'processing_override_reason': 'Page 2 is a scanned annexure; '
                              'its figures were confirmed by hand.'}, format='json')
        self.assertEqual(r.status_code, status.HTTP_200_OK, r.data)
        self.assertTrue(r.data['processing_incomplete'])
        self.assertIn('scanned annexure', r.data['processing_override_reason'])
        self.assertIn('1 partial', r.data['processing_snapshot']['description'])
        history = BaselineDecisionHistory.objects.filter(action='approved_provisional').get()
        self.assertIn('Decided with incomplete document processing', history.comments)
        data = build_report_data(self.sprint)
        self.assertFalse(data['processing']['complete'])
        self.assertIn('scanned annexure', data['processing']['baseline_override']['reason'])
        from apps.reports.rendering import _processing_lines
        self.assertTrue(_processing_lines(data)[0].startswith('INCOMPLETE DOCUMENT PROCESSING'))

    def test_fully_processed_documents_need_no_reason(self):
        self.fail_page_two = False
        self._process()
        self._pending_baseline()
        self._auth(self.admin)
        r = self.client.post(f'/api/v1/sprints/{self.sprint.id}/baseline/approve-provisional', {'comments': 'ok'},
                             format='json')
        self.assertEqual(r.status_code, status.HTTP_200_OK, r.data)
        self.assertFalse(r.data['processing_incomplete'])
        self.assertTrue(r.data['processing_snapshot']['complete'])

    def test_other_colleges_cannot_read_the_summary(self):
        self._auth(self.outsider)
        r = self.client.get(f'/api/v1/sprints/{self.sprint.id}/processing-summary')
        self.assertIn(r.status_code, (status.HTTP_403_FORBIDDEN, status.HTTP_404_NOT_FOUND))


@override_settings(CELERY_TASK_ALWAYS_EAGER=True, CELERY_TASK_EAGER_PROPAGATES=True)
class AutoConfirmTests(ExtractionJobTestBase):
    """Facts above FACT_AUTO_CONFIRM_MIN_CONFIDENCE are confirmed by the system at the end of extraction
    (apps/facts/auto_confirm.py) -- never ones in a conflict or whose number is not in their quote."""

    def setUp(self):
        super().setUp()
        on = override_settings(FACT_AUTO_CONFIRM_MIN_CONFIDENCE=0.8)
        on.enable()
        self.addCleanup(on.disable)
        self._auth(self.iqac)

    def _run(self, *facts):
        self.fact_openai_mock.return_value.extract_structured_data.return_value = {'facts': list(facts)}
        r = self.client.post(f'/api/v1/sprints/{self.sprint.id}/extraction-jobs/', {'document_id': str(self.document.id)})
        self.assertEqual(r.status_code, status.HTTP_201_CREATED, r.data)
        return {f.field_key: f for f in self.document.extracted_facts.select_related('reviewed_by')}

    def test_only_facts_strictly_above_the_threshold_are_confirmed_by_the_system(self):
        from apps.facts.auto_confirm import AUTO_CONFIRM_EMAIL
        facts = self._run(
            _valid_fact(),  # 0.96, "42" is in its quote
            _valid_fact(field_key='labs', field_name='Labs', value='7', source_snippet='There are 7 labs.',
                        confidence_score=0.8),
            _valid_fact(field_key='rooms', field_name='Rooms', value='9', source_snippet='There are 9 rooms.',
                        confidence_score=0.5),
        )
        auto = facts['total_faculty']
        self.assertEqual((auto.status, auto.reviewed_by.email), (ExtractedFact.Status.CONFIRMED, AUTO_CONFIRM_EMAIL))
        self.assertEqual(auto.review_history.get().reason, 'Auto-confirmed: AI confidence 0.96 is above 0.80.')
        self.assertEqual((facts['labs'].status, facts['rooms'].status), ('extracted', 'extracted'))
        listed = self.client.get(f'/api/v1/sprints/{self.sprint.id}/facts/').data
        rows = {f['field_key']: f for f in (listed['results'] if isinstance(listed, dict) else listed)}
        self.assertEqual((rows['total_faculty']['auto_confirmed'], rows['labs']['auto_confirmed']), (True, False))
        from apps.scoring.services import run_scoring_engine
        self.assertGreater(run_scoring_engine(self.sprint).evidence_count, 0)  # it counts in the CRI

    def test_a_number_not_in_its_own_quote_is_left_for_a_person(self):
        facts = self._run(
            _valid_fact(field_key='pct_revised', field_name='Percent revised', value='37', data_type='percentage',
                        source_snippet='Percentage of Programs where syllabus revision was carried out',
                        confidence_score=0.99),
            _valid_fact(field_key='doctors', field_name='Doctors', value='5',
                        source_snippet='Qualified Doctor (Full time) 12', confidence_score=0.95),
        )
        self.assertEqual({f.status for f in facts.values()}, {'extracted'})

    def test_a_fact_in_an_open_conflict_is_left_for_a_person(self):
        from apps.facts.auto_confirm import auto_confirm
        from apps.gaps.models import GapItem
        fact = ExtractedFact.objects.create(
            sprint=self.sprint, document=self.document, field_name='Faculty', field_key='faculty', value=42,
            data_type='number', pillar='aiqif_c4', source_page='1', source_snippet='42 faculty members.',
            confidence_score=0.95)
        GapItem.objects.create(sprint=self.sprint, gap_type=GapItem.GapType.CONFLICT, title='Two values',
                               source_fact=fact, pillar='aiqif_c4')
        result = auto_confirm(ExtractedFact.objects.filter(pk=fact.pk))
        self.assertEqual((result['confirmed'], result['skipped']), ([], {'open conflict': 1}))

    def test_reprocessing_updates_an_auto_confirmed_fact_but_never_a_persons(self):
        first = self._run(_valid_fact(), _valid_fact(field_key='labs', field_name='Labs', value='7',
                                                     source_snippet='There are 7 labs.', confidence_score=0.5))
        person = first['labs']
        person.status, person.reviewed_by = ExtractedFact.Status.CONFIRMED, self.iqac  # a person confirmed it
        person.save()
        again = self._run(
            _valid_fact(value='45', source_snippet='The institution has 45 faculty members.'),
            _valid_fact(field_key='labs', field_name='Labs', value='8', source_snippet='There are 8 labs.',
                        confidence_score=0.99),
        )
        auto = again['total_faculty']
        self.assertEqual((auto.id, auto.value, auto.status), (first['total_faculty'].id, 45, 'confirmed'))
        self.assertEqual(auto.review_history.count(), 2)  # confirmed automatically for 42, then again for 45
        self.assertEqual((again['labs'].value, again['labs'].reviewed_by), (7, self.iqac))  # untouched
        self.assertEqual(self.document.extracted_facts.count(), 2)

    def test_it_can_be_switched_off(self):
        with override_settings(FACT_AUTO_CONFIRM_MIN_CONFIDENCE=None):
            facts = self._run(_valid_fact())
        self.assertEqual(facts['total_faculty'].status, 'extracted')


@override_settings(CELERY_TASK_ALWAYS_EAGER=True, CELERY_TASK_EAGER_PROPAGATES=True)
class GapEvidenceUploadTests(ExtractionJobTestBase):
    """Upload Evidence from a Missing Document / Low Confidence gap (apps/gaps/evidence_upload.py)."""

    def setUp(self):
        super().setUp()
        self.admin = self._make_user('admin@test.edu', User.Role.INSTITUTION_ADMIN, self.institution)
        self.viewer = self._make_user('viewer@test.edu', User.Role.VIEWER, self.institution)
        self._auth(self.iqac)

    def _facts(self, *facts):
        self.fact_openai_mock.return_value.extract_structured_data.return_value = {'facts': list(facts)}

    #: One fixed PDF, so "the same file" really is byte-identical (generated PDFs carry a timestamp).
    EVIDENCE_BYTES = None

    def _evidence(self, name='evidence.pdf'):
        if GapEvidenceUploadTests.EVIDENCE_BYTES is None:
            GapEvidenceUploadTests.EVIDENCE_BYTES = make_pdf_bytes(
                ('Signed evidence: the institution has 45 faculty members.',))
        return SimpleUploadedFile(name, GapEvidenceUploadTests.EVIDENCE_BYTES, content_type='application/pdf')

    def _upload(self, gap, name='evidence.pdf', user=None):
        self._auth(user or self.admin)  # someone who may act on the gap (apps/facts/ownership.py)
        return self.client.post(f'/api/v1/gaps/{gap.id}/upload-evidence', {'file': self._evidence(name)},
                                format='multipart')

    def _events(self, gap):
        return list(gap.events.order_by('created_at').values_list('action', flat=True))

    def _missing_gap(self):
        from apps.gaps.models import GapItem
        from apps.gaps.services import generate_gaps_for_sprint
        generate_gaps_for_sprint(self.sprint)
        return GapItem.objects.filter(sprint=self.sprint, gap_type='missing_document', title='Missing AQAR Report').first() \
            or GapItem.objects.filter(sprint=self.sprint, gap_type='missing_document').first()

    def _low_confidence_gap(self):
        from apps.gaps.models import GapItem
        self._facts(_valid_fact(confidence_score=0.3, value='42'))
        self.client.post(f'/api/v1/sprints/{self.sprint.id}/extraction-jobs/', {'document_id': str(self.document.id)})
        return GapItem.objects.get(sprint=self.sprint, gap_type='low_confidence')

    # -- missing document ------------------------------------------------------------------------------
    def test_missing_document_upload_is_processed_in_this_sprint_but_resolves_nothing_by_itself(self):
        gap = self._missing_gap()
        self._facts(_valid_fact(value='45', source_snippet='Signed evidence: the institution has 45 faculty members.'))
        r = self._upload(gap)
        self.assertEqual(r.status_code, status.HTTP_201_CREATED, r.data)
        doc = Document.objects.get(id=r.data['document_id'])
        gap.refresh_from_db()
        self.assertEqual((doc.sprint, list(gap.evidence_documents.all())), (self.sprint, [doc]))
        self.assertEqual(ExtractionJob.objects.get(document=doc).status, ExtractionJob.Status.COMPLETED)
        self.assertEqual(gap.status, 'in_progress')  # never resolved by an upload
        self.assertEqual(self._events(gap)[-2:], ['evidence_uploaded', 'evidence_processed'])
        self.assertEqual(set(doc.extracted_facts.values_list('status', flat=True)), {'extracted'})
        state = self.client.get(f'/api/v1/gaps/{gap.id}').data['evidence_state']
        self.assertEqual(state['state'], 'awaiting_review')
        self.assertIn('Confirm its category', state['next_action'])
        # The reviewer confirms the document's category -> the existing rule closes the gap.
        t = gap.title.replace('Missing ', '')
        from apps.gaps.services import missing_document_type
        self._auth(self.admin)
        r = self.client.post(f'/api/v1/documents/{doc.id}/confirm-category',
                             {'category': missing_document_type(gap)}, format='json')
        self.assertEqual(r.status_code, status.HTTP_200_OK, r.data)
        gap.refresh_from_db()
        self.assertEqual(gap.status, 'resolved', t)

    def test_the_same_file_in_the_same_sprint_is_reused_not_duplicated(self):
        gap = self._missing_gap()
        first = self._upload(gap)
        again = self._upload(gap)
        self.assertEqual((again.status_code, again.data['reused']), (status.HTTP_201_CREATED, True))
        self.assertEqual(again.data['document_id'], first.data['document_id'])
        self.assertEqual(ExtractionJob.objects.filter(document_id=first.data['document_id']).count(), 1)

    def test_a_file_processed_in_another_sprint_is_processed_again_here(self):
        other = Sprint.objects.create(institution=self.institution, mode=Sprint.SprintMode.VERIFIED_CRI)
        first = self.client.post(f'/api/v1/sprints/{other.id}/upload-file/',
                                 {'file': self._evidence('renamed.pdf'), 'document_type': 'aqar_report'}, format='multipart')
        self.assertEqual(first.status_code, status.HTTP_201_CREATED, first.data)
        gap = self._missing_gap()
        r = self._upload(gap)
        doc = Document.objects.get(id=r.data['document_id'])
        self.assertEqual((r.data['reused'], doc.sprint), (False, self.sprint))
        self.assertTrue(ExtractionJob.objects.filter(document=doc, sprint=self.sprint).exists())

    def test_a_failed_extraction_is_recorded_and_retried_on_the_same_job(self):
        gap = self._missing_gap()
        self.analysis_service.extract_structured_data.side_effect = PermanentExtractionError('bad file')
        r = self._upload(gap)
        job = ExtractionJob.objects.get(document_id=r.data['document_id'])
        self.assertEqual(job.status, ExtractionJob.Status.FAILED)
        self.assertIn('evidence_failed', self._events(gap))
        self.assertEqual(self.client.get(f'/api/v1/gaps/{gap.id}').data['evidence_state']['state'], 'processing_failed')
        route_combined_analysis(self)  # the AI answers again
        self.client.post(f'/api/v1/extraction-jobs/{job.id}/process')
        self.assertEqual(ExtractionJob.objects.filter(document_id=r.data['document_id']).count(), 1)
        job.refresh_from_db()
        self.assertEqual(job.status, ExtractionJob.Status.COMPLETED)

    # -- low confidence -------------------------------------------------------------------------------
    def test_low_confidence_closes_only_when_the_new_value_is_reviewed(self):
        from apps.scoring.models import ScoringRun
        gap = self._low_confidence_gap()
        self._facts(_valid_fact(value='45', confidence_score=0.95,
                                source_snippet='Signed evidence: the institution has 45 faculty members.'))
        r = self._upload(gap)
        new = ExtractedFact.objects.get(document_id=r.data['document_id'], field_key='total_faculty')
        gap.refresh_from_db()
        self.assertEqual((new.status, gap.status), ('extracted', 'in_progress'))  # a separate suggestion
        listed = self.client.get(f'/api/v1/sprints/{self.sprint.id}/facts/').data
        row = next(f for f in (listed['results'] if isinstance(listed, dict) else listed) if f['id'] == str(new.id))
        self.assertEqual(row['evidence_for_gaps'][0]['id'], str(gap.id))
        runs = ScoringRun.objects.filter(sprint=self.sprint).count()
        self._auth(self.admin)
        r = self.client.post(f'/api/v1/facts/{new.id}/confirm', {'reason': 'Signed register.'})
        self.assertEqual(r.status_code, status.HTTP_200_OK, r.data)
        gap.refresh_from_db()
        self.assertEqual(gap.status, 'resolved')
        self.assertIn('Verified from uploaded evidence', gap.resolution)
        self.assertEqual(ScoringRun.objects.filter(sprint=self.sprint).count(), runs + 1)  # CRI recalculated

    def test_a_conflict_keeps_the_low_confidence_gap_open(self):
        from apps.gaps.models import GapItem
        gap = self._low_confidence_gap()
        self._facts(_valid_fact(value='45', confidence_score=0.95,
                                source_snippet='Signed evidence: the institution has 45 faculty members.'))
        r = self._upload(gap)
        new = ExtractedFact.objects.get(document_id=r.data['document_id'], field_key='total_faculty')
        GapItem.objects.create(sprint=self.sprint, gap_type='conflict', title='Two values', source_fact=gap.source_fact,
                               conflict_fact_b=new, pillar='aiqif_c4')
        self._auth(self.admin)
        self.client.post(f'/api/v1/facts/{new.id}/confirm', {'reason': 'x'})
        gap.refresh_from_db()
        self.assertEqual(gap.status, 'in_progress')

    # -- access --------------------------------------------------------------------------------------
    def test_only_people_who_may_act_and_only_eligible_gaps(self):
        from apps.gaps.models import GapItem
        gap = self._missing_gap()
        self.assertEqual(self._upload(gap, user=self.viewer).status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(self._upload(gap, user=self.outsider).status_code, status.HTTP_403_FORBIDDEN)
        from apps.facts.ownership import gap_owner_role
        if gap_owner_role(gap) not in ('', None, 'iqac_coordinator'):
            self.assertEqual(self._upload(gap, user=self.iqac).status_code, status.HTTP_403_FORBIDDEN)  # not its owner
        conflict = GapItem.objects.create(sprint=self.sprint, gap_type='conflict', title='x', pillar='aiqif_c4')
        self.assertEqual(self._upload(conflict).status_code, status.HTTP_400_BAD_REQUEST)
        self.assertFalse(gap.evidence_documents.exists())
