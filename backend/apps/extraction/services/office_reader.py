"""Text from non-PDF uploads, shaped like PDFPageReader's pages so the rest of
the pipeline (chunking, fact extraction with page citations) is unchanged.

These formats have no real pages, so a "page" here is a stable section the AI
can cite and a person can find again:
  .docx        consecutive paragraphs/tables grouped up to ~PAGE_CHARS characters
  .xlsx        one page per worksheet (long sheets split every SHEET_ROWS rows),
               each row as "A | B | C" with the sheet name as a heading
  .csv         like one worksheet
  .png/.jpg    one page, read by OCR (the same PaddleOCR as scanned PDF pages)

Legacy binary .doc / .xls and archives (.zip) are not read; the caller reports
that to the user instead of silently producing nothing.
"""
import csv
import io
import logging

logger = logging.getLogger(__name__)

PAGE_CHARS = 3000
SHEET_ROWS = 150
MAX_CELLS = 200_000  # safety cap on very large workbooks

OFFICE_EXTENSIONS = {'.docx', '.xlsx', '.csv'}
IMAGE_EXTENSIONS = {'.png', '.jpg', '.jpeg'}
UNREADABLE_NOTES = {
    '.doc': 'Old Word format (.doc) cannot be read. Save it as .docx or PDF and upload again.',
    '.xls': 'Old Excel format (.xls) cannot be read. Save it as .xlsx or CSV and upload again.',
    '.zip': 'ZIP archives are not unpacked. Upload the files inside it individually.',
}


def _group(blocks, limit=PAGE_CHARS):
    pages, current, size = [], [], 0
    for block in blocks:
        block = block.strip()
        if not block:
            continue
        if current and size + len(block) > limit:
            pages.append('\n'.join(current))
            current, size = [], 0
        current.append(block)
        size += len(block) + 1
    if current:
        pages.append('\n'.join(current))
    return pages


def read_docx(data):
    from docx import Document as DocxDocument  # python-docx

    doc = DocxDocument(io.BytesIO(data))
    blocks = []
    body = doc.element.body
    for child in body.iterchildren():
        tag = child.tag.rsplit('}', 1)[-1]
        if tag == 'p':
            text = ''.join(node.text or '' for node in child.iter() if node.tag.endswith('}t'))
            blocks.append(text)
        elif tag == 'tbl':
            for row in child.iter():
                if row.tag.endswith('}tr'):
                    cells = []
                    for cell in row.iterchildren():
                        if cell.tag.endswith('}tc'):
                            cells.append(' '.join((n.text or '').strip() for n in cell.iter() if n.tag.endswith('}t')).strip())
                    if any(cells):
                        blocks.append(' | '.join(cells))
    return _group(blocks)


def _rows_to_pages(title, rows):
    pages = []
    batch = []
    for row in rows:
        values = ['' if v is None else str(v).strip() for v in row]
        while values and not values[-1]:
            values.pop()
        if any(values):
            batch.append(' | '.join(values))
        if len(batch) >= SHEET_ROWS:
            pages.append(f'Sheet: {title}\n' + '\n'.join(batch))
            batch = []
    if batch:
        pages.append(f'Sheet: {title}\n' + '\n'.join(batch))
    return pages


def read_xlsx(data):
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        pages, cells = [], 0
        for ws in wb.worksheets:
            def limited(rows):
                nonlocal cells
                for row in rows:
                    cells += len(row)
                    if cells > MAX_CELLS:
                        logger.warning('office_reader.xlsx_truncated sheet=%s', ws.title)
                        return
                    yield row
            pages += _rows_to_pages(ws.title, limited(ws.iter_rows(values_only=True)))
        return pages
    finally:
        wb.close()


def read_csv(data):
    text = data.decode('utf-8-sig', errors='replace')
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=',;\t|')
    except csv.Error:
        dialect = csv.excel
    return _rows_to_pages('CSV', csv.reader(io.StringIO(text), dialect))


def image_to_bgr(data):
    """PNG/JPEG bytes -> BGR uint8 array (what PaddleOCR expects)."""
    import numpy as np
    from PIL import Image

    with Image.open(io.BytesIO(data)) as img:
        rgb = np.asarray(img.convert('RGB'))
    return np.ascontiguousarray(rgb[:, :, ::-1])
