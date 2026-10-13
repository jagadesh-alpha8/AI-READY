"""Import the intern Source Register (.xlsx) into the Source Library.

    rows = read_register(path)          # parse + check every row, no DB access
    summary = apply_register(rows, actor='...', commit=False)   # dry run

Errors skip a row; warnings import it with `review_flags` so nothing doubtful
becomes approved silently. Everything is imported as a DRAFT version -- a
version is activated only by a reviewer. Re-running is safe: drafts are
updated in place, approved (locked) versions are never touched.
"""
import datetime as dt
import re
from dataclasses import dataclass, field
from urllib.parse import urlsplit

from django.db import transaction

from ..models import Authority, Framework, SourceAuditEvent, SourceDocument, SourceFeed, SourceVersion

OFFICIAL_DOMAINS = ('.gov.in', '.nic.in', '.gov', 'nbaind.org', 'nirfindia.org', 'aicte-india.org', 'unesco.org',
                    'iso.org', 'w3.org', '1edtech.org', 'imsglobal.org', 'khronos.org', 'cert-in.org.in', 'ugc.ac.in')
PLACEHOLDER_TITLE = re.compile(r'\(one row per|\(identify|research-call pages|policy sources|integrity sources|'
                               r'\(official sources\)|roadmaps \(|requirements \(identify', re.I)
SHA256 = re.compile(r'^[0-9a-fA-F]{64}$')

COLLECTION = {label: value for value, label in SourceVersion.Collection.choices}
ROLES = {label: value for value, label in SourceDocument.Role.choices}
KINDS = set(Authority.Kind.values)
FEED_TYPES = set(SourceFeed.FeedType.values)
AUTHORITY_CODES = {'Ministry of Education': 'MOE', 'Academic Bank of Credits': 'ABC', 'Parliament of India': 'PARLIAMENT',
                   'MeitY/IndiaAI': 'MEITY-INDIAAI', 'ISO/IEC': 'ISO-IEC', 'Khronos Group': 'KHRONOS',
                   'Other (write in notes)': 'OTHER'}

REQUIRED_HEADERS = ('Row ID', 'Category', 'Authority', 'Authority Kind', 'Source Code', 'Official Title (exact)',
                    'Collection Status')


@dataclass
class RegisterRow:
    row_id: str
    data: dict
    errors: list = field(default_factory=list)
    warnings: list = field(default_factory=list)

    @property
    def ok(self):
        return not self.errors


def domain(url):
    return urlsplit(url).netloc.lower().removeprefix('www.') if url.startswith('http') else ''


def is_official(url):
    d = domain(url)
    return any(d.endswith(o) if o.startswith('.') else (d == o or d.endswith('.' + o)) for o in OFFICIAL_DOMAINS)


def authority_code(name):
    return AUTHORITY_CODES.get(name) or re.sub(r'[^A-Z0-9]+', '-', name.upper()).strip('-')[:40]


def _text(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _date(value, label, row):
    if value in (None, ''):
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    for fmt in ('%d-%b-%Y', '%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y'):
        try:
            return dt.datetime.strptime(str(value).strip(), fmt).date()
        except ValueError:
            pass
    row.warnings.append(f'{label} "{value}" is not a date - ignored.')
    return None


def _int(value):
    try:
        return int(float(value)) if value not in (None, '') else None
    except (TypeError, ValueError):
        return None


def read_register(path, sheet='Source Register'):
    """Parse and check every row. Returns [RegisterRow] for rows that have a title."""
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        if sheet not in wb.sheetnames:
            raise ValueError(f'No sheet named {sheet!r}.')
        rows = list(wb[sheet].iter_rows(values_only=True))
    finally:
        wb.close()  # read-only mode keeps the file open (and locked on Windows) until closed
    header_idx = next((i for i, r in enumerate(rows[:20]) if r and r[0] == 'Row ID'), None)
    if header_idx is None:
        raise ValueError(f'No "Row ID" header row found in sheet {sheet!r}.')
    headers = [_text(h) for h in rows[header_idx]]
    missing = [h for h in REQUIRED_HEADERS if h not in headers]
    if missing:
        raise ValueError(f'Sheet {sheet!r} is missing columns: {", ".join(missing)}')
    col = {h: i for i, h in enumerate(headers)}

    out, seen = [], {}
    for raw in rows[header_idx + 1:]:
        get = lambda h: raw[col[h]] if h in col and col[h] < len(raw) else None  # noqa: E731
        title = _text(get('Official Title (exact)'))
        if not title:
            continue
        row = RegisterRow(row_id=_text(get('Row ID')), data={})
        d = row.data
        d['title'] = title
        d['source_code'] = _text(get('Source Code'))  # kept exactly as the register writes it (e.g. NCrF-001)
        d['category'] = _text(get('Category'))
        d['authority'] = _text(get('Authority'))
        d['kind'] = _text(get('Authority Kind'))
        d['role'] = ROLES.get(_text(get('Source Role (suggested)')), '')
        d['document_type'] = _text(get('Document Type'))
        d['version_label'] = _text(get('Version / Edition')) or SourceVersion.UNVERSIONED
        d['issued_on'] = _date(get('Issued On'), 'Issued On', row)
        d['effective_from'] = _date(get('Effective From'), 'Effective From', row)
        d['reported_status'] = _text(get('Current Status'))
        d['applies_to'] = _text(get('Applies To'))
        d['programme_scope'] = _text(get('Programme / Trade Scope'))
        d['listing_url'] = _text(get('Listing Page URL'))
        d['direct_url'] = _text(get('Direct Document URL'))
        d['collection'] = COLLECTION.get(_text(get('Collection Status')))
        d['local_file'] = _text(get('Local File (folder/name)'))
        d['file_size_kb'] = _int(get('File Size (KB)'))
        d['pages'] = _int(get('Pages'))
        d['sha256'] = _text(get('SHA-256 Checksum')).lower()
        d['licensed'] = _text(get('Licensed Content?')).lower() == 'yes'
        d['cadence'] = _int(get('Review Cadence (months)')) or 12
        d['feed_type'] = _text(get('Watcher Feed Type'))
        d['collected_on'] = _date(get('Collected On'), 'Collected On', row)
        d['notes'] = _text(get('Notes'))
        _check(row, seen)
        out.append(row)
    return out


def _check(row, seen):
    d, err, warn = row.data, row.errors, row.warnings
    if not d['source_code']:
        err.append('Source Code is empty.')
    if not d['authority']:
        err.append('Authority is empty.')
    if d['kind'] not in KINDS:
        err.append(f'Authority Kind "{d["kind"]}" is not one of {sorted(KINDS)}.')
    if d['collection'] is None:
        err.append('Collection Status is empty or not a list value.')
    if d['sha256'] and not SHA256.match(d['sha256']):
        err.append('SHA-256 is not 64 hex characters.')
    key = (d['source_code'], d['version_label'])
    if d['source_code'] and key in seen:
        err.append(f'Duplicate of {seen[key]} (same Source Code and Version).')
    else:
        seen[key] = row.row_id

    for label, url in (('Listing Page URL', d['listing_url']), ('Direct Document URL', d['direct_url'])):
        if url and not url.startswith('http'):
            warn.append(f'{label} is not a URL - ignored.')
            d['listing_url' if label.startswith('Listing') else 'direct_url'] = ''
        elif url and not is_official(url):
            warn.append(f'{label} is on {domain(url)}, not an official domain.')
    drive = d['local_file'] if 'drive.google.com' in d['local_file'] else ''
    if d['licensed']:
        if drive:
            warn.append('Licensed standard: the Drive copy is NOT imported - delete it unless InGage holds a licence.')
        drive = ''
        if d['collection'] == SourceVersion.Collection.DOWNLOADED:
            d['collection'] = SourceVersion.Collection.LICENSED
            warn.append('Licensed standard marked Downloaded - imported as "Licensed - link only".')
    elif '/folders/' in drive:
        warn.append('Drive link is a folder, not the file.')
    d['drive_url'] = drive
    if d['local_file'] and not drive and not d['licensed']:
        d['notes'] = (d['notes'] + f' | Local file: {d["local_file"]}').strip(' |')

    if d['collection'] == SourceVersion.Collection.DOWNLOADED:
        missing = [label for label, k in (('direct URL', 'direct_url'), ('file link', 'drive_url'),
                                          ('SHA-256', 'sha256'), ('size', 'file_size_kb'))
                   if not d[k]]
        if d['version_label'] == SourceVersion.UNVERSIONED:
            missing.append('version')
        if not d['issued_on']:
            missing.append('issued date')
        if missing:
            warn.append('Downloaded but missing: ' + ', '.join(missing) + '.')
    if PLACEHOLDER_TITLE.search(d['title']):
        warn.append('Title is still a template placeholder.')
    if not d['applies_to']:
        warn.append('Applies To is empty.')
    if d['reported_status'] in ('', 'Not verified'):
        warn.append('Current Status not verified.')
    if d['role'] == '':
        warn.append('Source Role is empty or not a list value.')
    if d['feed_type'] and d['feed_type'] not in FEED_TYPES:
        warn.append(f'Watcher Feed Type "{d["feed_type"]}" unknown - no feed created.')


@dataclass
class ImportSummary:
    committed: bool
    rows: int = 0
    skipped: int = 0
    flagged: int = 0
    created: dict = field(default_factory=lambda: {'authority': 0, 'framework': 0, 'document': 0, 'version': 0, 'feed': 0})
    updated: dict = field(default_factory=lambda: {'document': 0, 'version': 0})
    locked: list = field(default_factory=list)


def _snapshot(obj, fields):
    return {f: (str(getattr(obj, f)) if getattr(obj, f) is not None else None) for f in fields}


DOC_FIELDS = ('title', 'official_url', 'licensed_content', 'review_cadence_months', 'category', 'source_role',
              'document_type', 'applies_to', 'programme_scope')
VERSION_FIELDS = ('issued_on', 'effective_from', 'checksum_sha256', 'direct_url', 'drive_url', 'file_size_kb', 'pages',
                  'collection_status', 'reported_status', 'collected_on', 'notes', 'register_row_id', 'review_flags')


def apply_register(rows, *, actor, commit=False, reason=''):
    """Write the checked rows. commit=False runs everything, then rolls back."""
    summary = ImportSummary(committed=commit)
    with transaction.atomic():
        for row in rows:
            summary.rows += 1
            if not row.ok:
                summary.skipped += 1
                continue
            if row.warnings:
                summary.flagged += 1
            _apply_row(row, actor, reason, summary)
        if not commit:
            transaction.set_rollback(True)
    return summary


def _audit(actor, action, obj, before, after, reason):
    SourceAuditEvent.objects.create(actor=actor, action=action, object_type=type(obj).__name__,
                                    object_id=str(obj.pk), before=before, after=after, reason=reason)


def _apply_row(row, actor, reason, summary):
    d = row.data
    auth_code = authority_code(d['authority'])
    authority, made = Authority.objects.get_or_create(code=auth_code, defaults={'name': d['authority'], 'kind': d['kind']})
    if made:
        summary.created['authority'] += 1
        _audit(actor, 'import.create', authority, None, {'code': auth_code, 'kind': d['kind']}, reason)
    cat_no = (re.match(r'(\d+)', d['category']) or [None, '0'])[1]
    framework, made = Framework.objects.get_or_create(
        code=f'{auth_code}-C{cat_no}',
        defaults={'authority': authority, 'name': f'{d["authority"]}: {d["category"]}', 'purpose': d['category'] or '-'})
    if made:
        summary.created['framework'] += 1
        _audit(actor, 'import.create', framework, None, {'code': framework.code}, reason)

    doc_values = {'title': d['title'], 'official_url': d['listing_url'], 'licensed_content': d['licensed'],
                  'review_cadence_months': d['cadence'], 'category': d['category'], 'source_role': d['role'],
                  'document_type': d['document_type'], 'applies_to': d['applies_to'],
                  'programme_scope': d['programme_scope']}
    doc = SourceDocument.objects.filter(source_code=d['source_code']).first()
    if doc is None:
        doc = SourceDocument.objects.create(framework=framework, source_code=d['source_code'], **doc_values)
        summary.created['document'] += 1
        _audit(actor, 'import.create', doc, None, _snapshot(doc, DOC_FIELDS), reason)
    else:
        before = _snapshot(doc, DOC_FIELDS)
        for k, v in doc_values.items():
            setattr(doc, k, v)
        if _snapshot(doc, DOC_FIELDS) != before:
            doc.save()
            summary.updated['document'] += 1
            _audit(actor, 'import.update', doc, before, _snapshot(doc, DOC_FIELDS), reason)

    version_values = {'issued_on': d['issued_on'], 'effective_from': d['effective_from'], 'checksum_sha256': d['sha256'],
                      'direct_url': d['direct_url'], 'drive_url': d['drive_url'], 'file_size_kb': d['file_size_kb'],
                      'pages': d['pages'], 'collection_status': d['collection'], 'reported_status': d['reported_status'],
                      'collected_on': d['collected_on'], 'notes': d['notes'], 'register_row_id': row.row_id,
                      'review_flags': list(row.warnings)}
    version = SourceVersion.objects.filter(document=doc, version_label=d['version_label']).first()
    if version is None:
        version = SourceVersion.objects.create(document=doc, version_label=d['version_label'], **version_values)
        summary.created['version'] += 1
        _audit(actor, 'import.create', version, None, _snapshot(version, VERSION_FIELDS), reason)
    elif version.is_locked:
        summary.locked.append(f'{row.row_id} {version} ({version.status}) - left unchanged')
    else:
        before = _snapshot(version, VERSION_FIELDS)
        for k, v in version_values.items():
            setattr(version, k, v)
        if _snapshot(version, VERSION_FIELDS) != before:
            version.save()
            summary.updated['version'] += 1
            _audit(actor, 'import.update', version, before, _snapshot(version, VERSION_FIELDS), reason)

    url = d['listing_url'] or d['direct_url']
    if url and d['feed_type'] in FEED_TYPES:
        feed, made = SourceFeed.objects.get_or_create(document=doc, endpoint_url=url,
                                                      defaults={'feed_type': d['feed_type']})
        if made:
            summary.created['feed'] += 1
            _audit(actor, 'import.create', feed, None, {'endpoint_url': url, 'feed_type': d['feed_type'],
                                                         'enabled': False}, reason)
