"""Stage 3.1: import the benchmark catalogue from the manager's Source Pack (.xlsx).

    data = read_catalogue(path)                      # parse + check, no DB access
    summary = apply_catalogue(data, actor=..., commit=False)   # dry run

Sheets used: "Read Me" (pillars), "Evidence Catalogue" (EV-xx), "Benchmark
Catalogue" (one row per benchmark, with L3/L2/L1 text). Everything lands as
DRAFT. A maturity text shared word-for-word by several benchmarks is the
Source Pack's generic template, so it is stored with is_placeholder=True --
real level definitions are expert work (stage 3.3), not something to infer.
Re-running is safe: drafts update in place; a benchmark past draft is never
touched.
"""
import re
from collections import Counter
from dataclasses import dataclass, field

from django.db import transaction

from ..models import Benchmark, Capability, EvidenceType, MaturityDefinition, Pillar, RequirementKind, SourceAuditEvent

PLACEHOLDER_MIN_SHARED = 3  # same level text on >= 3 benchmarks = template, not a definition
LEVEL_COLUMNS = (('L3', 'L3 Foundational'), ('L2', 'L2 Operating'), ('L1', 'L1 Advanced'))
CATALOGUE_HEADERS = ('Benchmark ID', 'Pillar ID', 'Benchmark Domain', 'Capability / Requirement Scope',
                     'Benchmark Class', 'Applicability', 'Source Basis', 'L3 Foundational', 'L2 Operating',
                     'L1 Advanced', 'Evidence Types')


def _t(v):
    return '' if v is None else str(v).strip()


def benchmark_kind(benchmark_class):
    c = benchmark_class.lower()
    if 'proposed' in c:
        return RequirementKind.PROPOSED
    if 'guidance' in c:
        return RequirementKind.GUIDANCE
    if 'mandatory' in c or 'authoritative' in c:
        return RequirementKind.MANDATORY
    return None


@dataclass
class Catalogue:
    version_label: str
    pillars: list = field(default_factory=list)
    evidence_types: list = field(default_factory=list)
    benchmarks: list = field(default_factory=list)
    errors: list = field(default_factory=list)
    warnings: list = field(default_factory=list)

    @property
    def ok(self):
        return not self.errors


def _rows(wb, name):
    if name not in wb.sheetnames:
        raise ValueError(f'Sheet {name!r} not found.')
    return [r for r in wb[name].iter_rows(values_only=True) if r and any(c not in (None, '') for c in r)]


def _table(rows, first_header):
    """(headers, data rows) of the table whose header row starts with first_header."""
    idx = next((i for i, r in enumerate(rows) if _t(r[0]) == first_header), None)
    if idx is None:
        raise ValueError(f'No table starting with {first_header!r}.')
    headers = [_t(h) for h in rows[idx]]
    return headers, rows[idx + 1:]


def read_catalogue(path, version_label=None):
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        readme, evidence, catalogue = _rows(wb, 'Read Me'), _rows(wb, 'Evidence Catalogue'), _rows(wb, 'Benchmark Catalogue')
    finally:
        wb.close()

    if not version_label:
        m = re.search(r'\bv(\d+(?:\.\d+)*)\b', _t(readme[0][0]) if readme else '')
        version_label = f'v{m.group(1)}' if m else 'v0'
    cat = Catalogue(version_label=version_label)

    headers, data = _table(readme, 'Pillar ID')
    col = {h: i for i, h in enumerate(headers)}
    for r in data:
        code = _t(r[0])
        if not re.fullmatch(r'P\d{2}', code):
            continue
        cat.pillars.append({'code': code, 'name': _t(r[col['Pillar']]),
                            'coverage': _t(r[col.get('Coverage statement', 2)]),
                            'catalogue_status': _t(r[col['Status']]) if 'Status' in col else ''})

    headers, data = _table(evidence, 'Evidence Type ID')
    col = {h: i for i, h in enumerate(headers)}
    for r in data:
        code = _t(r[0])
        if re.fullmatch(r'EV-\d{2}', code):
            cat.evidence_types.append({'code': code, 'name': _t(r[col['Evidence Family']]),
                                       'examples': _t(r[col['Examples']]),
                                       'validation_method': _t(r[col['Minimum Validation']]),
                                       'assessment_use': _t(r[col['Assessment Use']])})

    headers, data = _table(catalogue, 'Benchmark ID')
    missing = [h for h in CATALOGUE_HEADERS if h not in headers]
    if missing:
        raise ValueError(f'Benchmark Catalogue is missing columns: {", ".join(missing)}')
    col = {h: i for i, h in enumerate(headers)}
    pillar_codes = {p['code'] for p in cat.pillars}
    ev_codes = {e['code'] for e in cat.evidence_types}
    seen = set()
    for r in data:
        code = _t(r[col['Benchmark ID']])
        if not code:
            continue
        get = lambda h: _t(r[col[h]]) if h in col else ''  # noqa: E731
        b = {'code': code, 'pillar': get('Pillar ID'), 'title': get('Benchmark Domain'),
             'scope': get('Capability / Requirement Scope'), 'benchmark_class': get('Benchmark Class'),
             'applicability': get('Applicability'), 'source_basis': get('Source Basis'),
             'levels': {lvl: get(h) for lvl, h in LEVEL_COLUMNS},
             'evidence': [e.strip() for e in get('Evidence Types').split(';') if e.strip()],
             'scoring_dimensions': get('Scoring Dimensions'), 'activation_status': get('Activation Status')}
        b['kind'] = benchmark_kind(b['benchmark_class'])
        problems = []
        if code in seen:
            problems.append('duplicate Benchmark ID')
        seen.add(code)
        if b['pillar'] not in pillar_codes:
            problems.append(f'unknown pillar {b["pillar"]!r}')
        if b['kind'] is None:
            problems.append(f'unknown Benchmark Class {b["benchmark_class"]!r}')
        if not b['title']:
            problems.append('no Benchmark Domain')
        unknown_ev = [e for e in b['evidence'] if e not in ev_codes]
        if unknown_ev:
            problems.append(f'unknown evidence types {unknown_ev}')
        if problems:
            cat.errors.append(f'{code}: ' + '; '.join(problems))
            continue
        cat.benchmarks.append(b)

    # Template text shared by many benchmarks is a placeholder, not a definition.
    for lvl, _h in LEVEL_COLUMNS:
        counts = Counter(b['levels'][lvl] for b in cat.benchmarks if b['levels'][lvl])
        for b in cat.benchmarks:
            b.setdefault('placeholder', {})[lvl] = (not b['levels'][lvl]
                                                    or counts[b['levels'][lvl]] >= PLACEHOLDER_MIN_SHARED)
    n_placeholder = sum(all(b['placeholder'].values()) for b in cat.benchmarks)
    if n_placeholder:
        cat.warnings.append(f'{n_placeholder} of {len(cat.benchmarks)} benchmarks have only placeholder '
                            'L3/L2/L1 text (generic template shared across benchmarks) - imported as placeholders.')
    proposed_no_owner = sum(1 for b in cat.benchmarks if b['kind'] == RequirementKind.PROPOSED)
    if proposed_no_owner:
        cat.warnings.append(f'{proposed_no_owner} benchmarks are InGage-proposed and need an expert owner '
                            'before they can be released.')
    return cat


@dataclass
class CatalogueSummary:
    committed: bool
    created: Counter = field(default_factory=Counter)
    updated: Counter = field(default_factory=Counter)
    locked: list = field(default_factory=list)


def _audit(actor, action, obj, after, reason):
    SourceAuditEvent.objects.create(actor=actor, action=action, object_type=type(obj).__name__,
                                    object_id=str(obj.pk), after=after, reason=reason)


def _upsert(model, lookup, values, summary, key, actor, reason):
    obj = model.objects.filter(**lookup).first()
    if obj is None:
        obj = model.objects.create(**lookup, **values)
        summary.created[key] += 1
        _audit(actor, 'catalogue.create', obj, {k: str(v) for k, v in {**lookup, **values}.items()}, reason)
        return obj
    changed = {k: v for k, v in values.items() if getattr(obj, k) != v}
    if changed:
        for k, v in changed.items():
            setattr(obj, k, v)
        obj.save()
        summary.updated[key] += 1
        _audit(actor, 'catalogue.update', obj, {k: str(v) for k, v in changed.items()}, reason)
    return obj


def apply_catalogue(cat, *, actor, commit=False, reason=''):
    summary = CatalogueSummary(committed=commit)
    with transaction.atomic():
        pillars = {p['code']: _upsert(Pillar, {'code': p['code']},
                                      {k: p[k] for k in ('name', 'coverage', 'catalogue_status')},
                                      summary, 'pillar', actor, reason) for p in cat.pillars}
        evidence = {e['code']: _upsert(EvidenceType, {'code': e['code']},
                                       {k: e[k] for k in ('name', 'validation_method', 'examples', 'assessment_use')},
                                       summary, 'evidence_type', actor, reason) for e in cat.evidence_types}
        for b in cat.benchmarks:
            existing = Benchmark.objects.filter(benchmark_code=b['code']).first()
            if existing and existing.status != Benchmark.Status.DRAFT:
                summary.locked.append(f'{b["code"]} ({existing.status}) - left unchanged')
                continue
            pillar = pillars[b['pillar']]
            capability = _upsert(Capability, {'code': f'CAP-{b["code"]}'},
                                 {'name': b['title'], 'domain': pillar.name, 'definition': b['scope'] or b['title']},
                                 summary, 'capability', actor, reason)
            bench = _upsert(Benchmark, {'benchmark_code': b['code']}, {
                'capability': capability, 'pillar': pillar, 'title': b['title'], 'benchmark_kind': b['kind'],
                'benchmark_class': b['benchmark_class'][:120], 'version_label': cat.version_label,
                'applicability': b['applicability'], 'source_basis': b['source_basis'],
                'scoring_dimensions': b['scoring_dimensions'][:200], 'activation_status': b['activation_status'],
            }, summary, 'benchmark', actor, reason)
            wanted = {evidence[c].pk for c in b['evidence']}
            if set(bench.evidence_types.values_list('pk', flat=True)) != wanted:
                bench.evidence_types.set(wanted)
            for lvl, _h in LEVEL_COLUMNS:
                existing = MaturityDefinition.objects.filter(benchmark=bench, level=lvl).first()
                if existing and existing.edited_by_id:
                    summary.locked.append(f'{b["code"]} {lvl} - written by an expert, kept')
                    continue
                _upsert(MaturityDefinition, {'benchmark': bench, 'level': lvl},
                        {'definition': b['levels'][lvl] or '(not defined)', 'is_placeholder': b['placeholder'][lvl]},
                        summary, 'maturity_level', actor, reason)
        if not commit:
            transaction.set_rollback(True)
    return summary
