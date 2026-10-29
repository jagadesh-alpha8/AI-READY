import os
import tempfile

import openpyxl
from django.core.exceptions import ValidationError
from django.test import TestCase
from rest_framework.test import APITestCase

from apps.accounts.models import User
from apps.institutions.models import Institution

from .models import Authority, SourceAuditEvent, SourceDocument, SourceFeed, SourceVersion
from .services.register_import import apply_register, read_register

PASSWORD = 'Str0ng!DevPassw0rd'
SHA = 'a' * 64
HEADERS = ['Row ID', 'Assigned To', 'Category', 'Authority', 'Authority Kind', 'Source Role (suggested)', 'Source Code',
           'Official Title (exact)', 'Document Type', 'Version / Edition', 'Issued On', 'Effective From',
           'Current Status', 'Supersedes / Amends', 'Applies To', 'Programme / Trade Scope', 'Listing Page URL',
           'Direct Document URL', 'Official Domain?', 'Collection Status', 'Local File (folder/name)', 'File Size (KB)',
           'Pages', 'SHA-256 Checksum', 'Licensed Content?', 'Review Cadence (months)', 'Watcher Feed Type',
           'Related Benchmark IDs', 'Collected On', 'Notes']

GOOD = {'Row ID': 'S-001', 'Category': '2. Higher-Education Regulation', 'Authority': 'UGC',
        'Authority Kind': 'regulatory', 'Source Role (suggested)': 'Mandatory', 'Source Code': 'UGC-005',
        'Official Title (exact)': 'UGC Plagiarism Regulations, 2018', 'Document Type': 'Regulation',
        'Version / Edition': '2018', 'Issued On': '23-Jul-2018', 'Current Status': 'Current', 'Applies To': 'HEIs',
        'Listing Page URL': 'https://www.ugc.gov.in/Regulations',
        'Direct Document URL': 'https://www.ugc.gov.in/pdfnews/plagiarism.pdf', 'Collection Status': 'Downloaded',
        'Local File (folder/name)': 'https://drive.google.com/file/d/1AbCdEfGhIjKlMnOp/view', 'File Size (KB)': 300, 'Pages': 8,
        'SHA-256 Checksum': SHA, 'Licensed Content?': 'No', 'Review Cadence (months)': 3,
        'Watcher Feed Type': 'document_index'}


def make_register(rows):
    """A workbook laid out like the intern tracker: title rows, header on row 4."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Source Register'
    ws['A1'] = 'InGage Source Library - Source Register'
    for i, h in enumerate(HEADERS, start=1):
        ws.cell(row=4, column=i, value=h)
    for r, values in enumerate(rows, start=5):
        for i, h in enumerate(HEADERS, start=1):
            ws.cell(row=r, column=i, value=values.get(h))
    fd, path = tempfile.mkstemp(suffix='.xlsx')
    os.close(fd)
    wb.save(path)
    return path


class ReadRegisterTests(TestCase):
    def _rows(self, *rows):
        path = make_register(rows)
        self.addCleanup(os.remove, path)
        return read_register(path)

    def test_clean_row_has_no_errors_or_warnings(self):
        [row] = self._rows(GOOD)
        self.assertEqual((row.errors, row.warnings), ([], []))
        self.assertEqual(row.data['issued_on'].isoformat(), '2018-07-23')

    def test_rows_without_a_code_or_with_a_bad_checksum_are_skipped(self):
        rows = self._rows({**GOOD, 'Source Code': ''}, {**GOOD, 'Row ID': 'S-002', 'SHA-256 Checksum': 'abc'})
        self.assertIn('Source Code is empty.', rows[0].errors)
        self.assertIn('SHA-256 is not 64 hex characters.', rows[1].errors)

    def test_same_code_and_version_twice_is_a_duplicate(self):
        rows = self._rows(GOOD, {**GOOD, 'Row ID': 'S-002'})
        self.assertTrue(rows[0].ok)
        self.assertIn('Duplicate of S-001', rows[1].errors[0])

    def test_unofficial_site_and_drive_folder_are_flagged(self):
        [row] = self._rows({**GOOD, 'Listing Page URL': 'https://iqac.someuniv.edu/naac.pdf',
                            'Local File (folder/name)': 'https://drive.google.com/drive/folders/xyz'})
        self.assertTrue(row.ok)
        self.assertTrue(any('not an official domain' in w for w in row.warnings))
        self.assertIn('Drive link is a folder, not the file.', row.warnings)

    def test_licensed_standard_never_keeps_its_drive_copy(self):
        [row] = self._rows({**GOOD, 'Authority': 'ISO/IEC', 'Authority Kind': 'standards_body',
                            'Licensed Content?': 'Yes', 'Listing Page URL': 'https://www.iso.org/standard/1.html'})
        self.assertEqual(row.data['drive_url'], '')
        self.assertEqual(row.data['collection'], SourceVersion.Collection.LICENSED)
        self.assertTrue(any('Drive copy is NOT imported' in w for w in row.warnings))

    def test_missing_header_columns_are_rejected(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Source Register'
        ws.append(['Row ID', 'Title'])
        fd, path = tempfile.mkstemp(suffix='.xlsx')
        os.close(fd)
        wb.save(path)
        self.addCleanup(os.remove, path)
        with self.assertRaisesMessage(ValueError, 'missing columns'):
            read_register(path)


class ApplyRegisterTests(TestCase):
    def setUp(self):
        self.path = make_register([GOOD, {**GOOD, 'Row ID': 'S-002', 'Source Code': 'NBA-002', 'Authority': 'NBA',
                                          'Authority Kind': 'accreditation', 'Category': '1. Accreditation & Quality',
                                          'Listing Page URL': 'https://www.nbaind.org/Downloads/Documents'}])
        self.addCleanup(os.remove, self.path)

    def test_dry_run_saves_nothing_but_reports_exact_counts(self):
        summary = apply_register(read_register(self.path), actor='t', commit=False)
        self.assertEqual(summary.created['document'], 2)
        self.assertEqual(summary.created['version'], 2)
        self.assertEqual((SourceDocument.objects.count(), SourceAuditEvent.objects.count()), (0, 0))

    def test_commit_creates_drafts_disabled_feeds_and_audit_events(self):
        apply_register(read_register(self.path), actor='intern-import', commit=True)
        version = SourceVersion.objects.get(document__source_code='UGC-005')
        self.assertEqual((version.status, version.version_label, version.checksum_sha256), ('draft', '2018', SHA))
        self.assertEqual(Authority.objects.get(code='UGC').kind, 'regulatory')
        self.assertFalse(SourceFeed.objects.filter(enabled=True).exists())
        self.assertEqual(SourceFeed.objects.count(), 2)
        self.assertTrue(SourceAuditEvent.objects.filter(actor='intern-import', action='import.create').exists())

    def test_rerun_is_idempotent(self):
        apply_register(read_register(self.path), actor='t', commit=True)
        events = SourceAuditEvent.objects.count()
        summary = apply_register(read_register(self.path), actor='t', commit=True)
        self.assertEqual((summary.created['version'], summary.updated['version']), (0, 0))
        self.assertEqual(SourceAuditEvent.objects.count(), events)

    def test_reimport_never_changes_an_approved_version(self):
        apply_register(read_register(self.path), actor='t', commit=True)
        SourceVersion.objects.filter(document__source_code='UGC-005').update(status='active')
        path = make_register([{**GOOD, 'Pages': 99}])
        self.addCleanup(os.remove, path)
        summary = apply_register(read_register(path), actor='t', commit=True)
        self.assertEqual(len(summary.locked), 1)
        self.assertEqual(SourceVersion.objects.get(document__source_code='UGC-005').pages, 8)


class ImmutabilityTests(TestCase):
    def setUp(self):
        path = make_register([GOOD])
        self.addCleanup(os.remove, path)
        apply_register(read_register(path), actor='t', commit=True)
        self.version = SourceVersion.objects.get()

    def test_draft_can_be_edited(self):
        self.version.pages = 9
        self.version.save()

    def test_active_version_content_is_frozen(self):
        self.version.status = 'active'
        self.version.save()
        self.version.checksum_sha256 = 'b' * 64
        with self.assertRaisesMessage(ValidationError, 'cannot change'):
            self.version.save()

    def test_status_moves_only_forward(self):
        self.version.status = 'active'
        self.version.save()
        self.version.status = 'superseded'
        self.version.save()
        self.version.status = 'active'
        with self.assertRaisesMessage(ValidationError, 'cannot move'):
            self.version.save()

    def test_approved_version_cannot_be_deleted(self):
        self.version.status = 'active'
        self.version.save()
        with self.assertRaises(ValidationError):
            self.version.delete()

    def test_audit_log_is_append_only(self):
        event = SourceAuditEvent.objects.first()
        event.reason = 'edited'
        with self.assertRaises(ValidationError):
            event.save()
        with self.assertRaises(ValidationError):
            event.delete()


class SourceLibraryApiTests(APITestCase):
    def setUp(self):
        path = make_register([GOOD])
        self.addCleanup(os.remove, path)
        apply_register(read_register(path), actor='t', commit=True)
        self.institution = Institution.objects.create(name='MKCE', city='Karur', state='TN')

    def _login(self, role):
        user = User.objects.create_user(email=f'{role}@test.edu', username=role, password=PASSWORD,
                                        first_name='T', last_name=role, role=role, institution=self.institution)
        token = self.client.post('/api/v1/auth/login/', {'email': user.email, 'password': PASSWORD}).data['access_token']
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {token}')

    def test_ingage_staff_can_read_the_library(self):
        self._login(User.Role.CONSULTANT)
        listing = self.client.get('/api/v1/benchmarks/sources')
        self.assertEqual(listing.status_code, 200)
        self.assertEqual(listing.data[0]['source_code'], 'UGC-005')
        self.assertEqual(listing.data[0]['latest_version']['version_label'], '2018')
        detail = self.client.get('/api/v1/benchmarks/sources/UGC-005')
        self.assertEqual((detail.status_code, len(detail.data['versions']), len(detail.data['feeds'])), (200, 1, 1))
        summary = self.client.get('/api/v1/benchmarks/sources/summary')
        self.assertEqual((summary.data['documents'], summary.data['by_status']), (1, {'draft': 1}))

    def test_filters(self):
        self._login(User.Role.SUPER_ADMIN)
        self.assertEqual(len(self.client.get('/api/v1/benchmarks/sources?authority=ugc').data), 1)
        self.assertEqual(len(self.client.get('/api/v1/benchmarks/sources?authority=NBA').data), 0)
        self.assertEqual(len(self.client.get('/api/v1/benchmarks/sources?q=plagiarism').data), 1)

    def test_institution_users_cannot_see_the_library(self):
        for role in (User.Role.IQAC_COORDINATOR,):
            self._login(role)
            self.assertEqual(self.client.get('/api/v1/benchmarks/sources').status_code, 403, role)

    def test_anonymous_is_rejected(self):
        self.assertEqual(self.client.get('/api/v1/benchmarks/sources').status_code, 401)


# -- stage 1 step 3: file intake ----------------------------------------------

import hashlib  # noqa: E402
import shutil  # noqa: E402
from unittest.mock import patch  # noqa: E402

from django.test import override_settings  # noqa: E402

#: Benchmark Studio is read-only in the portal (settings.BENCHMARK_CATALOGUE_READ_ONLY, apps/benchmarks/permissions.py).
#: The API tests of the earlier authoring workflows run with the lock off: they cover the services the controlled
#: import and maintenance still use. ReadOnlyCatalogueTests checks the lock itself.
MAINTENANCE_MODE = override_settings(BENCHMARK_CATALOGUE_READ_ONLY=False)

from .models import Requirement, RequirementCandidate, SourceSection  # noqa: E402
from .services import approval, file_intake  # noqa: E402
from .services.requirement_drafting import (  # noqa: E402
    DraftingError, accept_candidate, build_sections, draft_section, reject_candidate,
)


def pdf_bytes(pages=('First page text.', 'Second page text.')):
    import pymupdf
    doc = pymupdf.open()
    for text in pages:
        doc.new_page().insert_text((72, 72), text)
    data = doc.tobytes()
    doc.close()
    return data


class _VersionMixin:
    def setUp(self):
        super().setUp()
        self.media = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.media, ignore_errors=True)
        override = override_settings(MEDIA_ROOT=self.media)
        override.enable()
        self.addCleanup(override.disable)
        self.content = pdf_bytes()
        self.sha = hashlib.sha256(self.content).hexdigest()
        path = make_register([{**GOOD, 'SHA-256 Checksum': self.sha}])
        self.addCleanup(os.remove, path)
        apply_register(read_register(path), actor='t', commit=True)
        self.version = SourceVersion.objects.get()

    def fetch_ok(self, file_id):
        return self.content, 'plagiarism.pdf'


class FileIntakeTests(_VersionMixin, TestCase):
    def test_matching_checksum_is_verified_and_stored(self):
        result = file_intake.ingest(self.version, fetcher=self.fetch_ok)
        self.version.refresh_from_db()
        self.assertEqual((result.status, self.version.intake_status), ('verified', 'verified'))
        self.assertTrue(os.path.exists(os.path.join(self.media, self.version.file_uri)))
        self.assertIn(self.sha, self.version.file_uri)
        self.assertEqual(self.version.pages, 2)  # corrected from the register's 8
        self.assertIn('pages recorded 8, actual 2', self.version.intake_note)

    def test_missing_checksum_is_computed(self):
        SourceVersion.objects.filter(pk=self.version.pk).update(checksum_sha256='')
        self.version.refresh_from_db()
        self.assertEqual(file_intake.ingest(self.version, fetcher=self.fetch_ok).status, 'computed')
        self.version.refresh_from_db()
        self.assertEqual(self.version.checksum_sha256, self.sha)

    def test_mismatch_is_not_stored(self):
        result = file_intake.ingest(self.version, fetcher=lambda fid: (b'other bytes', 'x.pdf'))
        self.version.refresh_from_db()
        self.assertEqual((result.status, self.version.file_uri), ('mismatch', ''))

    def test_folder_link_and_fetch_errors(self):
        SourceVersion.objects.filter(pk=self.version.pk).update(drive_url='https://drive.google.com/drive/folders/xyz')
        self.version.refresh_from_db()
        self.assertEqual(file_intake.ingest(self.version, fetcher=self.fetch_ok).status, 'skipped')
        SourceVersion.objects.filter(pk=self.version.pk).update(drive_url='https://drive.google.com/file/d/abcdefghijkl/view')
        self.version.refresh_from_db()

        def boom(file_id):
            raise RuntimeError('403 denied')
        self.assertEqual(file_intake.ingest(self.version, fetcher=boom).status, 'failed')

    def test_approved_version_is_never_touched(self):
        SourceVersion.objects.filter(pk=self.version.pk).update(status='active')
        self.version.refresh_from_db()
        result = file_intake.ingest(self.version, fetcher=self.fetch_ok)
        self.version.refresh_from_db()
        self.assertEqual((result.status, self.version.file_uri), ('skipped', ''))

    def test_drive_file_id_parsing(self):
        self.assertEqual(file_intake.drive_file_id('https://drive.google.com/file/d/1AbC_dEf-123456/view?usp=x'),
                         '1AbC_dEf-123456')
        self.assertIsNone(file_intake.drive_file_id('https://drive.google.com/drive/folders/1AbCdEf123456'))


# -- approval -------------------------------------------------------------------

class ApprovalTests(_VersionMixin, TestCase):
    def setUp(self):
        super().setUp()
        self.user = User.objects.create_user(email='c@ingage.ai', username='c', password=PASSWORD, role='consultant')

    def test_cannot_approve_before_intake(self):
        with self.assertRaisesMessage(approval.ApprovalError, 'run file intake'):
            approval.approve(self.version, self.user)

    def test_open_flags_need_acknowledgement(self):
        file_intake.ingest(self.version, fetcher=self.fetch_ok)
        self.version.refresh_from_db()
        self.version.review_flags = ['Applies To is empty.']
        self.version.save()
        with self.assertRaises(approval.ApprovalError) as ctx:
            approval.approve(self.version, self.user)
        self.assertEqual(ctx.exception.flags, ['Applies To is empty.'])
        approval.approve(self.version, self.user, acknowledge_flags=True, note='checked')
        self.version.refresh_from_db()
        self.assertEqual((self.version.status, self.version.approved_by), ('active', self.user))
        self.assertTrue(SourceAuditEvent.objects.filter(action='version.approved', actor='c@ingage.ai').exists())

    def test_new_edition_supersedes_the_active_one(self):
        file_intake.ingest(self.version, fetcher=self.fetch_ok)
        self.version.refresh_from_db()
        approval.approve(self.version, self.user)
        new = SourceVersion.objects.create(document=self.version.document, version_label='2024',
                                           collection_status='downloaded', file_uri='x', intake_status='verified')
        approval.approve(new, self.user)
        self.version.refresh_from_db()
        new.refresh_from_db()
        self.assertEqual((self.version.status, new.status, new.supersedes_id), ('superseded', 'active', self.version.pk))

    def test_withdraw_needs_a_reason(self):
        file_intake.ingest(self.version, fetcher=self.fetch_ok)
        self.version.refresh_from_db()
        approval.approve(self.version, self.user)
        with self.assertRaisesMessage(approval.ApprovalError, 'reason'):
            approval.withdraw(self.version, self.user, note=' ')
        approval.withdraw(self.version, self.user, note='Replaced by gazette correction')
        self.assertEqual(SourceVersion.objects.get(pk=self.version.pk).status, 'withdrawn')


# -- stage 2: sections, drafts, review -----------------------------------------

PAGES = [{'page_number': 1, 'text': 'Chapter 1. Every university shall establish an Academic Integrity Panel '
                                    'chaired by the Pro-Vice Chancellor to review plagiarism cases.'},
         {'page_number': 2, 'text': 'Institutions should conduct awareness programmes on academic integrity '
                                    'for students and faculty every year.'}]


class FakeAI:
    model = 'fake-model'

    def __init__(self, items):
        self.items = items

    def extract_structured_data(self, **kwargs):
        return {'requirements': self.items}


class DraftingTests(_VersionMixin, TestCase):
    def setUp(self):
        super().setUp()
        self.user = User.objects.create_user(email='e@ingage.ai', username='e', password=PASSWORD, role='consultant')
        build_sections(self.version, pages=PAGES)
        self.section = SourceSection.objects.get(version=self.version)

    def test_sections_keep_page_ranges(self):
        self.assertEqual((self.section.page_start, self.section.page_end), (1, 2))
        self.assertIn('[Page 2]', self.section.text)

    def test_only_drafts_quoting_the_source_are_kept(self):
        good = {'statement': 'The institution establishes an Academic Integrity Panel.', 'kind': 'mandatory',
                'page': 1, 'source_quote': 'Every university shall establish an Academic Integrity Panel',
                'confidence': 0.9, 'rationale': 'shall'}
        invented = {**good, 'source_quote': 'Every university shall appoint a Chief AI Officer within six months'}
        wrong_page = {**good, 'page': 7}
        kept, dropped = draft_section(self.section, FakeAI([good, invented, wrong_page]))
        self.assertEqual((kept, dropped), (1, 2))
        self.assertEqual(RequirementCandidate.objects.get().kind, 'mandatory')

    def _candidate(self):
        draft_section(self.section, FakeAI([{
            'statement': 'The institution runs yearly integrity awareness programmes.', 'kind': 'recommended',
            'page': 2, 'source_quote': 'Institutions should conduct awareness programmes on academic integrity',
            'confidence': 0.8, 'rationale': 'should'}]))
        return RequirementCandidate.objects.get()

    def test_accept_with_edit_creates_a_traceable_requirement(self):
        cand = self._candidate()
        req = accept_candidate(cand, self.user, statement='The institution runs academic-integrity awareness yearly.')
        self.assertEqual(req.code, 'UGC-005-R001')
        self.assertEqual((req.page, req.version, req.approved_by), (2, self.version, self.user))
        self.assertIn('awareness programmes', req.source_quote)
        cand.refresh_from_db()
        self.assertEqual(cand.status, 'accepted')
        with self.assertRaises(DraftingError):
            accept_candidate(cand, self.user)
        self.assertTrue(SourceAuditEvent.objects.filter(action='requirement.accepted').exists())

    def test_reject(self):
        cand = self._candidate()
        reject_candidate(cand, self.user, note='duplicate')
        cand.refresh_from_db()
        self.assertEqual((cand.status, Requirement.objects.count()), ('rejected', 0))

    def test_sections_are_kept_once_drafts_were_reviewed(self):
        reject_candidate(self._candidate(), self.user)
        with self.assertRaises(DraftingError):
            build_sections(self.version, pages=PAGES)


# -- rate limits and resuming an interrupted drafting run ---------------------------------------

from types import SimpleNamespace  # noqa: E402

from apps.extraction.exceptions import RecoverableExtractionError  # noqa: E402

from .services import requirement_drafting as _drafting  # noqa: E402

PANEL = {'statement': 'The institution establishes an Academic Integrity Panel.', 'kind': 'mandatory', 'page': 1,
         'source_quote': 'Every university shall establish an Academic Integrity Panel', 'confidence': 0.9,
         'rationale': 'shall'}
AWARENESS = {'statement': 'The institution runs yearly integrity awareness programmes.', 'kind': 'recommended',
             'page': 2, 'source_quote': 'Institutions should conduct awareness programmes on academic integrity',
             'confidence': 0.8, 'rationale': 'should'}


class _Http429(Exception):
    status_code = 429

    def __init__(self, message, retry_after=None):
        super().__init__(message)
        self.response = SimpleNamespace(headers={'retry-after': retry_after} if retry_after else {})


def _rate_limited(message="Quota exceeded ... Please retry in 4.549782715s. 'retryDelay': '4s'", retry_after=None):
    """What the OpenAI-compatible client raises on a 429: a recoverable error caused by the SDK's."""
    try:
        raise _Http429(message, retry_after)
    except _Http429 as cause:
        try:
            raise RecoverableExtractionError(f'OpenAI rate limit hit: {cause}') from cause
        except RecoverableExtractionError as exc:
            return exc


class PagedAI:
    """Answers per section (by its [Page N]); `fail` = errors to raise first, per page."""
    model = 'fake-model'

    def __init__(self, fail=None):
        self.fail = {page: list(errors) for page, errors in (fail or {}).items()}
        self.calls = []

    def extract_structured_data(self, *, user_content, **kwargs):
        page = 1 if '[Page 1]' in user_content else 2
        self.calls.append(page)
        if self.fail.get(page):
            raise self.fail[page].pop(0)
        return {'requirements': [PANEL if page == 1 else AWARENESS]}


class DraftingResumeTests(_VersionMixin, TestCase):
    def setUp(self):
        super().setUp()
        self.waits = []
        for p in (patch.object(_drafting, '_sleep', self.waits.append),
                  patch.object(_drafting, 'SECTION_MAX_CHARS', 10)):  # one section per page
            p.start()
            self.addCleanup(p.stop)
        build_sections(self.version, pages=PAGES)
        self.first, self.second = SourceSection.objects.filter(version=self.version).order_by('ordinal')

    def _run(self, ai):
        with patch('apps.extraction.services.ai_service.get_ai_service', return_value=ai):
            return _drafting.draft_version(self.version)

    def test_the_wait_follows_the_providers_hint(self):
        self.assertEqual(_drafting.rate_limit_wait(_rate_limited()), 5.549782715)  # Gemini: "retry in 4.5s" + 1
        self.assertEqual(_drafting.rate_limit_wait(_rate_limited("[{'retryDelay': '7s'}]")), 8)
        self.assertEqual(_drafting.rate_limit_wait(_rate_limited('slow down', retry_after='12')), 13)
        self.assertEqual(_drafting.rate_limit_wait(_rate_limited('no hint at all')), _drafting.RATE_LIMIT_DEFAULT_WAIT + 1)
        self.assertEqual(_drafting.rate_limit_wait(_rate_limited('retry in 3600s')), _drafting.RATE_LIMIT_MAX_WAIT)
        self.assertIsNone(_drafting.rate_limit_wait(RecoverableExtractionError('OpenAI request timed out')))

    def test_a_rate_limited_section_waits_and_is_retried(self):
        ai = PagedAI(fail={2: [_rate_limited(), _rate_limited()]})
        self.assertEqual(self._run(ai), (2, 0))
        self.assertEqual(ai.calls, [1, 2, 2, 2])
        self.assertEqual(len(self.waits), 2)
        self.assertEqual(RequirementCandidate.objects.filter(version=self.version).count(), 2)

    def test_other_errors_are_not_retried(self):
        ai = PagedAI(fail={1: [RecoverableExtractionError('OpenAI request timed out')]})
        with self.assertRaisesMessage(RecoverableExtractionError, 'timed out'):
            self._run(ai)
        self.assertEqual((ai.calls, self.waits), ([1], []))

    def test_a_rerun_keeps_finished_sections_and_drafts_only_the_rest(self):
        ai = PagedAI(fail={2: [_rate_limited()] * (_drafting.RATE_LIMIT_RETRIES + 1)})
        with self.assertRaisesMessage(RecoverableExtractionError, 'rate limit'):
            self._run(ai)  # the limit outlasted every retry
        self.assertEqual(len(self.waits), _drafting.RATE_LIMIT_RETRIES)
        self.first.refresh_from_db()
        self.second.refresh_from_db()
        self.assertIsNotNone(self.first.drafted_at)
        self.assertIsNone(self.second.drafted_at)
        kept = list(RequirementCandidate.objects.filter(section=self.first).values_list('pk', flat=True))
        self.assertEqual(len(kept), 1)

        again = PagedAI()
        self.assertEqual(self._run(again), (1, 0))
        self.assertEqual(again.calls, [2])  # only the section that was left
        self.assertEqual(list(RequirementCandidate.objects.filter(section=self.first).values_list('pk', flat=True)), kept)
        self.assertEqual(RequirementCandidate.objects.filter(version=self.version, status='pending').count(), 2)
        event = SourceAuditEvent.objects.filter(action='requirements.drafted').latest('created_at')
        self.assertEqual((event.after['sections_drafted'], event.after['sections_kept']), (1, 1))

    def test_a_finished_version_drafted_again_starts_over_and_keeps_reviewed_drafts(self):
        self._run(PagedAI())
        user = User.objects.create_user(email='r@ingage.ai', username='r', password=PASSWORD, role='consultant')
        reviewed = RequirementCandidate.objects.get(section=self.first)
        reject_candidate(reviewed, user, note='duplicate')
        old_pending = RequirementCandidate.objects.get(section=self.second).pk
        ai = PagedAI()
        self.assertEqual(self._run(ai), (2, 0))
        self.assertEqual(ai.calls, [1, 2])  # every section again, as before this change
        self.assertTrue(RequirementCandidate.objects.filter(pk=reviewed.pk, status='rejected').exists())
        self.assertFalse(RequirementCandidate.objects.filter(pk=old_pending).exists())


@override_settings(AI_PROCESSING_MODE='single')
class DraftingTaskTests(_VersionMixin, TestCase):
    def test_runs_only_when_queued(self):
        from .tasks import draft_source_requirements
        for i in range(3):  # what the (mocked) run leaves behind: the note counts the version's pending drafts
            RequirementCandidate.objects.create(version=self.version, statement=f'Draft {i}', kind='mandatory',
                                                source_quote='quoted words')
        with patch('apps.benchmarks.services.requirement_drafting.draft_version', return_value=(3, 1)) as run:
            draft_source_requirements.apply(args=[str(self.version.pk)])
            run.assert_not_called()  # not queued -> skipped
            SourceVersion.objects.filter(pk=self.version.pk).update(extraction_status='queued')
            draft_source_requirements.apply(args=[str(self.version.pk)])
        self.version.refresh_from_db()
        self.assertEqual(self.version.extraction_status, 'done')
        self.assertIn('3 draft requirement(s)', self.version.extraction_note)
        self.assertIn('1 AI suggestion(s) dropped', self.version.extraction_note)

    def test_failure_is_recorded(self):
        from .tasks import draft_source_requirements
        SourceVersion.objects.filter(pk=self.version.pk).update(extraction_status='queued')
        with patch('apps.benchmarks.services.requirement_drafting.draft_version', side_effect=DraftingError('no file')):
            draft_source_requirements.apply(args=[str(self.version.pk)])
        self.version.refresh_from_db()
        self.assertEqual((self.version.extraction_status, self.version.extraction_note), ('failed', 'no file'))


@MAINTENANCE_MODE
class ActionApiTests(_VersionMixin, APITestCase):
    def setUp(self):
        super().setUp()
        self.institution = Institution.objects.create(name='MKCE', city='Karur', state='TN')

    def _login(self, role):
        user = User.objects.create_user(email=f'{role}@x.edu', username=role, password=PASSWORD, role=role,
                                        institution=self.institution)
        token = self.client.post('/api/v1/auth/login/', {'email': user.email, 'password': PASSWORD}).data['access_token']
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {token}')

    def test_approve_flow_over_the_api(self):
        self._login('consultant')
        url = f'/api/v1/benchmarks/versions/{self.version.pk}'
        self.assertEqual(self.client.post(f'{url}/approve', {}).status_code, 400)  # no intake yet
        with patch('apps.benchmarks.services.file_intake.fetch_from_drive', side_effect=self.fetch_ok):
            r = self.client.post(f'{url}/intake')
        self.assertEqual(r.data['result'], 'verified')
        SourceVersion.objects.filter(pk=self.version.pk).update(review_flags=['x'])
        blocked = self.client.post(f'{url}/approve', {})
        self.assertEqual((blocked.status_code, blocked.data['flags']), (400, ['x']))
        ok = self.client.post(f'{url}/approve', {'acknowledge_flags': True, 'note': 'ok'}, format='json')
        self.assertEqual((ok.status_code, ok.data['status'], ok.data['approved_by']), (200, 'active', 'consultant@x.edu'))

    def test_drafting_and_review_over_the_api(self):
        self._login('super_admin')
        url = f'/api/v1/benchmarks/versions/{self.version.pk}'
        self.assertEqual(self.client.post(f'{url}/draft-requirements').status_code, 400)  # no stored file
        SourceVersion.objects.filter(pk=self.version.pk).update(file_uri='source_library/x.pdf')
        self.assertEqual(self.client.post(f'{url}/draft-requirements').status_code, 400)  # not approved yet
        SourceVersion.objects.filter(pk=self.version.pk).update(status='active')
        with patch('apps.benchmarks.tasks.draft_source_requirements.delay') as delay:
            self.assertEqual(self.client.post(f'{url}/draft-requirements').status_code, 202)
            self.assertEqual(self.client.post(f'{url}/draft-requirements').status_code, 409)  # already queued
            delay.assert_called_once()
        build_sections(self.version, pages=PAGES) if not self.version.sections.exists() else None
        section = SourceSection.objects.get(version=self.version)
        cand = RequirementCandidate.objects.create(version=self.version, section=section, statement='S', kind='mandatory',
                                                   page=1, source_quote='q', confidence=0.5)
        listed = self.client.get(f'{url}/candidates?status=pending')
        self.assertEqual(len(listed.data), 1)
        r = self.client.post(f'/api/v1/benchmarks/candidates/{cand.pk}/review',
                             {'decision': 'accept', 'statement': 'The institution does S.'}, format='json')
        self.assertEqual((r.status_code, r.data['status'], r.data['requirement_code']), (200, 'accepted', 'UGC-005-R001'))
        reqs = self.client.get('/api/v1/benchmarks/requirements?source_code=UGC-005')
        self.assertEqual(reqs.data[0]['statement'], 'The institution does S.')

    def test_actions_are_staff_only(self):
        self._login('iqac_coordinator')
        url = f'/api/v1/benchmarks/versions/{self.version.pk}'
        for path in ('intake', 'approve', 'draft-requirements'):
            self.assertEqual(self.client.post(f'{url}/{path}').status_code, 403, path)


class IncompleteDownloadTests(_VersionMixin, TestCase):
    def test_unfinished_browser_download_is_refused(self):
        result = file_intake.ingest(self.version, fetcher=lambda fid: (b'partial', 'Unconfirmed 807153.crdownload'))
        self.version.refresh_from_db()
        self.assertEqual((result.status, self.version.file_uri), ('failed', ''))
        self.assertIn('incomplete', result.note)


# -- stage 3.1: catalogue import ----------------------------------------------------

from .models import Benchmark, MaturityDefinition  # noqa: E402
from .services.catalogue_import import apply_catalogue, benchmark_kind, read_catalogue  # noqa: E402

TEMPLATE = {'L3': 'Policy/owner/process and baseline evidence exist.', 'L2': 'Operates across scope.',
            'L1': 'Institution-wide and measured.'}


def make_source_pack(benchmarks, title='InGage Benchmark Source Pack v0.2 | Expanded'):
    wb = openpyxl.Workbook()
    rm = wb.active
    rm.title = 'Read Me'
    rm.append([title])
    rm.append(['Purpose', 'x'])
    rm.append(['Pillar ID', 'Pillar', 'Coverage statement', 'Benchmark domains', 'Mandatory', 'Guidance',
               'Proposed', 'Status'])
    rm.append(['P01', 'Governance', 'Leadership', 2, 1, 0, 1, 'Catalogue built'])
    rm.append(['P08', 'AI transformation', 'AI', 1, 0, 0, 1, 'Catalogue built'])
    ev = wb.create_sheet('Evidence Catalogue')
    ev.append(['Evidence Type ID', 'Evidence Family', 'Examples', 'Minimum Validation', 'Assessment Use'])
    ev.append(['EV-01', 'Authority and policy', 'Policies', 'Authenticity and date', 'Mapped'])
    ev.append(['EV-10', 'AI system and model', 'Model cards', 'Reproducibility', 'Mapped'])
    bc = wb.create_sheet('Benchmark Catalogue')
    bc.append(['Benchmark ID', 'Pillar ID', 'Benchmark Domain', 'Capability / Requirement Scope', 'Benchmark Class',
               'Applicability', 'Source Basis', 'L3 Foundational', 'L2 Operating', 'L1 Advanced', 'Evidence Types',
               'Scoring Dimensions', 'Activation Status'])
    for b in benchmarks:
        bc.append(b)
    fd, path = tempfile.mkstemp(suffix='.xlsx')
    os.close(fd)
    wb.save(path)
    return path


MANDATORY_CLASS = 'Mandatory/authoritative source mapping plus maturity benchmark'
PROPOSED_CLASS = 'Proposed InGage maturity benchmark'


def bench_row(code, pillar='P01', cls=MANDATORY_CLASS, levels=None, evidence='EV-01', title=None):
    levels = levels or TEMPLATE
    return [code, pillar, title or f'Domain {code}', f'Scope {code}', cls, 'University; ITI', 'NAAC C6',
            levels['L3'], levels['L2'], levels['L1'], evidence, 'Compliance; maturity', 'Pending extraction']


class CatalogueImportTests(TestCase):
    def _cat(self, rows, **kw):
        path = make_source_pack(rows, **kw)
        self.addCleanup(os.remove, path)
        return read_catalogue(path)

    def test_kinds_from_benchmark_class(self):
        self.assertEqual(benchmark_kind(MANDATORY_CLASS), 'mandatory')
        self.assertEqual(benchmark_kind(PROPOSED_CLASS), 'proposed')
        self.assertEqual(benchmark_kind('Guidance-backed maturity benchmark'), 'guidance')
        self.assertIsNone(benchmark_kind('Something else'))

    def test_import_creates_the_catalogue_as_drafts(self):
        cat = self._cat([bench_row('GOV-01'), bench_row('GOV-02', cls=PROPOSED_CLASS),
                         bench_row('AIX-01', pillar='P08', evidence='EV-01; EV-10')])
        self.assertEqual((cat.version_label, len(cat.pillars), len(cat.evidence_types), len(cat.benchmarks)),
                         ('v0.2', 2, 2, 3))
        summary = apply_catalogue(cat, actor='t', commit=True)
        self.assertEqual(summary.created['benchmark'], 3)
        aix = Benchmark.objects.get(benchmark_code='AIX-01')
        self.assertEqual((aix.pillar.code, aix.status, aix.version_label), ('P08', 'draft', 'v0.2'))
        self.assertEqual(sorted(aix.evidence_types.values_list('code', flat=True)), ['EV-01', 'EV-10'])
        self.assertEqual(Benchmark.objects.get(benchmark_code='GOV-02').benchmark_kind, 'proposed')
        self.assertEqual(MaturityDefinition.objects.count(), 9)

    def test_shared_template_levels_are_marked_placeholders(self):
        real = {'L3': 'A named AI policy exists.', 'L2': 'Policy reviewed yearly.', 'L1': 'Outcomes audited.'}
        cat = self._cat([bench_row('GOV-01'), bench_row('GOV-02'), bench_row('GOV-03'),
                         bench_row('AIX-01', pillar='P08', levels=real)])
        apply_catalogue(cat, actor='t', commit=True)
        self.assertEqual(MaturityDefinition.objects.filter(benchmark__benchmark_code='GOV-01',
                                                           is_placeholder=True).count(), 3)
        self.assertFalse(MaturityDefinition.objects.filter(benchmark__benchmark_code='AIX-01',
                                                           is_placeholder=True).exists())
        self.assertTrue(any('placeholder' in w for w in cat.warnings))

    def test_bad_rows_are_skipped_with_a_reason(self):
        cat = self._cat([bench_row('GOV-01', pillar='P99'), bench_row('GOV-02', cls='??'),
                         bench_row('GOV-03', evidence='EV-77'), bench_row('GOV-04'), bench_row('GOV-04')])
        self.assertEqual([b['code'] for b in cat.benchmarks], ['GOV-04'])
        self.assertEqual(len(cat.errors), 4)

    def test_dry_run_and_idempotent_rerun_and_locked_benchmarks(self):
        cat = self._cat([bench_row('GOV-01')])
        apply_catalogue(cat, actor='t', commit=False)
        self.assertEqual(Benchmark.objects.count(), 0)
        apply_catalogue(cat, actor='t', commit=True)
        again = apply_catalogue(cat, actor='t', commit=True)
        self.assertEqual((dict(again.created), dict(again.updated)), ({}, {}))
        Benchmark.objects.update(status='active')
        locked = apply_catalogue(self._cat([bench_row('GOV-01', title='Renamed')]), actor='t', commit=True)
        self.assertEqual(len(locked.locked), 1)
        self.assertEqual(Benchmark.objects.get().title, 'Domain GOV-01')


class CatalogueApiTests(APITestCase):
    def setUp(self):
        path = make_source_pack([bench_row('GOV-01'), bench_row('GOV-02', cls=PROPOSED_CLASS),
                                 bench_row('AIX-01', pillar='P08')])
        self.addCleanup(os.remove, path)
        apply_catalogue(read_catalogue(path), actor='t', commit=True)
        self.institution = Institution.objects.create(name='MKCE', city='Karur', state='TN')

    def _login(self, role):
        user = User.objects.create_user(email=f'{role}@x.edu', username=role, password=PASSWORD, role=role,
                                        institution=self.institution)
        token = self.client.post('/api/v1/auth/login/', {'email': user.email, 'password': PASSWORD}).data['access_token']
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {token}')

    def test_catalogue_endpoints(self):
        self._login('consultant')
        listing = self.client.get('/api/v1/benchmarks/catalogue')
        self.assertEqual([b['benchmark_code'] for b in listing.data], ['AIX-01', 'GOV-01', 'GOV-02'])
        gov = listing.data[1]
        self.assertEqual(([lv['level'] for lv in gov['maturity_levels']], gov['levels_defined']),
                         (['L3', 'L2', 'L1'], False))
        self.assertEqual(len(self.client.get('/api/v1/benchmarks/catalogue?kind=proposed').data), 1)
        self.assertEqual(len(self.client.get('/api/v1/benchmarks/catalogue?pillar=P08').data), 1)
        self.assertEqual(self.client.get('/api/v1/benchmarks/catalogue/GOV-02').data['kind_label'], 'Proposed')
        pillars = self.client.get('/api/v1/benchmarks/pillars').data
        self.assertEqual((pillars[0]['code'], pillars[0]['benchmark_count'], pillars[0]['proposed_count']),
                         ('P01', 2, 1))
        self.assertEqual(len(self.client.get('/api/v1/benchmarks/evidence-types').data), 12)  # seeded EV-01..12 (migration 0014)

    def test_catalogue_is_staff_only(self):
        self._login('institution_admin')
        for url in ('/api/v1/benchmarks/catalogue', '/api/v1/benchmarks/pillars', '/api/v1/benchmarks/evidence-types'):
            self.assertEqual(self.client.get(url).status_code, 403, url)


# -- stage 3.2: linking requirements to benchmarks ----------------------------------

from django.utils import timezone as _tz  # noqa: E402

from .models import BenchmarkRequirement  # noqa: E402
from .services import linking  # noqa: E402


class _LinkFixture:
    def setUp(self):
        super().setUp()
        reg = make_register([GOOD])  # UGC-005, UGC Plagiarism Regulations, 2018
        pack = make_source_pack([
            bench_row('ACA-05', title='Assessment, examination and academic integrity'),
            bench_row('GOV-01', title='Strategic direction and institutional development'),
            bench_row('AI-04', pillar='P08', title='AI security, safety and adversarial resilience'),
        ])
        for p in (reg, pack):
            self.addCleanup(os.remove, p)
        apply_register(read_register(reg), actor='t', commit=True)
        apply_catalogue(read_catalogue(pack), actor='t', commit=True)
        Benchmark.objects.filter(benchmark_code='ACA-05').update(source_basis='NAAC C2; UGC plagiarism; NBA/DGT')
        Benchmark.objects.filter(benchmark_code='GOV-01').update(source_basis='UGC Institutional Development Plans; NAAC C6')
        Benchmark.objects.filter(benchmark_code='AI-04').update(source_basis='NIST AI RMF/GenAI Profile; CERT-In')
        self.user = User.objects.create_user(email='x@ingage.ai', username='x', password=PASSWORD, role='consultant')
        self.version = SourceVersion.objects.get()
        self.req = Requirement.objects.create(
            code='UGC-005-R001', version=self.version, statement='The institution checks every thesis for '
            'plagiarism with an approved similarity tool before evaluation.', kind='mandatory', page=3,
            source_quote='every thesis shall be checked for plagiarism using a similarity tool',
            approved_by=self.user, approved_at=_tz.now())


class LinkingTests(_LinkFixture, TestCase):
    def test_suggestions_rank_the_matching_benchmark_first_and_explain_why(self):
        suggestions = linking.suggest_benchmarks(self.req)
        self.assertEqual(suggestions[0]['benchmark'].benchmark_code, 'ACA-05')
        reasons = ' '.join(suggestions[0]['reasons'])
        self.assertIn('UGC', reasons)
        self.assertIn('plagiarism', reasons)
        self.assertNotIn('AI-04', [s['benchmark'].benchmark_code for s in suggestions])

    def test_link_update_unlink_are_audited(self):
        aca = Benchmark.objects.get(benchmark_code='ACA-05')
        linking.link(aca, self.req, self.user, relationship='primary')
        linking.link(aca, self.req, self.user, relationship='supporting', note='partly')
        self.assertEqual(BenchmarkRequirement.objects.get().relationship, 'supporting')
        self.assertNotIn('ACA-05', [s['benchmark'].benchmark_code for s in linking.suggest_benchmarks(self.req)])
        linking.unlink(aca, self.req, self.user)
        self.assertFalse(BenchmarkRequirement.objects.exists())
        actions = list(SourceAuditEvent.objects.filter(action__startswith='benchmark.').values_list('action', flat=True))
        self.assertEqual(sorted(actions), ['benchmark.link_changed', 'benchmark.linked', 'benchmark.unlinked'])

    def test_released_benchmark_and_withdrawn_source_are_refused(self):
        aca = Benchmark.objects.get(benchmark_code='ACA-05')
        with self.assertRaises(linking.LinkError):
            linking.link(aca, self.req, self.user, relationship='primary-ish')
        SourceVersion.objects.filter(pk=self.version.pk).update(status='withdrawn')
        self.req.refresh_from_db()
        with self.assertRaisesMessage(linking.LinkError, 'withdrawn'):
            linking.link(aca, self.req, self.user, relationship='primary')
        SourceVersion.objects.filter(pk=self.version.pk).update(status='active')
        self.req.refresh_from_db()
        aca.status = 'active'
        aca.save()
        with self.assertRaisesMessage(linking.LinkError, 'new draft'):
            linking.link(aca, self.req, self.user, relationship='primary')


@MAINTENANCE_MODE
class LinkingApiTests(_LinkFixture, APITestCase):
    def _login(self, role):
        user = User.objects.create_user(email=f'{role}@y.edu', username=f'{role}y', password=PASSWORD, role=role)
        token = self.client.post('/api/v1/auth/login/', {'email': user.email, 'password': PASSWORD}).data['access_token']
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {token}')

    def test_link_flow(self):
        self._login('consultant')
        sug = self.client.get('/api/v1/benchmarks/requirements/UGC-005-R001/suggestions')
        self.assertEqual(sug.data[0]['benchmark_code'], 'ACA-05')
        self.assertEqual(len(self.client.get('/api/v1/benchmarks/requirements?unlinked=1').data), 1)
        r = self.client.post('/api/v1/benchmarks/catalogue/ACA-05/requirements',
                             {'requirement_code': 'UGC-005-R001', 'relationship': 'primary'}, format='json')
        self.assertEqual((r.status_code, r.data[0]['requirement_code'], r.data[0]['linked_by']),
                         (201, 'UGC-005-R001', 'consultant@y.edu'))
        self.assertEqual(self.client.get('/api/v1/benchmarks/catalogue/ACA-05').data['requirement_count'], 1)
        reqs = self.client.get('/api/v1/benchmarks/requirements').data
        self.assertEqual(reqs[0]['benchmarks'], [{'benchmark_code': 'ACA-05', 'relationship': 'primary',
                                                  'title': 'Assessment, examination and academic integrity'}])
        self.assertEqual(len(self.client.get('/api/v1/benchmarks/requirements?unlinked=1').data), 0)
        bad = self.client.post('/api/v1/benchmarks/catalogue/ACA-05/requirements',
                               {'requirement_code': 'UGC-005-R001', 'relationship': 'nope'}, format='json')
        self.assertEqual(bad.status_code, 400)
        self.assertEqual(self.client.delete('/api/v1/benchmarks/catalogue/ACA-05/requirements/UGC-005-R001').status_code, 204)
        self.assertEqual(self.client.delete('/api/v1/benchmarks/catalogue/ACA-05/requirements/UGC-005-R001').status_code, 400)

    def test_staff_only(self):
        self._login('hod')
        self.assertEqual(self.client.get('/api/v1/benchmarks/catalogue/ACA-05/requirements').status_code, 403)
        self.assertEqual(self.client.get('/api/v1/benchmarks/requirements/UGC-005-R001/suggestions').status_code, 403)


class TitleAcronymTests(TestCase):
    def test_acronyms(self):
        self.assertEqual(linking.title_acronym('The Digital Personal Data Protection Act, 2023 (No. 22 of 2023)'), 'dpdp')
        self.assertEqual(linking.title_acronym('National Credit Framework (NCrF)'), 'ncf')
        self.assertEqual(linking.title_acronym('Policy'), '')


# -- stage 3.3: writing L3/L2/L1 -------------------------------------------------------

from .services import authoring  # noqa: E402

REAL_L3 = ('An academic integrity policy and a plagiarism-detection tool exist and are used for PhD theses; '
           'a named officer owns them.')
REAL_L2 = ('The integrity policy covers all programmes and assessments; panels decide cases on time and records '
           'are kept for every case.')
REAL_L1 = ('Integrity data is reviewed every year, trends drive training and assessment redesign, and outcomes '
           'are reported publicly.')


class _CataloguePack:
    def setUp(self):
        super().setUp()
        pack = make_source_pack([bench_row('ACA-05', title='Assessment and integrity'), bench_row('GOV-01'),
                                 bench_row('GOV-02', cls=PROPOSED_CLASS)])
        self.addCleanup(os.remove, pack)
        self.pack = pack
        apply_catalogue(read_catalogue(pack), actor='t', commit=True)
        self.user = User.objects.create_user(email='expert@ingage.ai', username='expert', password=PASSWORD,
                                             role='consultant')
        self.aca = Benchmark.objects.get(benchmark_code='ACA-05')


class AuthoringTests(_CataloguePack, TestCase):
    def test_saving_a_real_level_clears_the_placeholder_and_is_audited(self):
        md = authoring.save_level(self.aca, 'L3', self.user, {
            'definition': REAL_L3, 'non_examples': ['A policy document nobody has adopted', ' '],
            'minimum_evidence': ['Approved policy', 'Tool licence'], 'requires_manual_review': True})
        self.assertEqual((md.is_placeholder, md.edited_by, md.non_examples), (False, self.user,
                                                                             ['A policy document nobody has adopted']))
        event = SourceAuditEvent.objects.get(action='benchmark.level_saved')
        self.assertEqual((event.before['is_placeholder'], event.after['definition']), (True, REAL_L3))

    def test_template_short_and_duplicate_text_are_refused(self):
        with self.assertRaisesMessage(authoring.AuthoringError, 'template'):
            authoring.save_level(self.aca, 'L3', self.user, {'definition': TEMPLATE['L3'] + ' ' * 3 + ''})
        with self.assertRaisesMessage(authoring.AuthoringError, 'at least'):
            authoring.save_level(self.aca, 'L3', self.user, {'definition': 'Too short.'})
        authoring.save_level(self.aca, 'L3', self.user, {'definition': REAL_L3})
        with self.assertRaisesMessage(authoring.AuthoringError, 'own definition'):
            authoring.save_level(self.aca, 'L2', self.user, {'definition': REAL_L3})
        with self.assertRaisesMessage(authoring.AuthoringError, 'list'):
            authoring.save_level(self.aca, 'L2', self.user, {'definition': REAL_L2, 'non_examples': 'not a list'})

    def test_only_drafts_can_be_edited(self):
        self.aca.status = 'active'
        self.aca.save()
        with self.assertRaisesMessage(authoring.AuthoringError, 'new draft'):
            authoring.save_level(self.aca, 'L3', self.user, {'definition': REAL_L3})
        with self.assertRaises(authoring.AuthoringError):
            authoring.update_benchmark(self.aca, self.user, {'expert_owner': 'x'})

    def test_readiness_checklist(self):
        r = authoring.readiness(self.aca)
        checks = {c['key']: c for c in r['checks']}
        self.assertFalse(r['ready_for_review'])
        self.assertEqual(checks['levels']['detail'], '0/3 defined (others are placeholders)')
        self.assertEqual((checks['scoring']['severity'], checks['scoring']['ok']), ('required', False))
        self.assertIn('no approved weights', checks['scoring']['detail'])
        self.assertEqual((checks['confirmed']['severity'], checks['confirmed']['ok']), ('required', False))
        for lvl, text in (('L3', REAL_L3), ('L2', REAL_L2), ('L1', REAL_L1)):
            authoring.save_level(self.aca, lvl, self.user, {'definition': text})
        checks = {c['key']: c for c in authoring.readiness(self.aca)['checks']}
        self.assertTrue(checks['levels']['ok'])
        self.assertFalse(checks['requirements']['ok'])  # mandatory benchmark: a primary requirement is required
        gov02 = Benchmark.objects.get(benchmark_code='GOV-02')
        owner = {c['key']: c for c in authoring.readiness(gov02)['checks']}['owner']
        self.assertEqual((owner['severity'], owner['ok']), ('required', False))  # proposed -> owner required
        authoring.update_benchmark(gov02, self.user, {'expert_owner': 'Dr. A. Expert'})
        self.assertTrue({c['key']: c for c in authoring.readiness(gov02)['checks']}['owner']['ok'])

    def test_reimporting_the_catalogue_keeps_expert_levels(self):
        authoring.save_level(self.aca, 'L3', self.user, {'definition': REAL_L3})
        summary = apply_catalogue(read_catalogue(self.pack), actor='t', commit=True)
        self.assertIn('ACA-05 L3 - written by an expert, kept', summary.locked)
        self.assertEqual(MaturityDefinition.objects.get(benchmark=self.aca, level='L3').definition, REAL_L3)


@MAINTENANCE_MODE
class AuthoringApiTests(_CataloguePack, APITestCase):
    def setUp(self):
        super().setUp()
        token = self.client.post('/api/v1/auth/login/', {'email': 'expert@ingage.ai', 'password': PASSWORD}).data['access_token']
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {token}')

    def test_edit_flow(self):
        base = '/api/v1/benchmarks/catalogue/ACA-05'
        bad = self.client.put(f'{base}/levels/L3', {'definition': 'short'}, format='json')
        self.assertEqual(bad.status_code, 400)
        ok = self.client.put(f'{base}/levels/l3', {'definition': REAL_L3, 'non_examples': ['Unadopted draft policy']},
                             format='json')
        self.assertEqual((ok.status_code, ok.data['is_placeholder'], ok.data['edited_by']),
                         (200, False, 'expert@ingage.ai'))
        p = self.client.patch(base, {'expert_owner': 'Dr. A. Expert'}, format='json')
        self.assertEqual(p.data['expert_owner'], 'Dr. A. Expert')
        readiness = self.client.get(f'{base}/readiness').data
        self.assertEqual({c['key']: c['ok'] for c in readiness['checks']}['owner'], True)
        history = self.client.get(f'{base}/history').data
        self.assertEqual([h['action'] for h in history[:2]], ['benchmark.updated', 'benchmark.level_saved'])
        detail = self.client.get(base).data
        l3 = [lv for lv in detail['maturity_levels'] if lv['level'] == 'L3'][0]
        self.assertEqual(l3['non_examples'], ['Unadopted draft policy'])

    def test_staff_only(self):
        hod = User.objects.create_user(email='hod@x.edu', username='hodx', password=PASSWORD, role='iqac_coordinator')
        token = self.client.post('/api/v1/auth/login/', {'email': hod.email, 'password': PASSWORD}).data['access_token']
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {token}')
        self.assertEqual(self.client.put('/api/v1/benchmarks/catalogue/ACA-05/levels/L3', {}, format='json').status_code, 403)


# -- stage 3.4: evidence rules ---------------------------------------------------------------

from .models import EvidenceRule  # noqa: E402
from .services import evidence_rules  # noqa: E402

RULE = {'evidence_type': 'EV-01', 'title': 'Approved academic integrity policy', 'mandatory': True,
        'minimum_count': 1, 'minimum_currency_months': 36, 'validation_method': 'manual_review',
        'privacy_classification': 'internal'}


class EvidenceRuleTests(_CataloguePack, TestCase):
    def test_benchmark_rule_with_and_without_level(self):
        r1 = evidence_rules.create_rule(self.user, {**RULE, 'level': 'L3'}, benchmark=self.aca)
        r2 = evidence_rules.create_rule(self.user, {**RULE, 'title': 'Similarity reports for theses', 'level': ''},
                                        benchmark=self.aca)
        self.assertEqual((r1.level, r2.level, r1.minimum_currency_months), ('L3', '', 36))
        self.assertEqual(evidence_rules.level_coverage(self.aca), {'L3': 2, 'L2': 1, 'L1': 1})
        self.assertEqual(SourceAuditEvent.objects.filter(action='evidence_rule.created').count(), 2)

    def test_new_evidence_type_is_added_to_the_benchmark(self):
        self.assertFalse(self.aca.evidence_types.filter(code='EV-10').exists())
        evidence_rules.create_rule(self.user, {**RULE, 'evidence_type': 'ev-10', 'title': 'Model card for any AI grader'},
                                   benchmark=self.aca)
        self.assertTrue(self.aca.evidence_types.filter(code='EV-10').exists())
        self.assertTrue(SourceAuditEvent.objects.filter(action='benchmark.evidence_type_added').exists())

    def test_validation(self):
        for bad, msg in ((dict(RULE, evidence_type='EV-99'), 'EV code'), (dict(RULE, title='x'), 'title'),
                         (dict(RULE, minimum_count=0), 'between 1 and 100'),
                         (dict(RULE, minimum_currency_months=500), 'between 1 and 120'),
                         (dict(RULE, validation_method='magic'), 'validation_method'),
                         (dict(RULE, level='L9'), 'level')):
            with self.assertRaisesMessage(evidence_rules.EvidenceRuleError, msg):
                evidence_rules.create_rule(self.user, bad, benchmark=self.aca)

    def test_update_delete_and_released_benchmark(self):
        rule = evidence_rules.create_rule(self.user, RULE, benchmark=self.aca)
        evidence_rules.update_rule(rule, self.user, {'minimum_count': 3, 'minimum_currency_months': ''})
        rule.refresh_from_db()
        self.assertEqual((rule.minimum_count, rule.minimum_currency_months), (3, None))
        self.aca.status = 'active'
        self.aca.save()
        with self.assertRaisesMessage(evidence_rules.EvidenceRuleError, 'new draft'):
            evidence_rules.update_rule(rule, self.user, {'minimum_count': 2})
        with self.assertRaises(evidence_rules.EvidenceRuleError):
            evidence_rules.delete_rule(rule, self.user)
        self.aca.status = 'draft'
        self.aca.save()
        rule = EvidenceRule.objects.get(pk=rule.pk)  # fresh copy: its cached benchmark was the 'active' one
        evidence_rules.delete_rule(rule, self.user)
        self.assertFalse(EvidenceRule.objects.exists())
        self.assertTrue(SourceAuditEvent.objects.filter(action='evidence_rule.deleted').exists())

    def test_readiness_counts_evidence_rules(self):
        checks = {c['key']: c for c in authoring.readiness(self.aca)['checks']}
        self.assertEqual((checks['evidence_rules']['severity'], checks['evidence_rules']['ok']), ('required', False))
        evidence_rules.create_rule(self.user, RULE, benchmark=self.aca)  # any level -> covers all three
        checks = {c['key']: c for c in authoring.readiness(self.aca)['checks']}
        self.assertTrue(checks['evidence_rules']['ok'])


class RequirementEvidenceRuleTests(_LinkFixture, TestCase):
    def test_requirement_rules_follow_the_schema(self):
        rule = evidence_rules.create_rule(self.user, {**RULE, 'evidence_type': 'EV-01'}, requirement=self.req)
        self.assertEqual((rule.requirement, rule.benchmark, rule.level), (self.req, None, ''))
        with self.assertRaisesMessage(evidence_rules.EvidenceRuleError, 'already has a rule'):
            evidence_rules.create_rule(self.user, RULE, requirement=self.req)
        with self.assertRaisesMessage(evidence_rules.EvidenceRuleError, 'leave level empty'):
            evidence_rules.create_rule(self.user, {**RULE, 'evidence_type': 'EV-10', 'level': 'L3'}, requirement=self.req)

    def test_requirement_in_a_released_benchmark_is_locked(self):
        aca = Benchmark.objects.get(benchmark_code='ACA-05')
        linking.link(aca, self.req, self.user, relationship='primary')
        aca.status = 'active'
        aca.save()
        with self.assertRaisesMessage(evidence_rules.EvidenceRuleError, 'released'):
            evidence_rules.create_rule(self.user, RULE, requirement=self.req)


@MAINTENANCE_MODE
class EvidenceRuleApiTests(_CataloguePack, APITestCase):
    def setUp(self):
        super().setUp()
        token = self.client.post('/api/v1/auth/login/', {'email': 'expert@ingage.ai', 'password': PASSWORD}).data['access_token']
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {token}')

    def test_flow(self):
        url = '/api/v1/benchmarks/catalogue/ACA-05/evidence-rules'
        created = self.client.post(url, {**RULE, 'level': 'L2'}, format='json')
        self.assertEqual((created.status_code, created.data['evidence_type'], created.data['level']), (201, 'EV-01', 'L2'))
        self.assertEqual(self.client.post(url, {**RULE, 'title': 'x'}, format='json').status_code, 400)
        listing = self.client.get(url).data
        self.assertEqual((len(listing['benchmark_rules']), listing['level_coverage']), (1, {'L3': 0, 'L2': 1, 'L1': 0}))
        rid = created.data['id']
        patched = self.client.patch(f'/api/v1/benchmarks/evidence-rules/{rid}', {'mandatory': False}, format='json')
        self.assertFalse(patched.data['mandatory'])
        self.assertEqual(self.client.delete(f'/api/v1/benchmarks/evidence-rules/{rid}').status_code, 204)


# -- stage 3.5: applicability -----------------------------------------------------------------

from .models import ApplicabilityRule  # noqa: E402
from .services import applicability  # noqa: E402
from .services.applicability import Profile  # noqa: E402

ENGINEERING = Profile('engineering', 'autonomous', {'ug_engineering', 'mba_pgdm'}, {'AICTE', 'NBA', 'NAAC'})
POLYTECHNIC = Profile('polytechnic', 'standalone', {'diploma_engineering'}, {'AICTE'})
ARTS = Profile('arts_science', 'affiliated', {'ug_arts_science'}, {'UGC', 'NAAC'})
ITI = Profile('iti', 'standalone', {'iti_trade'}, {'DGT', 'NCVET'})


class ApplicabilityTests(_CataloguePack, TestCase):
    def setUp(self):
        super().setUp()
        self.gov01 = Benchmark.objects.get(benchmark_code='GOV-01')
        self.gov02 = Benchmark.objects.get(benchmark_code='GOV-02')

    def test_no_rules_means_review_not_applies(self):
        d = applicability.evaluate([], ENGINEERING)
        self.assertEqual(d.result, 'requires_review')
        self.assertIn('No applicability rules', d.reason)

    def test_first_matching_rule_decides_with_a_reason(self):
        # ACA-05: accredited-programme rule first, then HEIs in general, ITIs excluded
        applicability.create_rule(self.user, {'result': 'applies', 'priority': 10, 'approval_bodies': ['NBA'],
                                              'reason': 'NBA-accredited programmes must show integrity controls.'},
                                  benchmark=self.aca)
        applicability.create_rule(self.user, {'result': 'applies', 'priority': 20,
                                              'institution_types': ['engineering', 'polytechnic', 'arts_science', 'university']},
                                  benchmark=self.aca)
        applicability.create_rule(self.user, {'result': 'does_not_apply', 'priority': 30, 'institution_types': ['iti'],
                                              'reason': 'ITIs follow DGT assessment norms instead.'}, benchmark=self.aca)
        rules = list(self.aca.applicability_rules.all())
        eng = applicability.evaluate(rules, ENGINEERING)
        self.assertEqual((eng.result, eng.reason), ('applies', 'NBA-accredited programmes must show integrity controls.'))
        arts = applicability.evaluate(rules, ARTS)
        self.assertEqual(arts.result, 'applies')
        self.assertIn('institution type is', arts.reason)
        iti = applicability.evaluate(rules, ITI)
        self.assertEqual((iti.result, iti.reason), ('does_not_apply', 'ITIs follow DGT assessment norms instead.'))

    def test_unmatched_profile_gets_an_exclusion_reason(self):
        applicability.create_rule(self.user, {'result': 'applies', 'programme_types': ['iti_trade']}, benchmark=self.aca)
        d = applicability.evaluate(list(self.aca.applicability_rules.all()), ARTS)
        self.assertEqual(d.result, 'does_not_apply')
        self.assertIn('needs offers ITI craftsman trade', d.reason)

    def test_four_pilot_profiles_get_different_lists(self):
        """Developer guide, pilot done-criterion 3."""
        applicability.create_rule(self.user, {'result': 'applies', 'approval_bodies': ['AICTE']}, benchmark=self.aca)
        applicability.create_rule(self.user, {'result': 'applies', 'institution_types': ['iti']}, benchmark=self.gov01)
        applicability.create_rule(self.user, {'result': 'applies', 'approval_bodies': ['UGC', 'NBA']}, benchmark=self.gov02)

        def applying(profile):
            return frozenset(r['benchmark'].benchmark_code for r in applicability.resolve_benchmarks(profile)
                             if r['decision'].result == 'applies')
        lists = {name: applying(p) for name, p in
                 (('eng', ENGINEERING), ('poly', POLYTECHNIC), ('arts', ARTS), ('iti', ITI))}
        self.assertEqual(lists['eng'], {'ACA-05', 'GOV-02'})
        self.assertEqual(lists['poly'], {'ACA-05'})
        self.assertEqual(lists['arts'], {'GOV-02'})
        self.assertEqual(lists['iti'], {'GOV-01'})
        self.assertEqual(len(set(lists.values())), 4)

    def test_validation_and_guards(self):
        for bad, msg in (({'result': 'maybe'}, 'result'), ({'institution_types': ['spaceport']}, 'institution type'),
                         ({'autonomous_statuses': ['private']}, 'autonomous_statuses'),
                         ({'approval_bodies': ['FIFA']}, 'approval bodies'), ({'priority': 5000}, 'priority')):
            with self.assertRaisesMessage(applicability.ApplicabilityError, msg):
                applicability.create_rule(self.user, bad, benchmark=self.aca)
        rule = applicability.create_rule(self.user, {'result': 'optional'}, benchmark=self.aca)
        applicability.update_rule(rule, self.user, {'priority': 5, 'institution_types': ['iti']})
        rule.refresh_from_db()
        self.assertEqual((rule.priority, list(rule.institution_types.values_list('code', flat=True))), (5, ['iti']))
        self.aca.status = 'active'
        self.aca.save()
        rule = ApplicabilityRule.objects.get(pk=rule.pk)
        with self.assertRaisesMessage(applicability.ApplicabilityError, 'new draft'):
            applicability.delete_rule(rule, self.user)
        actions = set(SourceAuditEvent.objects.filter(action__startswith='applicability_rule').values_list('action', flat=True))
        self.assertEqual(actions, {'applicability_rule.created', 'applicability_rule.updated'})

    def test_readiness_needs_rules(self):
        checks = {c['key']: c for c in authoring.readiness(self.aca)['checks']}
        self.assertFalse(checks['applicability']['ok'])
        self.assertIn('Source Pack text', checks['applicability']['detail'])
        applicability.create_rule(self.user, {'result': 'applies'}, benchmark=self.aca)
        self.assertTrue({c['key']: c for c in authoring.readiness(self.aca)['checks']}['applicability']['ok'])


@MAINTENANCE_MODE
class ApplicabilityApiTests(_CataloguePack, APITestCase):
    def setUp(self):
        super().setUp()
        token = self.client.post('/api/v1/auth/login/', {'email': 'expert@ingage.ai', 'password': PASSWORD}).data['access_token']
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {token}')

    def test_flow(self):
        ref = self.client.get('/api/v1/benchmarks/applicability/reference').data
        self.assertIn('iti', [t['code'] for t in ref['institution_types']])
        self.assertIn('NCVET', ref['approval_bodies'])
        url = '/api/v1/benchmarks/catalogue/ACA-05/applicability-rules'
        r = self.client.post(url, {'result': 'does_not_apply', 'institution_types': ['iti'],
                                   'reason': 'Not for ITIs.'}, format='json')
        self.assertEqual((r.status_code, r.data['institution_types']), (201, ['iti']))
        self.assertEqual(self.client.post(url, {'result': 'nope'}, format='json').status_code, 400)
        # The preview reads a college's SAVED profile only (the one assessments use).
        from apps.assessments.models import BenchmarkProfile, Programme
        from .models import InstitutionType, ProgrammeType
        iti = Institution.objects.create(name='Govt ITI Karur', city='Karur', state='TN')
        bare = Institution.objects.create(name='No Profile College', city='Karur', state='TN')
        BenchmarkProfile.objects.create(institution=iti, institution_type=InstitutionType.objects.get(code='iti'))
        Programme.objects.create(institution=iti, programme_type=ProgrammeType.objects.get(code='iti_trade'),
                                 name='Electrician')
        url = '/api/v1/benchmarks/applicability/preview'
        colleges = {c['name']: c['has_profile'] for c in self.client.get(url).data}
        self.assertEqual((colleges['Govt ITI Karur'], colleges['No Profile College']), (True, False))
        preview = self.client.get(f'{url}?institution={iti.pk}').data
        self.assertEqual((preview['profile']['institution_type'], preview['profile']['programme_types']),
                         ('iti', ['iti_trade']))
        aca = [b for b in preview['benchmarks'] if b['benchmark_code'] == 'ACA-05'][0]
        self.assertEqual((aca['result'], aca['reason']), ('does_not_apply', 'Not for ITIs.'))
        self.assertEqual(preview['counts'], {'does_not_apply': 1, 'requires_review': 2})
        none = self.client.get(f'{url}?institution={bare.pk}').data
        self.assertEqual((none['profile'], none['benchmarks']), (None, []))  # no saved profile: nothing guessed
        self.assertIn('no saved benchmark profile', none['detail'])
        # A typed-in profile is not accepted any more.
        self.assertEqual(self.client.post(url, {'institution_type': 'iti'}, format='json').status_code, 405)
        rid = r.data['id']
        self.assertEqual(self.client.patch(f'/api/v1/benchmarks/applicability-rules/{rid}', {'priority': 1},
                                           format='json').data['priority'], 1)
        self.assertEqual(self.client.delete(f'/api/v1/benchmarks/applicability-rules/{rid}').status_code, 204)


# -- stage 3.6: releases ------------------------------------------------------------------

from django.core.exceptions import ValidationError as DjangoValidationError  # noqa: E402

from .models import BenchmarkRelease, BenchmarkReleaseItem  # noqa: E402
from .services import releases  # noqa: E402


def confirm_current(benchmark, user):
    """An InGage reviewer confirms the benchmark's current content (required before release submission)."""
    from .services import confirmation
    if confirmation.state(benchmark) != confirmation.CONFIRMED:
        confirmation.confirm(benchmark, user)


def approve_weights(benchmark, user):
    """A person approves the benchmark's composite scoring weights (required for release submission)."""
    from apps.assessments.services import set_score_rule
    set_score_rule(benchmark, user, {'compliance_weight': '0.25', 'maturity_weight': '0.25', 'evidence_weight': '0.25',
                                     'outcome_weight': '0.25', 'note': 'Approved by the benchmark owner (test).'})


def make_ready(benchmark, user, requirement):
    """Everything readiness requires: 3 levels, an evidence rule, an applicability rule, a primary
    requirement, approved scoring weights, and a reviewer's confirmation of that content."""
    for lvl, text in (('L3', REAL_L3), ('L2', REAL_L2), ('L1', REAL_L1)):
        authoring.save_level(benchmark, lvl, user, {'definition': text})
    evidence_rules.create_rule(user, RULE, benchmark=benchmark)
    applicability.create_rule(user, {'result': 'applies', 'institution_types': ['engineering']}, benchmark=benchmark)
    linking.link(benchmark, requirement, user, relationship='primary')
    approve_weights(benchmark, user)
    confirm_current(benchmark, user)


class ReleaseTests(_LinkFixture, TestCase):
    def setUp(self):
        super().setUp()
        self.aca = Benchmark.objects.get(benchmark_code='ACA-05')
        self.approver = User.objects.create_user(email='approver@ingage.ai', username='approver', password=PASSWORD,
                                                 role='super_admin')

    def _publish_first(self):
        make_ready(self.aca, self.user, self.req)
        rel = releases.create_release('Release 1.0', self.user, 'Pilot')
        releases.add_item(rel, self.aca, self.user)
        releases.submit(rel, self.user)
        releases.approve(rel, self.approver)
        return releases.publish(rel, self.approver)

    def test_not_ready_benchmarks_block_review_with_reasons(self):
        rel = releases.create_release('Release 1.0', self.user)
        with self.assertRaises(releases.ReleaseError) as ctx:
            releases.submit(rel, self.user)
        self.assertIn('no benchmarks', ctx.exception.problems[0])
        releases.add_item(rel, self.aca, self.user)
        with self.assertRaises(releases.ReleaseError) as ctx:
            releases.submit(rel, self.user)
        self.assertIn('ACA-05: L3, L2 and L1 written', ctx.exception.problems[0])

    def test_full_lifecycle_and_four_eyes(self):
        make_ready(self.aca, self.user, self.req)
        rel = releases.create_release('Release 1.0', self.user)
        releases.add_item(rel, self.aca, self.user, impact_summary='First version')
        warnings = releases.submit(rel, self.user)
        self.assertIn('20-30 pilot benchmarks', warnings[0])
        with self.assertRaisesMessage(releases.ReleaseError, 'four-eyes'):
            releases.approve(rel, self.user)
        releases.approve(rel, self.approver)
        releases.publish(rel, self.approver)
        rel.refresh_from_db()
        self.aca.refresh_from_db()
        item = rel.items.get()
        self.assertEqual((rel.status, self.aca.status, item.change_type), ('published', 'active', 'added'))
        self.assertEqual((len(rel.content_hash), len(item.snapshot_hash)), (64, 64))
        self.assertEqual(item.snapshot['levels']['L3']['definition'], REAL_L3)
        self.assertEqual(item.snapshot['requirements'][0]['code'], 'UGC-005-R001')
        self.assertEqual(releases.current_release(), rel)
        with self.assertRaisesMessage(authoring.AuthoringError, 'active'):
            authoring.save_level(self.aca, 'L3', self.user, {'definition': REAL_L3 + ' Changed.'})
        item.impact_summary = 'tampered'
        with self.assertRaises(DjangoValidationError):
            item.save()

    def test_benchmarks_are_frozen_while_in_review(self):
        make_ready(self.aca, self.user, self.req)
        rel = releases.create_release('Release 1.0', self.user)
        releases.add_item(rel, self.aca, self.user)
        releases.submit(rel, self.user)
        with self.assertRaisesMessage(authoring.AuthoringError, 'under review'):
            authoring.save_level(self.aca, 'L3', self.user, {'definition': REAL_L3 + ' Edited during review.'})
        with self.assertRaisesMessage(evidence_rules.EvidenceRuleError, 'under review'):
            evidence_rules.create_rule(self.user, {**RULE, 'title': 'Another rule'}, benchmark=self.aca)
        with self.assertRaisesMessage(releases.ReleaseError, 'Say what'):
            releases.send_back(rel, self.approver, '')
        releases.send_back(rel, self.approver, 'Tighten L3')
        authoring.save_level(self.aca, 'L3', self.user, {'definition': REAL_L3 + ' Edited after send-back.'})

    def test_second_release_carries_forward_and_keeps_history(self):
        first = self._publish_first()
        self.aca.refresh_from_db()  # publish made it active
        first_hash = first.items.get().snapshot_hash
        gov = Benchmark.objects.get(benchmark_code='GOV-01')
        make_ready(gov, self.user, self.req)
        second = releases.create_release('Release 1.1', self.user)
        self.assertEqual(second.items.get(benchmark=self.aca).change_type, 'unchanged')
        releases.add_item(second, gov, self.user)
        with self.assertRaisesMessage(releases.ReleaseError, 'carried forward'):
            releases.remove_item(second, self.aca, self.user)
        with self.assertRaisesMessage(releases.ReleaseError, 'reopen'):
            releases.add_item(second, self.aca, self.user)
        releases.reopen_benchmark(self.aca, self.user, 'Clarify L1 outcomes')
        self.aca.refresh_from_db()
        self.assertEqual((self.aca.status, self.aca.version_label), ('draft', 'v0.2.r1'))
        authoring.save_level(self.aca, 'L1', self.user, {'definition': REAL_L1 + ' Reported every semester.'})
        releases.add_item(second, self.aca, self.user)
        self.assertEqual(second.items.get(benchmark=self.aca).change_type, 'updated')
        changes = {d['benchmark_code']: d for d in releases.diff(second)}
        self.assertEqual((changes['ACA-05']['changed'], changes['GOV-01']['change_type']), (['levels'], 'added'))
        with self.assertRaisesMessage(releases.ReleaseError, 'Not ready'):
            releases.submit(second, self.user)  # the reopened version has not been confirmed yet
        confirm_current(self.aca, self.user)
        releases.submit(second, self.user)
        releases.approve(second, self.approver)
        releases.publish(second, self.approver)
        first.refresh_from_db()
        self.assertEqual(first.status, 'retired')
        self.assertEqual(first.items.get().snapshot_hash, first_hash)  # history untouched
        self.assertNotEqual(second.items.get(benchmark=self.aca).snapshot_hash, first_hash)
        self.assertEqual(releases.current_release(), second)

    def test_only_one_release_in_progress_and_unique_labels(self):
        releases.create_release('Release 1.0', self.user)
        with self.assertRaisesMessage(releases.ReleaseError, 'in progress'):
            releases.create_release('Release 2.0', self.user)
        with self.assertRaisesMessage(releases.ReleaseError, 'label'):
            releases.create_release('bad/label!', self.user)


@MAINTENANCE_MODE
class ReleaseApiTests(_LinkFixture, APITestCase):
    def _login(self, email):
        token = self.client.post('/api/v1/auth/login/', {'email': email, 'password': PASSWORD}).data['access_token']
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {token}')

    def test_api_flow(self):
        aca = Benchmark.objects.get(benchmark_code='ACA-05')
        make_ready(aca, self.user, self.req)
        User.objects.create_user(email='boss@ingage.ai', username='boss', password=PASSWORD, role='super_admin')
        self._login('x@ingage.ai')
        self.assertEqual(self.client.get('/api/v1/benchmarks/releases/current').status_code, 404)
        r = self.client.post('/api/v1/benchmarks/releases', {'release_label': 'Release 1.0'}, format='json')
        self.assertEqual((r.status_code, r.data['status']), (201, 'draft'))
        base = '/api/v1/benchmarks/releases/Release%201.0'
        r = self.client.post(f'{base}/items', {'benchmark_code': 'ACA-05'}, format='json')
        self.assertEqual(r.data['items'][0]['change_type'], 'added')
        self.assertEqual(self.client.post(f'{base}/submit').data['status'], 'review')
        refused = self.client.post(f'{base}/approve')
        self.assertIn('four-eyes', refused.data['detail'])
        self._login('boss@ingage.ai')
        self.assertEqual(self.client.post(f'{base}/approve', {'note': 'ok'}, format='json').data['status'], 'approved')
        pub = self.client.post(f'{base}/publish').data
        self.assertEqual((pub['status'], pub['published_by'], len(pub['content_hash'])), ('published', 'boss@ingage.ai', 64))
        self.assertEqual(self.client.get('/api/v1/benchmarks/releases/current').data['release_label'], 'Release 1.0')
        self.assertEqual(self.client.get(f'{base}/diff').data[0]['change_type'], 'added')
        reopened = self.client.post('/api/v1/benchmarks/catalogue/ACA-05/reopen', {'reason': 'fix'}, format='json')
        self.assertEqual((reopened.data['status'], reopened.data['version_label']), ('draft', 'v0.2.r1'))

    def test_staff_only(self):
        User.objects.create_user(email='hod@z.edu', username='hodz', password=PASSWORD, role='iqac_coordinator')
        self._login('hod@z.edu')
        self.assertEqual(self.client.get('/api/v1/benchmarks/releases').status_code, 403)


# -- Stage 4: daily watcher ---------------------------------------------------------------
from .models import SourceChangeEvent, SourceSnapshot, WatcherRun  # noqa: E402
from .services import watcher  # noqa: E402
from .services.watcher import FetchError, Fetched  # noqa: E402


def index_page(*links, footer='Last updated: 12 Oct 2026 10:41'):
    rows = ''.join(f'<li><a href="{href}">{text}</a></li>' for text, href in links)
    return (f'<!doctype html><html><body><h1>Regulations</h1><ul>{rows}</ul>'
            f'<a href="/about">About</a><p>{footer}</p><script>var t=1;</script></body></html>').encode()


def fake(*responses):
    """A fetcher that returns (or raises) the given responses in order."""
    queue = list(responses)

    def _fetch(url):
        item = queue.pop(0)
        if isinstance(item, Exception):
            raise item
        content, ctype = item if isinstance(item, tuple) else (item, 'text/html; charset=utf-8')
        return Fetched(url=url, status=200, content=content, content_type=ctype, headers={'Content-Type': ctype})
    return _fetch


PLAG = ('UGC Plagiarism Regulations, 2018', '/pdfnews/plagiarism.pdf')
NEW = ('UGC (Minimum Standards) Regulations amendment 2026', '/pdfnews/minimum-standards-2026.pdf')


class _WatcherFixture(_LinkFixture):
    def setUp(self):
        super().setUp()
        self.media = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.media, ignore_errors=True)
        override = override_settings(MEDIA_ROOT=self.media)
        override.enable()
        self.addCleanup(override.disable)
        self.feed = SourceFeed.objects.get(document__source_code='UGC-005', feed_type='document_index')
        self.feed.enabled = True
        self.feed.save()


@MAINTENANCE_MODE
class WatcherTests(_WatcherFixture, TestCase):
    def test_first_check_is_a_baseline_and_noise_is_ignored(self):
        r = watcher.check_feed(self.feed, fetcher=fake(index_page(PLAG)))
        self.assertEqual((r['result'], r['events']), ('baseline', 0))
        snap = SourceSnapshot.objects.get()
        self.assertEqual(snap.normalized['links'], [[PLAG[0], 'https://www.ugc.gov.in/pdfnews/plagiarism.pdf']])
        self.assertTrue(os.path.exists(os.path.join(self.media, snap.content_uri)))
        # Different date, script and non-document link: still "unchanged".
        r = watcher.check_feed(self.feed, fetcher=fake(index_page(PLAG, footer='Last updated: 13 Oct 2026 02:00')))
        self.assertEqual(r['result'], 'unchanged')
        self.assertEqual((SourceSnapshot.objects.count(), SourceChangeEvent.objects.count()), (1, 0))
        self.feed.refresh_from_db()
        self.assertEqual(self.feed.last_result, 'unchanged')

    def test_new_document_link_raises_an_event_with_impact_and_changes_nothing_else(self):
        aca = Benchmark.objects.get(benchmark_code='ACA-05')
        BenchmarkRequirement.objects.create(benchmark=aca, requirement=self.req, relationship='primary')
        before = (Requirement.objects.count(), aca.status, SourceVersion.objects.get().status)
        watcher.check_feed(self.feed, fetcher=fake(index_page(PLAG)))
        r = watcher.check_feed(self.feed, fetcher=fake(index_page(PLAG, NEW)))
        self.assertEqual((r['result'], r['events']), ('changed', 1))
        event = SourceChangeEvent.objects.get()
        self.assertEqual((event.change_kind, event.status), ('new_document', 'triaged'))
        self.assertGreaterEqual(event.materiality_score, 55)  # "Regulations" is a key term
        self.assertEqual(event.details['added'][0][1], 'https://www.ugc.gov.in/pdfnews/minimum-standards-2026.pdf')
        self.assertEqual(event.impact['requirements'], ['UGC-005-R001'])
        self.assertEqual(event.impact['benchmarks'], ['ACA-05'])
        self.assertIsNotNone(event.previous_snapshot)
        self.assertIsNotNone(event.current_snapshot)
        aca.refresh_from_db()
        self.assertEqual((Requirement.objects.count(), aca.status, SourceVersion.objects.get().status), before)

    def test_removed_link_is_a_withdrawal(self):
        watcher.check_feed(self.feed, fetcher=fake(index_page(PLAG, NEW)))
        watcher.check_feed(self.feed, fetcher=fake(index_page(NEW)))
        event = SourceChangeEvent.objects.get()
        self.assertEqual(event.change_kind, 'withdrawal')
        self.assertEqual(event.details['removed'][0][0], PLAG[0])

    def test_web_page_text_change(self):
        self.feed.feed_type = 'web_page'
        self.feed.save()
        page = '<html><body><p>Directions under section 70B</p><p>Visitors: {}</p></body></html>'
        watcher.check_feed(self.feed, fetcher=fake(page.format(10).encode()))
        self.assertEqual(watcher.check_feed(self.feed, fetcher=fake(page.format(99).encode()))['result'], 'unchanged')
        changed = page.replace('</body>', '<p>New guideline on incident reporting issued</p></body>')
        watcher.check_feed(self.feed, fetcher=fake(changed.format(99).encode()))
        event = SourceChangeEvent.objects.get()
        self.assertEqual(event.change_kind, 'content_change')
        self.assertEqual(event.details['added_lines'], ['New guideline on incident reporting issued'])

    def test_changed_file_is_a_new_version(self):
        watcher.check_feed(self.feed, fetcher=fake((b'%PDF-1.4 one', 'application/pdf')))
        watcher.check_feed(self.feed, fetcher=fake((b'%PDF-1.4 two', 'application/pdf')))
        self.assertEqual(SourceChangeEvent.objects.get().change_kind, 'new_version')

    def test_page_gone_raises_one_withdrawal_not_one_per_day(self):
        watcher.check_feed(self.feed, fetcher=fake(index_page(PLAG)))
        gone = FetchError('HTTP 404', status=404, kind='http')
        self.assertEqual(watcher.check_feed(self.feed, fetcher=fake(gone))['result'], 'failed')
        watcher.check_feed(self.feed, fetcher=fake(gone))
        event = SourceChangeEvent.objects.get()
        self.assertEqual((event.change_kind, event.current_snapshot), ('withdrawal', None))
        self.feed.refresh_from_db()
        self.assertEqual((self.feed.last_result, self.feed.last_error), ('failed', 'HTTP 404'))

    def test_blocked_access_is_an_access_change(self):
        watcher.check_feed(self.feed, fetcher=fake(index_page(PLAG)))
        watcher.check_feed(self.feed, fetcher=fake(FetchError('HTTP 403', status=403, kind='http')))
        self.assertEqual(SourceChangeEvent.objects.get().change_kind, 'access_change')

    def test_dry_run_saves_nothing(self):
        r = watcher.check_feed(self.feed, fetcher=fake(index_page(PLAG)), dry_run=True)
        self.assertEqual(r['result'], 'baseline')
        self.feed.refresh_from_db()
        self.assertEqual((SourceSnapshot.objects.count(), self.feed.last_checked_at), (0, None))

    def test_run_checks_only_enabled_feeds_and_records_failures(self):
        other = SourceFeed.objects.create(document=self.feed.document, feed_type='web_page', enabled=True,
                                          endpoint_url='https://www.ugc.gov.in/other')
        SourceFeed.objects.create(document=self.feed.document, feed_type='web_page', enabled=False,
                                  endpoint_url='https://www.ugc.gov.in/off')

        def fetcher(url):
            if url == other.endpoint_url:
                raise FetchError('ConnectTimeout', kind='network')
            return fake(index_page(PLAG))(url)
        run, results = watcher.run_watcher(trigger='manual', fetcher=fetcher)
        self.assertEqual((run.feeds_checked, run.status), (2, 'partial_failure'))
        self.assertEqual(run.failure_detail[0]['url'], 'https://www.ugc.gov.in/other')
        self.assertIsNotNone(run.completed_at)

    def test_manual_upload_feeds_are_never_fetched(self):
        self.feed.feed_type = 'manual_upload'
        self.feed.save()
        self.assertEqual(watcher.check_feed(self.feed, fetcher=fake())['result'], 'skipped')

    def test_review_is_audited_and_needs_a_real_decision(self):
        watcher.check_feed(self.feed, fetcher=fake(index_page(PLAG)))
        watcher.check_feed(self.feed, fetcher=fake(index_page(PLAG, NEW)))
        event = SourceChangeEvent.objects.get()
        with self.assertRaises(ValueError):
            watcher.review_event(event, self.user, 'triaged')
        watcher.review_event(event, self.user, 'accepted', 'New regulation; start a new source version.')
        event.refresh_from_db()
        self.assertEqual((event.status, event.reviewed_by), ('accepted', self.user))
        self.assertTrue(SourceAuditEvent.objects.filter(action='watcher.event_reviewed', object_id=str(event.pk)).exists())

    def test_task_runs_and_releases_its_lock(self):
        from django.core.cache import cache

        from .tasks import run_source_watcher
        cache.add('benchmarks:watcher:lock', 'other', 60)
        self.assertIsNone(run_source_watcher(trigger='manual'))  # another run holds the lock
        cache.delete('benchmarks:watcher:lock')
        real = watcher.run_watcher
        with patch.object(watcher, 'run_watcher', lambda **kw: real(fetcher=fake(index_page(PLAG)), **kw)):
            run_id = run_source_watcher(trigger='manual')
        self.assertEqual(WatcherRun.objects.get(pk=run_id).status, 'complete')
        self.assertEqual(SourceSnapshot.objects.count(), 1)
        self.assertFalse(cache.get('benchmarks:watcher:lock'))


@MAINTENANCE_MODE
class WatcherApiTests(_WatcherFixture, APITestCase):
    def _login(self, email):
        token = self.client.post('/api/v1/auth/login/', {'email': email, 'password': PASSWORD}).data['access_token']
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {token}')

    def test_feeds_events_and_review(self):
        self._login('x@ingage.ai')
        feeds = self.client.get('/api/v1/benchmarks/watcher/feeds?enabled=true').data
        self.assertEqual([f['source_code'] for f in feeds], ['UGC-005'])
        off = self.client.post(f'/api/v1/benchmarks/watcher/feeds/{self.feed.pk}/enable', {'enabled': False,
                               'note': 'pause'}, format='json')
        self.assertFalse(off.data['enabled'])
        self.assertTrue(SourceAuditEvent.objects.filter(action='watcher.feed_disabled').exists())
        watcher.check_feed(self.feed, fetcher=fake(index_page(PLAG)))
        watcher.check_feed(self.feed, fetcher=fake(index_page(PLAG, NEW)))
        summary = self.client.get('/api/v1/benchmarks/watcher/summary').data
        self.assertEqual((summary['open_events'], summary['feeds_enabled']), (1, 0))
        [event] = self.client.get('/api/v1/benchmarks/watcher/events').data
        url = f'/api/v1/benchmarks/watcher/events/{event["id"]}/review'
        self.assertEqual(self.client.post(url, {'status': 'ignored'}, format='json').status_code, 400)  # note needed
        done = self.client.post(url, {'status': 'ignored', 'note': 'Duplicate of an existing file.'}, format='json')
        self.assertEqual((done.status_code, done.data['status'], done.data['reviewed_by']), (200, 'ignored', 'x@ingage.ai'))
        self.assertEqual(self.client.get('/api/v1/benchmarks/watcher/events').data, [])
        self.assertEqual(len(self.client.get('/api/v1/benchmarks/watcher/events?status=all').data), 1)

    def test_run_now_needs_an_enabled_feed(self):
        self._login('x@ingage.ai')
        SourceFeed.objects.update(enabled=False)
        self.assertEqual(self.client.post('/api/v1/benchmarks/watcher/runs').status_code, 400)

    def test_college_users_cannot_see_the_watcher(self):
        inst = Institution.objects.create(name='MKCE', city='Karur', state='TN')
        User.objects.create_user(email='c@test.edu', username='c', password=PASSWORD, role='iqac_coordinator',
                                 institution=inst)
        self._login('c@test.edu')
        self.assertEqual(self.client.get('/api/v1/benchmarks/watcher/feeds').status_code, 403)


# -- Stage 4b: benchmark candidates from watcher changes ---------------------------------
from .models import BenchmarkCandidate  # noqa: E402
from .services import candidates as candidate_svc  # noqa: E402

NEW2 = ('UGC Regulations on Academic Integrity amendment 2027', '/pdfnews/integrity-2027.pdf')


class CandidateTests(_WatcherFixture, TestCase):
    def setUp(self):
        super().setUp()
        self.aca = Benchmark.objects.get(benchmark_code='ACA-05')
        BenchmarkRequirement.objects.create(benchmark=self.aca, requirement=self.req, relationship='primary')
        self.legal = User.objects.create_user(email='legal@ingage.ai', username='legal', password=PASSWORD,
                                              role='consultant')

    def _change(self, *links):
        watcher.check_feed(self.feed, fetcher=fake(index_page(PLAG)))
        return watcher.check_feed(self.feed, fetcher=fake(index_page(PLAG, *links)))

    def test_material_change_creates_one_review_candidate_per_linked_benchmark(self):
        r = self._change(NEW)
        self.assertEqual(r['candidates'], 1)
        c = BenchmarkCandidate.objects.get()
        self.assertEqual((c.candidate_type, c.status, c.benchmark, c.created_by),
                         ('update_benchmark', 'review', self.aca, 'watcher'))
        self.assertEqual(c.required_roles, ['accreditation', 'legal'])  # UGC = regulatory authority
        self.assertIn('minimum-standards-2026.pdf', c.source_rationale)
        self.aca.refresh_from_db()
        self.assertEqual(self.aca.status, 'draft')  # nothing about the benchmark changed

    def test_an_open_candidate_is_not_duplicated(self):
        self._change(NEW)
        r = watcher.check_feed(self.feed, fetcher=fake(index_page(PLAG, NEW, NEW2)))
        self.assertEqual((r['events'], r['candidates']), (1, 0))
        self.assertEqual(BenchmarkCandidate.objects.count(), 1)

    def test_unlinked_source_gets_a_review_source_candidate(self):
        BenchmarkRequirement.objects.all().delete()
        self._change(NEW)
        c = BenchmarkCandidate.objects.get()
        self.assertEqual((c.candidate_type, c.benchmark), ('source_update', None))

    def test_low_authority_is_rejected_automatically(self):
        self.feed.authority_trust_score = 20
        self.feed.save()
        self._change(NEW)
        self.assertEqual(BenchmarkCandidate.objects.get().status, 'rejected')

    def test_minor_changes_make_no_candidate(self):
        self.feed.feed_type = 'web_page'
        self.feed.save()
        page = '<html><body>' + ''.join(f'<p>Line number {i} of the regulation text</p>' for i in range(60)) + '{}</body></html>'
        watcher.check_feed(self.feed, fetcher=fake(page.format('').encode()))
        r = watcher.check_feed(self.feed, fetcher=fake(page.format('<p>Office closed on Friday</p>').encode()))
        self.assertEqual((r['events'], r['candidates']), (1, 0))

    def test_review_rules_and_revision(self):
        self._change(NEW)
        c = BenchmarkCandidate.objects.get()
        with self.assertRaisesMessage(candidate_svc.CandidateError, 'needs reviews from'):
            candidate_svc.review(c, self.user, 'ai_governance', 'approve')
        candidate_svc.review(c, self.user, 'accreditation', 'approve')
        with self.assertRaisesMessage(candidate_svc.CandidateError, 'different person'):
            candidate_svc.review(c, self.user, 'legal', 'approve')
        with self.assertRaisesMessage(candidate_svc.CandidateError, 'Say why'):
            candidate_svc.review(c, self.legal, 'legal', 'reject')
        candidate_svc.review(c, self.legal, 'legal', 'approve', 'Amendment changes the integrity rules.')
        c.refresh_from_db()
        self.assertEqual(c.status, 'approved')
        Benchmark.objects.filter(pk=self.aca.pk).update(status='active', version_label='v0.2')
        c.benchmark.refresh_from_db()
        candidate_svc.reopen_for_revision(c, self.user)
        self.aca.refresh_from_db()
        self.assertEqual((self.aca.status, self.aca.version_label), ('draft', 'v0.2.r1'))

    def test_a_reject_rejects_and_needs_research_goes_back(self):
        self._change(NEW)
        c = BenchmarkCandidate.objects.get()
        candidate_svc.review(c, self.user, 'accreditation', 'needs_research', 'Find the gazette copy first.')
        c.refresh_from_db()
        self.assertEqual(c.status, 'research')
        candidate_svc.submit(c, self.user)
        candidate_svc.review(c, self.legal, 'legal', 'reject', 'Not applicable to HEIs.')
        c.refresh_from_db()
        self.assertEqual(c.status, 'rejected')

    def test_watcher_run_counts_candidates(self):
        watcher.check_feed(self.feed, fetcher=fake(index_page(PLAG)))
        run, _ = watcher.run_watcher(trigger='manual', fetcher=fake(index_page(PLAG, NEW)))
        self.assertEqual(run.candidates_created, 1)


@MAINTENANCE_MODE
class CandidateApiTests(_WatcherFixture, APITestCase):
    def _login(self, email):
        token = self.client.post('/api/v1/auth/login/', {'email': email, 'password': PASSWORD}).data['access_token']
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {token}')

    def test_manual_candidate_and_listing(self):
        self._login('x@ingage.ai')
        r = self.client.post('/api/v1/benchmarks/benchmark-candidates', {
            'candidate_type': 'new_benchmark', 'title': 'Generative-AI use policy for assessment',
            'source_rationale': 'UNESCO 2023 guidance recommends an institutional GenAI policy.'}, format='json')
        self.assertEqual((r.status_code, r.data['status']), (201, 'research'))
        self.assertIn('ai_governance', r.data['required_roles'])
        cid = r.data['id']
        self.assertEqual(self.client.post(f'/api/v1/benchmarks/benchmark-candidates/{cid}/submit').data['status'],
                         'review')
        self.assertEqual(len(self.client.get('/api/v1/benchmarks/benchmark-candidates').data), 1)
        inst = Institution.objects.create(name='C', city='X', state='TN')
        User.objects.create_user(email='c@college.edu', username='c', password=PASSWORD, role='iqac_coordinator',
                                 institution=inst)
        self._login('c@college.edu')
        self.assertEqual(self.client.get('/api/v1/benchmarks/benchmark-candidates').status_code, 403)


# -- AI-QIF v1.0 framework import ---------------------------------------------------

from decimal import Decimal  # noqa: E402

from .models import Framework, Pillar  # noqa: E402
from .services import authoring as _authoring  # noqa: E402
from .services.aiqif_import import apply_aiqif, read_aiqif  # noqa: E402

_LV = ('L3 — Foundational', 'L2 — Good / Recommended', 'L1 — Advanced / Aspirational')


def _aiqif_workbook(rows=None, crosswalk=None, source_refs=None):
    """A small AI-QIF master: criterion C1 (60 = 30 + 30) and C2 (40, one KI) -- total 100, not 1000.
    source_refs: rows of the optional 'AI-QIF Source References' sheet
    (indicator, source code, version, page, quote, relationship)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'AI Benchmark Master'
    ws.append(['AI-QIF Benchmark ID', 'Criterion ID', 'Criterion', 'Criterion Weight (of 1000)', 'KI ID',
               'Key Indicator', 'KI Provisional Weight', 'What AI Readiness Capability Is Being Benchmarked?',
               'Why It Matters for AI Readiness', *_LV, 'Applicability', 'Target Rule', 'Status'])
    tmpl = ('capability is defined and owned.', 'capability is systematic and measured.', 'capability is advanced.')
    for code, crit, cname, cw, ki, kw in rows or [
            ('AIQIF-C1.1', 'AIQIF-C1', 'Strategy', 60, 'AI Vision', 30),
            ('AIQIF-C1.2', 'AIQIF-C1', 'Strategy', 60, 'AI Governance', 30),
            ('AIQIF-C2.1', 'AIQIF-C2', 'Responsible AI', 40, 'Responsible AI Policy', 40)]:
        ws.append([code, crit, cname, cw, 1, ki, kw, f'Ability for {ki}.', 'Turns AI into a capability.',
                   *[f'{ki}: {t}' for t in tmpl], 'Institution-specific.', 'Target engine picks the level.',
                   'Proposed v1.0 - requires expert validation'])
    m = wb.create_sheet('AI-QIF Maturity')
    m.append(['Dimension', *_LV])
    m.append(['Strategy', 'Intent defined.', 'Integrated in planning.', 'Dynamic strategy.'])
    c = wb.create_sheet('NAAC → AI-QIF Crosswalk')
    c.append(['NAAC Metric ID', 'NAAC Criterion', 'NAAC Key Indicator', 'NAAC Requirement', 'NAAC Source Framework',
              'AI-QIF Mapping', 'AI-QIF Key Indicator', 'Relationship', 'How NAAC Evidence Can Be Reused',
              'AI-Specific Additional Evidence Needed'])
    for metric, ki in crosswalk or [('1.1.1', 'AIQIF-C1.1 / C1.2'), ('6.2.1', 'AIQIF-C1.1')]:
        c.append([metric, 'Criterion 1', '1.1', f'Requirement {metric}', 'NAAC', 'AIQIF-C1', ki, 'Supporting', 'x', 'y'])
    p = wb.create_sheet('AI-QIF Principles')
    p.append(['Principle', 'Rule'])
    p.append(['NAAC relationship', 'NAAC is supporting evidence, not AI readiness.'])
    a = wb.create_sheet('Scoring Architecture')
    a.append(['Layer', 'Definition', 'Authority'])
    a.append(['AI-QIF Criteria', '9 domains totaling 1000', 'InGage methodology'])
    if source_refs is not None:
        r = wb.create_sheet('AI-QIF Source References')
        r.append(['AI-QIF Benchmark ID', 'Source Code', 'Source Version', 'Page', 'Exact Quote', 'Relationship'])
        for row in source_refs:
            r.append(list(row))
    fd, path = tempfile.mkstemp(suffix='.xlsx')
    os.close(fd)
    wb.save(path)
    return path


def _rows(response):
    data = response.json()
    return data.get('results', data) if isinstance(data, dict) else data


class AiQifImportTests(TestCase):
    def setUp(self):
        self.path = _aiqif_workbook()
        self.addCleanup(os.remove, self.path)

    def test_reads_criteria_weights_and_crosswalk(self):
        data = read_aiqif(self.path)
        self.assertTrue(data.ok)
        self.assertEqual([(c['code'], c['weight'], c['kis']) for c in data.criteria],
                         [('AIQIF-C1', Decimal(60), 2), ('AIQIF-C2', Decimal(40), 1)])
        by_code = {b['code']: b for b in data.benchmarks}
        # "AIQIF-C1.1 / C1.2" supports both key indicators.
        self.assertEqual([x['metric'] for x in by_code['AIQIF-C1.1']['crosswalk']], ['1.1.1', '6.2.1'])
        self.assertEqual(by_code['AIQIF-C1.2']['crosswalk'][0]['shared_with'], ['AIQIF-C1.1'])
        self.assertTrue(any('add up to 100' in w for w in data.warnings))  # not 1000
        self.assertTrue(any('no NAAC crosswalk' in w and 'AIQIF-C2.1' in w for w in data.warnings))
        # Same level wording on every KI once its name is removed = template.
        self.assertTrue(all(all(b['placeholder'].values()) for b in data.benchmarks))

    def test_kis_not_adding_up_to_their_criterion_are_reported(self):
        path = _aiqif_workbook(rows=[('AIQIF-C1.1', 'AIQIF-C1', 'Strategy', 60, 'AI Vision', 25)])
        self.addCleanup(os.remove, path)
        self.assertTrue(any('AIQIF-C1: key indicator weights add up to 25' in w for w in read_aiqif(path).warnings))

    def test_dry_run_saves_nothing(self):
        apply_aiqif(read_aiqif(self.path), actor='test', commit=False)
        self.assertFalse(Benchmark.objects.exists())
        self.assertFalse(Pillar.objects.exists())

    def test_commit_creates_drafts_with_weights_levels_and_framework(self):
        apply_aiqif(read_aiqif(self.path), actor='test', commit=True)
        b = Benchmark.objects.get(benchmark_code='AIQIF-C1.1')
        self.assertEqual((b.status, b.benchmark_kind, b.version_label), ('draft', 'proposed', 'AI-QIF v1.0'))
        self.assertEqual((b.weight, b.pillar.code, b.pillar.weight), (Decimal(30), 'AIQIF-C1', Decimal(60)))
        self.assertEqual(b.framework.code, 'AIQIF-V1.0')
        self.assertIn('NAAC is supporting evidence', b.framework.purpose)
        self.assertEqual(b.target_guidance, 'Target engine picks the level.')
        self.assertEqual(len(b.naac_crosswalk), 2)
        l3 = b.maturity_levels.get(level='L3')
        self.assertTrue(l3.is_placeholder)
        self.assertEqual(l3.definition, 'AI Vision: capability is defined and owned.')
        self.assertEqual(l3.observable_requirements, ['Strategy: Intent defined.'])
        self.assertFalse(_authoring.readiness(b)['ready_for_review'])  # experts still have to confirm

    def test_rerun_keeps_expert_work_and_withdraws_old_drafts(self):
        apply_aiqif(read_aiqif(self.path), actor='test', commit=True)
        expert = User.objects.create_user(email='x@ingage.ai', username='x', password='Str0ng!Pass', role='consultant')
        b = Benchmark.objects.get(benchmark_code='AIQIF-C1.1')
        _authoring.save_level(b, 'L3', expert, {'definition': 'A written, approved AI vision with named owners '
                                                              'and a funded roadmap for the next three years.'})
        _authoring.update_benchmark(b, expert, {'target_guidance': 'Expert target note.'})
        old = Benchmark.objects.create(benchmark_code='ACA-01', title='Old', capability=b.capability, pillar=b.pillar,
                                       benchmark_kind='guidance', version_label='v0.2')
        summary = apply_aiqif(read_aiqif(self.path), actor='test', commit=True, withdraw_others=True)
        b.refresh_from_db()
        old.refresh_from_db()
        self.assertTrue(b.maturity_levels.get(level='L3').definition.startswith('A written, approved'))
        self.assertEqual(b.target_guidance, 'Expert target note.')
        self.assertIn('AIQIF-C1.1 L3 - written by an expert, kept', summary.locked)
        self.assertEqual(old.status, 'withdrawn')
        self.assertEqual(summary.withdrawn, ['ACA-01'])
        self.assertEqual(Framework.objects.filter(code='AIQIF-V1.0').count(), 1)


class AiQifApiTests(APITestCase):
    def setUp(self):
        path = _aiqif_workbook()
        self.addCleanup(os.remove, path)
        apply_aiqif(read_aiqif(path), actor='test', commit=True)
        b = Benchmark.objects.get(benchmark_code='AIQIF-C1.1')
        old_pillar = Pillar.objects.create(code='P01', name='Old pillar')
        Benchmark.objects.create(benchmark_code='ACA-01', title='Old', capability=b.capability, pillar=old_pillar,
                                 benchmark_kind='guidance', version_label='v0.2', status='withdrawn')
        self.client.force_authenticate(User.objects.create_user(email='c@ingage.ai', username='c',
                                                                password='Str0ng!Pass', role='consultant'))

    def test_catalogue_hides_withdrawn_and_shows_weights(self):
        rows = _rows(self.client.get('/api/v1/benchmarks/catalogue'))
        self.assertEqual([r['benchmark_code'] for r in rows], ['AIQIF-C1.1', 'AIQIF-C1.2', 'AIQIF-C2.1'])
        self.assertEqual((rows[0]['weight'], rows[0]['pillar_weight'], rows[0]['framework']),
                         ('30.00', '60.00', 'AIQIF-V1.0'))
        self.assertEqual(rows[0]['naac_crosswalk'][0]['metric'], '1.1.1')
        all_rows = _rows(self.client.get('/api/v1/benchmarks/catalogue?status=all'))
        self.assertIn('ACA-01', [r['benchmark_code'] for r in all_rows])

    def test_pillars_are_the_aiqif_criteria(self):
        pillars = self.client.get('/api/v1/benchmarks/pillars').json()
        self.assertEqual([(p['code'], p['weight'], p['benchmark_count']) for p in pillars],
                         [('AIQIF-C1', '60.00', 2), ('AIQIF-C2', '40.00', 1)])


# -- AI drafting of benchmark content (fully automatic) ---------------------------------

from .models import ApplicabilityRule, EvidenceRule  # noqa: E402
from .services import ai_drafting  # noqa: E402


def _good_draft():
    def level(lv, text):
        return {'level': lv, 'definition': text, 'measurement_rule': f'{lv}: all listed evidence is accepted.',
                'observable_requirements': [f'{lv} observable'], 'minimum_evidence': [f'{lv} record'],
                'outcome_expectations': [f'{lv} outcome'], 'non_examples': [f'{lv} non-example'],
                'requires_manual_review': False}

    def rule(lv, ev, mandatory=True):
        return {'level': lv, 'evidence_type': ev, 'title': f'{lv} {ev} evidence', 'description': 'What to submit.',
                'scope_rule': 'Whole institution', 'mandatory': mandatory, 'minimum_count': 1,
                'minimum_currency_months': 24, 'validation_method': 'manual_review',
                'privacy_classification': 'internal'}

    return {
        'levels': [
            level('L3', 'A governing-body approved AI vision names owners and an initial portfolio of use cases.'),
            level('L2', 'The AI vision drives an annual plan with KPIs that are reviewed each semester and acted on.'),
            level('L1', 'AI strategy is adapted yearly from outcome data and peer benchmarks, with sustained impact.'),
        ],
        'evidence_rules': [rule('L3', 'EV-01'), rule('L2', 'EV-01'), rule('L1', 'EV-12'), rule('L2', 'EV-03', False)],
        'applicability_rules': [{'result': 'applies', 'priority': 100, 'institution_types': [], 'programme_types': [],
                                 'autonomous_statuses': [], 'approval_bodies': [], 'reason': 'Every institution.'}],
        'target_guidance': 'Most colleges should target L2 within two years.',
        'rationale': 'Built from the AI-QIF template and NAAC crosswalk.',
    }


class FakeDraftAI:
    model = 'fake-model-1'

    def __init__(self, draft=None, error=None):
        self.draft, self.error, self.calls = draft, error, []

    def extract_structured_data(self, **kwargs):
        self.calls.append(kwargs)
        if self.error:
            raise self.error
        return self.draft


class AiDraftingTests(TestCase):
    def setUp(self):
        path = _aiqif_workbook()
        self.addCleanup(os.remove, path)
        apply_aiqif(read_aiqif(path), actor='test', commit=True)
        self.b = Benchmark.objects.get(benchmark_code='AIQIF-C1.1')

    def test_good_draft_is_saved_labelled_and_ready_for_release(self):
        ai = FakeDraftAI(_good_draft())
        result = ai_drafting.draft_benchmark(self.b, ai)
        self.assertTrue(result.ok, result.detail)
        self.b.refresh_from_db()
        levels = {m.level: m for m in self.b.maturity_levels.all()}
        self.assertTrue(all(not m.is_placeholder and m.ai_model == 'fake-model-1' for m in levels.values()))
        self.assertTrue(levels['L3'].definition.startswith('A governing-body approved'))
        self.assertEqual(levels['L3'].non_examples, ['L3 non-example'])
        self.assertEqual(EvidenceRule.objects.filter(benchmark=self.b, ai_model='fake-model-1').count(), 4)
        self.assertEqual(ApplicabilityRule.objects.filter(benchmark=self.b, ai_model='fake-model-1').count(), 1)
        # The owner is only suggested: a person names the expert owner.
        self.assertEqual(self.b.expert_owner, '')
        # Complete except for what people do: name the owner, approve the scoring weights, confirm.
        r = _authoring.readiness(self.b)
        self.assertEqual([c['key'] for c in r['checks'] if c['severity'] == 'required' and not c['ok']],
                         ['owner', 'scoring', 'confirmed'])
        self.assertFalse(r['ready_for_review'])
        # The AI was given the real catalogues and the indicator.
        prompt = ai.calls[0]['user_content']
        self.assertIn('AIQIF-C1.1', prompt)
        self.assertIn('EV-01', prompt)
        self.assertIn('1.1.1', prompt)  # NAAC crosswalk
        drafter = ai_drafting.ai_user()
        self.assertFalse(drafter.is_active)
        self.assertFalse(drafter.has_usable_password())
        self.assertTrue(SourceAuditEvent.objects.filter(action='benchmark.ai_drafted', object_id=str(self.b.pk)).exists())

    def test_draft_missing_a_mandatory_rule_saves_nothing(self):
        bad = _good_draft()
        bad['evidence_rules'] = [r for r in bad['evidence_rules'] if r['level'] != 'L1']
        result = ai_drafting.draft_benchmark(self.b, FakeDraftAI(bad))
        self.assertFalse(result.ok)
        self.assertIn('L1', result.detail)
        self.assertTrue(all(m.is_placeholder for m in self.b.maturity_levels.all()))
        self.assertFalse(EvidenceRule.objects.filter(benchmark=self.b).exists())

    def test_template_text_or_identical_levels_are_refused(self):
        bad = _good_draft()
        bad['levels'][1]['definition'] = bad['levels'][0]['definition']
        self.assertFalse(ai_drafting.draft_benchmark(self.b, FakeDraftAI(bad)).ok)
        self.assertFalse(EvidenceRule.objects.filter(benchmark=self.b).exists())

    def test_ai_failure_changes_nothing(self):
        result = ai_drafting.draft_benchmark(self.b, FakeDraftAI(error=RuntimeError('router down')))
        self.assertFalse(result.ok)
        self.assertIn('router down', result.detail)

    def test_rerun_replaces_ai_work_but_keeps_what_a_person_wrote(self):
        ai_drafting.draft_benchmark(self.b, FakeDraftAI(_good_draft()))
        person = User.objects.create_user(email='e@ingage.ai', username='e', password='Str0ng!Pass', role='consultant')
        _authoring.save_level(self.b, 'L2', person, {'definition': 'An expert-written L2 that names the review '
                                                                   'board, KPIs and the corrective-action log.'})
        from .services import evidence_rules
        evidence_rules.create_rule(person, {'evidence_type': 'EV-09', 'title': 'Expert data rule', 'level': 'L2'},
                                   benchmark=self.b)
        result = ai_drafting.draft_benchmark(self.b, FakeDraftAI(_good_draft()))
        self.assertTrue(result.ok, result.detail)
        self.assertIn('L2 level', result.kept_human)
        self.assertTrue(self.b.maturity_levels.get(level='L2').definition.startswith('An expert-written'))
        self.assertEqual(EvidenceRule.objects.filter(benchmark=self.b, ai_model='').count(), 1)  # the expert's
        self.assertEqual(EvidenceRule.objects.filter(benchmark=self.b).exclude(ai_model='').count(), 4)  # not 8

    def test_only_drafts_are_written(self):
        Benchmark.objects.filter(pk=self.b.pk).update(status='active')
        self.b.refresh_from_db()
        ai = FakeDraftAI(_good_draft())
        self.assertFalse(ai_drafting.draft_benchmark(self.b, ai).ok)
        self.assertEqual(ai.calls, [])


class AiDraftingFixTests(AiDraftingTests.__bases__[0]):
    def setUp(self):
        path = _aiqif_workbook()
        self.addCleanup(os.remove, path)
        apply_aiqif(read_aiqif(path), actor='test', commit=True)
        self.b = Benchmark.objects.get(benchmark_code='AIQIF-C1.1')

    def test_a_condition_listing_every_option_is_stored_as_everyone(self):
        from .models import APPROVAL_BODIES, AUTONOMOUS_STATUSES, InstitutionType, ProgrammeType
        draft = _good_draft()
        draft['applicability_rules'][0].update({
            'institution_types': list(InstitutionType.objects.values_list('code', flat=True)),
            'programme_types': list(ProgrammeType.objects.values_list('code', flat=True)),
            'autonomous_statuses': [c for c, _ in AUTONOMOUS_STATUSES],
            'approval_bodies': list(APPROVAL_BODIES)})
        self.assertTrue(ai_drafting.draft_benchmark(self.b, FakeDraftAI(draft)).ok)
        rule = ApplicabilityRule.objects.get(benchmark=self.b)
        self.assertEqual((rule.approval_bodies, rule.autonomous_statuses), ([], []))
        self.assertFalse(rule.institution_types.exists() or rule.programme_types.exists())

    def test_a_narrow_condition_is_kept(self):
        draft = _good_draft()
        draft['applicability_rules'][0].update({'result': 'requires_review', 'institution_types': ['iti']})
        self.assertTrue(ai_drafting.draft_benchmark(self.b, FakeDraftAI(draft)).ok)
        rule = ApplicabilityRule.objects.get(benchmark=self.b)
        self.assertEqual(list(rule.institution_types.values_list('code', flat=True)), ['iti'])

    def test_a_failed_call_is_retried_on_another_router(self):
        from unittest.mock import MagicMock, patch
        from apps.extraction.services.router_pool import Lease, Router
        r1, r2 = Router(index=1, base_url='https://r/1/v1', api_key='k'), Router(index=2, base_url='https://r/2/v1', api_key='k')
        pool = MagicMock()
        pool.acquire.side_effect = [Lease(router=r1), Lease(router=r2)]
        outcomes = iter([ai_drafting.DraftResult('AIQIF-C1.1', False, 'AI call failed: 502'),
                         ai_drafting.DraftResult('AIQIF-C1.1', True, 'written')])
        with patch('apps.extraction.services.router_pool.get_router_pool', return_value=pool), \
                patch.object(ai_drafting, 'draft_many', side_effect=lambda bs: [next(outcomes)]):
            result = ai_drafting.draft_with_router(self.b)
        self.assertTrue(result.ok)
        self.assertEqual(pool.acquire.call_args_list[1].kwargs['avoid'], r1.label)
        self.assertEqual(pool.release.call_count, 2)


# -- Add Source (Studio) ---------------------------------------------------------------

from django.conf import settings  # noqa: E402
from django.core.files.uploadedfile import SimpleUploadedFile  # noqa: E402

from .models import Framework as _Framework  # noqa: E402


@override_settings(MEDIA_ROOT=tempfile.mkdtemp(prefix='source-create-'))
@MAINTENANCE_MODE
class AddSourceApiTests(APITestCase):
    URL = '/api/v1/benchmarks/sources'

    def setUp(self):
        self.staff = User.objects.create_user(email='c@ingage.ai', username='c', password='Str0ng!Pass', role='consultant')
        self.client.force_authenticate(self.staff)
        Authority.objects.create(code='UGC', name='University Grants Commission', kind='regulatory')

    def _post(self, **fields):
        data = {'authority': 'UGC', 'category': '2. Governance', 'source_code': 'UGC-900',
                'title': 'UGC Guidelines on AI in Higher Education', 'source_role': 'mandatory',
                'applies_to': 'All UGC-recognised institutions', 'official_url': 'https://www.ugc.gov.in/page'}
        data.update(fields)
        data = {k: v for k, v in data.items() if v is not None}
        return self.client.post(self.URL, data, format='multipart')

    def test_only_ingage_staff_can_add(self):
        college = User.objects.create_user(email='p@college.edu', username='p', password='Str0ng!Pass',
                                           role='iqac_coordinator')
        self.client.force_authenticate(college)
        self.assertEqual(self._post().status_code, 403)
        self.assertFalse(SourceDocument.objects.exists())

    def test_link_only_source_starts_as_a_draft_found_link_only(self):
        r = self._post()
        self.assertEqual(r.status_code, 201, r.data)
        doc = SourceDocument.objects.get(source_code='UGC-900')
        v = doc.versions.get()
        self.assertEqual((v.status, v.collection_status, v.version_label), ('draft', 'link_only', 'unversioned'))
        self.assertEqual(doc.framework.code, 'UGC-C2')
        self.assertEqual(v.review_flags, [])
        self.assertEqual(r.data['source_code'], 'UGC-900')
        self.assertTrue(SourceAuditEvent.objects.filter(action='source.create', object_id=str(doc.pk)).exists())

    def test_uploaded_pdf_is_stored_with_checksum_and_marked_downloaded(self):
        pdf = SimpleUploadedFile('ugc-ai.pdf', b'%PDF-1.4 small test file', content_type='application/pdf')
        r = self._post(file=pdf, version_label='2025', issued_on='2025-03-01')
        self.assertEqual(r.status_code, 201, r.data)
        v = SourceVersion.objects.get(document__source_code='UGC-900')
        self.assertEqual((v.collection_status, v.intake_status), ('downloaded', 'computed'))
        self.assertEqual(len(v.checksum_sha256), 64)
        self.assertTrue(v.file_uri.startswith('source_library/'))
        self.assertTrue(os.path.exists(os.path.join(settings.MEDIA_ROOT, v.file_uri)))

    def test_licensed_standard_is_link_only_and_refuses_a_file(self):
        self.assertEqual(self._post(licensed_content='true', source_code='ISO-1').status_code, 201)
        self.assertEqual(SourceVersion.objects.get(document__source_code='ISO-1').collection_status, 'licensed_link_only')
        r = self._post(licensed_content='true', source_code='ISO-2',
                       file=SimpleUploadedFile('std.pdf', b'%PDF-1.4 x', content_type='application/pdf'))
        self.assertEqual(r.status_code, 400)
        self.assertIn('file', r.data['errors'])

    def test_nothing_known_yet_is_to_find_and_flagged_like_an_import(self):
        r = self._post(official_url='', source_role='', applies_to='')
        self.assertEqual(r.status_code, 201, r.data)
        v = SourceVersion.objects.get(document__source_code='UGC-900')
        self.assertEqual(v.collection_status, 'to_find')
        self.assertIn('Applies To is empty.', v.review_flags)

    def test_unofficial_domain_is_saved_but_flagged(self):
        self._post(official_url='https://example.com/ugc')
        v = SourceVersion.objects.get(document__source_code='UGC-900')
        self.assertTrue(any('not an official domain' in f for f in v.review_flags))

    def test_validation_returns_every_problem_and_saves_nothing(self):
        r = self._post(source_code='x', title='AI', official_url='ugc.gov.in', issued_on='2025-13-40',
                       file=SimpleUploadedFile('notes.exe', b'MZ', content_type='application/octet-stream'))
        self.assertEqual(r.status_code, 400)
        self.assertEqual(set(r.data['errors']), {'source_code', 'title', 'official_url', 'issued_on', 'file'})
        self.assertFalse(SourceDocument.objects.exists())

    def test_duplicate_code_is_refused(self):
        self._post()
        r = self._post(title='Another title for the same code')
        self.assertEqual(r.status_code, 400)
        self.assertIn('already exists', r.data['errors']['source_code'])
        self.assertEqual(SourceDocument.objects.count(), 1)

    def test_new_authority_is_created_with_its_framework(self):
        r = self._post(authority='', authority_name='National Board of Accreditation', authority_kind='accreditation',
                       source_code='NBA-900', category='4. Accreditation')
        self.assertEqual(r.status_code, 201, r.data)
        auth = Authority.objects.get(name='National Board of Accreditation')
        self.assertEqual(auth.kind, 'accreditation')
        self.assertTrue(_Framework.objects.filter(code=f'{auth.code}-C4', authority=auth).exists())
        bad = self._post(authority='', authority_name='', authority_kind='nope', source_code='NBA-901')
        self.assertEqual(set(bad.data['errors']) & {'authority_name', 'authority_kind'}, {'authority_name', 'authority_kind'})

    def test_reference_lists_the_form_choices(self):
        data = self.client.get('/api/v1/benchmarks/sources/reference').data
        self.assertIn('UGC', [a['code'] for a in data['authorities']])
        self.assertIn('.pdf', data['file_extensions'])
        self.assertEqual(data['max_file_mb'], 25)
        self.assertIn('mandatory', [r['value'] for r in data['source_roles']])


# -- Confirming a draft benchmark (separate from publishing) ------------------------------

from .models import BenchmarkRelease  # noqa: E402
from .services import confirmation as _confirmation  # noqa: E402
from .services import releases as _releases  # noqa: E402


@MAINTENANCE_MODE
class BenchmarkConfirmationTests(APITestCase):
    def setUp(self):
        path = _aiqif_workbook()
        self.addCleanup(os.remove, path)
        apply_aiqif(read_aiqif(path), actor='test', commit=True)
        self.b = Benchmark.objects.get(benchmark_code='AIQIF-C1.1')
        self.staff = User.objects.create_user(email='s@ingage.ai', username='s', password='Str0ng!Pass', role='consultant')
        ai_drafting.draft_benchmark(self.b, FakeDraftAI(_good_draft()))  # makes the content complete
        authoring.update_benchmark(self.b, self.staff, {'expert_owner': 'Dr. A. Expert'})  # a person names the owner
        self.b.refresh_from_db()

    def _events(self, action):
        return SourceAuditEvent.objects.filter(object_id=str(self.b.pk), action=action).order_by('created_at')

    def test_an_incomplete_draft_cannot_be_confirmed(self):
        other = Benchmark.objects.get(benchmark_code='AIQIF-C1.2')  # still template levels
        with self.assertRaisesMessage(_confirmation.ConfirmationError, 'Finish these before confirming'):
            _confirmation.confirm(other, self.staff)
        self.assertEqual(_confirmation.state(other), 'none')

    def test_confirm_marks_it_confirmed_records_it_and_publishes_nothing(self):
        _confirmation.confirm(self.b, self.staff, note='Checked against AI-QIF')
        self.b.refresh_from_db()
        self.assertEqual(_confirmation.state(self.b), 'confirmed')
        self.assertEqual((self.b.status, self.b.confirmed_by), ('draft', self.staff))
        event = self._events('benchmark.confirmed').get()
        self.assertEqual(event.after['confirmed_hash'], self.b.confirmed_hash)
        self.assertEqual(event.after['version_label'], self.b.version_label)
        self.assertEqual(event.reason, 'Checked against AI-QIF')
        # Not published: still a draft, no release, nothing for assessments.
        self.assertFalse(BenchmarkRelease.objects.exists())
        self.assertIsNone(_releases.current_release())
        check = {c['key']: c for c in _authoring.readiness(self.b)['checks']}['confirmed']
        self.assertTrue(check['ok'])
        self.assertEqual(check['severity'], 'required')

    def test_an_edit_after_confirming_needs_a_new_confirmation(self):
        _confirmation.confirm(self.b, self.staff)
        _authoring.save_level(self.b, 'L2', self.staff, {'definition': 'An expert rewrite of L2 naming the review '
                                                                     'board, the KPI set and the action log.'})
        self.b.refresh_from_db()
        self.assertEqual(_confirmation.state(self.b), 'changed')
        _confirmation.confirm(self.b, self.staff)
        self.b.refresh_from_db()
        self.assertEqual(_confirmation.state(self.b), 'confirmed')
        events = list(self._events('benchmark.confirmed'))
        self.assertEqual(len(events), 2)  # history kept
        self.assertEqual(events[1].before['confirmed_hash'], events[0].after['confirmed_hash'])

    def test_confirming_again_without_changes_is_refused(self):
        _confirmation.confirm(self.b, self.staff)
        with self.assertRaisesMessage(_confirmation.ConfirmationError, 'already confirmed'):
            _confirmation.confirm(self.b, self.staff)

    def test_withdraw_needs_a_reason_and_is_recorded(self):
        _confirmation.confirm(self.b, self.staff)
        with self.assertRaisesMessage(_confirmation.ConfirmationError, 'reason'):
            _confirmation.withdraw_confirmation(self.b, self.staff, '')
        _confirmation.withdraw_confirmation(self.b, self.staff, 'Evidence rules need another look')
        self.b.refresh_from_db()
        self.assertEqual(_confirmation.state(self.b), 'none')
        self.assertEqual(self._events('benchmark.confirmation_withdrawn').get().reason, 'Evidence rules need another look')

    def test_only_drafts_and_reopening_starts_unconfirmed(self):
        _confirmation.confirm(self.b, self.staff)
        Benchmark.objects.filter(pk=self.b.pk).update(status='active')
        self.b.refresh_from_db()
        self.assertIsNone(_confirmation.state(self.b))
        with self.assertRaisesMessage(_confirmation.ConfirmationError, 'only a draft'):
            _confirmation.confirm(self.b, self.staff)
        _releases.reopen_benchmark(self.b, self.staff, 'Next version')
        self.b.refresh_from_db()
        self.assertEqual((self.b.status, _confirmation.state(self.b)), ('draft', 'none'))
        self.assertTrue(self._events('benchmark.confirmed').exists())  # the earlier confirmation stays in history

    def test_api_confirm_and_catalogue_shows_it(self):
        college = User.objects.create_user(email='p@c.edu', username='p', password='Str0ng!Pass', role='iqac_coordinator')
        self.client.force_authenticate(college)
        self.assertEqual(self.client.post('/api/v1/benchmarks/catalogue/AIQIF-C1.1/confirm', {}).status_code, 403)
        self.client.force_authenticate(self.staff)
        r = self.client.post('/api/v1/benchmarks/catalogue/AIQIF-C1.1/confirm', {'note': 'ok'}, format='json')
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual(r.data['confirmation']['state'], 'confirmed')
        rows = _rows(self.client.get('/api/v1/benchmarks/catalogue'))
        states = {x['benchmark_code']: x['confirmation']['state'] for x in rows}
        self.assertEqual((states['AIQIF-C1.1'], states['AIQIF-C1.2']), ('confirmed', 'none'))
        bad = self.client.post('/api/v1/benchmarks/catalogue/AIQIF-C1.2/confirm', {}, format='json')
        self.assertEqual(bad.status_code, 400)
        self.assertIn('Finish these', bad.data['detail'])
        undo = self.client.post('/api/v1/benchmarks/catalogue/AIQIF-C1.1/unconfirm', {'reason': 'recheck'}, format='json')
        self.assertEqual(undo.data['confirmation']['state'], 'none')


class TidyApplicabilityTests(AiDraftingFixTests):
    def test_exceptions_run_first_and_applies_becomes_the_catch_all(self):
        from .services.applicability import Profile, evaluate
        draft = _good_draft()
        draft['applicability_rules'] = [
            {'result': 'applies', 'priority': 1, 'institution_types': ['engineering', 'arts_science', 'university'],
             'programme_types': ['ug_engineering', 'ug_arts_science'], 'autonomous_statuses': [],
             'approval_bodies': ['UGC', 'AICTE'], 'reason': 'Higher education.'},
            {'result': 'requires_review', 'priority': 2, 'institution_types': ['iti'], 'programme_types': [],
             'autonomous_statuses': [], 'approval_bodies': [], 'reason': 'ITIs need a review.'}]
        self.assertTrue(ai_drafting.draft_benchmark(self.b, FakeDraftAI(draft)).ok)
        rules = list(ApplicabilityRule.objects.filter(benchmark=self.b).order_by('priority'))
        self.assertEqual([r.result for r in rules], ['requires_review', 'applies'])
        catch_all = rules[1]
        self.assertEqual((catch_all.approval_bodies, catch_all.institution_types.count()), ([], 0))
        self.assertTrue(all(r.ai_model == 'fake-model-1' for r in rules))
        nothing_recorded = evaluate(rules, Profile('engineering'))
        self.assertEqual(nothing_recorded.result, 'applies')
        self.assertEqual(evaluate(rules, Profile('iti')).result, 'requires_review')


# -- release gates: confirmation, scoring weights, applicability validation, prefill ----------

from decimal import Decimal as _D  # noqa: E402

from rest_framework.test import APIClient as _APIClient  # noqa: E402

from .services import confirmation as _conf  # noqa: E402
from .services.release_gates import score_rule_status, validate_applicability  # noqa: E402


def _approve_weights(benchmark, user, weights=('0.25', '0.25', '0.25', '0.25')):
    from apps.assessments.models import BenchmarkScoreRule, ScoringModel
    from django.utils import timezone
    model = ScoringModel.objects.get(active=True)
    c, m, e, o = (_D(w) for w in weights)
    return BenchmarkScoreRule.objects.create(benchmark=benchmark, scoring_model=model, compliance_weight=c,
                                             maturity_weight=m, evidence_weight=e, outcome_weight=o,
                                             approved_by=user, approved_at=timezone.now())


@MAINTENANCE_MODE
class ReleaseGateTests(_LinkFixture, TestCase):
    def setUp(self):
        super().setUp()
        self.aca = Benchmark.objects.get(benchmark_code='ACA-05')
        self.gov = Benchmark.objects.get(benchmark_code='GOV-01')

    # confirmation is a required, human step
    def test_only_an_active_ingage_reviewer_can_confirm(self):
        for lvl, text in (('L3', REAL_L3), ('L2', REAL_L2), ('L1', REAL_L1)):
            authoring.save_level(self.aca, lvl, self.user, {'definition': text})
        evidence_rules.create_rule(self.user, RULE, benchmark=self.aca)
        applicability.create_rule(self.user, {'result': 'applies'}, benchmark=self.aca)
        linking.link(self.aca, self.req, self.user, relationship='primary')
        college = User.objects.create_user(email='p@c.edu', username='pc', password=PASSWORD, role='iqac_coordinator')
        system = User.objects.create_user(email='bot@system.ingage.local', username='bot', password=PASSWORD,
                                          role='consultant', is_active=False)
        for who in (college, system):
            with self.assertRaisesMessage(_conf.ConfirmationError, 'InGage reviewer'):
                _conf.confirm(self.aca, who)
        rel = releases.create_release('Release 1.0', self.user)
        releases.add_item(rel, self.aca, self.user)
        with self.assertRaises(releases.ReleaseError) as ctx:
            releases.submit(rel, self.user)
        self.assertIn('Content confirmed by an InGage reviewer', ctx.exception.problems[0])
        _conf.confirm(self.aca, self.user)
        _approve_weights(self.aca, self.user)
        releases.submit(rel, self.user)

    # scoring weights
    def test_any_release_is_blocked_without_approved_weights(self):
        from apps.assessments.models import BenchmarkScoreRule
        make_ready(self.aca, self.user, self.req)
        BenchmarkScoreRule.objects.filter(benchmark=self.aca).delete()  # weights not (or no longer) approved
        for composite in (True, False):
            rel = releases.create_release(f'Release 1.{int(composite)}', self.user, requires_composite=composite)
            releases.add_item(rel, self.aca, self.user)
            with self.assertRaises(releases.ReleaseError) as ctx:
                releases.submit(rel, self.user)
            self.assertTrue(any(p.startswith('ACA-05:') and 'Approved scoring weights' in p
                                for p in ctx.exception.problems), ctx.exception.problems)
            rel.delete()
        _approve_weights(self.aca, self.user)
        self.assertTrue(score_rule_status(self.aca)[0])
        rel = releases.create_release('Release 1.0', self.user)
        releases.add_item(rel, self.aca, self.user)
        self.assertTrue(releases.scoring_summary(rel)['fully_scoreable'])
        warnings = releases.submit(rel, self.user)
        self.assertTrue(any(w.startswith('Pilot size (guidance only, does not block submission)') for w in warnings))

    def test_weights_not_adding_up_to_one_are_not_valid(self):
        _approve_weights(self.aca, self.user, ('0.5', '0.5', '0.5', '0'))
        ok, detail = score_rule_status(self.aca)
        self.assertFalse(ok)
        self.assertIn('add up to 1.5000, not 1', detail)

    # applicability validation
    def test_validation_finds_overridden_rules_and_unmatched_profiles(self):
        self.assertFalse(validate_applicability(self.aca)['has_rules'])
        applicability.create_rule(self.user, {'result': 'applies', 'priority': 10,
                                              'institution_types': ['engineering', 'polytechnic']}, benchmark=self.aca)
        applicability.create_rule(self.user, {'result': 'requires_review', 'priority': 20,
                                              'institution_types': ['polytechnic']}, benchmark=self.aca)
        v = validate_applicability(self.aca)
        rules = {r['priority']: r for r in v['rules']}
        self.assertEqual(rules[20]['decides'], 0)
        self.assertIn('always overridden by earlier rule(s) #10', rules[20]['note'])
        self.assertFalse(v['ok'])
        self.assertIn('arts_science', v['unmatched']['by_institution_type'])  # matches no rule -> does not apply
        self.assertIsNone(v['requires_review'])
        check = {c['key']: c for c in authoring.readiness(self.aca)['checks']}['applicability_coverage']
        self.assertEqual((check['severity'], check['ok']), ('required', False))  # a dead rule blocks release

    def test_validation_reports_requires_review_profiles(self):
        applicability.create_rule(self.user, {'result': 'requires_review', 'priority': 10,
                                              'institution_types': ['iti']}, benchmark=self.aca)
        applicability.create_rule(self.user, {'result': 'applies', 'priority': 20}, benchmark=self.aca)
        v = validate_applicability(self.aca)
        self.assertTrue(v['ok'])
        self.assertEqual(list(v['requires_review']['by_institution_type']), ['iti'])
        self.assertIsNone(v['unmatched'])

    # prefill
    def test_prefill_adds_only_ready_benchmarks_once_and_lists_the_rest(self):
        make_ready(self.aca, self.user, self.req)
        rel = releases.create_release('Release 1.0', self.user, prefill_ready=True)
        self.assertEqual(list(rel.items.values_list('benchmark__benchmark_code', flat=True)), ['ACA-05'])
        releases.prefill(rel, self.user)  # again: nothing duplicated
        self.assertEqual(rel.items.count(), 1)
        rows = {r['code']: r for r in releases.eligibility(rel)}
        self.assertTrue(rows['ACA-05']['included'])
        self.assertFalse(rows['GOV-01']['included'])
        self.assertIn('L3, L2 and L1 written for this benchmark', rows['GOV-01']['missing'])

    def test_api_create_with_options_patch_and_validation(self):
        make_ready(self.aca, self.user, self.req)
        client = _APIClient()
        client.force_authenticate(self.user)
        r = client.post('/api/v1/benchmarks/releases', {'release_label': 'Release 1.0', 'prefill': True,
                                                         'requires_composite': True}, format='json')
        self.assertEqual(r.status_code, 201, r.data)
        self.assertEqual((r.data['requires_composite'], r.data['item_count']), (True, 1))
        self.assertEqual((r.data['scoring']['missing'], r.data['scoring']['fully_scoreable']), ([], True))
        self.assertTrue(any(row['code'] == 'GOV-01' and not row['included'] for row in r.data['eligibility']))
        p = client.patch('/api/v1/benchmarks/releases/Release%201.0', {'requires_composite': False}, format='json')
        self.assertEqual((p.status_code, p.data['requires_composite']), (200, False))
        self.assertEqual(p.data['problems'], [])
        v = client.get('/api/v1/benchmarks/catalogue/ACA-05/applicability-validation')
        self.assertEqual((v.status_code, v.data['has_rules']), (200, True))


# -- an approved "new benchmark" candidate -> a draft benchmark (candidate_svc.create_benchmark_draft) --------

from .models import Capability, EvidenceType  # noqa: E402


def _no_ai(*_args, **_kwargs):
    raise AssertionError('This flow must never call the AI service.')


@patch('apps.extraction.services.ai_service.get_ai_service', _no_ai)
@MAINTENANCE_MODE
class NewBenchmarkDraftTests(APITestCase):
    """Approve -> create draft (explicitly, once) -> complete by hand -> confirm -> release -> published.
    The AI service is patched to fail, so every step here is proven to work without AI."""

    URL = '/api/v1/benchmarks/benchmark-candidates'

    def setUp(self):
        path = _aiqif_workbook()
        self.addCleanup(os.remove, path)
        apply_aiqif(read_aiqif(path), actor='test', commit=True)
        self.staff = User.objects.create_user(email='x@ingage.ai', username='x', password=PASSWORD, role='consultant')
        self.second = User.objects.create_user(email='y@ingage.ai', username='y', password=PASSWORD, role='super_admin')
        EvidenceType.objects.get_or_create(code='EV-01', defaults={'name': 'Policy', 'validation_method': 'Read it.'})
        # A legacy pillar outside AI-QIF must never be offered.
        legacy = Pillar.objects.create(code='P01', name='Institutional governance')
        Benchmark.objects.create(benchmark_code='GOV-01', title='Old', pillar=legacy, benchmark_kind='mandatory',
                                 version_label='v0.2', capability=Capability.objects.create(
                                     code='CAP-GOV-01', name='Old', domain=legacy.name, definition='Old.'))
        self._login(self.staff)
        r = self.client.post(self.URL, {'candidate_type': 'new_benchmark', 'title': 'Generative-AI use policy for assessment',
                                        'source_rationale': 'UNESCO 2023 guidance: institutions need a GenAI policy.'},
                             format='json')
        self.cid = r.data['id']
        self.assertEqual(r.data['proposed_kind'], 'proposed')

    def _login(self, user):
        token = self.client.post('/api/v1/auth/login/', {'email': user.email, 'password': PASSWORD}).data['access_token']
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {token}')

    def _candidate(self):
        return BenchmarkCandidate.objects.get(pk=self.cid)

    def _approve_all(self):
        self.client.post(f'{self.URL}/{self.cid}/submit')
        c = self._candidate()
        for role, user in zip(c.required_roles, (self.staff, self.second)):
            candidate_svc.review(c, user, role, 'approve')
        self.assertEqual(self._candidate().status, 'approved')

    def _create(self, **body):
        payload = {'pillar_code': 'AIQIF-C1', 'capability_code': 'CAP-AIQIF-C1.1', **body}
        return self.client.post(f'{self.URL}/{self.cid}/create-draft', payload, format='json')

    def test_approval_alone_creates_nothing(self):
        before = Benchmark.objects.count()
        self._approve_all()
        self.assertEqual((Benchmark.objects.count(), self._candidate().benchmark_id), (before, None))
        options = self.client.get(f'{self.URL}/{self.cid}/draft-options').data
        self.assertTrue(options['can_create'])
        self.assertEqual([c['code'] for c in options['criteria']], ['AIQIF-C1', 'AIQIF-C2'])  # not P01
        c1 = options['criteria'][0]
        self.assertEqual(c1['next_code'], 'AIQIF-C1.3')
        self.assertEqual([c['code'] for c in c1['capabilities']], ['CAP-AIQIF-C1.1', 'CAP-AIQIF-C1.2'])
        self.assertEqual((options['benchmark_kind'], options['title']),
                         ('proposed', 'Generative-AI use policy for assessment'))

    def test_only_a_fully_approved_new_benchmark_candidate_can_create(self):
        self.assertEqual(self._create().status_code, 400)  # still in research
        self.client.post(f'{self.URL}/{self.cid}/submit')
        c = self._candidate()
        candidate_svc.review(c, self.staff, c.required_roles[0], 'approve')
        r = self._create()
        self.assertEqual(r.status_code, 400)
        self.assertIn('Every required expert role', r.data['detail'])
        self.assertFalse(self.client.get(f'{self.URL}/{self.cid}/draft-options').data['can_create'])
        update = candidate_svc.create(self.staff, {'candidate_type': 'update_benchmark', 'benchmark_code': 'AIQIF-C1.1',
                                                   'title': 'Revise', 'source_rationale': 'x'})
        self.assertEqual(self.client.post(f'{self.URL}/{update.pk}/create-draft', {}, format='json').status_code, 400)

    def test_creates_one_traceable_empty_draft_and_repeats_return_it(self):
        self._approve_all()
        count = Benchmark.objects.count()
        r = self._create(title='Generative-AI policy for assessment', target_guidance='Aim for L2 in year one.')
        self.assertEqual((r.status_code, r.data['benchmark_code'], r.data['benchmark_status'], r.data['status']),
                         (201, 'AIQIF-C1.3', 'draft', 'approved'))  # approved, not published
        b = Benchmark.objects.get(benchmark_code='AIQIF-C1.3')
        self.assertEqual((b.pillar.code, b.capability.code, b.framework.code, b.version_label, b.benchmark_kind),
                         ('AIQIF-C1', 'CAP-AIQIF-C1.1', 'AIQIF-V1.0', 'AI-QIF v1.0', 'proposed'))
        self.assertEqual((b.title, b.target_guidance), ('Generative-AI policy for assessment', 'Aim for L2 in year one.'))
        self.assertIn(self.cid, b.source_basis)
        self.assertIn('UNESCO 2023 guidance', b.source_basis)
        # Empty for staff: placeholder levels only, no evidence, applicability, owner or weight.
        self.assertEqual(sorted((m.level, m.definition, m.is_placeholder) for m in b.maturity_levels.all()),
                         [(lv, '(not defined)', True) for lv in ('L1', 'L2', 'L3')])
        self.assertEqual((b.evidence_types.count(), b.evidence_rules.count(), b.applicability_rules.count(),
                          b.expert_owner, b.weight), (0, 0, 0, '', None))
        self.assertEqual(SourceAuditEvent.objects.get(action='benchmark.created_from_candidate').after['candidate'],
                         self.cid)
        self.assertEqual(SourceAuditEvent.objects.get(action='candidate.draft_created').object_id, self.cid)
        # Repeated clicks: the same draft, never a second one.
        for _ in range(2):
            again = self._create(pillar_code='AIQIF-C2', capability_code='CAP-AIQIF-C2.1')
            self.assertEqual((again.status_code, again.data['benchmark_code']), (200, 'AIQIF-C1.3'))
        self.assertEqual(Benchmark.objects.count(), count + 1)
        self.assertEqual(SourceAuditEvent.objects.filter(action='candidate.draft_created').count(), 1)
        self.assertEqual(self.client.get(f'{self.URL}/{self.cid}/draft-options').data['existing_benchmark'], 'AIQIF-C1.3')
        self.assertEqual(self.client.post(f'{self.URL}/{self.cid}/open-revision').status_code, 400)

    def test_criterion_and_capability_must_be_existing_and_match(self):
        self._approve_all()
        self.assertIn('AI-QIF criteria', self._create(pillar_code='P01', capability_code='CAP-GOV-01').data['detail'])
        self.assertIn('under AIQIF-C1', self._create(capability_code='CAP-AIQIF-C2.1').data['detail'])
        self.assertIn('AI-QIF criteria', self._create(pillar_code='AIQIF-C10').data['detail'])  # no new criterion
        self.assertFalse(Benchmark.objects.filter(benchmark_code__startswith='AIQIF-C1.3').exists())

    def test_codes_follow_the_criterion_and_are_never_reused(self):
        Benchmark.objects.filter(benchmark_code='AIQIF-C1.2').update(status='withdrawn')
        c11 = Benchmark.objects.get(benchmark_code='AIQIF-C1.1')
        Benchmark.objects.create(benchmark_code='AIQIF-C1.7', title='Retired', pillar=c11.pillar, framework=c11.framework,
                                 capability=c11.capability, benchmark_kind='proposed', version_label='AI-QIF v1.0',
                                 status='withdrawn')
        self.assertEqual(candidate_svc.next_code(c11.pillar), 'AIQIF-C1.8')
        self._approve_all()
        self.assertEqual(self._create().data['benchmark_code'], 'AIQIF-C1.8')

    def test_staff_only(self):
        self._approve_all()
        inst = Institution.objects.create(name='C', city='X', state='TN')
        outsider = User.objects.create_user(email='c@college.edu', username='c', password=PASSWORD,
                                            role='iqac_coordinator', institution=inst)
        self._login(outsider)
        self.assertEqual(self.client.get(f'{self.URL}/{self.cid}/draft-options').status_code, 403)
        self.assertEqual(self._create().status_code, 403)
        self.assertIsNone(self._candidate().benchmark_id)

    def test_completed_by_hand_confirmed_and_released_then_the_candidate_is_published(self):
        self._approve_all()
        code = self._create().data['benchmark_code']
        b = Benchmark.objects.get(benchmark_code=code)
        ready = self.client.get(f'/api/v1/benchmarks/catalogue/{code}/readiness').data
        self.assertFalse(ready['ready_for_review'])
        checks = {c['key']: c for c in ready['checks']}
        self.assertEqual((checks['levels']['ok'], checks['levels']['detail']), (False, '0/3 defined (others are placeholders)'))
        self.assertEqual(checks['requirements']['severity'], 'recommended')  # optional for a proposed benchmark
        self.assertEqual(checks['owner']['severity'], 'required')

        rel = releases.create_release('Release 2.0', self.staff)
        releases.add_item(rel, b, self.staff)
        with self.assertRaises(releases.ReleaseError) as ctx:
            releases.submit(rel, self.staff)  # empty draft: readiness blocks it
        self.assertTrue(any(code in p for p in ctx.exception.problems))

        # Staff complete it in the existing Studio flows -- written by people, no AI.
        for lvl, text in (('L3', REAL_L3), ('L2', REAL_L2), ('L1', REAL_L1)):
            authoring.save_level(b, lvl, self.staff, {'definition': text})
        evidence_rules.create_rule(self.staff, RULE, benchmark=b)  # also chooses evidence type EV-01
        applicability.create_rule(self.staff, {'result': 'applies', 'institution_types': ['engineering']}, benchmark=b)
        authoring.update_benchmark(b, self.staff, {'expert_owner': 'Dr. A. Expert'})
        approve_weights(b, self.staff)
        with self.assertRaises(releases.ReleaseError) as ctx:
            releases.submit(rel, self.staff)  # everything but the reviewer's confirmation
        self.assertTrue(any('Content confirmed' in p for p in ctx.exception.problems))
        self.assertFalse(b.requirement_links.exists())  # no requirement needed for a proposed benchmark
        confirm_current(b, self.second)
        self.assertTrue(authoring.readiness(b)['ready_for_review'])
        releases.submit(rel, self.staff)
        self.assertEqual(self._candidate().status, 'approved')  # submitted / approved release: not yet
        releases.approve(rel, self.second)
        self.assertEqual(self._candidate().status, 'approved')
        releases.publish(rel, self.second)
        b.refresh_from_db()
        self.assertEqual((b.status, rel.items.get(benchmark=b).change_type), ('active', 'added'))
        self.assertEqual(self._candidate().status, 'published')


# -- AI preparation across the Studio (services/studio_ai.py) --------------------------------------------

from unittest.mock import patch as _patch  # noqa: E402

from .models import SourceChangeEvent, StudioAIJob  # noqa: E402
from .services import studio_ai  # noqa: E402


class ScriptedAI:
    """Answers by schema name; an Exception value is raised (a provider / router failure)."""
    model = 'fake-model-2'

    def __init__(self, answers):
        self.answers, self.calls = answers, []

    def extract_structured_data(self, **kwargs):
        self.calls.append(kwargs['schema_name'])
        answer = self.answers[kwargs['schema_name']]
        if isinstance(answer, Exception):
            raise answer
        return answer


@MAINTENANCE_MODE
class StudioAITests(_LinkFixture, APITestCase):
    QUOTE = 'every thesis shall be checked for plagiarism'

    def setUp(self):
        super().setUp()
        path = _aiqif_workbook()
        self.addCleanup(os.remove, path)
        apply_aiqif(read_aiqif(path), actor='test', commit=True)
        EvidenceType.objects.get_or_create(code='EV-01', defaults={'name': 'Policy', 'validation_method': 'Read it.'})
        EvidenceType.objects.get_or_create(code='EV-03', defaults={'name': 'Minutes', 'validation_method': 'Read it.'})
        EvidenceType.objects.get_or_create(code='EV-12', defaults={'name': 'Outcomes', 'validation_method': 'Read it.'})
        self.reviewers = [User.objects.create_user(email=f'r{i}@ingage.ai', username=f'r{i}', password=PASSWORD,
                                                   role='super_admin') for i in range(4)]
        self.ai = ScriptedAI({})
        for target, value in (('apps.extraction.services.ai_service.get_ai_service', self.ai),
                              ('apps.benchmarks.services.studio_ai._dispatch', None)):
            p = _patch(target, return_value=value) if value is not None else _patch(target)
            p.start()
            self.addCleanup(p.stop)
        self.client.force_authenticate(self.user)

    def _run(self, kind, target_id, answer=None, name=None):
        if answer is not None:
            self.ai.answers[name] = answer
        job = studio_ai.start(kind, target_id, self.user)
        studio_ai.run_job(job.pk)
        job.refresh_from_db()
        return job

    def _fit(self, **new):
        return {'decision': 'new', 'rationale': 'Not covered by any indicator.', 'confidence': 0.7,
                'matches': [{'benchmark_code': 'ACA-05', 'relationship': 'supporting', 'rationale': 'Integrity.',
                             'confidence': 0.6},
                            {'benchmark_code': 'NOT-REAL', 'relationship': 'primary', 'rationale': 'x', 'confidence': 1}],
                'new_indicator': {'pillar_code': 'AIQIF-C1', 'capability_code': 'CAP-AIQIF-C1.1',
                                  'title': 'Thesis similarity screening with AI-text detection',
                                  'benchmark_kind': 'proposed', 'capability_statement': 'Screens theses.',
                                  'how_measured': 'Share of theses screened.', 'rationale': 'Distinct capability.',
                                  'source_quote': self.QUOTE, **new}}

    # -- requirements -> indicators ------------------------------------------------------------------
    def test_fit_keeps_only_verified_items_and_a_proposal_becomes_a_draft_after_expert_approval(self):
        job = self._run('requirement_fit', 'UGC-005-R001', self._fit(), 'requirement_indicator_fit')
        self.assertEqual(job.status, 'ready', job.error)
        self.assertEqual([m['benchmark_code'] for m in job.output['matches']], ['ACA-05'])
        self.assertEqual(job.output['rejected'][0]['why'], 'unknown indicator')
        self.assertFalse(job.output['new_indicator']['blocked'])
        self.assertEqual(job.model, 'fake-model-2')

        r = self.client.post(f'/api/v1/benchmarks/ai-jobs/{job.pk}/propose', {'note': 'Worth a look'}, format='json')
        self.assertEqual((r.status_code, r.data['status']), (200, 'applied'), r.data)
        c = BenchmarkCandidate.objects.get(candidate_type='new_benchmark')
        self.assertEqual((c.status, c.title, c.ai_proposal['requirement_code']),
                         ('review', 'Thesis similarity screening with AI-text detection', 'UGC-005-R001'))
        self.assertIn(self.QUOTE, c.source_rationale)
        self.assertIsNone(c.benchmark)  # nothing exists before the experts approve

        for role, person in zip(c.required_roles, self.reviewers):
            candidate_svc.review(c, person, role, 'approve')
            c.refresh_from_db()
        bench = c.benchmark
        self.assertEqual((c.status, bench.status, bench.pillar.code, bench.capability.code),
                         ('approved', 'draft', 'AIQIF-C1', 'CAP-AIQIF-C1.1'))
        self.assertEqual(bench.title, 'Thesis similarity screening with AI-text detection')
        self.assertEqual(list(bench.requirement_links.values_list('requirement__code', 'relationship')),
                         [('UGC-005-R001', 'primary')])
        drafting = StudioAIJob.objects.get(kind='benchmark_content', target_id=bench.benchmark_code)
        self.assertEqual(drafting.status, 'queued')  # AI prepares its levels next; never published by itself
        self.assertFalse(bench.release_items.exists())

    def test_a_duplicate_or_unquoted_proposal_is_not_offered(self):
        dup = self._run('requirement_fit', 'UGC-005-R001',
                        self._fit(title='Assessment, examination and academic integrity'), 'requirement_indicator_fit')
        self.assertTrue(dup.output['new_indicator']['blocked'])
        r = self.client.post(f'/api/v1/benchmarks/ai-jobs/{dup.pk}/propose', {}, format='json')
        self.assertEqual(r.status_code, 400)
        self.assertIn('ACA-05', r.data['detail'])
        bad = self._run('requirement_fit', 'UGC-005-R001', self._fit(source_quote='words that are not in the source'),
                        'requirement_indicator_fit')
        self.assertIsNone(bad.output['new_indicator'])
        self.assertIn('quote not found', bad.output['rejected'][-1]['why'])

    # -- levels / evidence / applicability --------------------------------------------------------------
    def _content(self):
        return {**_good_draft(), 'owner_suggestion': 'Accreditation lead', 'confidence': 0.8,
                'source_references': ['UGC-005-R001', 'MADE-UP-R9']}

    def test_content_is_ai_labelled_keeps_human_work_and_still_needs_confirmation(self):
        b = Benchmark.objects.get(benchmark_code='AIQIF-C1.1')
        linking.link(b, self.req, self.user, relationship='primary')
        authoring.save_level(b, 'L1', self.user, {'definition': REAL_L1})  # a person wrote L1
        authoring.update_benchmark(b, self.user, {'target_guidance': 'Aim for L2 by 2027 (owner).'})
        job = self._run('benchmark_content', 'AIQIF-C1.1', self._content(), 'benchmark_content')
        self.assertEqual(job.status, 'ready', job.error)
        b.refresh_from_db()
        levels = {m.level: m for m in b.maturity_levels.all()}
        self.assertEqual((levels['L3'].ai_model, levels['L1'].ai_model, levels['L1'].definition),
                         ('fake-model-2', '', REAL_L1))
        self.assertEqual(b.target_guidance, 'Aim for L2 by 2027 (owner).')
        self.assertEqual(set(job.output['kept_human']), {'L1 level', 'target guidance'})
        self.assertEqual((job.output['owner_suggestion'], job.confidence), ('Accreditation lead', 0.8))
        self.assertEqual(job.citations, [{'requirement': 'UGC-005-R001'}])  # only codes really linked
        self.assertEqual(b.expert_owner, '')
        self.assertFalse(authoring.readiness(b)['ready_for_review'])  # a reviewer must still confirm
        self.assertTrue(SourceAuditEvent.objects.filter(action='ai.ready', object_id=job.target_id).exists())

    def test_only_drafts_are_prepared_and_failures_can_be_retried(self):
        Benchmark.objects.filter(benchmark_code='AIQIF-C1.2').update(status='active')
        r = self.client.post('/api/v1/benchmarks/ai-jobs', {'kind': 'benchmark_content', 'target_id': 'AIQIF-C1.2'},
                             format='json')
        self.assertEqual(r.status_code, 400)
        job = self._run('benchmark_content', 'AIQIF-C1.1', RuntimeError('router down'), 'benchmark_content')
        self.assertEqual(job.status, 'failed')
        self.assertIn('router down', job.error)
        self.assertTrue(all(m.is_placeholder for m in Benchmark.objects.get(benchmark_code='AIQIF-C1.1')
                            .maturity_levels.all()))  # nothing applied
        self.assertEqual(self.client.post(f'/api/v1/benchmarks/ai-jobs/{job.pk}/dismiss', {}).status_code, 400)
        r = self.client.post(f'/api/v1/benchmarks/ai-jobs/{job.pk}/retry', {})
        self.assertEqual((r.status_code, r.data['status']), (202, 'queued'))
        self.ai.answers['benchmark_content'] = {**self._content(), 'levels': []}  # fails a check
        studio_ai.run_job(r.data['id'])
        again = StudioAIJob.objects.get(pk=r.data['id'])
        self.assertEqual(again.status, 'failed')
        self.assertIn('failed a check', again.error)

    # -- sources ---------------------------------------------------------------------------------------
    def test_source_details_need_a_quote_from_the_source_and_apply_only_what_is_chosen(self):
        doc = self.version.document
        text = ('UNIVERSITY GRANTS COMMISSION (Promotion of Academic Integrity and Prevention of Plagiarism in Higher '
                'Educational Institutions) Regulations, 2018. Notified on 23rd July, 2018. ' * 4)
        answer = {'title': {'value': 'UGC Plagiarism Regulations 2018',
                            'quote': 'Prevention of Plagiarism in Higher Educational Institutions) Regulations, 2018'},
                  'issued_on': {'value': '2018-07-23', 'quote': 'Notified on 23rd July, 2018'},
                  'document_type': {'value': 'circular', 'quote': 'This circular is addressed to all'},
                  **{f: {'value': None, 'quote': None} for f in ('authority_name', 'source_role', 'applies_to',
                                                                 'programme_scope', 'version_label', 'effective_from')},
                  'confidence': 0.9}
        with _patch('apps.benchmarks.services.studio_ai._source_text', return_value=(text, {'version': 'v1'})):
            SourceDocument.objects.filter(pk=doc.pk).update(official_url='https://www.ugc.gov.in/x')
            job = self._run('source_metadata', doc.source_code, answer, 'source_details')
        self.assertEqual(job.status, 'ready', job.error)
        self.assertEqual(sorted(job.output['fields']), ['issued_on', 'title'])
        self.assertEqual(job.output['rejected'][0]['field'], 'document_type')
        r = self.client.post(f'/api/v1/benchmarks/ai-jobs/{job.pk}/apply', {'fields': ['title']}, format='json')
        self.assertEqual(r.status_code, 200, r.data)
        doc.refresh_from_db()
        self.assertEqual(doc.title, 'UGC Plagiarism Regulations 2018')
        self.assertTrue(SourceAuditEvent.objects.filter(action='source.ai_details_applied').exists())

    def test_version_comparison_lists_added_and_removed_lines(self):
        from .models import SourceSection
        new = SourceVersion.objects.create(document=self.version.document, version_label='2026 amendment',
                                           supersedes=self.version)
        SourceSection.objects.create(version=self.version, ordinal=1, page_start=1, page_end=1,
                                     text='Theses are checked.\nOld clause.')
        SourceSection.objects.create(version=new, ordinal=1, page_start=1, page_end=1,
                                     text='Theses are checked.\nAI-text detection is required.')
        data = self.client.get(f'/api/v1/benchmarks/versions/{new.pk}/compare').data
        self.assertEqual((data['added'], data['removed']), (['AI-text detection is required.'], ['Old clause.']))

    def test_requirement_drafts_show_duplicates_of_approved_requirements(self):
        from .models import RequirementCandidate
        from .serializers import RequirementCandidateSerializer
        draft = RequirementCandidate.objects.create(
            version=self.version, statement=self.req.statement, kind='mandatory', page=3,
            source_quote=self.req.source_quote, confidence=0.9)
        data = RequirementCandidateSerializer(draft).data
        self.assertEqual(data['duplicates'][0]['code'], 'UGC-005-R001')

    # -- watcher -----------------------------------------------------------------------------------------
    def _event(self):
        return SourceChangeEvent.objects.create(
            document=self.version.document, change_kind='content_change', materiality_score=80,
            summary='2 lines added', details={'added_lines': ['Regulation 7: AI-detection reports are required for '
                                                              'every thesis.']})

    def test_change_analysis_quotes_the_change_and_prepares_a_candidate(self):
        linking.link(Benchmark.objects.get(benchmark_code='ACA-05'), self.req, self.user, relationship='primary')
        event = self._event()
        answer = {'summary': 'Detection reports become mandatory.', 'impact_level': 'high',
                  'changes': [{'description': 'New duty', 'quote': 'AI-detection reports are required'},
                              {'description': 'Invented', 'quote': 'a sentence nobody wrote'}],
                  'affected_requirements': ['UGC-005-R001'], 'affected_benchmarks': ['ACA-05', 'NOPE'],
                  'impact_rationale': 'Changes the evidence for ACA-05.', 'recommendation': 'update_benchmark',
                  'proposed_title': 'Add AI-detection evidence to ACA-05', 'confidence': 0.75}
        job = self._run('change_analysis', event.pk, answer, 'source_change_analysis')
        self.assertEqual(job.status, 'ready', job.error)
        self.assertEqual((len(job.output['changes']), job.output['affected_benchmarks']), (1, ['ACA-05']))
        r = self.client.post(f'/api/v1/benchmarks/ai-jobs/{job.pk}/candidate', {}, format='json')
        self.assertEqual(r.status_code, 200, r.data)
        c = BenchmarkCandidate.objects.get(change_event=event, created_by=self.user.email)
        self.assertEqual((c.candidate_type, c.benchmark.benchmark_code, c.status), ('update_benchmark', 'ACA-05', 'review'))
        self.assertIn('AI-detection reports are required', c.source_rationale)

    def test_material_changes_are_queued_for_analysis(self):
        event = self._event()
        studio_ai.queue_change_analysis([event])
        self.assertEqual(StudioAIJob.objects.get(kind='change_analysis').target_id, str(event.pk))

    # -- releases ----------------------------------------------------------------------------------------
    def test_release_notes_are_applied_to_a_draft_only(self):
        aca = Benchmark.objects.get(benchmark_code='ACA-05')
        make_ready(aca, self.user, self.req)
        rel = releases.create_release('Release 1.0', self.user)
        releases.add_item(rel, aca, self.user)
        answer = {'version_notes': 'First pilot release.', 'change_summary': 'Adds ACA-05.',
                  'impact_summary': 'Colleges evidence thesis screening.', 'readiness_notes': '',
                  'item_impacts': [{'benchmark_code': 'ACA-05', 'impact_summary': 'New: thesis screening evidence.'}],
                  'confidence': 0.6}
        job = self._run('release_notes', 'Release 1.0', answer, 'release_notes')
        self.assertEqual(job.status, 'ready', job.error)
        self.client.post(f'/api/v1/benchmarks/ai-jobs/{job.pk}/apply', {}, format='json')
        rel.refresh_from_db()
        self.assertTrue(rel.release_notes.startswith('First pilot release.'))
        self.assertEqual(rel.items.get().impact_summary, 'New: thesis screening evidence.')
        self.assertEqual(rel.status, 'draft')  # approving / publishing stay human steps
        releases.submit(rel, self.user)
        with self.assertRaises(studio_ai.StudioAIError):
            studio_ai.start('release_notes', 'Release 1.0', self.user)

    def test_college_users_cannot_use_studio_ai(self):
        college = User.objects.create_user(email='c@c.edu', username='cc', password=PASSWORD, role='iqac_coordinator')
        self.client.force_authenticate(college)
        r = self.client.post('/api/v1/benchmarks/ai-jobs', {'kind': 'benchmark_content', 'target_id': 'AIQIF-C1.1'},
                             format='json')
        self.assertEqual(r.status_code, 403)


# -- provenance: indicator origin vs supporting source references ----------------------------

from .models import Capability, StudioAIJob  # noqa: E402
from .services import provenance  # noqa: E402
from .services.requirement_drafting import accept_candidate as _accept  # noqa: E402

QUOTE = 'Every institution shall publish a policy on the responsible use of artificial intelligence'


class _ProvenanceFixture:
    """A Source Library version with extracted text (UGC-005, page 3 holds QUOTE) and a reviewer."""

    def setUp(self):
        super().setUp()
        reg = make_register([GOOD])
        self.addCleanup(os.remove, reg)
        apply_register(read_register(reg), actor='t', commit=True)
        self.version = SourceVersion.objects.get()  # UGC-005 @ 2018, 8 pages
        SourceSection.objects.create(version=self.version, ordinal=1, page_start=3, page_end=3,
                                     text=f'Clause 4. {QUOTE}, reviewed every year.')
        self.reviewer = User.objects.create_user(email='rev@ingage.ai', username='rev', password=PASSWORD,
                                                 role='consultant')

    def _import(self, source_refs=None, linked_by=None):
        path = _aiqif_workbook(source_refs=source_refs)
        self.addCleanup(os.remove, path)
        data = read_aiqif(path)
        self.assertEqual(data.errors, [])
        return data, apply_aiqif(data, actor='import_aiqif', commit=True, linked_by=linked_by,
                                 reason=f'AI-QIF import from {path}')


class ProvenanceImportTests(_ProvenanceFixture, TestCase):
    def test_every_indicator_gets_the_workbook_as_its_origin_and_no_invented_source(self):
        data, _ = self._import()
        self.assertTrue(any('no source document linked' in w for w in data.warnings))
        for code, row in (('AIQIF-C1.1', 2), ('AIQIF-C1.2', 3), ('AIQIF-C2.1', 4)):
            b = Benchmark.objects.get(benchmark_code=code)
            p = provenance.provenance(b)
            o = p['origin']
            self.assertEqual((o['type'], o['recorded'], o['workbook'], o['sheet'], o['row'], o['indicator_code'],
                              o['framework_version']),
                             ('framework_import', 'row', data.workbook, 'AI Benchmark Master', row, code, 'AI-QIF v1.0'))
            self.assertTrue(o['imported_at'])
            self.assertEqual(o['workbook_source_refs'], [])
            # The workbook defined the indicator; no source document is linked, and none is made up.
            self.assertEqual((p['source_references'], p['source_document_linked']), ([], False))
        self.assertEqual(Requirement.objects.count(), 0)

    def test_reimporting_the_same_workbook_changes_nothing(self):
        path = _aiqif_workbook()
        self.addCleanup(os.remove, path)
        apply_aiqif(read_aiqif(path), actor='a', commit=True)
        first = Benchmark.objects.get(benchmark_code='AIQIF-C1.1').origin
        summary = apply_aiqif(read_aiqif(path), actor='b', commit=True)
        self.assertEqual(summary.updated['benchmark'], 0)
        self.assertEqual(Benchmark.objects.get(benchmark_code='AIQIF-C1.1').origin, first)

    def test_an_older_import_without_a_stored_origin_reads_its_audit_event(self):
        data, _ = self._import()
        Benchmark.objects.update(origin_type='', origin={})
        o = provenance.origin_of(Benchmark.objects.get(benchmark_code='AIQIF-C1.1'))
        self.assertEqual((o['type'], o['recorded'], o['workbook'], o['sheet'], o['row'], o['imported_by']),
                         ('framework_import', 'audit', data.workbook, 'AI Benchmark Master', None, 'import_aiqif'))

    def test_an_indicator_with_no_recorded_origin_says_so(self):
        b = Benchmark.objects.create(
            benchmark_code='X-01', title='Hand-made', benchmark_kind='guidance', version_label='v1',
            capability=Capability.objects.create(code='CAP-X', name='x', domain='d', definition='d'),
            pillar=Pillar.objects.create(code='PX', name='px'))
        self.assertEqual(provenance.origin_of(b), {'type': '', 'label': 'Origin not recorded', 'recorded': ''})

    def test_workbook_source_mappings_are_linked_as_the_workbook_says_and_never_guessed(self):
        refs = [('AIQIF-C1.1', 'UGC-005', '2018', 3, QUOTE, 'primary'),
                ('AIQIF-C1.2', 'UGC-005', '', 3, 'words that are not in the source', 'supporting'),
                ('AIQIF-C2.1', 'NOPE-001', '', '', 'anything', 'supporting')]
        data, summary = self._import(source_refs=refs)
        # The workbook is authoritative: no reviewer is needed, but only a real version and a real quote link.
        self.assertEqual(summary.source_refs, {'linked': 1, 'not linked': 2})
        statuses = {b.benchmark_code: b.origin['workbook_source_refs'][0]['status'] for b in Benchmark.objects.all()}
        self.assertTrue(statuses['AIQIF-C1.1'].startswith('linked as UGC-005-R'))
        self.assertIn('the quote is not on page 3', statuses['AIQIF-C1.2'])
        self.assertIn('NOPE-001 is not in the Source Library', statuses['AIQIF-C2.1'])
        ref = provenance.source_references(Benchmark.objects.get(benchmark_code='AIQIF-C1.1'))
        self.assertEqual(len(ref), 1)
        self.assertEqual((ref[0]['source_code'], ref[0]['source_version'], ref[0]['page'], ref[0]['source_quote'],
                          ref[0]['relationship'], ref[0]['requirement_origin'], ref[0]['accepted_by'],
                          ref[0]['quote_in_source_text']),
                         ('UGC-005', '2018', 3, QUOTE, 'primary', 'manager_workbook', None, True))
        self.assertEqual((ref[0]['workbook_ref']['workbook'], ref[0]['workbook_ref']['sheet'], ref[0]['workbook_ref']['row'],
                          ref[0]['workbook_ref']['workbook_sha256']),
                         (data.workbook, 'AI-QIF Source References', 2, data.workbook_sha256))
        for code in ('AIQIF-C1.2', 'AIQIF-C2.1'):  # no mapping -> nothing linked, nothing invented
            self.assertEqual(provenance.source_references(Benchmark.objects.get(benchmark_code=code)), [])
        # Re-running links nothing twice.
        self._import(source_refs=refs)
        self.assertEqual(Requirement.objects.count(), 1)

    def test_bad_source_reference_rows_are_errors(self):
        path = _aiqif_workbook(source_refs=[('AIQIF-C9.9', 'UGC-005', '', 'x', QUOTE, 'main')])
        self.addCleanup(os.remove, path)
        err = read_aiqif(path).errors[0]
        for part in ('row 2', 'AIQIF-C9.9 is not a key indicator', "relationship 'main'", "page 'x'"):
            self.assertIn(part, err)


class ProvenanceLinkTests(_ProvenanceFixture, TestCase):
    def setUp(self):
        super().setUp()
        self._import()
        self.c11, self.c12 = (Benchmark.objects.get(benchmark_code=c) for c in ('AIQIF-C1.1', 'AIQIF-C1.2'))

    def _accepted(self):
        cand = RequirementCandidate.objects.create(version=self.version, statement='Publish an AI-use policy.',
                                                   kind='mandatory', page=3, source_quote=QUOTE, confidence=0.9,
                                                   model='test-model')
        return _accept(cand, self.reviewer)

    def test_an_accepted_requirement_keeps_source_version_page_and_quote_on_every_indicator(self):
        req = self._accepted()
        linking.link(self.c11, req, self.reviewer, relationship='primary')
        linking.link(self.c12, req, self.reviewer, relationship='supporting')
        for b, rel in ((self.c11, 'primary'), (self.c12, 'supporting')):
            ref = provenance.source_references(b)[0]
            self.assertEqual((ref['requirement_code'], ref['source_code'], ref['source_version'], ref['page'],
                              ref['source_quote'], ref['relationship'], ref['requirement_origin'], ref['ai_model']),
                             (req.code, 'UGC-005', '2018', 3, QUOTE, rel, 'ai_draft_accepted', 'test-model'))
        # The origin is still the workbook: a supporting source is not where the indicator came from.
        self.assertEqual(provenance.origin_of(self.c11)['type'], 'framework_import')

    def test_one_indicator_can_have_several_sources(self):
        first = self._accepted()
        second = provenance.add_source_reference(self.c11, self.reviewer, version=self.version, page=3,
                                                 source_quote='reviewed every year', statement='Review it yearly.',
                                                 relationship='supporting')
        linking.link(self.c11, first, self.reviewer, relationship='primary')
        codes = [r['requirement_code'] for r in provenance.source_references(self.c11)]
        self.assertEqual(sorted(codes), sorted([first.code, second.code]))

    def test_ai_drafts_never_become_source_mappings(self):
        cand = RequirementCandidate.objects.create(version=self.version, statement='Publish an AI-use policy.',
                                                   kind='mandatory', page=3, source_quote=QUOTE, confidence=0.99)
        StudioAIJob.objects.create(kind='requirement_fit', target_type='Requirement', target_id='UGC-005-R001',
                                   status='ready', model='m', output={'matches': [
                                       {'benchmark_code': 'AIQIF-C1.1', 'relationship': 'primary',
                                        'rationale': 'same topic', 'confidence': 0.95}]})
        p = provenance.provenance(self.c11)
        self.assertEqual((p['source_references'], Requirement.objects.count(), BenchmarkRequirement.objects.count()),
                         ([], 0, 0))
        self.assertNotIn('ai_suggestions', p)  # AI output is not shown as provenance at all
        _accept(cand, self.reviewer)  # even an accepted draft links to nothing by itself
        self.assertEqual(provenance.source_references(self.c11), [])

    def test_a_reviewer_source_reference_must_quote_the_source_exactly(self):
        with self.assertRaisesMessage(provenance.ProvenanceError, 'not on page 4'):
            provenance.add_source_reference(self.c11, self.reviewer, version=self.version, page=4,
                                            source_quote=QUOTE, statement='x', relationship='primary')
        with self.assertRaisesMessage(provenance.ProvenanceError, 'Page 9 is not in'):
            provenance.add_source_reference(self.c11, self.reviewer, version=self.version, page=9,
                                            source_quote=QUOTE, statement='x', relationship='primary')
        self.assertEqual(Requirement.objects.count(), 0)
        req = provenance.add_source_reference(self.c11, self.reviewer, version=self.version, page=3,
                                              source_quote=QUOTE, statement='Publish an AI-use policy.',
                                              relationship='primary', note='Cited by the manager')
        actions = list(SourceAuditEvent.objects.filter(created_at__gte=req.approved_at)
                       .order_by('created_at').values_list('action', flat=True))
        self.assertEqual(actions, ['requirement.entered', 'benchmark.linked'])

    def test_manual_origin_only_when_none_is_recorded_and_it_is_confirmed_content(self):
        with self.assertRaisesMessage(provenance.ProvenanceError, 'already has a recorded origin'):
            provenance.set_manual_origin(self.c11, self.reviewer, description='Board resolution 2026/14')
        Benchmark.objects.filter(pk=self.c11.pk).update(origin_type='', origin={})
        SourceAuditEvent.objects.filter(object_id=str(self.c11.pk), action='catalogue.create').update(reason='by hand')
        self.c11.refresh_from_db()
        before = releases.snapshot(self.c11)
        provenance.set_manual_origin(self.c11, self.reviewer, description='Board resolution 2026/14',
                                     reference='Minutes p.2')
        o = provenance.origin_of(self.c11)
        self.assertEqual((o['type'], o['description'], o['reference'], o['entered_by']),
                         ('manual', 'Board resolution 2026/14', 'Minutes p.2', 'rev@ingage.ai'))
        self.assertNotEqual(releases.snapshot(self.c11), before)  # an edit the reviewer must confirm again
        with self.assertRaisesMessage(provenance.ProvenanceError, 'reason'):
            provenance.clear_manual_origin(self.c11, self.reviewer, reason='')
        provenance.clear_manual_origin(self.c11, self.reviewer, reason='Entered on the wrong indicator')
        actions = (SourceAuditEvent.objects.filter(object_id=str(self.c11.pk), action__startswith='benchmark.origin')
                   .order_by('created_at').values_list('action', flat=True))
        self.assertEqual(list(actions), ['benchmark.origin_set', 'benchmark.origin_cleared'])

    def test_a_released_indicator_changes_provenance_only_in_a_new_draft(self):
        Benchmark.objects.filter(pk=self.c11.pk).update(status='active')
        self.c11.refresh_from_db()
        with self.assertRaisesMessage(provenance.ProvenanceError, 'new draft'):
            provenance.add_source_reference(self.c11, self.reviewer, version=self.version, page=3,
                                            source_quote=QUOTE, statement='x', relationship='primary')


class ProvenanceReleaseTests(_LinkFixture, TestCase):
    """A published release keeps the origin and source references it was published with."""

    def setUp(self):
        super().setUp()
        self.aca = Benchmark.objects.get(benchmark_code='ACA-05')
        self.approver = User.objects.create_user(email='approver@ingage.ai', username='approver', password=PASSWORD,
                                                 role='super_admin')
        make_ready(self.aca, self.user, self.req)
        rel = releases.create_release('Release 1.0', self.user)
        releases.add_item(rel, self.aca, self.user)
        releases.submit(rel, self.user)
        releases.approve(rel, self.approver)
        self.release = releases.publish(rel, self.approver)

    def test_history_keeps_the_links_after_later_draft_edits(self):
        item = self.release.items.get()
        kept = provenance.release_provenance(item)
        self.assertEqual([r['requirement_code'] for r in kept['source_references']], ['UGC-005-R001'])
        self.assertEqual(kept['source_references'][0]['source_quote'], self.req.source_quote)
        self.assertEqual(kept['origin']['label'], 'Origin not recorded')  # Source Pack fixture: none was recorded
        self.assertEqual(kept['frozen_in']['release_label'], 'Release 1.0')

        self.aca.refresh_from_db()
        releases.reopen_benchmark(self.aca, self.user, 'next version')
        linking.unlink(self.aca, self.req, self.user)
        other = Requirement.objects.create(code='UGC-005-R002', version=self.version, statement='Another.',
                                           kind='guidance', page=4, source_quote='another clause',
                                           approved_by=self.user, approved_at=_tz.now())
        linking.link(self.aca, other, self.user, relationship='supporting')
        live = [r['requirement_code'] for r in provenance.provenance(self.aca)['source_references']]
        self.assertEqual(live, ['UGC-005-R002'])
        item.refresh_from_db()
        self.assertEqual([r['requirement_code'] for r in provenance.release_provenance(item)['source_references']],
                         ['UGC-005-R001'])

    def test_a_release_published_before_provenance_was_frozen_still_shows_its_links(self):
        item = self.release.items.get()
        legacy = {k: v for k, v in item.snapshot.items() if k != 'provenance'}
        BenchmarkReleaseItem.objects.filter(pk=item.pk).update(snapshot=legacy)  # (bypasses the publish lock)
        item.refresh_from_db()
        kept = provenance.release_provenance(item)
        self.assertEqual((kept['source_references'][0]['requirement_code'], kept['source_references'][0]['page']),
                         ('UGC-005-R001', 3))
        self.assertIn('not frozen', kept['origin']['label'])


@MAINTENANCE_MODE
class ProvenanceApiTests(_ProvenanceFixture, APITestCase):
    def setUp(self):
        super().setUp()
        self._import()
        self.institution = Institution.objects.create(name='MKCE', city='Karur', state='TN')

    def _as(self, role):
        user = User.objects.create_user(email=f'{role}@x.edu', username=f'u-{role}', password=PASSWORD, role=role,
                                        institution=None if role in User.CROSS_INSTITUTION_ROLES else self.institution)
        self.client.force_authenticate(user)
        return user

    def _add(self, **over):
        body = {'version_id': str(self.version.pk), 'page': 3, 'source_quote': QUOTE,
                'statement': 'Publish an AI-use policy.', 'relationship': 'primary', **over}
        return self.client.post('/api/v1/benchmarks/catalogue/AIQIF-C1.1/source-references', body, format='json')

    def test_only_ingage_staff_read_or_change_provenance(self):
        for role in ('institution_admin', 'iqac_coordinator', 'faculty'):
            self._as(role)
            self.assertEqual(self.client.get('/api/v1/benchmarks/catalogue/AIQIF-C1.1/provenance').status_code, 403)
            self.assertEqual(self._add().status_code, 403)
            r = self.client.put('/api/v1/benchmarks/catalogue/AIQIF-C1.1/origin', {'description': 'x'}, format='json')
            self.assertEqual(r.status_code, 403)
        self.client.force_authenticate(None)
        self.assertIn(self._add().status_code, (401, 403))
        self.assertEqual(Requirement.objects.count(), 0)

    def test_staff_see_origin_add_a_source_and_navigate_both_ways(self):
        self._as('consultant')
        p = self.client.get('/api/v1/benchmarks/catalogue/AIQIF-C1.1/provenance').data
        self.assertEqual((p['origin']['type'], p['source_document_linked'], p['releases']),
                         ('framework_import', False, []))
        self.assertEqual(self._add(page=4).status_code, 400)  # not on that page: refused, nothing saved
        r = self._add()
        self.assertEqual(r.status_code, 201)
        code = r.data['requirement_code']
        self.assertEqual(r.data['source_references'][0]['requirement_code'], code)
        # Sources -> indicators, and requirement -> indicators.
        doc = self.client.get('/api/v1/benchmarks/sources').data[0]
        self.assertEqual([i['benchmark_code'] for i in doc['linked_indicators']], ['AIQIF-C1.1'])
        detail = self.client.get('/api/v1/benchmarks/sources/UGC-005').data
        self.assertEqual(detail['linked_indicators'][0]['links'], [{'requirement_code': code, 'relationship': 'primary'}])
        req = self.client.get(f'/api/v1/benchmarks/requirements?code={code}').data[0]
        self.assertEqual((req['benchmarks'][0]['benchmark_code'], req['benchmarks'][0]['title'],
                          req['requirement_origin']), ('AIQIF-C1.1', 'AI Vision', 'entered_by_reviewer'))
        by_indicator = self.client.get('/api/v1/benchmarks/requirements?benchmark=AIQIF-C1.1').data
        self.assertEqual([x['code'] for x in by_indicator], [code])

    def test_a_release_the_indicator_is_not_in_is_a_404(self):
        self._as('super_admin')
        r = self.client.get('/api/v1/benchmarks/catalogue/AIQIF-C1.1/provenance?release=Release 9')
        self.assertEqual(r.status_code, 404)


# -- editing and deleting a source (services/source_edit.py) ---------------------------------

@MAINTENANCE_MODE
class SourceEditApiTests(APITestCase):
    def setUp(self):
        reg = make_register([GOOD])  # UGC-005 with one draft version "2018"
        self.addCleanup(os.remove, reg)
        apply_register(read_register(reg), actor='t', commit=True)
        self.doc = SourceDocument.objects.get()
        self.version = self.doc.versions.get()
        self.staff = User.objects.create_user(email='c@ingage.ai', username='c', password=PASSWORD, role='consultant')
        self.client.force_authenticate(self.staff)
        self.url = '/api/v1/benchmarks/sources/UGC-005'

    def test_staff_edit_the_source_and_its_draft_version_and_it_is_audited(self):
        r = self.client.patch(self.url, {'title': 'UGC Plagiarism Regulations 2018 (amended)', 'pages': '9',
                                         'version_label': '2018 rev 1', 'notes': 'Checked'}, format='json')
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual((r.data['title'], r.data['latest_version']['version_label'], r.data['latest_version']['pages']),
                         ('UGC Plagiarism Regulations 2018 (amended)', '2018 rev 1', 9))
        events = {e.action: e for e in SourceAuditEvent.objects.filter(action__in=['source.updated', 'version.updated'])}
        self.assertEqual(events['source.updated'].before, {'title': 'UGC Plagiarism Regulations, 2018'})
        self.assertEqual(events['version.updated'].after['version_label'], '2018 rev 1')
        self.assertEqual(r.data['source_code'], 'UGC-005')  # the code never changes

    def test_bad_values_are_refused_field_by_field(self):
        r = self.client.patch(self.url, {'title': 'x', 'official_url': 'not a url', 'pages': 'many'}, format='json')
        self.assertEqual(r.status_code, 400)
        self.assertEqual(set(r.data['errors']), {'title', 'official_url', 'pages'})
        self.doc.refresh_from_db()
        self.assertEqual(self.doc.title, 'UGC Plagiarism Regulations, 2018')

    def test_an_approved_version_is_frozen_except_its_notes(self):
        self.version.status = SourceVersion.Status.ACTIVE
        self.version.save()
        r = self.client.patch(self.url, {'version_label': '2019'}, format='json')
        self.assertEqual(r.status_code, 400)
        self.assertIn('cannot change', r.data['errors']['version_label'])
        self.assertEqual(self.client.patch(self.url, {'notes': 'Still current'}, format='json').status_code, 200)

    def test_delete_needs_a_reason_and_removes_the_source_with_its_drafts(self):
        RequirementCandidate.objects.create(version=self.version, statement='s', kind='mandatory', page=1,
                                            source_quote='q', confidence=0.5)
        self.assertEqual(self.client.get('/api/v1/benchmarks/sources').data[0]['delete_blockers'], [])
        r = self.client.delete(self.url, {'reason': ''}, format='json')
        self.assertEqual(r.status_code, 400)
        r = self.client.delete(self.url, {'reason': 'Added twice by mistake'}, format='json')
        self.assertEqual(r.status_code, 204)
        self.assertFalse(SourceDocument.objects.exists())
        self.assertFalse(RequirementCandidate.objects.exists())
        event = SourceAuditEvent.objects.get(action='source.deleted')
        self.assertEqual((event.before['source_code'], event.before['pending_ai_drafts_removed'], event.reason),
                         ('UGC-005', 1, 'Added twice by mistake'))

    def test_a_source_something_depends_on_cannot_be_deleted(self):
        Requirement.objects.create(code='UGC-005-R001', version=self.version, statement='s', kind='mandatory', page=1,
                                   source_quote='q', approved_by=self.staff, approved_at=_tz.now())
        blockers = self.client.get('/api/v1/benchmarks/sources').data[0]['delete_blockers']
        self.assertEqual(blockers, ['1 accepted requirement(s) come from it.'])
        r = self.client.delete(self.url, {'reason': 'cleanup'}, format='json')
        self.assertEqual(r.status_code, 400)
        self.assertIn('cannot be deleted', r.data['detail'])
        self.assertTrue(SourceDocument.objects.exists())

    def test_the_stored_file_goes_only_when_no_other_version_uses_it(self):
        media = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, media, True)
        with override_settings(MEDIA_ROOT=media):
            rel, _sha = file_intake.store(b'%PDF-1.4 test', 'x.pdf')
            SourceVersion.objects.filter(pk=self.version.pk).update(file_uri=rel)
            with self.captureOnCommitCallbacks(execute=True):
                self.assertEqual(self.client.delete(self.url, {'reason': 'wrong file'}, format='json').status_code, 204)
            self.assertFalse(os.path.exists(os.path.join(media, rel)))

    def test_only_ingage_staff_edit_or_delete(self):
        inst = Institution.objects.create(name='MKCE', city='Karur', state='TN')
        self.client.force_authenticate(User.objects.create_user(email='a@x.edu', username='a', password=PASSWORD,
                                                                role='iqac_coordinator', institution=inst))
        self.assertEqual(self.client.patch(self.url, {'title': 'Changed title'}, format='json').status_code, 403)
        self.assertEqual(self.client.delete(self.url, {'reason': 'x'}, format='json').status_code, 403)
        self.assertTrue(SourceDocument.objects.filter(title='UGC Plagiarism Regulations, 2018').exists())

    def test_category_and_type_choices_are_listed_once(self):
        first = SourceDocument.objects.get()
        SourceDocument.objects.create(framework=first.framework, source_code='UGC-006', title='Another UGC regulation',
                                      category=first.category, document_type=first.document_type)
        ref = self.client.get('/api/v1/benchmarks/sources/reference').data
        self.assertEqual((ref['categories'], ref['document_types']), ([first.category], [first.document_type]))
        summary = self.client.get('/api/v1/benchmarks/sources/summary').data
        self.assertEqual(summary['by_category'], {first.category: 2})


# -- read-only Benchmark Studio: the manager workbooks are the standard catalogue -------------

from django.core.management import CommandError, call_command  # noqa: E402

from .models import WatcherRun  # noqa: E402
from .services.baseline import label_for, publish_workbook_baseline  # noqa: E402


class ReadOnlyCatalogueTests(_ProvenanceFixture, APITestCase):
    """With the default setting, staff can only read; every write the Studio API had is refused on the server."""

    def setUp(self):
        super().setUp()
        self._import()
        self.client.force_authenticate(self.reviewer)

    def test_staff_can_view_search_and_inspect(self):
        for url in ('/api/v1/benchmarks/sources', '/api/v1/benchmarks/sources/UGC-005', '/api/v1/benchmarks/catalogue',
                    '/api/v1/benchmarks/catalogue?q=Vision', '/api/v1/benchmarks/catalogue/AIQIF-C1.1',
                    '/api/v1/benchmarks/catalogue/AIQIF-C1.1/provenance', '/api/v1/benchmarks/pillars',
                    '/api/v1/benchmarks/applicability/preview'):
            self.assertEqual(self.client.get(url).status_code, 200, url)

    def test_every_change_to_master_data_is_refused(self):
        v = str(self.version.pk)
        writes = [
            ('post', '/api/v1/benchmarks/sources', {'source_code': 'NEW-001', 'title': 'A new source'}),
            ('patch', '/api/v1/benchmarks/sources/UGC-005', {'title': 'Changed title'}),
            ('delete', '/api/v1/benchmarks/sources/UGC-005', {'reason': 'x'}),
            ('post', f'/api/v1/benchmarks/versions/{v}/intake', {}),
            ('post', f'/api/v1/benchmarks/versions/{v}/approve', {'acknowledge_flags': True}),
            ('post', f'/api/v1/benchmarks/versions/{v}/draft-requirements', {}),
            ('put', '/api/v1/benchmarks/catalogue/AIQIF-C1.1/levels/L3', {'definition': 'x' * 80}),
            ('patch', '/api/v1/benchmarks/catalogue/AIQIF-C1.1', {'expert_owner': 'Someone'}),
            ('post', '/api/v1/benchmarks/catalogue/AIQIF-C1.1/confirm', {}),
            ('post', '/api/v1/benchmarks/catalogue/AIQIF-C1.1/requirements', {'requirement_code': 'X'}),
            ('post', '/api/v1/benchmarks/catalogue/AIQIF-C1.1/evidence-rules', {'title': 'x'}),
            ('post', '/api/v1/benchmarks/catalogue/AIQIF-C1.1/applicability-rules', {'result': 'applies'}),
            ('put', '/api/v1/benchmarks/catalogue/AIQIF-C1.1/origin', {'description': 'x'}),
            ('post', '/api/v1/benchmarks/catalogue/AIQIF-C1.1/source-references', {'version_id': v}),
            ('post', '/api/v1/benchmarks/releases', {'release_label': 'R1'}),
            ('post', '/api/v1/benchmarks/watcher/runs', {}),
            ('post', '/api/v1/benchmarks/benchmark-candidates', {'candidate_type': 'new_benchmark', 'title': 'x'}),
            ('post', '/api/v1/benchmarks/ai-jobs', {'kind': 'requirement_fit', 'target_id': 'X'}),
        ]
        before = (SourceDocument.objects.get().title, Benchmark.objects.get(benchmark_code='AIQIF-C1.1').expert_owner)
        for method, url, body in writes:
            r = getattr(self.client, method)(url, body, format='json')
            self.assertEqual(r.status_code, 403, f'{method} {url}')
            self.assertIn('read-only', str(r.data['detail']), url)
        self.assertEqual((SourceDocument.objects.get().title, Benchmark.objects.get(benchmark_code='AIQIF-C1.1').expert_owner),
                         before)
        self.assertEqual((Requirement.objects.count(), BenchmarkRelease.objects.count(), WatcherRun.objects.count(),
                          BenchmarkCandidate.objects.count(), StudioAIJob.objects.count()), (0, 0, 0, 0, 0))

    def test_the_watcher_and_ai_drafting_are_switched_off(self):
        from .tasks import run_source_watcher
        self.assertIsNone(run_source_watcher())
        self.assertEqual(WatcherRun.objects.count(), 0)
        for command in ('run_source_watcher', 'ai_draft_benchmarks'):
            with self.assertRaisesMessage(CommandError, 'read-only'):
                call_command(command)

    def test_a_missing_mapping_is_reported_not_invented(self):
        p = self.client.get('/api/v1/benchmarks/catalogue/AIQIF-C2.1/provenance').data
        self.assertEqual((p['source_references'], p['source_document_linked'], p['origin']['type']),
                         ([], False, 'framework_import'))


class WorkbookBaselineTests(_ProvenanceFixture, TestCase):
    """The controlled import freezes the standard catalogue as the release assessments use -- no portal workflow."""

    def setUp(self):
        super().setUp()
        refs = [('AIQIF-C1.1', 'UGC-005', '2018', 3, QUOTE, 'primary')]
        self.data, _ = self._import(source_refs=refs)
        self.old = BenchmarkRelease.objects.create(release_label='Old 1.0', status='published')

    def test_baseline_freezes_the_catalogue_without_changing_indicator_status(self):
        statuses = dict(Benchmark.objects.values_list('benchmark_code', 'status'))
        plan, release = publish_workbook_baseline(self.data.workbook, self.data.workbook_sha256, 'import', commit=False)
        self.assertIsNone(release)  # dry run
        self.assertEqual((plan['indicators'], plan['current'], BenchmarkRelease.objects.count()), (3, 'Old 1.0', 1))
        _plan, release = publish_workbook_baseline(self.data.workbook, self.data.workbook_sha256, 'import', commit=True)
        self.assertEqual((release.kind, release.status, release.workbook, release.workbook_sha256, release.release_label),
                         ('workbook_baseline', 'published', self.data.workbook, self.data.workbook_sha256,
                          label_for(self.data.workbook_sha256)))
        self.assertEqual(releases.current_release(), release)
        self.old.refresh_from_db()
        self.assertEqual(self.old.status, 'retired')  # kept, not deleted
        self.assertEqual(dict(Benchmark.objects.values_list('benchmark_code', 'status')), statuses)  # still drafts
        item = release.items.get(benchmark__benchmark_code='AIQIF-C1.1')
        refs = item.snapshot['provenance']['source_references']
        self.assertEqual((refs[0]['source_code'], refs[0]['page'], refs[0]['requirement_origin']),
                         ('UGC-005', 3, 'manager_workbook'))
        empty = release.items.get(benchmark__benchmark_code='AIQIF-C2.1').snapshot['provenance']['source_references']
        self.assertEqual(empty, [])
        # The same workbook again: nothing to publish.
        plan, again = publish_workbook_baseline(self.data.workbook, self.data.workbook_sha256, 'import', commit=True)
        self.assertEqual((plan['unchanged'], again, BenchmarkRelease.objects.count()), (True, None, 2))

    def test_assessments_run_on_the_baseline_and_record_the_workbook(self):
        from apps.assessments.models import BenchmarkProfile
        from apps.assessments.services import create_cycle
        from .models import InstitutionType
        _plan, release = publish_workbook_baseline(self.data.workbook, self.data.workbook_sha256, 'import', commit=True)
        college = Institution.objects.create(name='MKCE', city='Karur', state='TN')
        BenchmarkProfile.objects.create(institution=college, institution_type=InstitutionType.objects.get(code='engineering'))
        cycle = create_cycle(college, self.reviewer, 'AI readiness 2026-27', '2026-10-01')
        self.assertEqual((cycle.release, cycle.benchmarks.count()), (release, 3))
        self.assertEqual((cycle.release.workbook, cycle.release.workbook_sha256),
                         (self.data.workbook, self.data.workbook_sha256))

    def test_the_import_command_previews_then_publishes_the_baseline(self):
        from io import StringIO
        path = _aiqif_workbook()
        self.addCleanup(os.remove, path)
        out = StringIO()
        call_command('import_aiqif', path, '--baseline', stdout=out)
        text = out.getvalue()
        self.assertIn('DRY RUN', text)
        self.assertIn('baseline (dry run): would publish', text)
        self.assertEqual(BenchmarkRelease.objects.filter(kind='workbook_baseline').count(), 0)
        out = StringIO()
        call_command('import_aiqif', path, '--commit', '--baseline', stdout=out)
        self.assertIn('baseline: published', out.getvalue())
        self.assertEqual(BenchmarkRelease.objects.filter(kind='workbook_baseline', status='published').count(), 1)
